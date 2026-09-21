import openpyxl
import pathlib
import sys

sys.stdout.reconfigure(errors='replace')

ROOT = pathlib.Path('D:/behaviour analysis')
fpath = ROOT / 'Nifty211_All_Stocks_12_Quarters_WinRate_Ranking_Master.xlsx'

wb = openpyxl.load_workbook(fpath, data_only=True)
ws1 = wb['All 211 Stocks 12Q Ranking']
ws2 = wb['Qualified Stocks (69)']

print("==========================================================================================")
print("TOP 15 STOCKS IN 12-QUARTER WINDOW RANKED BY WIN RATE %")
print("==========================================================================================")

for r in range(17, 32):
    row_vals = [ws1.cell(row=r, column=c).value for c in range(1, 13)]
    rank, sym, strat, status, n_q, win_str, wr, w_c, l_c, avg_r, margin, pnl = row_vals
    print(f"Rank #{rank:<2} | {sym:<12} | {strat:<13} | {status:<20} | 12Q Win Rate: {wr:<7} | Wins: {w_c}/{w_c+l_c} | 12-Qtr Net P&L: {pnl}")

print("\nTOP 15 QUALIFIED STOCKS (>=12 QTRS) IN 12-QUARTER WINDOW RANKED BY WIN RATE %")
print("=" * 90)
for r in range(17, 32):
    row_vals = [ws2.cell(row=r, column=c).value for c in range(1, 13)]
    rank, sym, strat, status, n_q, win_str, wr, w_c, l_c, avg_r, margin, pnl = row_vals
    print(f"Rank #{rank:<2} | {sym:<12} | {strat:<13} | {n_q:<12} | 12Q Win Rate: {wr:<7} | Wins: {w_c}/{w_c+l_c} | 12-Qtr Net P&L: {pnl}")

print("==========================================================================================")
