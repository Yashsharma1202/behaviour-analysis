import os
import pathlib
import sys
import bisect
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(errors='replace')

ROOT = pathlib.Path('D:/behaviour analysis')
OI_DIR = ROOT / 'OI_DATA'
PROC = ROOT / 'processed'
PRICE_CACHE = PROC / 'price_cache'

SYMBOLS = sorted([e.name.strip().upper() for e in os.scandir(OI_DIR) if e.is_dir()])

print("==========================================================================================")
print("BUILDING QUARTERLY RESULTS PORTFOLIO & RANKINGS MASTER WORKBOOK")
print("==========================================================================================")
print(f"Target Universe: {len(SYMBOLS)} Stocks (Strictly Corporate Quarterly Results Data)")

LOT_SIZES = {
    "RELIANCE": 250, "TCS": 175, "INFY": 400, "HDFCBANK": 550, "ICICIBANK": 700,
    "BHARTIARTL": 950, "ITC": 1600, "SBIN": 1500, "LTIM": 150, "LT": 300,
    "HINDUNILVR": 300, "AXISBANK": 625, "KOTAKBANK": 400, "BAJFINANCE": 125,
    "M&M": 350, "MARUTI": 100, "SUNPHARMA": 350, "TATASTEEL": 5500,
    "NTPC": 1500, "POWERGRID": 1800, "TITAN": 175, "ADANIENT": 300,
    "ADANIPORTS": 625, "ULTRACEMCO": 100, "ASIANPAINT": 200, "COALINDIA": 2100,
    "BAJAJ-AUTO": 125, "JSWSTEEL": 675, "TATAMOTORS": 1425, "HCLTECH": 350,
    "GRASIM": 250, "HEROMOTOCO": 150, "EICHERMOT": 175, "CIPLA": 650,
    "HDFCLIFE": 1100, "SBILIFE": 375, "DRREDDY": 125, "BRITANNIA": 200,
    "APOLLOHOSP": 125, "TATACONSUM": 450, "HINDALCO": 1400, "BPCL": 1800,
    "INDUSINDBK": 500, "DIVISLAB": 200, "BAJAJFINSV": 500, "NESTLEIND": 200,
    "WIPRO": 1500, "ONGC": 3750, "TECHM": 600, "ASHOKLEY": 5000, "ADANIGREEN": 500
}

SECTOR_MAP = {
    "HDFCBANK": "Banking & Finance", "ICICIBANK": "Banking & Finance", "AXISBANK": "Banking & Finance",
    "KOTAKBANK": "Banking & Finance", "SBIN": "Banking & Finance", "INDUSINDBK": "Banking & Finance",
    "BANKBARODA": "Banking & Finance", "PNB": "Banking & Finance", "CANBK": "Banking & Finance",
    "IDFCFIRSTB": "Banking & Finance", "AUBANK": "Banking & Finance", "FEDERALBNK": "Banking & Finance",
    "RBLBANK": "Banking & Finance", "INDIANB": "Banking & Finance", "UNIONBANK": "Banking & Finance",
    
    "BAJFINANCE": "NBFC & Fintech", "BAJAJFINSV": "NBFC & Fintech", "SHRIRAMFIN": "NBFC & Fintech",
    "JIOFIN": "NBFC & Fintech", "CHOLAFIN": "NBFC & Fintech", "MUTHOOTFIN": "NBFC & Fintech",
    "REC": "NBFC & Fintech", "PFC": "NBFC & Fintech", "SBICARD": "NBFC & Fintech",
    "M&MFIN": "NBFC & Fintech", "LICHSGFIN": "NBFC & Fintech", "IRFC": "NBFC & Fintech",
    "CREDITACC": "NBFC & Fintech", "MANAPPURAM": "NBFC & Fintech", "PAYTM": "NBFC & Fintech",
    
    "HDFCLIFE": "Insurance", "SBILIFE": "Insurance", "ICICIPRULI": "Insurance", "ICICIGI": "Insurance",
    "GICRE": "Insurance", "NIACL": "Insurance", "STARHEALTH": "Insurance",
    
    "TCS": "IT & Tech Services", "INFY": "IT & Tech Services", "HCLTECH": "IT & Tech Services",
    "TECHM": "IT & Tech Services", "WIPRO": "IT & Tech Services", "LTIM": "IT & Tech Services",
    "PERSISTENT": "IT & Tech Services", "COFORGE": "IT & Tech Services", "MPHASIS": "IT & Tech Services",
    "TATAELXSI": "IT & Tech Services", "KPITTECH": "IT & Tech Services", "LTTS": "IT & Tech Services",
    "OFSS": "IT & Tech Services", "CYIENT": "IT & Tech Services", "KFINTECH": "IT & Tech Services",
    
    "MARUTI": "Automobile & Auto Ancillaries", "M&M": "Automobile & Auto Ancillaries",
    "BAJAJ-AUTO": "Automobile & Auto Ancillaries", "EICHERMOT": "Automobile & Auto Ancillaries",
    "HEROMOTOCO": "Automobile & Auto Ancillaries", "TATAMOTORS": "Automobile & Auto Ancillaries",
    "ASHOKLEY": "Automobile & Auto Ancillaries", "TVSMOTOR": "Automobile & Auto Ancillaries",
    "BHARATFORG": "Automobile & Auto Ancillaries", "MOTHERSON": "Automobile & Auto Ancillaries",
    "BOSCHLTD": "Automobile & Auto Ancillaries", "BALKRISIND": "Automobile & Auto Ancillaries",
    "MRF": "Automobile & Auto Ancillaries", "APOLLOTYRE": "Automobile & Auto Ancillaries",
    "UNOMINDA": "Automobile & Auto Ancillaries", "TIINDIA": "Automobile & Auto Ancillaries",
    "ATHERENERG": "Automobile & Auto Ancillaries", "HYUNDAI": "Automobile & Auto Ancillaries",
    
    "TATASTEEL": "Metals & Mining", "JSWSTEEL": "Metals & Mining", "HINDALCO": "Metals & Mining",
    "COALINDIA": "Metals & Mining", "JINDALSTEL": "Metals & Mining", "NMDC": "Metals & Mining",
    "VEDL": "Metals & Mining", "NATIONALUM": "Metals & Mining", "SAIL": "Metals & Mining",
    "APLAPOLLO": "Metals & Mining", "HINDZINC": "Metals & Mining",
    
    "RELIANCE": "Oil, Gas & Energy", "ONGC": "Oil, Gas & Energy", "NTPC": "Oil, Gas & Energy",
    "POWERGRID": "Oil, Gas & Energy", "BPCL": "Oil, Gas & Energy", "IOC": "Oil, Gas & Energy",
    "HPCL": "Oil, Gas & Energy", "GAIL": "Oil, Gas & Energy", "ADANIGREEN": "Oil, Gas & Energy",
    "ADANIPOWER": "Oil, Gas & Energy", "TATAPOWER": "Oil, Gas & Energy", "NHPC": "Oil, Gas & Energy",
    "SJVN": "Oil, Gas & Energy", "OIL": "Oil, Gas & Energy", "SUZLON": "Oil, Gas & Energy",
    "INOXWIND": "Oil, Gas & Energy", "IREDA": "Oil, Gas & Energy", "WAAREEENER": "Oil, Gas & Energy",
    
    "SUNPHARMA": "Pharma & Healthcare", "CIPLA": "Pharma & Healthcare", "DRREDDY": "Pharma & Healthcare",
    "APOLLOHOSP": "Pharma & Healthcare", "DIVISLAB": "Pharma & Healthcare", "TORNTPHARM": "Pharma & Healthcare",
    "MANKIND": "Pharma & Healthcare", "LUPIN": "Pharma & Healthcare", "ZYDUSLIFE": "Pharma & Healthcare",
    "ALKEM": "Pharma & Healthcare", "BIOCON": "Pharma & Healthcare", "MAXHEALTH": "Pharma & Healthcare",
    "SYNGENE": "Pharma & Healthcare", "FORTIS": "Pharma & Healthcare", "LAURUSLABS": "Pharma & Healthcare",
    "SAGILITY": "Pharma & Healthcare",
    
    "HINDUNILVR": "FMCG & Consumer Goods", "ITC": "FMCG & Consumer Goods", "NESTLEIND": "FMCG & Consumer Goods",
    "TATACONSUM": "FMCG & Consumer Goods", "BRITANNIA": "FMCG & Consumer Goods", "DABUR": "FMCG & Consumer Goods",
    "GODREJCP": "FMCG & Consumer Goods", "MARICO": "FMCG & Consumer Goods", "COLPAL": "FMCG & Consumer Goods",
    "VBL": "FMCG & Consumer Goods", "PGHH": "FMCG & Consumer Goods", "UBL": "FMCG & Consumer Goods",
    "MCDOWELL-N": "FMCG & Consumer Goods", "GODFRYPHLP": "FMCG & Consumer Goods",
    
    "TITAN": "Consumer Durables & Retail", "TRENT": "Consumer Durables & Retail",
    "ASIANPAINT": "Consumer Durables & Retail", "BERGEPAINT": "Consumer Durables & Retail",
    "PIDILITIND": "Consumer Durables & Retail", "HAVELLS": "Consumer Durables & Retail",
    "DIXON": "Consumer Durables & Retail", "VOLTAS": "Consumer Durables & Retail",
    "POLYCAB": "Consumer Durables & Retail", "CROMPTON": "Consumer Durables & Retail",
    "KEI": "Consumer Durables & Retail", "WHIRLPOOL": "Consumer Durables & Retail",
    "DMART": "Consumer Durables & Retail", "KALYANKJIL": "Consumer Durables & Retail",
    
    "L&T": "Capital Goods & Infrastructure", "LT": "Capital Goods & Infrastructure",
    "SIEMENS": "Capital Goods & Infrastructure", "ABB": "Capital Goods & Infrastructure",
    "BEL": "Capital Goods & Infrastructure", "HAL": "Capital Goods & Infrastructure",
    "BHEL": "Capital Goods & Infrastructure", "CGPOWER": "Capital Goods & Infrastructure",
    "CUMMINSIND": "Capital Goods & Infrastructure", "TITAGARH": "Capital Goods & Infrastructure",
    "RAILTEL": "Capital Goods & Infrastructure", "RITES": "Capital Goods & Infrastructure",
    "MAZDOCK": "Capital Goods & Infrastructure", "COCHINSHIP": "Capital Goods & Infrastructure",
    "BDL": "Capital Goods & Infrastructure", "IRCON": "Capital Goods & Infrastructure",
    
    "DLF": "Realty & Construction", "LODHA": "Realty & Construction", "GODREJPROP": "Realty & Construction",
    "OBERREALTY": "Realty & Construction", "PHOENIXLTD": "Realty & Construction", "PRESTIGE": "Realty & Construction"
}

pre_candidates = [1, 2, 3, 4, 5, 7, 8]
post_candidates = [1, 2, 3, 4, 5, 8]

def assign_financial_quarter(row):
    rel = str(row.get('relatingTo', '')).lower()
    if 'first' in rel or 'q1' in rel: return 'Q1'
    elif 'second' in rel or 'q2' in rel: return 'Q2'
    elif 'third' in rel or 'q3' in rel: return 'Q3'
    elif 'fourth' in rel or 'q4' in rel: return 'Q4'
        
    to_dt = row.get('to_dt', None)
    if pd.notnull(to_dt) and hasattr(to_dt, 'month'):
        m = to_dt.month
        if m in (4, 5, 6): return 'Q1'
        elif m in (7, 8, 9): return 'Q2'
        elif m in (10, 11, 12): return 'Q3'
        elif m in (1, 2, 3): return 'Q4'

    bc_dt = row.get('dt', None)
    if pd.notnull(bc_dt) and hasattr(bc_dt, 'month'):
        m = bc_dt.month
        if m in (7, 8, 9): return 'Q1'
        elif m in (10, 11, 12): return 'Q2'
        elif m in (1, 2, 3): return 'Q3'
        else: return 'Q4'
        
    return 'Q1'

def get_financial_quarter_display_name(q_code):
    names = {
        "Q1": "Q1 (April - June Results)",
        "Q2": "Q2 (July - September Results)",
        "Q3": "Q3 (October - December Results)",
        "Q4": "Q4 (January - March Results)"
    }
    return names.get(q_code, "N/A")

quarter_data_store = {
    "ALL": [], "Q1": [], "Q2": [], "Q3": [], "Q4": []
}

stock_quarterly_returns = {}

for idx, sym in enumerate(SYMBOLS, 1):
    stock_folder = ROOT / sym
    fr_path = stock_folder / 'financial_results.csv'
    price_path = PRICE_CACHE / f"{sym}.csv"
    
    df_prices = pd.DataFrame()
    if price_path.exists():
        try:
            df_prices = pd.read_csv(price_path)
            df_prices['date'] = pd.to_datetime(df_prices['date'])
            df_prices = df_prices.sort_values('date').reset_index(drop=True)
        except Exception:
            pass
            
    if df_prices.empty:
        continue
        
    df_fr_clean = pd.DataFrame()
    first_q_name = "N/A"
    
    if fr_path.exists() and fr_path.stat().st_size > 10:
        try:
            df_fr = pd.read_csv(fr_path, dtype=str).fillna('')
            if not df_fr.empty and 'broadCastDate' in df_fr.columns:
                df_fr['dt'] = pd.to_datetime(df_fr['broadCastDate'], format='mixed', errors='coerce')
                if 'toDate' in df_fr.columns:
                    df_fr['to_dt'] = pd.to_datetime(df_fr['toDate'], format='mixed', errors='coerce')
                else:
                    df_fr['to_dt'] = None
                df_fr = df_fr.dropna(subset=['dt']).sort_values('dt', ascending=True)
                
                if not df_fr.empty:
                    earliest_row = df_fr.iloc[0]
                    first_q_code = assign_financial_quarter(earliest_row)
                    first_q_name = get_financial_quarter_display_name(first_q_code)
                
                df_fr = df_fr.sort_values('dt', ascending=False)
                df_fr['year'] = df_fr['dt'].dt.year
                df_fr['q_type'] = df_fr.apply(assign_financial_quarter, axis=1)
                
                df_fr_clean = df_fr.drop_duplicates(subset=['year', 'q_type']).copy()
        except Exception:
            pass
            
    n_quarters_total = len(df_fr_clean)
    trading_status = "QUALIFIED" if n_quarters_total >= 12 else "AVOID BUT MONITOR IT"
    sector_name = SECTOR_MAP.get(sym, "Diversified / Others")
    valid_dates = df_prices['date'].tolist()
    if not valid_dates:
        continue
        
    lot_size = LOT_SIZES.get(sym, max(100, int(1000000 / float(df_prices['adj'].iloc[-1]))))
    last_price = float(df_prices['adj'].iloc[-1])
    margin_for_stock = round(0.20 * lot_size * last_price, 2)
    
    stock_quarterly_returns[sym] = {}

    for q_t in ["ALL", "Q1", "Q2", "Q3", "Q4"]:
        if q_t == "ALL":
            df_q_sub = df_fr_clean
        else:
            df_q_sub = df_fr_clean[df_fr_clean['q_type'] == q_t] if not df_fr_clean.empty else pd.DataFrame()
            
        tot_q_count = len(df_q_sub)
        
        if tot_q_count == 0:
            quarter_data_store[q_t].append({
                "sym": sym, "sector": sector_name,
                "first_q_name": first_q_name,
                "trading_status": trading_status,
                "n_quarters_total": f"{n_quarters_total} Quarters",
                "n_quarters_qtype": "0 Quarters", "window": "N/A", "strategy": "N/A",
                "win_rate": 0.0, "wins": 0, "losses": 0, "win_ratio_str": "N/A", "avg_ret": 0.0,
                "margin": margin_for_stock, "net_pnl": 0.0, "is_available": False
            })
            continue

        best_pnl = -999999999
        best_eval = None

        for pre in pre_candidates:
            for post in post_candidates:
                emp_long_wins = 0; emp_short_wins = 0
                rets_long = []; rets_short = []
                pnl_long_total = 0.0; pnl_short_total = 0.0
                q_ret_map_long = {}; q_ret_map_short = {}

                for _, row in df_q_sub.iterrows():
                    dt = row['dt']
                    q_key = f"{row['year']}_{row['q_type']}"
                    pos = bisect.bisect_left(valid_dates, dt)
                    if pos == 0: r_i = 0
                    elif pos >= len(valid_dates): r_i = len(valid_dates) - 1
                    else:
                        r_i = pos if (valid_dates[pos] - dt) < (dt - valid_dates[pos - 1]) else (pos - 1)

                    e_i = max(0, r_i - pre)
                    x_i = min(len(valid_dates) - 1, r_i + post)
                    pe = float(df_prices.loc[e_i, 'adj']); px = float(df_prices.loc[x_i, 'adj'])
                    if pe > 0:
                        rl = (px - pe) / pe * 100
                        rs = (pe - px) / pe * 100
                        rets_long.append(rl); rets_short.append(rs)
                        q_ret_map_long[q_key] = rl
                        q_ret_map_short[q_key] = rs
                        
                        b_to_l = round(lot_size * pe, 2); s_to_l = round(lot_size * px, 2)
                        cost_l = round((b_to_l + s_to_l) * 0.0005, 2)
                        net_pnl_l = round(s_to_l - b_to_l - cost_l, 2)
                        pnl_long_total += net_pnl_l
                        
                        b_to_s = round(lot_size * px, 2); s_to_s = round(lot_size * pe, 2)
                        cost_s = round((b_to_s + s_to_s) * 0.0005, 2)
                        net_pnl_s = round(s_to_s - b_to_s - cost_s, 2)
                        pnl_short_total += net_pnl_s
                        
                        if net_pnl_l > 0: emp_long_wins += 1
                        if net_pnl_s > 0: emp_short_wins += 1

                wr_long = (emp_long_wins / tot_q_count) * 100
                wr_short = (emp_short_wins / tot_q_count) * 100
                
                if wr_long >= wr_short:
                    opt_strat = "FUTURE LONG"
                    opt_wr = round(wr_long, 2)
                    opt_wins = emp_long_wins
                    opt_losses = tot_q_count - emp_long_wins
                    opt_avg_ret = round(np.mean(rets_long), 2) if rets_long else 0.0
                    opt_pnl = round(pnl_long_total, 2)
                    opt_q_map = q_ret_map_long
                else:
                    opt_strat = "FUTURE SHORT"
                    opt_wr = round(wr_short, 2)
                    opt_wins = emp_short_wins
                    opt_losses = tot_q_count - emp_short_wins
                    opt_avg_ret = round(np.mean(rets_short), 2) if rets_short else 0.0
                    opt_pnl = round(pnl_short_total, 2)
                    opt_q_map = q_ret_map_short

                if opt_pnl > best_pnl:
                    best_pnl = opt_pnl
                    best_eval = {
                        "pre": pre, "post": post, "window": f"T-{pre} to T+{post}",
                        "strategy": opt_strat, "win_rate": opt_wr,
                        "wins": opt_wins, "losses": opt_losses,
                        "avg_ret": opt_avg_ret, "pnl": opt_pnl,
                        "q_map": opt_q_map
                    }

        window_str = best_eval["window"]
        opt_strat = best_eval["strategy"]
        opt_wr = best_eval["win_rate"]
        opt_wins = best_eval["wins"]
        opt_losses = best_eval["losses"]
        opt_avg_ret = best_eval["avg_ret"]
        opt_pnl = best_eval["pnl"]
        
        win_ratio_str = f"{opt_wins} Wins / {tot_q_count} Qtrs"

        if q_t == "ALL":
            stock_quarterly_returns[sym] = {
                "strategy": opt_strat,
                "q_map": best_eval["q_map"],
                "window": window_str
            }

        quarter_data_store[q_t].append({
            "sym": sym, "sector": sector_name,
            "first_q_name": first_q_name,
            "trading_status": trading_status,
            "n_quarters_total": f"{n_quarters_total} Quarters",
            "n_quarters_qtype": f"{tot_q_count} Quarters", "window": window_str, "strategy": opt_strat,
            "win_rate": opt_wr, "wins": opt_wins, "losses": opt_losses, "win_ratio_str": win_ratio_str,
            "avg_ret": opt_avg_ret, "margin": margin_for_stock, "net_pnl": opt_pnl, "is_available": True
        })

# ALL QUARTERS COMBINED DataFrame
df_all = pd.DataFrame(quarter_data_store["ALL"])
df_qual = df_all[(df_all['is_available'] == True) & (df_all['trading_status'] == 'QUALIFIED')].sort_values(['net_pnl', 'win_rate', 'avg_ret'], ascending=[False, False, False])

top_10_stocks = df_qual.head(10)['sym'].tolist()

all_q_keys = set()
for sym in top_10_stocks:
    all_q_keys.update(stock_quarterly_returns[sym]["q_map"].keys())

sorted_q_keys = sorted(list(all_q_keys), key=lambda x: (int(x.split('_')[0]), x.split('_')[1]))

portfolio_sim_25l = []
portfolio_sim_50l = []

capital_25l = 2500000.0  # ₹25 Lakhs
capital_50l = 5000000.0  # ₹50 Lakhs

alloc_25l_per_trade = capital_25l / 10.0
alloc_50l_per_trade = capital_50l / 10.0

curr_bal_25l = capital_25l
curr_bal_50l = capital_50l

eq_curve_25l = [capital_25l]
eq_curve_50l = [capital_50l]

for q_k in sorted_q_keys:
    yr, q_type = q_k.split('_')
    q_disp = f"{yr} {q_type} Results"
    
    q_pnl_25l = 0.0
    q_pnl_50l = 0.0
    
    n_active_trades = 0
    trade_details = []
    
    for sym in top_10_stocks:
        q_map = stock_quarterly_returns[sym]["q_map"]
        if q_k in q_map:
            ret_pct = q_map[q_k]
            n_active_trades += 1
            
            pos_val_25l = alloc_25l_per_trade * 5.0
            pos_val_50l = alloc_50l_per_trade * 5.0
            
            trade_pnl_25l = round(pos_val_25l * (ret_pct / 100.0) - (pos_val_25l * 0.0005), 2)
            trade_pnl_50l = round(pos_val_50l * (ret_pct / 100.0) - (pos_val_50l * 0.0005), 2)
            
            q_pnl_25l += trade_pnl_25l
            q_pnl_50l += trade_pnl_50l
            trade_details.append(f"{sym}: {ret_pct:+.2f}%")

    start_25l = curr_bal_25l
    start_50l = curr_bal_50l
    
    curr_bal_25l += q_pnl_25l
    curr_bal_50l += q_pnl_50l
    
    eq_curve_25l.append(curr_bal_25l)
    eq_curve_50l.append(curr_bal_50l)
    
    ret_q_pct_25l = round((q_pnl_25l / start_25l) * 100.0, 2)
    ret_q_pct_50l = round((q_pnl_50l / start_50l) * 100.0, 2)
    
    portfolio_sim_25l.append({
        "quarter": q_disp,
        "active_trades": n_active_trades,
        "start_bal": start_25l,
        "q_pnl": q_pnl_25l,
        "q_ret_pct": ret_q_pct_25l,
        "end_bal": curr_bal_25l,
        "trades_summary": ", ".join(trade_details)
    })

    portfolio_sim_50l.append({
        "quarter": q_disp,
        "active_trades": n_active_trades,
        "start_bal": start_50l,
        "q_pnl": q_pnl_50l,
        "q_ret_pct": ret_q_pct_50l,
        "end_bal": curr_bal_50l,
        "trades_summary": ", ".join(trade_details)
    })

def get_max_drawdown_pct(eq_list):
    arr = np.array(eq_list)
    pk = np.maximum.accumulate(arr)
    dd = (arr - pk) / pk * 100.0
    return round(abs(float(np.min(dd))), 2)

tot_pnl_25l = round(curr_bal_25l - capital_25l, 2)
tot_ret_25l = round((tot_pnl_25l / capital_25l) * 100.0, 2)
max_dd_25l = get_max_drawdown_pct(eq_curve_25l)

tot_pnl_50l = round(curr_bal_50l - capital_50l, 2)
tot_ret_50l = round((tot_pnl_50l / capital_50l) * 100.0, 2)
max_dd_50l = get_max_drawdown_pct(eq_curve_50l)

out_path_q_master = ROOT / "Nifty211_Quarterly_Results_Portfolio_Master.xlsx"
out_path_q_clean = ROOT / "Nifty211_Quarterly_Performance_Master_Clean.xlsx"
out_path_q_dynamic = ROOT / "Nifty211_Dynamic_Quarterly_Performance_Master.xlsx"

wb = openpyxl.Workbook()

title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
title_fill = PatternFill("solid", fgColor="1F4E79")
card_val_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
hdr_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1F4E79")
subhdr_fill = PatternFill("solid", fgColor="2F5597")
border_thin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                     top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
card_fill = PatternFill("solid", fgColor="F2F4F7")
pos_fill = PatternFill("solid", fgColor="E2EFDA")
neg_fill = PatternFill("solid", fgColor="FCE4D6")
amb_fill = PatternFill("solid", fgColor="FFF2CC")
pos_font = Font(name="Calibri", size=11, bold=True, color="276A3C")
neg_font = Font(name="Calibri", size=11, bold=True, color="9C0006")
amb_font = Font(name="Calibri", size=11, bold=True, color="B25900")

# SHEET 1: PORTFOLIO ALLOCATOR
ws_p = wb.active
ws_p.title = "PORTFOLIO ALLOCATOR (25L & 50L)"
ws_p.views.sheetView[0].showGridLines = True

ws_p.merge_cells("A1:N1")
ws_p["A1"] = "REAL PORTFOLIO CAPITAL ALLOCATOR — CORPORATE QUARTERLY RESULTS SIMULATION"
ws_p["A1"].font = title_font; ws_p["A1"].fill = title_fill
ws_p["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws_p.row_dimensions[1].height = 40

port_cards = [
    ("₹25 LAKHS PORTFOLIO NET P&L", f"₹{tot_pnl_25l:,.2f}\n(+{tot_ret_25l}%)", "A3:C4", "A3"),
    ("₹25L MAX DRAWDOWN %", f"-{max_dd_25l}%", "D3:F4", "D3"),
    ("₹50 LAKHS PORTFOLIO NET P&L", f"₹{tot_pnl_50l:,.2f}\n(+{tot_ret_50l}%)", "G3:I4", "G3"),
    ("₹50L MAX DRAWDOWN %", f"-{max_dd_50l}%", "J3:L4", "J3"),
    ("TOP 10 QUALIFIED STOCKS", f"{', '.join(top_10_stocks[:5])}\n{', '.join(top_10_stocks[5:])}", "M3:N4", "M3")
]
for title, val, merge_range, top_left in port_cards:
    ws_p.merge_cells(merge_range)
    ws_p[top_left] = f"{title}\n{val}"
    ws_p[top_left].font = card_val_font
    ws_p[top_left].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws_p[top_left].fill = card_fill; ws_p[top_left].border = border_thin

ws_p.cell(row=6, column=1, value="QUARTERLY RESULTS TIMELINE EQUITY GROWTH (TOP 10 QUALIFIED STOCKS PORTFOLIO)").font = hdr_font
ws_p.cell(row=6, column=1).fill = hdr_fill
ws_p.merge_cells("A6:N6")

p_headers = [
    "Quarter Results Timeline", "Active Trades Count",
    "₹25L Start Capital (₹)", "₹25L Quarter P&L (₹)", "₹25L Return %", "₹25L Ending Equity (₹)",
    "₹50L Start Capital (₹)", "₹50L Quarter P&L (₹)", "₹50L Return %", "₹50L Ending Equity (₹)",
    "Individual Stock Earnings Trade Details"
]

ws_p.row_dimensions[7].height = 28
c_i = 1
for h in p_headers:
    if h == "Individual Stock Earnings Trade Details":
        ws_p.merge_cells("K7:N7")
        cell = ws_p.cell(row=7, column=11, value=h)
        cell.font = hdr_font; cell.fill = subhdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    else:
        cell = ws_p.cell(row=7, column=c_i, value=h)
        cell.font = hdr_font; cell.fill = subhdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c_i += 1

for r_idx, (r25, r50) in enumerate(zip(portfolio_sim_25l, portfolio_sim_50l), start=8):
    ws_p.cell(row=r_idx, column=1, value=r25['quarter']).border = border_thin
    ws_p.cell(row=r_idx, column=1).alignment = Alignment(horizontal="center")
    ws_p.cell(row=r_idx, column=1).font = Font(bold=True)
    
    ws_p.cell(row=r_idx, column=2, value=r25['active_trades']).border = border_thin
    ws_p.cell(row=r_idx, column=2).alignment = Alignment(horizontal="center")
    
    ws_p.cell(row=r_idx, column=3, value=f"₹{r25['start_bal']:,.2f}").border = border_thin
    ws_p.cell(row=r_idx, column=3).alignment = Alignment(horizontal="right")
    
    c_pnl25 = ws_p.cell(row=r_idx, column=4, value=f"₹{r25['q_pnl']:,.2f}")
    c_pnl25.border = border_thin; c_pnl25.alignment = Alignment(horizontal="right")
    if r25['q_pnl'] >= 0: c_pnl25.fill = pos_fill; c_pnl25.font = pos_font
    else: c_pnl25.fill = neg_fill; c_pnl25.font = neg_font
    
    c_ret25 = ws_p.cell(row=r_idx, column=5, value=f"{r25['q_ret_pct']}%")
    c_ret25.border = border_thin; c_ret25.alignment = Alignment(horizontal="right")
    if r25['q_ret_pct'] >= 0: c_ret25.fill = pos_fill; c_ret25.font = pos_font
    else: c_ret25.fill = neg_fill; c_ret25.font = neg_font
    
    ws_p.cell(row=r_idx, column=6, value=f"₹{r25['end_bal']:,.2f}").border = border_thin
    ws_p.cell(row=r_idx, column=6).alignment = Alignment(horizontal="right")
    ws_p.cell(row=r_idx, column=6).font = Font(bold=True)
    
    ws_p.cell(row=r_idx, column=7, value=f"₹{r50['start_bal']:,.2f}").border = border_thin
    ws_p.cell(row=r_idx, column=7).alignment = Alignment(horizontal="right")
    
    c_pnl50 = ws_p.cell(row=r_idx, column=8, value=f"₹{r50['q_pnl']:,.2f}")
    c_pnl50.border = border_thin; c_pnl50.alignment = Alignment(horizontal="right")
    if r50['q_pnl'] >= 0: c_pnl50.fill = pos_fill; c_pnl50.font = pos_font
    else: c_pnl50.fill = neg_fill; c_pnl50.font = neg_font
    
    c_ret50 = ws_p.cell(row=r_idx, column=9, value=f"{r50['q_ret_pct']}%")
    c_ret50.border = border_thin; c_ret50.alignment = Alignment(horizontal="right")
    if r50['q_ret_pct'] >= 0: c_ret50.fill = pos_fill; c_ret50.font = pos_font
    else: c_ret50.fill = neg_fill; c_ret50.font = neg_font
    
    ws_p.cell(row=r_idx, column=10, value=f"₹{r50['end_bal']:,.2f}").border = border_thin
    ws_p.cell(row=r_idx, column=10).alignment = Alignment(horizontal="right")
    ws_p.cell(row=r_idx, column=10).font = Font(bold=True)
    
    ws_p.merge_cells(start_row=r_idx, start_column=11, end_row=r_idx, end_column=14)
    ws_p.cell(row=r_idx, column=11, value=r25['trades_summary']).border = border_thin
    ws_p.cell(row=r_idx, column=11).alignment = Alignment(horizontal="left")

for col in ws_p.columns:
    col_letter = get_column_letter(col[0].column)
    ws_p.column_dimensions[col_letter].width = 24

# SHEETS 2-6: ALL QUARTERS + Q1, Q2, Q3, Q4
quarter_names = {
    "ALL": "ALL QUARTERS COMBINED",
    "Q1": "Q1 (April - June Results)",
    "Q2": "Q2 (July - September Results)",
    "Q3": "Q3 (October - December Results)",
    "Q4": "Q4 (January - March Results)"
}

for q_t in ["ALL", "Q1", "Q2", "Q3", "Q4"]:
    sheet_title = quarter_names[q_t]
    ws = wb.create_sheet(title=sheet_title)
    ws.views.sheetView[0].showGridLines = True
    
    raw_list = quarter_data_store[q_t]
    df_qtype = pd.DataFrame(raw_list)
    df_avail = df_qtype[df_qtype['is_available'] == True].sort_values(['net_pnl', 'win_rate', 'avg_ret'], ascending=[False, False, False])
    df_unavail = df_qtype[df_qtype['is_available'] == False].sort_values('sym', ascending=True)
    df_sorted = pd.concat([df_avail, df_unavail], ignore_index=True)
    
    ws.merge_cells("A1:O1")
    ws["A1"] = f"QUARTERLY RESULTS PERFORMANCE MASTER — {sheet_title.upper()}"
    ws["A1"].font = title_font; ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    top_stock = df_avail.iloc[0]['sym'] if not df_avail.empty else "N/A"
    top_pnl = df_avail.iloc[0]['net_pnl'] if not df_avail.empty else 0.0
    avg_ret_q = round(df_avail['avg_ret'].mean(), 2) if not df_avail.empty else 0.0
    avg_wr_q = round(df_avail['win_rate'].mean(), 2) if not df_avail.empty else 0.0
    tot_pnl_q = round(df_avail['net_pnl'].sum(), 2) if not df_avail.empty else 0.0
    
    cards = [
        ("TOTAL STOCKS RANKED", f"{len(df_sorted)} Stocks ({len(df_avail)} Avail)", "A3:C4", "A3"),
        ("TOP RANKED STOCK (NET P&L)", f"{top_stock} (₹{top_pnl:,.2f})", "D3:F4", "D3"),
        ("AVERAGE QUARTER RETURN %", f"{avg_ret_q}%", "G3:I4", "G3"),
        ("AVERAGE QUARTER WIN RATE %", f"{avg_wr_q}%", "J3:L4", "J3"),
        ("TOTAL QUARTER NET P&L", f"₹{tot_pnl_q:,.2f}", "M3:O4", "M3")
    ]
    for title, val, merge_range, top_left in cards:
        ws.merge_cells(merge_range)
        ws[top_left] = f"{title}\n{val}"
        ws[top_left].font = card_val_font
        ws[top_left].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws[top_left].fill = card_fill; ws[top_left].border = border_thin

    rank_start_row = 6
    ws.cell(row=rank_start_row, column=1, value=f"ALL {len(df_sorted)} STOCKS RANKED BY {sheet_title.upper()} NET REALISED P&L").font = hdr_font
    ws.cell(row=rank_start_row, column=1).fill = hdr_fill
    ws.merge_cells(start_row=rank_start_row, start_column=1, end_row=rank_start_row, end_column=15)

    headers = [
        "Rank", "Stock Symbol", "Industry Sector",
        "Listing Start Quarter Name",
        "Strategy", "Action Status", "Total History Quarters", "Quarter Data Count",
        "Position Window", "Win Rate %", "Win Ratio", "Losing Quarters",
        "Average Return %", "Margin Capital (₹)", "Net Realised P&L (₹)"
    ]
    ws.row_dimensions[rank_start_row+1].height = 28
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=rank_start_row+1, column=c_idx, value=h)
        cell.font = hdr_font; cell.fill = subhdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, r in enumerate(df_sorted.to_dict('records'), start=rank_start_row+2):
        ws.cell(row=r_idx, column=1, value=r_idx - rank_start_row - 1).border = border_thin
        ws.cell(row=r_idx, column=1).alignment = Alignment(horizontal="center")
        
        ws.cell(row=r_idx, column=2, value=r['sym']).border = border_thin
        ws.cell(row=r_idx, column=2).font = Font(bold=True)
        
        ws.cell(row=r_idx, column=3, value=r['sector']).border = border_thin
        ws.cell(row=r_idx, column=3).alignment = Alignment(horizontal="center")
        
        ws.cell(row=r_idx, column=4, value=r['first_q_name']).border = border_thin
        ws.cell(row=r_idx, column=4).alignment = Alignment(horizontal="center")

        c_st = ws.cell(row=r_idx, column=5, value=r['strategy'])
        c_st.border = border_thin; c_st.alignment = Alignment(horizontal="center")
        if r['strategy'] == 'FUTURE LONG': c_st.fill = pos_fill; c_st.font = pos_font
        elif r['strategy'] == 'FUTURE SHORT': c_st.fill = neg_fill; c_st.font = neg_font
        else: c_st.fill = card_fill; c_st.font = Font(color="7F7F7F", italic=True)

        c_act = ws.cell(row=r_idx, column=6, value=r['trading_status'])
        c_act.border = border_thin; c_act.alignment = Alignment(horizontal="center")
        if r['trading_status'] == 'QUALIFIED': c_act.fill = pos_fill; c_act.font = pos_font
        else: c_act.fill = amb_fill; c_act.font = amb_font
        
        ws.cell(row=r_idx, column=7, value=r['n_quarters_total']).border = border_thin
        ws.cell(row=r_idx, column=7).alignment = Alignment(horizontal="center")
        
        ws.cell(row=r_idx, column=8, value=r['n_quarters_qtype']).border = border_thin
        ws.cell(row=r_idx, column=8).alignment = Alignment(horizontal="center")
        
        ws.cell(row=r_idx, column=9, value=r['window']).border = border_thin
        ws.cell(row=r_idx, column=9).alignment = Alignment(horizontal="center")
        
        wr_val_str = f"{r['win_rate']}%" if r['is_available'] else "N/A"
        c_wr = ws.cell(row=r_idx, column=10, value=wr_val_str)
        c_wr.border = border_thin; c_wr.alignment = Alignment(horizontal="right")
        c_wr.font = Font(bold=True)
        if r['is_available']:
            if r['win_rate'] >= 65.0: c_wr.fill = pos_fill; c_wr.font = pos_font
            elif r['win_rate'] < 50.0: c_wr.fill = neg_fill; c_wr.font = neg_font
        else:
            c_wr.alignment = Alignment(horizontal="center"); c_wr.font = Font(color="7F7F7F", italic=True)
        
        c_w = ws.cell(row=r_idx, column=11, value=r['win_ratio_str'])
        c_w.border = border_thin; c_w.alignment = Alignment(horizontal="center")
        if not r['is_available']: c_w.font = Font(color="7F7F7F", italic=True)
        
        c_l = ws.cell(row=r_idx, column=12, value=r['losses'] if r['is_available'] else "N/A")
        c_l.border = border_thin; c_l.alignment = Alignment(horizontal="center")
        if not r['is_available']: c_l.font = Font(color="7F7F7F", italic=True)
        
        ret_val_str = f"{r['avg_ret']}%" if r['is_available'] else "N/A"
        c_ret = ws.cell(row=r_idx, column=13, value=ret_val_str)
        c_ret.border = border_thin; c_ret.alignment = Alignment(horizontal="right")
        c_ret.font = Font(bold=True)
        if r['is_available']:
            if r['avg_ret'] >= 0: c_ret.fill = pos_fill; c_ret.font = pos_font
            else: c_ret.fill = neg_fill; c_ret.font = neg_font
        else:
            c_ret.alignment = Alignment(horizontal="center"); c_ret.font = Font(color="7F7F7F", italic=True)
        
        ws.cell(row=r_idx, column=14, value=f"₹{r['margin']:,.2f}").border = border_thin
        ws.cell(row=r_idx, column=14).alignment = Alignment(horizontal="right")
        
        pnl_val_str = f"₹{r['net_pnl']:,.2f}" if r['is_available'] else "N/A"
        c_p = ws.cell(row=r_idx, column=15, value=pnl_val_str)
        c_p.border = border_thin; c_p.alignment = Alignment(horizontal="right")
        if r['is_available']:
            if r['net_pnl'] >= 0: c_p.fill = pos_fill; c_p.font = pos_font
            else: c_p.fill = neg_fill; c_p.font = neg_font
        else:
            c_p.alignment = Alignment(horizontal="center"); c_p.font = Font(color="7F7F7F", italic=True)

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = 24

saved_paths = []
target_list = [
    out_path_q_master, out_path_q_clean, out_path_q_dynamic
]
for target_path in target_list:
    try:
        wb.save(target_path)
        print(f"  • Successfully generated: {target_path.name}")
        saved_paths.append(target_path)
    except PermissionError:
        print(f"  • Warning: {target_path.name} is currently open and locked by Excel!")

print("==========================================================================================")
print(f"QUARTERLY RESULTS PORTFOLIO & RANKINGS MASTER CREATED SUCCESSFULLY! ({len(saved_paths)} files saved)")
print("==========================================================================================")
