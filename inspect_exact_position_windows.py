import openpyxl, pandas as pd, sys, pathlib

sys.stdout.reconfigure(errors='replace')

b_dir = pathlib.Path('D:/behaviour analysis')

master_path = b_dir / 'Nifty50_12_Quarters_Futures_0.05PCT_Cost_Master.xlsx'
wb = openpyxl.load_workbook(master_path, data_only=True)

ws = wb['Q3 2023-24']

print("==========================================================================================")
print("INSPECTING EXACT POSITION TAKING WINDOW VALUES IN MASTER WORKBOOK")
print("==========================================================================================")

print("Sample Rows from Q3 2023-24:")
print(f"Header: Col 5 ({ws.cell(row=7, column=5).value}), Col 6 ({ws.cell(row=7, column=6).value}), Col 7 ({ws.cell(row=7, column=7).value}), Col 8 ({ws.cell(row=7, column=8).value}), Col 23 ({ws.cell(row=7, column=23).value})")

for r in range(8, 20):
    strat    = ws.cell(row=r, column=5).value
    window   = ws.cell(row=r, column=6).value
    entry_d  = ws.cell(row=r, column=7).value
    exit_d   = ws.cell(row=r, column=8).value
    result_d = ws.cell(row=r, column=23).value
    print(f"Row {r:2d}: Strategy={strat:15s} | Window={window:12s} | Entry={entry_d} | Exit={exit_d} | Result Date={result_d}")

print("==========================================================================================")
