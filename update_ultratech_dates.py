import urllib.request
import urllib.parse
import http.cookiejar
import json
import pandas as pd
import pathlib
import sys
import warnings
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore')
sys.stdout.reconfigure(errors='replace')

ROOT = pathlib.Path('D:/behaviour analysis')
stock_folder = ROOT / 'ULTRACEMCO'
stock_folder.mkdir(parents=True, exist_ok=True)

print("==========================================================================================")
print("ULTRATECH CEMENT (ULTRACEMCO) LIVE RESULT DATE & ANNOUNCEMENTS FETCH")
print("==========================================================================================")
print(f"Current System Time: {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M:%S')}")
print("------------------------------------------------------------------------------------------")

cj = http.cookiejar.CookieJar()
opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
headers = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
    'Accept': '*/*',
    'Accept-Language': 'en-US,en;q=0.9',
    'Referer': 'https://www.nseindia.com/'
}

# 1. Prime session with ULTRACEMCO quote page
try:
    req = urllib.request.Request('https://www.nseindia.com/get-quotes/equity?symbol=ULTRACEMCO', headers=headers)
    with opener.open(req, timeout=12) as r:
        r.read()
except Exception as e:
    print(f"Warning priming NSE session: {e}")

# 2. Fetch Corporate Announcements for ULTRACEMCO
ann_url = 'https://www.nseindia.com/api/corporate-announcements?index=equities&symbol=ULTRACEMCO'
anns = []
try:
    with opener.open(urllib.request.Request(ann_url, headers=headers), timeout=15) as r:
        anns = json.loads(r.read().decode('utf-8'))
    print(f"Total live corporate announcements fetched for ULTRACEMCO: {len(anns)}")
except Exception as e:
    print(f"Error fetching corporate announcements: {e}")

bm_list = []
for a in anns:
    if not isinstance(a, dict): continue
    desc = str(a.get('desc', ''))
    att = str(a.get('attchmntText', ''))
    dt = a.get('an_dt') or a.get('sort_date')
    file_name = a.get('attchmntFile', '')
    combo = (desc + ' ' + att).lower()
    
    if any(k in combo for k in ['board meeting', 'financial result', 'results', 'audited', 'unaudited']):
        bm_list.append({
            'symbol': 'ULTRACEMCO',
            'announcement_date': dt,
            'purpose': desc,
            'details': att,
            'attachment': f"https://nsearchives.nseindia.com/corporate/{file_name}" if file_name else ""
        })

df_ann = pd.DataFrame(anns)
df_bm = pd.DataFrame(bm_list)

# 3. Fetch Financial Results Comparison for ULTRACEMCO
res_url = 'https://www.nseindia.com/api/results-comparision?symbol=ULTRACEMCO'
res_raw = []
try:
    with opener.open(urllib.request.Request(res_url, headers=headers), timeout=15) as r:
        res_data = json.loads(r.read().decode('utf-8'))
        res_raw = (res_data.get('resCmpData') or []) if isinstance(res_data, dict) else []
    print(f"Total financial results records fetched for ULTRACEMCO: {len(res_raw)}")
except Exception as e:
    print(f"Error fetching financial results comparison: {e}")

fr_rows = []
for r in res_raw:
    if not isinstance(r, dict): continue
    fr_rows.append({
        'broadCastDate': r.get('re_create_dt', ''),
        'relatingTo': f"Qtr ended {r.get('re_to_dt', '')}",
        'consolidated': r.get('re_res_type', 'Audited'),
        'audited': 'Audited',
        'fromDate': r.get('re_from_dt', ''),
        'toDate': r.get('re_to_dt', ''),
        'xbrl': ''
    })

df_fr = pd.DataFrame(fr_rows)

# Save local files in ULTRACEMCO folder
if not df_ann.empty:
    df_ann.to_csv(stock_folder / 'announcements.csv', index=False)
if not df_bm.empty:
    df_bm.to_csv(stock_folder / 'board_meetings.csv', index=False)
if not df_fr.empty:
    df_fr.to_csv(stock_folder / 'financial_results.csv', index=False)

print("\n------------------------------------------------------------------------------------------")
print("ULTRATECH CEMENT (ULTRACEMCO) ANNOUNCED RESULT DATES & BOARD MEETINGS:")
print("------------------------------------------------------------------------------------------")

if not df_bm.empty:
    print("\nRecent Board Meetings & Financial Result Announcements for ULTRACEMCO:")
    print(df_bm[['announcement_date', 'purpose', 'details']].head(10).to_string(index=False))

if not df_fr.empty:
    df_fr['_dt'] = pd.to_datetime(df_fr['broadCastDate'], errors='coerce')
    df_fr_sorted = df_fr.sort_values('_dt', ascending=False).drop(columns=['_dt'])
    print("\nHistorical & Latest Broadcast Financial Results for ULTRACEMCO:")
    print(df_fr_sorted.head(10).to_string(index=False))

# Update Master Workbooks for ULTRACEMCO
print("\nUpdating Master Workbooks with latest ULTRACEMCO dates...")

# 1. Update Verified_All_Quarters_Result_Dates_211_Stocks.xlsx
master_path = ROOT / 'Verified_All_Quarters_Result_Dates_211_Stocks.xlsx'
if master_path.exists():
    df_master = pd.read_excel(master_path)
    # Remove existing ULTRACEMCO entries to avoid duplicates
    df_master = df_master[df_master['Stock Symbol'] != 'ULTRACEMCO'].copy()
    
    # Add updated ULTRACEMCO rows
    new_rows = []
    for _, r in df_fr.iterrows():
        raw_dt = pd.to_datetime(r.get('broadCastDate', ''), errors='coerce')
        if pd.notna(raw_dt):
            new_rows.append({
                'Stock Symbol': 'ULTRACEMCO',
                'Result Announcement Date': raw_dt.strftime('%Y-%m-%d'),
                'Financial Period / Description': r.get('relatingTo', 'Financial Results'),
                'Audit Status': r.get('audited', 'Audited'),
                'Period From Date': r.get('fromDate', ''),
                'Period To Date': r.get('toDate', '')
            })
    df_new = pd.DataFrame(new_rows)
    df_combined = pd.concat([df_master, df_new], ignore_index=True)
    df_combined['_dt'] = pd.to_datetime(df_combined['Result Announcement Date'], errors='coerce')
    df_combined = df_combined.sort_values(['_dt', 'Stock Symbol'], ascending=[False, True]).drop(columns=['_dt'])
    df_combined['Sr. No.'] = range(1, len(df_combined) + 1)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All Quarters Result Dates"
    ws.views.sheetView[0].showGridLines = True
    
    hdr_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    border_thin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                         top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
    
    headers = ["Sr. No.", "Stock Symbol", "Result Announcement Date", "Financial Period / Description", "Audit Status", "Period From Date", "Period To Date"]
    ws.row_dimensions[1].height = 26
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for r_idx, r in enumerate(df_combined.to_dict('records'), start=2):
        for c_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=r.get(h, ''))
            if h in ["Sr. No.", "Result Announcement Date", "Audit Status", "Period From Date", "Period To Date"]:
                cell.alignment = Alignment(horizontal="center")
            cell.border = border_thin
            
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    wb.save(master_path)
    print(f"Updated Master Excel: {master_path}")

print("==========================================================================================")
print("ULTRATECH CEMENT (ULTRACEMCO) LIVE RESULT DATE UPDATE COMPLETE!")
print("==========================================================================================")
