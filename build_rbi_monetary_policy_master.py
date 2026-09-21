import os
import pathlib
import sys
import bisect
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(errors='replace')

ROOT = pathlib.Path('D:/behaviour analysis')
OI_DIR = ROOT / 'OI_DATA'
PROC = ROOT / 'processed'
PRICE_CACHE = PROC / 'price_cache'
RBI_CSV = ROOT / 'rbi_monetary_policy_dates.csv'

SYMBOLS = sorted([e.name.strip().upper() for e in os.scandir(OI_DIR) if e.is_dir()])

print("==========================================================================================")
print("BUILDING RBI MONETARY POLICY EVENT BACKTEST MASTER WORKBOOK")
print("==========================================================================================")
print(f"Target Universe: {len(SYMBOLS)} Stocks")

LOT_SIZES = {
    "RELIANCE": 250, "TCS": 175, "INFY": 400, "HDFCBANK": 550, "ICICIBANK": 700,
    "BHARTIARTL": 950, "ITC": 1600, "SBIN": 1500, "LTIM": 150, "LT": 300,
    "HINDUNILVR": 300, "AXISBANK": 625, "KOTAKBANK": 400, "BAJFINANCE": 125,
    "M&M": 350, "MARUTI": 100, "SUNPHARMA": 350, "TATASTEEL": 5500,
    "NTPC": 1500, "POWERGRID": 1800, "TITAN": 175, "ADANIENT": 300,
    "ADANIPORTS": 625, "ULTRACEMCO": 100, "ASIANPAINT": 200, "COALINDIA": 2100,
    "BAJAJ-AUTO": 125, "JSWSTEEL": 675, "TATAMOTORS": 1425, "HCLTECH": 350,
    "GRASIM": 250, "HEROMOTOCO": 150, "EICHERMOT": 175, "CIPLA": 650,
    "HDFCLIFE": 1100, "SBILIFE": 375, "DRREDDY": 125, "BRITANNIA": 200,
    "APOLLOHOSP": 125, "TATACONSUM": 450, "HINDALCO": 1400, "BPCL": 1800,
    "INDUSINDBK": 500, "DIVISLAB": 200, "BAJAJFINSV": 500, "NESTLEIND": 200,
    "WIPRO": 1500, "ONGC": 3750, "TECHM": 600, "ASHOKLEY": 5000, "ADANIGREEN": 500
}

SECTOR_MAP = {
    "HDFCBANK": "Banking & Finance", "ICICIBANK": "Banking & Finance", "AXISBANK": "Banking & Finance",
    "KOTAKBANK": "Banking & Finance", "SBIN": "Banking & Finance", "INDUSINDBK": "Banking & Finance",
    "BANKBARODA": "Banking & Finance", "PNB": "Banking & Finance", "CANBK": "Banking & Finance",
    "IDFCFIRSTB": "Banking & Finance", "AUBANK": "Banking & Finance", "FEDERALBNK": "Banking & Finance",
    "RBLBANK": "Banking & Finance", "INDIANB": "Banking & Finance", "UNIONBANK": "Banking & Finance",
    
    "BAJFINANCE": "NBFC & Fintech", "BAJAJFINSV": "NBFC & Fintech", "SHRIRAMFIN": "NBFC & Fintech",
    "JIOFIN": "NBFC & Fintech", "CHOLAFIN": "NBFC & Fintech", "MUTHOOTFIN": "NBFC & Fintech",
    "REC": "NBFC & Fintech", "PFC": "NBFC & Fintech", "SBICARD": "NBFC & Fintech",
    "M&MFIN": "NBFC & Fintech", "LICHSGFIN": "NBFC & Fintech", "IRFC": "NBFC & Fintech",
    "CREDITACC": "NBFC & Fintech", "MANAPPURAM": "NBFC & Fintech", "PAYTM": "NBFC & Fintech",
    
    "HDFCLIFE": "Insurance", "SBILIFE": "Insurance", "ICICIPRULI": "Insurance", "ICICIGI": "Insurance",
    "GICRE": "Insurance", "NIACL": "Insurance", "STARHEALTH": "Insurance",
    
    "TCS": "IT & Tech Services", "INFY": "IT & Tech Services", "HCLTECH": "IT & Tech Services",
    "TECHM": "IT & Tech Services", "WIPRO": "IT & Tech Services", "LTIM": "IT & Tech Services",
    "PERSISTENT": "IT & Tech Services", "COFORGE": "IT & Tech Services", "MPHASIS": "IT & Tech Services",
    "TATAELXSI": "IT & Tech Services", "KPITTECH": "IT & Tech Services", "LTTS": "IT & Tech Services",
    "OFSS": "IT & Tech Services", "CYIENT": "IT & Tech Services", "KFINTECH": "IT & Tech Services",
    
    "MARUTI": "Automobile & Auto Ancillaries", "M&M": "Automobile & Auto Ancillaries",
    "BAJAJ-AUTO": "Automobile & Auto Ancillaries", "EICHERMOT": "Automobile & Auto Ancillaries",
    "HEROMOTOCO": "Automobile & Auto Ancillaries", "TATAMOTORS": "Automobile & Auto Ancillaries",
    "ASHOKLEY": "Automobile & Auto Ancillaries", "TVSMOTOR": "Automobile & Auto Ancillaries",
    "BHARATFORG": "Automobile & Auto Ancillaries", "MOTHERSON": "Automobile & Auto Ancillaries",
    "BOSCHLTD": "Automobile & Auto Ancillaries", "BALKRISIND": "Automobile & Auto Ancillaries",
    "MRF": "Automobile & Auto Ancillaries", "APOLLOTYRE": "Automobile & Auto Ancillaries",
    "UNOMINDA": "Automobile & Auto Ancillaries", "TIINDIA": "Automobile & Auto Ancillaries",
    "ATHERENERG": "Automobile & Auto Ancillaries", "HYUNDAI": "Automobile & Auto Ancillaries",
    
    "TATASTEEL": "Metals & Mining", "JSWSTEEL": "Metals & Mining", "HINDALCO": "Metals & Mining",
    "COALINDIA": "Metals & Mining", "JINDALSTEL": "Metals & Mining", "NMDC": "Metals & Mining",
    "VEDL": "Metals & Mining", "NATIONALUM": "Metals & Mining", "SAIL": "Metals & Mining",
    "APLAPOLLO": "Metals & Mining", "HINDZINC": "Metals & Mining",
    
    "RELIANCE": "Oil, Gas & Energy", "ONGC": "Oil, Gas & Energy", "NTPC": "Oil, Gas & Energy",
    "POWERGRID": "Oil, Gas & Energy", "BPCL": "Oil, Gas & Energy", "IOC": "Oil, Gas & Energy",
    "HPCL": "Oil, Gas & Energy", "GAIL": "Oil, Gas & Energy", "ADANIGREEN": "Oil, Gas & Energy",
    "ADANIPOWER": "Oil, Gas & Energy", "TATAPOWER": "Oil, Gas & Energy", "NHPC": "Oil, Gas & Energy",
    "SJVN": "Oil, Gas & Energy", "OIL": "Oil, Gas & Energy", "SUZLON": "Oil, Gas & Energy",
    "INOXWIND": "Oil, Gas & Energy", "IREDA": "Oil, Gas & Energy", "WAAREEENER": "Oil, Gas & Energy",
    
    "SUNPHARMA": "Pharma & Healthcare", "CIPLA": "Pharma & Healthcare", "DRREDDY": "Pharma & Healthcare",
    "APOLLOHOSP": "Pharma & Healthcare", "DIVISLAB": "Pharma & Healthcare", "TORNTPHARM": "Pharma & Healthcare",
    "MANKIND": "Pharma & Healthcare", "LUPIN": "Pharma & Healthcare", "ZYDUSLIFE": "Pharma & Healthcare",
    "ALKEM": "Pharma & Healthcare", "BIOCON": "Pharma & Healthcare", "MAXHEALTH": "Pharma & Healthcare",
    "SYNGENE": "Pharma & Healthcare", "FORTIS": "Pharma & Healthcare", "LAURUSLABS": "Pharma & Healthcare",
    "SAGILITY": "Pharma & Healthcare",
    
    "HINDUNILVR": "FMCG & Consumer Goods", "ITC": "FMCG & Consumer Goods", "NESTLEIND": "FMCG & Consumer Goods",
    "TATACONSUM": "FMCG & Consumer Goods", "BRITANNIA": "FMCG & Consumer Goods", "DABUR": "FMCG & Consumer Goods",
    "GODREJCP": "FMCG & Consumer Goods", "MARICO": "FMCG & Consumer Goods", "COLPAL": "FMCG & Consumer Goods",
    "VBL": "FMCG & Consumer Goods", "PGHH": "FMCG & Consumer Goods", "UBL": "FMCG & Consumer Goods",
    "MCDOWELL-N": "FMCG & Consumer Goods", "GODFRYPHLP": "FMCG & Consumer Goods",
    
    "TITAN": "Consumer Durables & Retail", "TRENT": "Consumer Durables & Retail",
    "ASIANPAINT": "Consumer Durables & Retail", "BERGEPAINT": "Consumer Durables & Retail",
    "PIDILITIND": "Consumer Durables & Retail", "HAVELLS": "Consumer Durables & Retail",
    "DIXON": "Consumer Durables & Retail", "VOLTAS": "Consumer Durables & Retail",
    "POLYCAB": "Consumer Durables & Retail", "CROMPTON": "Consumer Durables & Retail",
    "KEI": "Consumer Durables & Retail", "WHIRLPOOL": "Consumer Durables & Retail",
    "DMART": "Consumer Durables & Retail", "KALYANKJIL": "Consumer Durables & Retail",
    
    "L&T": "Capital Goods & Infrastructure", "LT": "Capital Goods & Infrastructure",
    "SIEMENS": "Capital Goods & Infrastructure", "ABB": "Capital Goods & Infrastructure",
    "BEL": "Capital Goods & Infrastructure", "HAL": "Capital Goods & Infrastructure",
    "BHEL": "Capital Goods & Infrastructure", "CGPOWER": "Capital Goods & Infrastructure",
    "CUMMINSIND": "Capital Goods & Infrastructure", "TITAGARH": "Capital Goods & Infrastructure",
    "RAILTEL": "Capital Goods & Infrastructure", "RITES": "Capital Goods & Infrastructure",
    "MAZDOCK": "Capital Goods & Infrastructure", "COCHINSHIP": "Capital Goods & Infrastructure",
    "BDL": "Capital Goods & Infrastructure", "IRCON": "Capital Goods & Infrastructure",
    
    "DLF": "Realty & Construction", "LODHA": "Realty & Construction", "GODREJPROP": "Realty & Construction",
    "OBERREALTY": "Realty & Construction", "PHOENIXLTD": "Realty & Construction", "PRESTIGE": "Realty & Construction"
}

df_rbi = pd.read_csv(RBI_CSV)
df_rbi['dt'] = pd.to_datetime(df_rbi['Date'])
df_rbi = df_rbi.sort_values('dt').reset_index(drop=True)

df_rbi['Category'] = 'Status Quo'
df_rbi.loc[df_rbi['Action'].str.contains('Hike', case=False, na=False), 'Category'] = 'Rate Hike'
df_rbi.loc[df_rbi['Action'].str.contains('Cut', case=False, na=False), 'Category'] = 'Rate Cut'

categories = {
    "ALL": df_rbi,
    "HIKE": df_rbi[df_rbi['Category'] == 'Rate Hike'],
    "CUT": df_rbi[df_rbi['Category'] == 'Rate Cut'],
    "STATUS_QUO": df_rbi[df_rbi['Category'] == 'Status Quo']
}

pre_candidates = [1, 2, 3, 4, 5, 6, 7, 8]
post_candidates = [1, 2, 3, 4, 5, 6, 7, 8]

category_data_store = {
    "ALL": [], "HIKE": [], "CUT": [], "STATUS_QUO": []
}

def calculate_max_drawdown(equity_curve):
    if not equity_curve or len(equity_curve) < 2:
        return 0.0
    arr = np.array(equity_curve)
    peak = np.maximum.accumulate(arr)
    # prevent div by zero
    peak_safe = np.where(peak == 0, 1.0, peak)
    dd = (arr - peak) / np.abs(peak_safe) * 100
    return round(abs(float(np.min(dd))), 2)

for idx, sym in enumerate(SYMBOLS, 1):
    stock_folder = ROOT / sym
    price_path = PRICE_CACHE / f"{sym}.csv"
    
    df_prices = pd.DataFrame()
    if price_path.exists():
        try:
            df_prices = pd.read_csv(price_path)
            df_prices['date'] = pd.to_datetime(df_prices['date'])
            df_prices = df_prices.sort_values('date').reset_index(drop=True)
        except Exception:
            pass
            
    if df_prices.empty:
        continue
        
    sector_name = SECTOR_MAP.get(sym, "Diversified / Others")
    valid_dates = df_prices['date'].tolist()
    if not valid_dates:
        continue
        
    lot_size = LOT_SIZES.get(sym, max(100, int(1000000 / float(df_prices['adj'].iloc[-1]))))
    last_price = float(df_prices['adj'].iloc[-1])
    margin_for_stock = round(0.20 * lot_size * last_price, 2)

    for cat_key, cat_df in categories.items():
        tot_events = len(cat_df)
        if tot_events == 0:
            category_data_store[cat_key].append({
                "sym": sym, "sector": sector_name,
                "n_events": "0 Events", "window": "N/A", "strategy": "N/A",
                "win_rate": 0.0, "wins": 0, "losses": 0, "win_ratio_str": "N/A",
                "avg_ret": 0.0, "max_dd": 0.0, "margin": margin_for_stock, "net_pnl": 0.0, "is_available": False
            })
            continue

        best_pnl = -999999999
        best_eval = None

        for pre in pre_candidates:
            for post in post_candidates:
                emp_long_wins = 0; emp_short_wins = 0
                rets_long = []; rets_short = []
                pnl_long_total = 0.0; pnl_short_total = 0.0
                equity_long = [0.0]; equity_short = [0.0]

                valid_event_count = 0
                for _, rbi_row in cat_df.iterrows():
                    dt = rbi_row['dt']
                    pos = bisect.bisect_left(valid_dates, dt)
                    if pos == 0: r_i = 0
                    elif pos >= len(valid_dates): r_i = len(valid_dates) - 1
                    else:
                        r_i = pos if (valid_dates[pos] - dt) < (dt - valid_dates[pos - 1]) else (pos - 1)

                    e_i = max(0, r_i - pre)
                    x_i = min(len(valid_dates) - 1, r_i + post)
                    pe = float(df_prices.loc[e_i, 'adj']); px = float(df_prices.loc[x_i, 'adj'])
                    
                    if pe > 0:
                        valid_event_count += 1
                        rl = (px - pe) / pe * 100
                        rs = (pe - px) / pe * 100
                        rets_long.append(rl); rets_short.append(rs)

                        b_to_l = round(lot_size * pe, 2); s_to_l = round(lot_size * px, 2)
                        cost_l = round((b_to_l + s_to_l) * 0.0005, 2)
                        net_pnl_l = round(s_to_l - b_to_l - cost_l, 2)
                        pnl_long_total += net_pnl_l
                        equity_long.append(pnl_long_total)

                        b_to_s = round(lot_size * px, 2); s_to_s = round(lot_size * pe, 2)
                        cost_s = round((b_to_s + s_to_s) * 0.0005, 2)
                        net_pnl_s = round(s_to_s - b_to_s - cost_s, 2)
                        pnl_short_total += net_pnl_s
                        equity_short.append(pnl_short_total)

                        if net_pnl_l > 0: emp_long_wins += 1
                        if net_pnl_s > 0: emp_short_wins += 1

                if valid_event_count == 0:
                    continue

                wr_long = (emp_long_wins / valid_event_count) * 100
                wr_short = (emp_short_wins / valid_event_count) * 100

                if wr_long >= wr_short:
                    opt_strat = "FUTURE LONG"
                    opt_wr = round(wr_long, 2)
                    opt_wins = emp_long_wins
                    opt_losses = valid_event_count - emp_long_wins
                    opt_avg_ret = round(np.mean(rets_long), 2) if rets_long else 0.0
                    opt_pnl = round(pnl_long_total, 2)
                    opt_dd = calculate_max_drawdown(equity_long)
                else:
                    opt_strat = "FUTURE SHORT"
                    opt_wr = round(wr_short, 2)
                    opt_wins = emp_short_wins
                    opt_losses = valid_event_count - emp_short_wins
                    opt_avg_ret = round(np.mean(rets_short), 2) if rets_short else 0.0
                    opt_pnl = round(pnl_short_total, 2)
                    opt_dd = calculate_max_drawdown(equity_short)

                if opt_pnl > best_pnl:
                    best_pnl = opt_pnl
                    best_eval = {
                        "pre": pre, "post": post, "window": f"T-{pre} to T+{post}",
                        "strategy": opt_strat, "win_rate": opt_wr,
                        "wins": opt_wins, "losses": opt_losses, "event_count": valid_event_count,
                        "avg_ret": opt_avg_ret, "max_dd": opt_dd, "pnl": opt_pnl
                    }

        if best_eval is None:
            category_data_store[cat_key].append({
                "sym": sym, "sector": sector_name,
                "n_events": f"{tot_events} Events", "window": "N/A", "strategy": "N/A",
                "win_rate": 0.0, "wins": 0, "losses": 0, "win_ratio_str": "N/A",
                "avg_ret": 0.0, "max_dd": 0.0, "margin": margin_for_stock, "net_pnl": 0.0, "is_available": False
            })
        else:
            v_cnt = best_eval["event_count"]
            w_str = f"{best_eval['wins']} Wins / {v_cnt} Events"
            trading_status = "QUALIFIED" if v_cnt >= 10 and best_eval["win_rate"] >= 60.0 else "AVOID BUT MONITOR IT"
            category_data_store[cat_key].append({
                "sym": sym, "sector": sector_name,
                "trading_status": trading_status,
                "n_events": f"{v_cnt} Events",
                "window": best_eval["window"],
                "strategy": best_eval["strategy"],
                "win_rate": best_eval["win_rate"],
                "wins": best_eval["wins"],
                "losses": best_eval["losses"],
                "win_ratio_str": w_str,
                "avg_ret": best_eval["avg_ret"],
                "max_dd": best_eval["max_dd"],
                "margin": margin_for_stock,
                "net_pnl": best_eval["pnl"],
                "is_available": True
            })

out_path_master = ROOT / "Nifty211_RBI_Monetary_Policy_Performance_Master.xlsx"

wb = openpyxl.Workbook()

title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
title_fill = PatternFill("solid", fgColor="1F4E79")
card_val_font = Font(name="Calibri", size=15, bold=True, color="1F4E79")
hdr_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1F4E79")
subhdr_fill = PatternFill("solid", fgColor="2F5597")
border_thin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                     top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
card_fill = PatternFill("solid", fgColor="F2F4F7")
pos_fill = PatternFill("solid", fgColor="E2EFDA")
neg_fill = PatternFill("solid", fgColor="FCE4D6")
amb_fill = PatternFill("solid", fgColor="FFF2CC")
pos_font = Font(name="Calibri", size=11, bold=True, color="276A3C")
neg_font = Font(name="Calibri", size=11, bold=True, color="9C0006")
amb_font = Font(name="Calibri", size=11, bold=True, color="B25900")

sheet_map = {
    "ALL": "All RBI Policy Events",
    "HIKE": "Rate Hike Meetings",
    "CUT": "Rate Cut Meetings",
    "STATUS_QUO": "Status Quo Meetings"
}

for s_idx, (cat_key, sheet_title) in enumerate(sheet_map.items()):
    if s_idx == 0:
        ws = wb.active
        ws.title = sheet_title
    else:
        ws = wb.create_sheet(title=sheet_title)

    ws.views.sheetView[0].showGridLines = True
    
    raw_list = category_data_store[cat_key]
    df_cat = pd.DataFrame(raw_list)
    df_avail = df_cat[df_cat['is_available'] == True].sort_values(['net_pnl', 'win_rate', 'avg_ret'], ascending=[False, False, False])
    df_unavail = df_cat[df_cat['is_available'] == False].sort_values('sym', ascending=True)
    df_sorted = pd.concat([df_avail, df_unavail], ignore_index=True)

    # Title Banner (15 Columns: A to O)
    ws.merge_cells("A1:O1")
    ws["A1"] = f"RBI MONETARY POLICY EVENT BACKTEST MASTER — {sheet_title.upper()}"
    ws["A1"].font = title_font; ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    # KPI Cards
    top_stock = df_avail.iloc[0]['sym'] if not df_avail.empty else "N/A"
    top_pnl = df_avail.iloc[0]['net_pnl'] if not df_avail.empty else 0.0
    avg_ret_c = round(df_avail['avg_ret'].mean(), 2) if not df_avail.empty else 0.0
    avg_wr_c = round(df_avail['win_rate'].mean(), 2) if not df_avail.empty else 0.0
    tot_pnl_c = round(df_avail['net_pnl'].sum(), 2) if not df_avail.empty else 0.0

    cards = [
        ("TOTAL STOCKS ANALYZED", f"{len(df_sorted)} Stocks ({len(df_avail)} Avail)", "A3:C4", "A3"),
        ("TOP PERFORMER", f"{top_stock} (₹{top_pnl:,.2f})", "D3:F4", "D3"),
        ("AVERAGE EVENT RETURN %", f"{avg_ret_c}%", "G3:I4", "G3"),
        ("AVERAGE EVENT WIN RATE %", f"{avg_wr_c}%", "J3:L4", "J3"),
        ("TOTAL EVENT NET P&L", f"₹{tot_pnl_c:,.2f}", "M3:O4", "M3")
    ]
    for title, val, merge_range, top_left in cards:
        ws.merge_cells(merge_range)
        ws[top_left] = f"{title}\n{val}"
        ws[top_left].font = card_val_font
        ws[top_left].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws[top_left].fill = card_fill; ws[top_left].border = border_thin

    # Main Table Header
    rank_start_row = 6
    ws.cell(row=rank_start_row, column=1, value=f"ALL {len(df_sorted)} STOCKS RANKED BY {sheet_title.upper()} PERFORMANCE").font = hdr_font
    ws.cell(row=rank_start_row, column=1).fill = hdr_fill
    ws.merge_cells(start_row=rank_start_row, start_column=1, end_row=rank_start_row, end_column=15)

    headers = [
        "Rank", "Stock Symbol", "Industry Sector", "Action Status",
        "Strategy", "Policy Events Count", "Position Window", "Win Rate %",
        "Win Ratio", "Losing Events", "Average Return %", "Max Drawdown %",
        "Margin Capital (₹)", "Net Realised P&L (₹)", "Status / Notes"
    ]
    ws.row_dimensions[rank_start_row+1].height = 28
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=rank_start_row+1, column=c_idx, value=h)
        cell.font = hdr_font; cell.fill = subhdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, r in enumerate(df_sorted.to_dict('records'), start=rank_start_row+2):
        ws.cell(row=r_idx, column=1, value=r_idx - rank_start_row - 1).border = border_thin
        ws.cell(row=r_idx, column=1).alignment = Alignment(horizontal="center")
        
        ws.cell(row=r_idx, column=2, value=r['sym']).border = border_thin
        ws.cell(row=r_idx, column=2).font = Font(bold=True)
        
        ws.cell(row=r_idx, column=3, value=r['sector']).border = border_thin
        ws.cell(row=r_idx, column=3).alignment = Alignment(horizontal="center")

        # Action Status
        c_act = ws.cell(row=r_idx, column=4, value=r.get('trading_status', 'AVOID BUT MONITOR IT'))
        c_act.border = border_thin; c_act.alignment = Alignment(horizontal="center")
        if r.get('trading_status') == 'QUALIFIED': c_act.fill = pos_fill; c_act.font = pos_font
        else: c_act.fill = amb_fill; c_act.font = amb_font

        # Strategy
        c_st = ws.cell(row=r_idx, column=5, value=r['strategy'])
        c_st.border = border_thin; c_st.alignment = Alignment(horizontal="center")
        if r['strategy'] == 'FUTURE LONG': c_st.fill = pos_fill; c_st.font = pos_font
        elif r['strategy'] == 'FUTURE SHORT': c_st.fill = neg_fill; c_st.font = neg_font
        else: c_st.fill = card_fill; c_st.font = Font(color="7F7F7F", italic=True)

        # Policy Events Count
        ws.cell(row=r_idx, column=6, value=r['n_events']).border = border_thin
        ws.cell(row=r_idx, column=6).alignment = Alignment(horizontal="center")
        
        # Position Window
        ws.cell(row=r_idx, column=7, value=r['window']).border = border_thin
        ws.cell(row=r_idx, column=7).alignment = Alignment(horizontal="center")
        
        # Win Rate %
        wr_val_str = f"{r['win_rate']}%" if r['is_available'] else "N/A"
        c_wr = ws.cell(row=r_idx, column=8, value=wr_val_str)
        c_wr.border = border_thin; c_wr.alignment = Alignment(horizontal="right")
        c_wr.font = Font(bold=True)
        if r['is_available']:
            if r['win_rate'] >= 65.0: c_wr.fill = pos_fill; c_wr.font = pos_font
            elif r['win_rate'] < 50.0: c_wr.fill = neg_fill; c_wr.font = neg_font
        else:
            c_wr.alignment = Alignment(horizontal="center"); c_wr.font = Font(color="7F7F7F", italic=True)
        
        # Win Ratio
        c_w = ws.cell(row=r_idx, column=9, value=r['win_ratio_str'])
        c_w.border = border_thin; c_w.alignment = Alignment(horizontal="center")
        if not r['is_available']: c_w.font = Font(color="7F7F7F", italic=True)
        
        # Losing Events
        c_l = ws.cell(row=r_idx, column=10, value=r['losses'] if r['is_available'] else "N/A")
        c_l.border = border_thin; c_l.alignment = Alignment(horizontal="center")
        if not r['is_available']: c_l.font = Font(color="7F7F7F", italic=True)
        
        # Average Return %
        ret_val_str = f"{r['avg_ret']}%" if r['is_available'] else "N/A"
        c_ret = ws.cell(row=r_idx, column=11, value=ret_val_str)
        c_ret.border = border_thin; c_ret.alignment = Alignment(horizontal="right")
        c_ret.font = Font(bold=True)
        if r['is_available']:
            if r['avg_ret'] >= 0: c_ret.fill = pos_fill; c_ret.font = pos_font
            else: c_ret.fill = neg_fill; c_ret.font = neg_font
        else:
            c_ret.alignment = Alignment(horizontal="center"); c_ret.font = Font(color="7F7F7F", italic=True)
            
        # Max Drawdown %
        dd_val_str = f"-{r['max_dd']}%" if r['is_available'] else "N/A"
        c_dd = ws.cell(row=r_idx, column=12, value=dd_val_str)
        c_dd.border = border_thin; c_dd.alignment = Alignment(horizontal="right")
        if r['is_available']:
            if r['max_dd'] <= 15.0: c_dd.fill = pos_fill; c_dd.font = pos_font
            else: c_dd.fill = neg_fill; c_dd.font = neg_font
        else:
            c_dd.alignment = Alignment(horizontal="center"); c_dd.font = Font(color="7F7F7F", italic=True)
        
        # Margin Capital (₹)
        ws.cell(row=r_idx, column=13, value=f"₹{r['margin']:,.2f}").border = border_thin
        ws.cell(row=r_idx, column=13).alignment = Alignment(horizontal="right")
        
        # Net Realised P&L (₹)
        pnl_val_str = f"₹{r['net_pnl']:,.2f}" if r['is_available'] else "N/A"
        c_p = ws.cell(row=r_idx, column=14, value=pnl_val_str)
        c_p.border = border_thin; c_p.alignment = Alignment(horizontal="right")
        if r['is_available']:
            if r['net_pnl'] >= 0: c_p.fill = pos_fill; c_p.font = pos_font
            else: c_p.fill = neg_fill; c_p.font = neg_font
        else:
            c_p.alignment = Alignment(horizontal="center"); c_p.font = Font(color="7F7F7F", italic=True)

        # Status / Notes
        c_n = ws.cell(row=r_idx, column=15, value="High Conviction" if r.get('trading_status') == 'QUALIFIED' else "Monitor")
        c_n.border = border_thin; c_n.alignment = Alignment(horizontal="center")

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = 24

try:
    wb.save(out_path_master)
    print(f"  • Successfully generated: {out_path_master.name}")
except PermissionError:
    alt_path = ROOT / "Nifty211_RBI_Monetary_Policy_Performance_Master_v1.xlsx"
    wb.save(alt_path)
    print(f"  • Note: Saved as {alt_path.name}")

print("==========================================================================================")
print("RBI MONETARY POLICY BACKTEST MASTER CREATED SUCCESSFULLY!")
print("==========================================================================================")
