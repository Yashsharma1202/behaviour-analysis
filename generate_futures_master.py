import pandas as pd
import pathlib
import sys
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

sys.stdout.reconfigure(encoding='utf-8')

# ------------------- Configuration -------------------
# Root directory containing per‑stock OI data (futures_parquet subfolders)
OI_ROOT = pathlib.Path('D:/behaviour analysis/OI_DATA')
# Where the master workbook will be saved
OUTPUT_PATH = pathlib.Path('D:/behaviour analysis/Futures_Master_Only.xlsx')

# ------------------- Helper functions -------------------

def get_close_price(df: pd.DataFrame) -> float:
    """Return the closing price for the day.
    Prefer the row with Time == '15:29:59' if present; otherwise use the last row.
    """
    if 'Time' in df.columns:
        row = df[df['Time'] == '15:29:59']
        if not row.empty:
            return float(row.iloc[0]['Close'])
    return float(df.iloc[-1]['Close'])

def load_futures_price(stock: str, date_str: str):
    """Load futures close price (and optional expiry) for a given stock/date.
    Returns (price, expiry) where expiry may be an empty string.
    """
    fut_path = OI_ROOT / stock / 'futures_parquet' / f"{date_str}.parquet"
    if not fut_path.exists():
        return None, ''
    try:
        df = pd.read_parquet(fut_path)
    except Exception as e:
        print(f"⚠️ Failed to read {fut_path}: {e}")
        return None, ''
    price = get_close_price(df)
    expiry = ''
    if 'ExpiryDate' in df.columns:
        # Use the first unique expiry date if present
        uniq = df['ExpiryDate'].dropna().unique()
        if len(uniq) > 0:
            expiry = str(uniq[0])
    return price, expiry

# ------------------- Build master data -------------------
records = []
for stock_dir in OI_ROOT.iterdir():
    if not stock_dir.is_dir():
        continue
    stock = stock_dir.name
    futures_folder = stock_dir / 'futures_parquet'
    if not futures_folder.is_dir():
        continue
    for parquet_file in sorted(futures_folder.glob('*.parquet')):
        date_str = parquet_file.stem  # e.g., 2023-01-15
        price, expiry = load_futures_price(stock, date_str)
        if price is None:
            continue
        records.append({
            'Symbol': stock,
            'Date': date_str,
            'Futures Close Price (₹)': price,
            'Futures Expiry Date': expiry,
            # Placeholder columns for future enrichment
            'Position': '',
            'Entry Date': '',
            'Exit Date': '',
            'Expected Return (%)': '',
            'Return We Get (%)': '',
            'Booked P&L (₹)': '',
            'Quarterly Result Date': ''
        })

if not records:
    print('❌ No futures data found – nothing to write.')
    sys.exit(1)

master_df = pd.DataFrame(records)

# ------------------- Write to Excel with styling -------------------
wb = Workbook()
ws = wb.active
ws.title = 'Futures_Master'

# Header row with styling (dark teal background, white bold text)
headers = list(master_df.columns)
ws.append(headers)
header_font = Font(bold=True, color='FFFFFFFF')
header_fill = PatternFill(start_color='FF006688', end_color='FF006688', fill_type='solid')
for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center')

# Data rows
for row in master_df.itertuples(index=False, name=None):
    ws.append(row)

# Auto‑size columns for readability
for col in ws.columns:
    max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
    ws.column_dimensions[col[0].column_letter].width = max_len + 3

# Safe save (handles occasional permission errors)
def safe_save(wb_obj, path):
    try:
        wb_obj.save(path)
        return True
    except PermissionError:
        for i in range(2, 6):
            versioned = path.with_name(f"{path.stem}_v{i}{path.suffix}")
            try:
                wb_obj.save(versioned)
                print(f'⚠️ Permission error on {path.name}; saved as {versioned.name}')
                return True
            except PermissionError:
                continue
        print(f'❌ Could not save {path.name} after multiple attempts.')
        return False

if safe_save(wb, OUTPUT_PATH):
    print(f'✅ Futures‑only master Excel created at: {OUTPUT_PATH}')
    print(f'Rows written: {len(master_df)}')
else:
    print('❌ Failed to write the master workbook.')
