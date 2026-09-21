import openpyxl, pandas as pd, numpy as np, pathlib, sys

sys.stdout.reconfigure(errors='replace')

BASE_DIR = pathlib.Path('D:/behaviour analysis')
FUT_V13 = BASE_DIR / 'Nifty50_12_Quarters_Futures_OI_Master_v13.xlsx'
EQ_V2   = BASE_DIR / 'Nifty50_12_Quarters_Equity_Master_v2.xlsx'

wb_fut = openpyxl.load_workbook(FUT_V13, data_only=True)
ws_fut_exec = wb_fut['Exec_12Q_Combined_Summary']

wb_eq = openpyxl.load_workbook(EQ_V2, data_only=True)
ws_eq_exec = wb_eq['Exec_12Q_Combined_Summary']

print("==========================================================================================")
print("FINAL AUDIT & VERIFICATION OF ALL MASTER EXCEL WORKBOOKS")
print("==========================================================================================")

print(f"Futures Master Sheet Count: {len(wb_fut.sheetnames)}")
print(f"Equity Master Sheet Count : {len(wb_eq.sheetnames)}")

print("\n--- FUTURES 12-QUARTER DRAWDOWN VERIFIED SUMMARY ---")
for r in range(8, 20):
    qtr = ws_fut_exec.cell(row=r, column=1).value
    trades = ws_fut_exec.cell(row=r, column=2).value
    pnl = float(ws_fut_exec.cell(row=r, column=7).value or 0.0)
    dd_val = float(ws_fut_exec.cell(row=r, column=10).value or 0.0)
    dd_pct = float(ws_fut_exec.cell(row=r, column=11).value or 0.0)
    print(f"  {qtr}: Trades={trades}, Net PnL=₹{pnl:,.2f}, Max DD=₹{dd_val:,.2f}, Max DD%={dd_pct:.2%}")

print("\n--- EQUITY 12-QUARTER DRAWDOWN VERIFIED SUMMARY ---")
for r in range(8, 20):
    qtr = ws_eq_exec.cell(row=r, column=1).value
    trades = ws_eq_exec.cell(row=r, column=2).value
    pnl = float(ws_eq_exec.cell(row=r, column=7).value or 0.0)
    dd_val = float(ws_eq_exec.cell(row=r, column=10).value or 0.0)
    dd_pct = float(ws_eq_exec.cell(row=r, column=11).value or 0.0)
    print(f"  {qtr}: Trades={trades}, Net PnL=₹{pnl:,.2f}, Max DD=₹{dd_val:,.2f}, Max DD%={dd_pct:.2%}")

print("==========================================================================================")
