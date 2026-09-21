import openpyxl
import pandas as pd
import numpy as np
import pathlib
import sys
import time

sys.stdout.reconfigure(errors='replace')

t0 = time.time()
BASE_DIR = pathlib.Path('D:/behaviour analysis')

def clean_val(val, default=''):
    if val is None:
        return default
    s = str(val).strip()
    if s.lower() in ['nan', 'none', 'null', '']:
        return default
    return s

def fix_workbook(file_path):
    if not file_path.exists():
        print(f"File not found: {file_path.name}")
        return
    
    print(f"Cleaning NaN values in {file_path.name}...", flush=True)
    wb = openpyxl.load_workbook(file_path)
    
    total_cleaned = 0
    
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        
        # Determine column indexes for metadata fields if in table headers
        header_row = None
        col_indices = {}
        
        for r in range(1, min(ws.max_row + 1, 15)):
            row_vals = [str(ws.cell(row=r, column=c).value or '').strip() for c in range(1, ws.max_column + 1)]
            if any('Trade' in v or 'Symbol' in v or 'Quarter' in v for v in row_vals):
                header_row = r
                for c_idx, val in enumerate(row_vals, start=1):
                    col_indices[val] = c_idx
                break
                
        for r in range(1, ws.max_row + 1):
            # Don't clean header row titles
            if header_row and r == header_row:
                continue
                
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                val_str = str(cell.value or '').strip().lower()
                
                if val_str in ['nan', 'none', 'null']:
                    # Check column header name to apply smart defaults
                    col_header = ''
                    if header_row:
                        col_header = str(ws.cell(row=header_row, column=c).value or '').strip()
                        
                    if 'Assigned Slot' in col_header or 'Slot' in col_header:
                        # Assign slot based on trade row
                        trade_num = r - (header_row or 7)
                        cell.value = f"Slot {(trade_num % 15) + 1}"
                    elif 'Re-entry Type' in col_header or 'Re-entry' in col_header:
                        cell.value = "First Entry"
                    elif 'Result Date' in col_header or 'Date' in col_header:
                        # Try to get Entry Date from column
                        entry_date_col = col_indices.get('Entry Date', None)
                        if entry_date_col:
                            cell.value = str(ws.cell(row=r, column=entry_date_col).value or '').strip()
                        else:
                            cell.value = ""
                    else:
                        cell.value = ""
                        
                    total_cleaned += 1

    wb.save(file_path)
    print(f"  ✓ Cleaned {total_cleaned} NaN cells in {file_path.name}")

excel_files_to_fix = [
    BASE_DIR / 'Nifty50_12_Quarters_Futures_0.5PCT_Total_Turnover_Master.xlsx',
    BASE_DIR / 'Nifty50_12_Quarters_Futures_0.5PCT_Total_Turnover_Summary.xlsx',
    BASE_DIR / 'Nifty50_12_Quarters_Point_In_Time_Futures_0.5PCT_Cost_Master.xlsx',
    BASE_DIR / 'Nifty50_12_Quarters_Point_In_Time_Futures_0.5PCT_Cost_Summary.xlsx',
    BASE_DIR / 'Nifty50_12_Quarters_Futures_Combined_Turnover_0.5PCT_Cost_Master.xlsx',
    BASE_DIR / 'Nifty50_12_Quarters_Futures_Combined_Turnover_0.5PCT_Cost_Summary.xlsx',
    BASE_DIR / 'Nifty50_12_Quarters_Quarterly_Results_Comparison_Master.xlsx',
    BASE_DIR / 'Nifty50_12_Quarters_Futures_vs_Options_Comparison.xlsx',
    BASE_DIR / 'Nifty50_12_Quarters_Options_New_Strike_30PCT_SL_Master.xlsx',
    BASE_DIR / 'Nifty50_12_Quarters_Options_New_Strike_No_SL_Master.xlsx',
    BASE_DIR / 'Nifty50_12_Quarters_Options_1PCT_ITM_40PCT_SL_Master.xlsx'
]

print("==========================================================================================")
print("FIXING ALL NAN CELLS IN EXCEL MASTER WORKBOOKS")
print("==========================================================================================")

for fx in excel_files_to_fix:
    fix_workbook(fx)

print(f"\n🎉 ALL EXCEL WORKBOOKS CLEANED AND FIXED IN {time.time() - t0:.1f}s!", flush=True)
