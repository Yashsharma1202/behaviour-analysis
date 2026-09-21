import os
import glob
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import datetime as dt

ROOT = "D:/behaviour analysis"
REPORTS_DIR = os.path.join(ROOT, "12_Quarters_Reports")
OUTPUT_FILE = os.path.join(REPORTS_DIR, "12_Quarters_Consolidated.xlsx")

QUARTERS = [
    ("Q3_2023_24", "Q3 2023-24"),
    ("Q4_2023_24", "Q4 2023-24"),
    ("Q1_2024_25", "Q1 2024-25"),
    ("Q2_2024_25", "Q2 2024-25"),
    ("Q3_2024_25", "Q3 2024-25"),
    ("Q4_2024_25", "Q4 2024-25"),
    ("Q1_2025_26", "Q1 2025-26"),
    ("Q2_2025_26", "Q2 2025-26"),
    ("Q3_2025_26", "Q3 2025-26"),
    ("Q4_2025_26", "Q4 2025-26"),
    ("Q1_2026_27", "Q1 2026-27"),
    ("Q2_2026_27", "Q2 2026-27")
]

# Load Sector mapping
try:
    import sys
    sys.path.append(ROOT)
    import expectation_bar as EB
    SYM_SECTOR = EB.SYM_SECTOR
except Exception:
    SYM_SECTOR = {}

# Load Company mappings from dates file
SYM_COMPANY = {}
try:
    dates_file = os.path.join(ROOT, "Nifty50Stocks_QtyResultDates.xlsx")
    if os.path.exists(dates_file):
        wb_dates = openpyxl.load_workbook(dates_file, read_only=True)
        for sn in wb_dates.sheetnames:
            ws_d = wb_dates[sn]
            for r in range(2, ws_d.max_row + 1):
                comp = ws_d.cell(r, 1).value
                sym = ws_d.cell(r, 2).value
                if sym and comp:
                    SYM_COMPANY[str(sym).strip().upper()] = str(comp).strip()
        wb_dates.close()
except Exception as e:
    print(f"Error loading company names: {e}")

def get_company_name(sym):
    return SYM_COMPANY.get(sym, sym)

def get_sector(sym):
    return SYM_SECTOR.get(sym, "Diversified")

def format_date_flow(d):
    if pd.isna(d):
        return ""
    ts = pd.Timestamp(d)
    return ts.strftime("%d-%b")

def run_slot_scheduling(trades_df):
    """
    Greedy slot scheduling algorithm:
    Sort trades chronologically by Entry Date, then Exit Date.
    Assign trades to the first free slot, compounding capital.
    """
    if trades_df.empty:
        return [], trades_df
    
    # Standardize columns
    trades = trades_df.copy()
    trades["Entry Date"] = pd.to_datetime(trades["Entry Date"])
    trades["Exit Date"] = pd.to_datetime(trades["Exit Date"])
    trades = trades.sort_values(by=["Entry Date", "Exit Date"]).reset_index(drop=True)
    
    # Overwrite and cast to string to avoid datetime64 dtype errors
    trades["Sector"] = trades["Symbol"].map(get_sector).astype(str)
    trades["Company Name"] = trades["Symbol"].map(get_company_name).astype(str)
    trades["Assigned Slot"] = ""
    trades["Re-entry Type"] = ""
    trades["Realised Profit (Rs.)"] = 0.0
    
    slots = []  # list of dicts: {"Slot ID": "Slot X", "End Date": ts, "Trades": [...], "Current Capital": float}
    
    for idx, row in trades.iterrows():
        entry_dt = row["Entry Date"]
        exit_dt = row["Exit Date"]
        ret = float(row["Realised Return (%)"])
        entry_price = float(row["Entry Price"])
        
        # Find first free slot
        assigned_slot = None
        for s in slots:
            if s["End Date"] <= entry_dt:
                assigned_slot = s
                break
                
        if assigned_slot is None:
            # Create a new slot
            slot_idx = len(slots) + 1
            s_name = f"Slot {slot_idx}"
            new_slot = {
                "Slot ID": s_name,
                "End Date": exit_dt,
                "Trades": [],
                "Current Capital": entry_price,
                "Start Capital": entry_price
            }
            slots.append(new_slot)
            assigned_slot = new_slot
        else:
            # Update slot end date
            assigned_slot["End Date"] = exit_dt
            
        # Capital compounding logic:
        # If it's the first trade in the slot, the capital is the entry price.
        # Otherwise, the capital is the compounded value from the previous trade's exit.
        trade_start_cap = assigned_slot["Current Capital"]
        trade_exit_cap = trade_start_cap * (1.0 + ret)
        assigned_slot["Current Capital"] = trade_exit_cap
        
        re_entry_type = "First Entry" if len(assigned_slot["Trades"]) == 0 else "Re-entry"
        
        trade_info = {
            "Symbol": row["Symbol"],
            "Company Name": get_company_name(row["Symbol"]),
            "Sector": get_sector(row["Symbol"]),
            "Strategy": row["Strategy"],
            "Entry Date": entry_dt,
            "Exit Date": exit_dt,
            "Entry Price": entry_price,
            "Exit Price": row["Exit Price"],
            "Expected Return (%)": row["Expected Return (%)"],
            "Realised Return (%)": ret,
            "Realised Profit (Rs.)": trade_exit_cap - trade_start_cap,
            "Re-entry Type": re_entry_type,
            "Assigned Slot": assigned_slot["Slot ID"]
        }
        
        assigned_slot["Trades"].append(trade_info)
        trades.loc[idx, "Assigned Slot"] = assigned_slot["Slot ID"]
        trades.loc[idx, "Re-entry Type"] = re_entry_type
        trades.loc[idx, "Realised Profit (Rs.)"] = trade_exit_cap - trade_start_cap
        trades.loc[idx, "Company Name"] = trade_info["Company Name"]
        trades.loc[idx, "Sector"] = trade_info["Sector"]
        
    # Compile slot summaries
    slots_summary = []
    for s in slots:
        net_prof = s["Current Capital"] - s["Start Capital"]
        comp_ret = (s["Current Capital"] / s["Start Capital"] - 1) if s["Start Capital"] > 0 else 0
        
        flow_path_parts = []
        for t in s["Trades"]:
            ent_str = format_date_flow(t["Entry Date"])
            ex_str = format_date_flow(t["Exit Date"])
            flow_path_parts.append(f"{t['Symbol']} ({ent_str} to {ex_str})")
        
        flow_path = " -> Re-entry: ".join(flow_path_parts) if flow_path_parts else ""
        
        slots_summary.append({
            "Slot ID": s["Slot ID"],
            "Start Capital (Rs.)": s["Start Capital"],
            "Final Value (Rs.)": s["Current Capital"],
            "Net Profit (Rs.)": net_prof,
            "Compounded Return (%)": comp_ret,
            "Number of Re-entries": len(s["Trades"]) - 1,
            "Trade Flow Path": flow_path
        })
        
    return slots_summary, trades

def main():
    print("Starting 12 Quarters Consolidation...")
    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)
    
    all_trades = []
    
    # Styles
    font_family = "Segoe UI"
    navy_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    light_fill = PatternFill(start_color='F5F5F5', end_color='F5F5F5', fill_type='solid')
    
    border_side = Side(style='medium', color='112233')
    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='F0F0F0'),
        bottom=Side(style='thin', color='F0F0F0')
    )
    
    for fn, q_name in QUARTERS:
        file_path = os.path.join(REPORTS_DIR, f"{fn}_Combined_Best_Capital_Utilisation.xlsx")
        if not os.path.exists(file_path):
            print(f"Warning: File {file_path} not found. Skipping...")
            continue
            
        print(f"Processing Quarter: {q_name}...")
        
        # Load trades from workbook
        wb_q = openpyxl.load_workbook(file_path, data_only=True)
        ws_q_trades = wb_q["Detailed_Trades"]
        
        raw_rows = list(ws_q_trades.values)
        if len(raw_rows) < 3:
            print(f"Warning: {q_name} Detailed_Trades has too few rows. Skipping...")
            wb_q.close()
            continue
            
        # Parse rows
        header_row_idx = 2
        for r_idx, r in enumerate(raw_rows[:5]):
            if any(isinstance(val, str) and "Symbol" in val for val in r):
                header_row_idx = r_idx
                break
        
        headers = [str(h).strip() for h in raw_rows[header_row_idx]]
        data_rows = raw_rows[header_row_idx+1:]
        
        df_raw = pd.DataFrame(data_rows, columns=headers).dropna(subset=["Symbol"])
        
        # Standardize strategy names
        strategy_col = next((c for c in df_raw.columns if "Strategy" in c), "Strategy")
        df_raw["Strategy"] = df_raw[strategy_col].astype(str).str.strip().str.upper()
        
        # Map prices and returns
        entry_price_col = next((c for c in df_raw.columns if "Entry Price" in c or "Buy Price" in c), "Entry Price")
        exit_price_col = next((c for c in df_raw.columns if "Exit Price" in c or "Sell Price" in c), "Exit Price")
        realised_ret_col = next((c for c in df_raw.columns if "Realised Return" in c), "Realised Return (%)")
        expected_ret_col = next((c for c in df_raw.columns if "Expected Return" in c), "Expected Return (%)")
        
        df_raw["Entry Price"] = pd.to_numeric(df_raw[entry_price_col], errors="coerce").fillna(0.0)
        df_raw["Exit Price"] = pd.to_numeric(df_raw[exit_price_col], errors="coerce").fillna(0.0)
        df_raw["Realised Return (%)"] = pd.to_numeric(df_raw[realised_ret_col], errors="coerce").fillna(0.0)
        df_raw["Expected Return (%)"] = pd.to_numeric(df_raw[expected_ret_col], errors="coerce").fillna(0.0)
        
        # Run slot scheduling
        slots_sum, scheduled_trades = run_slot_scheduling(df_raw)
        
        # Total metrics for quarter
        fund_utilised = sum(s["Start Capital (Rs.)"] for s in slots_sum)
        final_value = sum(s["Final Value (Rs.)"] for s in slots_sum)
        net_profit = final_value - fund_utilised
        comp_return = (final_value / fund_utilised - 1) if fund_utilised > 0 else 0.0
        total_trades = len(scheduled_trades)
        total_reentries = sum(s["Number of Re-entries"] for s in slots_sum)
        
        # Save quarter detailed trades to master list
        for idx, row in scheduled_trades.iterrows():
            all_trades.append({
                "Quarter": q_name,
                "Symbol": row["Symbol"],
                "Company Name": get_company_name(row["Symbol"]),
                "Sector": get_sector(row["Symbol"]),
                "Strategy": row["Strategy"],
                "Entry Date": pd.to_datetime(row["Entry Date"]).strftime("%Y-%m-%d"),
                "Exit Date": pd.to_datetime(row["Exit Date"]).strftime("%Y-%m-%d"),
                "Buy Price (Rs.)": row["Entry Price"],
                "Exit Price (Rs.)": row["Exit Price"],
                "Expected Return (%)": row["Expected Return (%)"],
                "Realised Return (%)": row["Realised Return (%)"],
                "Realised Profit/Loss (Rs.)": row["Realised Profit (Rs.)"],
                "Re-entry Type": row["Re-entry Type"],
                "Assigned Slot": row["Assigned Slot"]
            })
            
        # Create sheet for this quarter
        ws_q = wb_out.create_sheet(title=q_name)
        ws_q.views.sheetView[0].showGridLines = True
        
        # 1. Summary Block
        ws_q.merge_cells("A1:G1")
        title_cell = ws_q.cell(row=1, column=1)
        title_cell.value = f"Nifty 50 Capital Utilisation Summary - {q_name} (Walk-Forward)"
        title_cell.font = Font(name=font_family, size=14, bold=True, color='FFFFFF')
        title_cell.fill = navy_fill
        title_cell.alignment = Alignment(horizontal='center', vertical='center')
        ws_q.row_dimensions[1].height = 40
        
        # Summary headers
        sum_headers = ["Quarter", "Fund Utilised", "Final Value", "Net Profit", "Compounded Return", "Total Trades", "Total Re-entries"]
        for col_idx, h in enumerate(sum_headers, 1):
            cell = ws_q.cell(row=2, column=col_idx, value=h)
            cell.font = Font(name=font_family, size=10, bold=True, color='FFFFFF')
            cell.fill = navy_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = Border(bottom=border_side, top=border_side, left=border_side, right=border_side)
        ws_q.row_dimensions[2].height = 25
        
        # Summary row
        sum_vals = [q_name, fund_utilised, final_value, net_profit, comp_return, total_trades, total_reentries]
        for col_idx, val in enumerate(sum_vals, 1):
            cell = ws_q.cell(row=3, column=col_idx, value=val)
            cell.font = Font(name=font_family, size=10, bold=True)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center')
            if col_idx in [2, 3, 4]:
                cell.number_format = '₹#,##0.00'
            elif col_idx == 5:
                cell.number_format = '0.00%'
            elif col_idx in [6, 7]:
                cell.number_format = '#,##0'
        ws_q.row_dimensions[3].height = 22
        
        # Spacer rows
        ws_q.row_dimensions[4].height = 10
        ws_q.row_dimensions[5].height = 10
        ws_q.row_dimensions[6].height = 10
        
        # 2. Slot Details Block
        slot_headers = ["Slot ID", "Start Capital (Rs.)", "Final Value (Rs.)", "Net Profit (Rs.)", "Compounded Return (%)", "Number of Re-entries", "Trade Flow Path"]
        for col_idx, h in enumerate(slot_headers, 1):
            cell = ws_q.cell(row=7, column=col_idx, value=h)
            cell.font = Font(name=font_family, size=10, bold=True, color='FFFFFF')
            cell.fill = navy_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = Border(bottom=border_side, top=border_side, left=border_side, right=border_side)
        ws_q.row_dimensions[7].height = 25
        
        # Slot values
        for r_idx, s in enumerate(slots_sum, 8):
            ws_q.row_dimensions[r_idx].height = 20
            vals = [
                s["Slot ID"],
                s["Start Capital (Rs.)"],
                s["Final Value (Rs.)"],
                s["Net Profit (Rs.)"],
                s["Compounded Return (%)"],
                s["Number of Re-entries"],
                s["Trade Flow Path"]
            ]
            for col_idx, val in enumerate(vals, 1):
                cell = ws_q.cell(row=r_idx, column=col_idx, value=val)
                cell.font = Font(name=font_family, size=9)
                cell.border = thin_border
                if col_idx == 7:
                    cell.alignment = Alignment(horizontal='left', vertical='center')
                else:
                    cell.alignment = Alignment(horizontal='center')
                
                if col_idx in [2, 3, 4]:
                    cell.number_format = '₹#,##0.00'
                elif col_idx == 5:
                    cell.number_format = '0.00%'
                elif col_idx == 6:
                    cell.number_format = '#,##0'
                    
        # Auto-fit columns
        for col in ws_q.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row == 1:
                    continue
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws_q.column_dimensions[col_letter].width = max(max_len + 3, 12)
        # Cap path column width
        ws_q.column_dimensions["G"].width = 85
        wb_q.close()
        
    # Create master Detailed_Trades sheet
    ws_master = wb_out.create_sheet(title="Detailed_Trades")
    ws_master.views.sheetView[0].showGridLines = True
    
    # Headers
    master_headers = ["Quarter", "Symbol", "Company Name", "Sector", "Strategy", "Entry Date", "Exit Date", "Buy Price (Rs.)", "Exit Price (Rs.)", "Expected Return (%)", "Realised Return (%)", "Realised Profit/Loss (Rs.)", "Re-entry Type", "Assigned Slot"]
    for col_idx, h in enumerate(master_headers, 1):
        cell = ws_master.cell(row=1, column=col_idx, value=h)
        cell.font = Font(name=font_family, size=10, bold=True, color='FFFFFF')
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = Border(bottom=border_side, top=border_side, left=border_side, right=border_side)
    ws_master.row_dimensions[1].height = 25
    
    # Fill values
    for r_idx, t in enumerate(all_trades, 2):
        ws_master.row_dimensions[r_idx].height = 18
        row_vals = [t[h] for h in master_headers]
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws_master.cell(row=r_idx, column=col_idx, value=val)
            cell.font = Font(name=font_family, size=9)
            cell.border = thin_border
            
            # Formats & Alignments
            if col_idx in [1, 2, 5, 6, 7, 13, 14]:
                cell.alignment = Alignment(horizontal='center')
            else:
                cell.alignment = Alignment(horizontal='left')
                
            if col_idx in [8, 9, 12]:
                cell.number_format = '₹#,##0.00'
            elif col_idx in [10, 11]:
                cell.number_format = '0.00%'
                
    # Auto-fit master columns
    for col in ws_master.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws_master.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    wb_out.save(OUTPUT_FILE)
    print(f"Consolidation complete! Saved workbook to: {OUTPUT_FILE}")

if __name__ == "__main__":
    main()
