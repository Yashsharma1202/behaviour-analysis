import sys
import os
import glob
import pathlib
import datetime
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

ROOT = pathlib.Path('D:/behaviour analysis')
PRICE_CACHE = ROOT / 'processed' / 'price_cache'

REPORTS_DIR = ROOT / 'NSE_17_Holidays_Past_4Years_Reports'
REPORTS_DIR.mkdir(parents=True, exist_ok=True)

MAIN_HOLIDAY_DIR = ROOT / 'Nifty50_Futures_Holiday_TradeLogs_Past_4Years'
MAIN_HOLIDAY_DIR.mkdir(parents=True, exist_ok=True)

MASTER_EXCEL_OUT = ROOT / 'Nifty50_Futures_All_Holidays_Past_4Years_Master_TradeLog_v4.xlsx'

print("==========================================================================================")
print("FAST EMPIRICAL BEHAVIORAL ENGINE — OPTIMAL LONG VS SHORT HOLIDAY STRATEGY BACKTEST")
print("GENERATING 17 HOLIDAY DIRECTORIES × 4 YEARLY EXCEL FILES + CONSOLIDATED REPORTS (2022-2025)")
print("==========================================================================================")

nifty50_specs = {
    'RELIANCE': {'name': 'Reliance Industries Ltd', 'lot': 250},
    'HDFCBANK': {'name': 'HDFC Bank Ltd', 'lot': 550},
    'ICICIBANK': {'name': 'ICICI Bank Ltd', 'lot': 700},
    'INFY': {'name': 'Infosys Ltd', 'lot': 400},
    'TCS': {'name': 'Tata Consultancy Services Ltd', 'lot': 175},
    'AXISBANK': {'name': 'Axis Bank Ltd', 'lot': 625},
    'SBIN': {'name': 'State Bank of India', 'lot': 750},
    'BHARTIARTL': {'name': 'Bharti Airtel Ltd', 'lot': 475},
    'BAJFINANCE': {'name': 'Bajaj Finance Ltd', 'lot': 125},
    'KOTAKBANK': {'name': 'Kotak Mahindra Bank Ltd', 'lot': 400},
    'LT': {'name': 'Larsen & Toubro Ltd', 'lot': 150},
    'M&M': {'name': 'Mahindra & Mahindra Ltd', 'lot': 350},
    'MARUTI': {'name': 'Maruti Suzuki India Ltd', 'lot': 50},
    'TATACONSUM': {'name': 'Tata Consumer Products Ltd', 'lot': 900},
    'SUNPHARMA': {'name': 'Sun Pharmaceutical Industries Ltd', 'lot': 350},
    'TATAMOTORS': {'name': 'Tata Motors Ltd', 'lot': 550},
    'TATASTEEL': {'name': 'Tata Steel Ltd', 'lot': 5500},
    'HINDUNILVR': {'name': 'Hindustan Unilever Ltd', 'lot': 300},
    'NTPC': {'name': 'NTPC Ltd', 'lot': 1500},
    'POWERGRID': {'name': 'Power Grid Corporation of India Ltd', 'lot': 1900},
    'ITC': {'name': 'ITC Ltd', 'lot': 1600},
    'COALINDIA': {'name': 'Coal India Ltd', 'lot': 2100},
    'JSWSTEEL': {'name': 'JSW Steel Ltd', 'lot': 675},
    'HINDALCO': {'name': 'Hindalco Industries Ltd', 'lot': 1400},
    'GRASIM': {'name': 'Grasim Industries Ltd', 'lot': 250},
    'EICHERMOT': {'name': 'Eicher Motors Ltd', 'lot': 175},
    'BAJAJ-AUTO': {'name': 'Bajaj Auto Ltd', 'lot': 75},
    'CIPLA': {'name': 'Cipla Ltd', 'lot': 650},
    'DRREDDY': {'name': 'Dr Reddys Laboratories Ltd', 'lot': 125},
    'APOLLOHOSP': {'name': 'Apollo Hospitals Enterprise Ltd', 'lot': 125},
    'ASIANPAINT': {'name': 'Asian Paints Ltd', 'lot': 200},
    'TITAN': {'name': 'Titan Company Ltd', 'lot': 175},
    'DIVISLAB': {'name': 'Divis Laboratories Ltd', 'lot': 200},
    'NESTLEIND': {'name': 'Nestle India Ltd', 'lot': 250},
    'BRITANNIA': {'name': 'Britannia Industries Ltd', 'lot': 200},
    'ADANIENT': {'name': 'Adani Enterprises Ltd', 'lot': 300},
    'ADANIPORTS': {'name': 'Adani Ports & SEZ Ltd', 'lot': 400},
    'ONGC': {'name': 'Oil & Natural Gas Corp Ltd', 'lot': 3850},
    'BPCL': {'name': 'Bharat Petroleum Corp Ltd', 'lot': 1800},
    'BEL': {'name': 'Bharat Electronics Ltd', 'lot': 2850},
    'ULTRACEMCO': {'name': 'UltraTech Cement Ltd', 'lot': 100},
    'TRENT': {'name': 'Trent Ltd', 'lot': 100},
    'HCLTECH': {'name': 'HCL Technologies Ltd', 'lot': 350},
    'TECHM': {'name': 'Tech Mahindra Ltd', 'lot': 600},
    'WIPRO': {'name': 'Wipro Ltd', 'lot': 1500},
    'SHRIRAMFIN': {'name': 'Shriram Finance Ltd', 'lot': 600},
    'BAJAJFINSV': {'name': 'Bajaj Finserv Ltd', 'lot': 500},
    'INDIGO': {'name': 'InterGlobe Aviation Ltd', 'lot': 150},
    'JIOFIN': {'name': 'Jio Financial Services Ltd', 'lot': 2400},
    'HDFCLIFE': {'name': 'HDFC Life Insurance Co Ltd', 'lot': 1100}
}

all_holidays = [
    {'folder': '01_Republic_Day', 'sheet': 'Republic Day', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Republic_Day.xlsx', 'dates': {2022: '2022-01-26', 2023: '2023-01-26', 2024: '2024-01-26', 2025: '2025-01-26'}},
    {'folder': '02_Mahashivratri', 'sheet': 'Mahashivratri', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Mahashivratri.xlsx', 'dates': {2022: '2022-03-01', 2023: '2023-02-18', 2024: '2024-03-08', 2025: '2025-02-26'}},
    {'folder': '03_Holi_Festival', 'sheet': 'Holi Festival', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Holi_Festival.xlsx', 'dates': {2022: '2022-03-18', 2023: '2023-03-07', 2024: '2024-03-25', 2025: '2025-03-14'}},
    {'folder': '04_Shri_Ram_Navami', 'sheet': 'Shri Ram Navami', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Shri_Ram_Navami.xlsx', 'dates': {2022: '2022-04-10', 2023: '2023-03-30', 2024: '2024-04-17', 2025: '2025-04-06'}},
    {'folder': '05_Shri_Mahavir_Jayanti', 'sheet': 'Shri Mahavir Jayanti', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Shri_Mahavir_Jayanti.xlsx', 'dates': {2022: '2022-04-14', 2023: '2023-04-04', 2024: '2024-04-21', 2025: '2025-04-10'}},
    {'folder': '06_Good_Friday', 'sheet': 'Good Friday', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Good_Friday.xlsx', 'dates': {2022: '2022-04-15', 2023: '2023-04-07', 2024: '2024-03-29', 2025: '2025-04-18'}},
    {'folder': '07_Dr_Ambedkar_Jayanti', 'sheet': 'Dr Ambedkar Jayanti', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Dr_Ambedkar_Jayanti.xlsx', 'dates': {2022: '2022-04-14', 2023: '2023-04-14', 2024: '2024-04-14', 2025: '2025-04-14'}},
    {'folder': '08_Maharashtra_Day', 'sheet': 'Maharashtra Day', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Maharashtra_Day.xlsx', 'dates': {2022: '2022-05-01', 2023: '2023-05-01', 2024: '2024-05-01', 2025: '2025-05-01'}},
    {'folder': '09_Bakri_Id', 'sheet': 'Bakri Id (Id-Ul-Adha)', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Bakri_Id.xlsx', 'dates': {2022: '2022-07-10', 2023: '2023-06-29', 2024: '2024-06-17', 2025: '2025-06-07'}},
    {'folder': '10_Muharram', 'sheet': 'Muharram', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Muharram.xlsx', 'dates': {2022: '2022-08-09', 2023: '2023-07-29', 2024: '2024-07-17', 2025: '2025-07-06'}},
    {'folder': '11_Independence_Day', 'sheet': 'Independence Day', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Independence_Day.xlsx', 'dates': {2022: '2022-08-15', 2023: '2023-08-15', 2024: '2024-08-15', 2025: '2025-08-15'}},
    {'folder': '12_Ganesh_Chaturthi', 'sheet': 'Ganesh Chaturthi', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Ganesh_Chaturthi.xlsx', 'dates': {2022: '2022-08-31', 2023: '2023-09-19', 2024: '2024-09-07', 2025: '2025-08-27'}},
    {'folder': '13_Mahatma_Gandhi_Jayanti', 'sheet': 'Mahatma Gandhi Jayanti', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Mahatma_Gandhi_Jayanti.xlsx', 'dates': {2022: '2022-10-02', 2023: '2023-10-02', 2024: '2024-10-02', 2025: '2025-10-02'}},
    {'folder': '14_Dussehra', 'sheet': 'Dussehra Dasera', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Dussehra.xlsx', 'dates': {2022: '2022-10-05', 2023: '2023-10-24', 2024: '2024-10-12', 2025: '2025-10-02'}},
    {'folder': '15_Diwali_Laxmi_Pujan', 'sheet': 'Diwali Laxmi Pujan', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Diwali_Laxmi_Pujan.xlsx', 'dates': {2022: '2022-10-24', 2023: '2023-11-12', 2024: '2024-11-01', 2025: '2025-10-20'}},
    {'folder': '16_Gurunanak_Jayanti', 'sheet': 'Gurunanak Jayanti', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Gurunanak_Jayanti.xlsx', 'dates': {2022: '2022-11-08', 2023: '2023-11-27', 2024: '2024-11-15', 2025: '2025-11-05'}},
    {'folder': '17_Christmas_Year_End', 'sheet': 'Christmas Year End', 'filename': 'Nifty50_Futures_Holiday_TradeLog_Christmas.xlsx', 'dates': {2022: '2022-12-25', 2023: '2023-12-25', 2024: '2024-12-25', 2025: '2025-12-25'}}
]

# Preload Price Data for All Nifty 50 Stocks into RAM
stock_data_cache = {}
for sym in nifty50_specs:
    csv_file = PRICE_CACHE / f"{sym}.csv"
    if csv_file.exists():
        df = pd.read_csv(csv_file)
        df['date'] = pd.to_datetime(df['date'])
        df = df.sort_values('date').reset_index(drop=True)
        dates_list = df['date'].tolist()
        adj_list = df['adj'].astype(float).tolist()
        stock_data_cache[sym] = (dates_list, adj_list)

candidate_windows = [(2,1), (3,2), (4,2), (5,3), (6,4), (7,5), (4,5), (5,5), (3,5), (6,2), (2,4), (5,2)]

def fast_backtest(dates_list, adj_list, dates_dict, n_back, m_fwd, direction='LONG', lot=100):
    trades = []
    for yr, h_date_str in dates_dict.items():
        target_dt = pd.to_datetime(h_date_str)
        # Find index in trading calendar
        idx = 0
        for i, d in enumerate(dates_list):
            if d <= target_dt:
                idx = i
            else:
                break
        entry_idx = max(0, idx - n_back)
        exit_idx = min(len(dates_list) - 1, idx + m_fwd)
        
        entry_p = adj_list[entry_idx]
        exit_p = adj_list[exit_idx]
        
        if direction == 'LONG':
            ret = ((exit_p - entry_p) / entry_p) * 100.0
            pnl = (exit_p - entry_p) * lot
        else:
            ret = ((entry_p - exit_p) / entry_p) * 100.0
            pnl = (entry_p - exit_p) * lot
            
        trades.append({
            'year': yr,
            'entry_date': dates_list[entry_idx].strftime('%Y-%m-%d'),
            'entry_price': entry_p,
            'exit_date': dates_list[exit_idx].strftime('%Y-%m-%d'),
            'exit_price': exit_p,
            'ret_pct': ret,
            'pnl': pnl,
            'status': "WIN" if pnl > 0 else "LOSS"
        })
        
    wins = sum(1 for t in trades if t['status'] == 'WIN')
    wr = (wins / len(trades) * 100.0) if trades else 0.0
    tot_pnl = sum(t['pnl'] for t in trades)
    avg_ret = np.mean([t['ret_pct'] for t in trades]) if trades else 0.0
    
    return {
        'n': n_back, 'm': m_fwd, 'direction': direction,
        'wr': wr, 'tot_pnl': tot_pnl, 'avg_ret': avg_ret,
        'trades': trades
    }

def find_best_behavior_strategy_fast(dates_list, adj_list, dates_dict, lot=100):
    best_res = None
    max_score = -999999999
    
    for direction in ['LONG', 'SHORT']:
        for n, m in candidate_windows:
            res = fast_backtest(dates_list, adj_list, dates_dict, n, m, direction, lot)
            if res:
                score = res['tot_pnl'] + (res['wr'] * 1000)
                if score > max_score:
                    max_score = score
                    best_res = res
    return best_res

# Openpyxl Styles Setup
font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_win = Font(name="Calibri", size=10, bold=True, color="385723")
font_loss = Font(name="Calibri", size=10, bold=True, color="C00000")

title_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
win_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
loss_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

border_thin = Border(
    left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9")
)
align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

all_master_trades = []

headers_log = [
    "Trade ID", "Rank", "Stock Symbol", "Company Name", "Holiday Event", "Holiday Date",
    "Optimal Window", "Strategy", "Entry Date", "Entry Price (₹)", "Exit Date", "Exit Price (₹)",
    "Return (%)", "Futures Lot Size", "Gross PnL / Lot (₹)", "Trade Status", "20% Margin Req. (₹)", "Return on Margin (ROC %)"
]

headers_sum = [
    "Rank", "Stock Symbol", "Company Name", "Optimal Window", "Optimal Direction", "4-Year Win Rate (%)",
    "2022 PnL (₹)", "2023 PnL (₹)", "2024 PnL (₹)", "2025 PnL (₹)", "Total 4-Year Gross PnL / Lot (₹)"
]

years = [2022, 2023, 2024, 2025]

# Build Reports across all 17 holidays
for h in all_holidays:
    holiday_subfolder = REPORTS_DIR / h['folder']
    holiday_subfolder.mkdir(parents=True, exist_ok=True)
    
    holiday_trades_by_year = {yr: [] for yr in years}
    trade_counter = 1
    
    for sym, spec in nifty50_specs.items():
        if sym not in stock_data_cache:
            continue
        dates_list, adj_list = stock_data_cache[sym]
        best_strat = find_best_behavior_strategy_fast(dates_list, adj_list, h['dates'], spec['lot'])
        if not best_strat:
            continue
            
        n_days = best_strat['n']
        m_days = best_strat['m']
        direction = best_strat['direction']
        
        for tr in best_strat['trades']:
            yr = tr['year']
            margin = tr['entry_price'] * spec['lot'] * 0.20
            roc_pct = (tr['pnl'] / margin) * 100.0 if margin > 0 else 0.0
            
            t_obj = {
                'id': f"TRD-{yr}-{trade_counter:03d}",
                'year': yr,
                'sym': sym,
                'name': spec['name'],
                'holiday': h['sheet'],
                'h_date': h['dates'][yr],
                'window': f"T-{n_days} to T+{m_days}",
                'strat': f"FUTURE {direction}",
                'entry_date': tr['entry_date'],
                'entry_price': tr['entry_price'],
                'exit_date': tr['exit_date'],
                'exit_price': tr['exit_price'],
                'ret_pct': tr['ret_pct'],
                'lot': spec['lot'],
                'pnl': tr['pnl'],
                'status': tr['status'],
                'margin': margin,
                'roc_pct': roc_pct
            }
            holiday_trades_by_year[yr].append(t_obj)
            all_master_trades.append(t_obj)
            
        trade_counter += 1

    # 1. GENERATE THE 4 INDIVIDUAL YEARLY EXCEL FILES FOR THIS HOLIDAY
    for yr in years:
        wb_yr = openpyxl.Workbook()
        ws_yr = wb_yr.active
        ws_yr.title = f"Trade Log {yr}"
        
        ws_yr.merge_cells("A1:R1")
        title_cell = ws_yr["A1"]
        title_cell.value = f"NIFTY 50 FUTURES — {h['sheet'].upper()} ({yr} BEHAVIORAL TRADE LOG)"
        title_cell.font = font_title
        title_cell.fill = title_fill
        title_cell.alignment = align_center
        ws_yr.row_dimensions[1].height = 40
        
        ws_yr.merge_cells("A2:R2")
        ws_yr["A2"] = " "
        
        ws_yr.append(headers_log)
        ws_yr.row_dimensions[3].height = 28
        for col_idx in range(1, len(headers_log) + 1):
            cell = ws_yr.cell(row=3, column=col_idx)
            cell.font = font_header
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = border_thin
            
        yr_trades = sorted(holiday_trades_by_year[yr], key=lambda x: x['pnl'], reverse=True)
        
        row_idx = 4
        for rank, t in enumerate(yr_trades, 1):
            ws_yr.append([
                t['id'], f"Rank {rank}", t['sym'], t['name'], t['holiday'], t['h_date'],
                t['window'], t['strat'], t['entry_date'], t['entry_price'], t['exit_date'], t['exit_price'],
                t['ret_pct'] / 100.0, t['lot'], t['pnl'], t['status'], t['margin'], t['roc_pct'] / 100.0
            ])
            ws_yr.row_dimensions[row_idx].height = 20
            
            ws_yr.cell(row=row_idx, column=1).alignment = align_center
            ws_yr.cell(row=row_idx, column=2).alignment = align_center
            ws_yr.cell(row=row_idx, column=2).font = Font(name="Calibri", size=10, bold=True, color="1F4E78")
            ws_yr.cell(row=row_idx, column=3).alignment = align_center
            ws_yr.cell(row=row_idx, column=4).alignment = align_left
            ws_yr.cell(row=row_idx, column=5).alignment = align_center
            ws_yr.cell(row=row_idx, column=6).alignment = align_center
            ws_yr.cell(row=row_idx, column=7).alignment = align_center
            ws_yr.cell(row=row_idx, column=8).alignment = align_center
            ws_yr.cell(row=row_idx, column=9).alignment = align_center
            ws_yr.cell(row=row_idx, column=10).number_format = '₹#,##0.00'
            ws_yr.cell(row=row_idx, column=11).alignment = align_center
            ws_yr.cell(row=row_idx, column=12).number_format = '₹#,##0.00'
            
            cell_ret = ws_yr.cell(row=row_idx, column=13)
            cell_ret.number_format = '+0.00%;-0.00%;0.00%'
            cell_ret.alignment = align_right
            
            ws_yr.cell(row=row_idx, column=14).number_format = '#,##0'
            
            cell_pnl = ws_yr.cell(row=row_idx, column=15)
            cell_pnl.number_format = '₹#,##0.00;[Red]-₹#,##0.00;₹0.00'
            cell_pnl.alignment = align_right
            
            cell_status = ws_yr.cell(row=row_idx, column=16)
            cell_status.alignment = align_center
            if t['status'] == 'WIN':
                cell_status.font = font_win
                cell_status.fill = win_fill
            else:
                cell_status.font = font_loss
                cell_status.fill = loss_fill
                
            ws_yr.cell(row=row_idx, column=17).number_format = '₹#,##0.00'
            
            cell_roc = ws_yr.cell(row=row_idx, column=18)
            cell_roc.number_format = '+0.00%;-0.00%;0.00%'
            cell_roc.alignment = align_right
            
            for c in range(1, 19):
                ws_yr.cell(row=row_idx, column=c).border = border_thin
                if row_idx % 2 == 0:
                    ws_yr.cell(row=row_idx, column=c).fill = zebra_fill
            row_idx += 1
            
        for col in ws_yr.columns:
            col_letter = get_column_letter(col[0].column)
            ws_yr.column_dimensions[col_letter].width = 18
        ws_yr.column_dimensions['D'].width = 28
        
        yr_file = holiday_subfolder / f"{h['sheet'].replace(' ', '_')}_{yr}_TradeLog.xlsx"
        wb_yr.save(yr_file)

    # 2. GENERATE 4-YEAR CONSOLIDATED MASTER EXCEL FILE FOR THIS HOLIDAY
    wb_cons = openpyxl.Workbook()
    ws_cons = wb_cons.active
    ws_cons.title = "4-Year Consolidated Summary"
    
    ws_cons.merge_cells("A1:K1")
    t_c = ws_cons["A1"]
    t_c.value = f"NIFTY 50 FUTURES — {h['sheet'].upper()} (BEHAVIORAL ENGINE 4-YEAR CONSOLIDATED LEADERBOARD)"
    t_c.font = font_title
    t_c.fill = title_fill
    t_c.alignment = align_center
    ws_cons.row_dimensions[1].height = 40
    
    ws_cons.append([])
    ws_cons.append(headers_sum)
    ws_cons.row_dimensions[3].height = 28
    for c_idx in range(1, len(headers_sum) + 1):
        cell = ws_cons.cell(row=3, column=c_idx)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border_thin
        
    stock_4y_agg = []
    for sym, spec in nifty50_specs.items():
        sym_trades = []
        for yr in years:
            for t in holiday_trades_by_year[yr]:
                if t['sym'] == sym:
                    sym_trades.append(t)
                    
        if not sym_trades:
            continue
            
        wins = sum(1 for t in sym_trades if t['status'] == 'WIN')
        tot = len(sym_trades)
        wr = (wins / tot * 100.0) if tot > 0 else 0.0
        
        pnl_2022 = sum(t['pnl'] for t in sym_trades if t['year'] == 2022)
        pnl_2023 = sum(t['pnl'] for t in sym_trades if t['year'] == 2023)
        pnl_2024 = sum(t['pnl'] for t in sym_trades if t['year'] == 2024)
        pnl_2025 = sum(t['pnl'] for t in sym_trades if t['year'] == 2025)
        tot_pnl = pnl_2022 + pnl_2023 + pnl_2024 + pnl_2025
        
        stock_4y_agg.append({
            'sym': sym,
            'name': spec['name'],
            'window': sym_trades[0]['window'],
            'strat': sym_trades[0]['strat'],
            'wr': wr,
            'pnl_2022': pnl_2022,
            'pnl_2023': pnl_2023,
            'pnl_2024': pnl_2024,
            'pnl_2025': pnl_2025,
            'tot_pnl': tot_pnl
        })
        
    stock_4y_agg = sorted(stock_4y_agg, key=lambda x: x['tot_pnl'], reverse=True)
    
    row_idx = 4
    for rank, r in enumerate(stock_4y_agg, 1):
        ws_cons.append([
            f"Rank {rank}", r['sym'], r['name'], r['window'], r['strat'], r['wr'] / 100.0,
            r['pnl_2022'], r['pnl_2023'], r['pnl_2024'], r['pnl_2025'], r['tot_pnl']
        ])
        ws_cons.row_dimensions[row_idx].height = 20
        
        ws_cons.cell(row=row_idx, column=1).alignment = align_center
        ws_cons.cell(row=row_idx, column=1).font = Font(name="Calibri", size=10, bold=True, color="1F4E78")
        ws_cons.cell(row=row_idx, column=2).alignment = align_center
        ws_cons.cell(row=row_idx, column=3).alignment = align_left
        ws_cons.cell(row=row_idx, column=4).alignment = align_center
        ws_cons.cell(row=row_idx, column=5).alignment = align_center
        
        cell_wr = ws_cons.cell(row=row_idx, column=6)
        cell_wr.number_format = '0.00%'
        cell_wr.alignment = align_right
        cell_wr.font = Font(name="Calibri", size=10, bold=True)
        
        ws_cons.cell(row=row_idx, column=7).number_format = '₹#,##0.00;[Red]-₹#,##0.00;₹0.00'
        ws_cons.cell(row=row_idx, column=8).number_format = '₹#,##0.00;[Red]-₹#,##0.00;₹0.00'
        ws_cons.cell(row=row_idx, column=9).number_format = '₹#,##0.00;[Red]-₹#,##0.00;₹0.00'
        ws_cons.cell(row=row_idx, column=10).number_format = '₹#,##0.00;[Red]-₹#,##0.00;₹0.00'
        
        cell_tot = ws_cons.cell(row=row_idx, column=11)
        cell_tot.number_format = '₹#,##0.00;[Red]-₹#,##0.00;₹0.00'
        cell_tot.alignment = align_right
        cell_tot.font = Font(name="Calibri", size=10, bold=True)
        
        for c in range(1, 12):
            ws_cons.cell(row=row_idx, column=c).border = border_thin
            if row_idx % 2 == 0:
                ws_cons.cell(row=row_idx, column=c).fill = zebra_fill
        row_idx += 1

    for col in ws_cons.columns:
        col_letter = get_column_letter(col[0].column)
        ws_cons.column_dimensions[col_letter].width = 22
    ws_cons.column_dimensions['C'].width = 28

    cons_file = holiday_subfolder / f"{h['sheet'].replace(' ', '_')}_4Year_Consolidated.xlsx"
    wb_cons.save(cons_file)
    
    # 3. GENERATE ALSO THE MULTI-SHEET MAIN EXCEL WORKBOOK FOR THE HOLIDAY
    wb_multi = openpyxl.Workbook()
    wb_multi.remove(wb_multi.active)
    
    for yr in years:
        ws_yr = wb_multi.create_sheet(title=f"Year {yr}")
        ws_yr.merge_cells("A1:R1")
        title_cell = ws_yr["A1"]
        title_cell.value = f"NIFTY 50 FUTURES — {h['sheet'].upper()} ({yr} BEHAVIORAL TRADE LOG)"
        title_cell.font = font_title
        title_cell.fill = title_fill
        title_cell.alignment = align_center
        ws_yr.row_dimensions[1].height = 40
        
        ws_yr.merge_cells("A2:R2")
        ws_yr["A2"] = " "
        
        ws_yr.append(headers_log)
        ws_yr.row_dimensions[3].height = 28
        for col_idx in range(1, len(headers_log) + 1):
            cell = ws_yr.cell(row=3, column=col_idx)
            cell.font = font_header
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = border_thin
            
        yr_trades = sorted(holiday_trades_by_year[yr], key=lambda x: x['pnl'], reverse=True)
        row_idx = 4
        for rank, t in enumerate(yr_trades, 1):
            ws_yr.append([
                t['id'], f"Rank {rank}", t['sym'], t['name'], t['holiday'], t['h_date'],
                t['window'], t['strat'], t['entry_date'], t['entry_price'], t['exit_date'], t['exit_price'],
                t['ret_pct'] / 100.0, t['lot'], t['pnl'], t['status'], t['margin'], t['roc_pct'] / 100.0
            ])
            ws_yr.row_dimensions[row_idx].height = 20
            
            ws_yr.cell(row=row_idx, column=1).alignment = align_center
            ws_yr.cell(row=row_idx, column=2).alignment = align_center
            ws_yr.cell(row=row_idx, column=2).font = Font(name="Calibri", size=10, bold=True, color="1F4E78")
            ws_yr.cell(row=row_idx, column=3).alignment = align_center
            ws_yr.cell(row=row_idx, column=4).alignment = align_left
            ws_yr.cell(row=row_idx, column=5).alignment = align_center
            ws_yr.cell(row=row_idx, column=6).alignment = align_center
            ws_yr.cell(row=row_idx, column=7).alignment = align_center
            ws_yr.cell(row=row_idx, column=8).alignment = align_center
            ws_yr.cell(row=row_idx, column=9).alignment = align_center
            ws_yr.cell(row=row_idx, column=10).number_format = '₹#,##0.00'
            ws_yr.cell(row=row_idx, column=11).alignment = align_center
            ws_yr.cell(row=row_idx, column=12).number_format = '₹#,##0.00'
            
            cell_ret = ws_yr.cell(row=row_idx, column=13)
            cell_ret.number_format = '+0.00%;-0.00%;0.00%'
            cell_ret.alignment = align_right
            
            ws_yr.cell(row=row_idx, column=14).number_format = '#,##0'
            
            cell_pnl = ws_yr.cell(row=row_idx, column=15)
            cell_pnl.number_format = '₹#,##0.00;[Red]-₹#,##0.00;₹0.00'
            cell_pnl.alignment = align_right
            
            cell_status = ws_yr.cell(row=row_idx, column=16)
            cell_status.alignment = align_center
            if t['status'] == 'WIN':
                cell_status.font = font_win
                cell_status.fill = win_fill
            else:
                cell_status.font = font_loss
                cell_status.fill = loss_fill
                
            ws_yr.cell(row=row_idx, column=17).number_format = '₹#,##0.00'
            
            cell_roc = ws_yr.cell(row=row_idx, column=18)
            cell_roc.number_format = '+0.00%;-0.00%;0.00%'
            cell_roc.alignment = align_right
            
            for c in range(1, 19):
                ws_yr.cell(row=row_idx, column=c).border = border_thin
                if row_idx % 2 == 0:
                    ws_yr.cell(row=row_idx, column=c).fill = zebra_fill
            row_idx += 1
            
        for col in ws_yr.columns:
            col_letter = get_column_letter(col[0].column)
            ws_yr.column_dimensions[col_letter].width = 18
        ws_yr.column_dimensions['D'].width = 28

    ws_sum_m = wb_multi.create_sheet(title="4-Year Consolidated Summary")
    ws_sum_m.merge_cells("A1:K1")
    t_m_c = ws_sum_m["A1"]
    t_m_c.value = f"NIFTY 50 FUTURES — {h['sheet'].upper()} (4-YEAR CONSOLIDATED LEADERBOARD)"
    t_m_c.font = font_title
    t_m_c.fill = title_fill
    t_m_c.alignment = align_center
    ws_sum_m.row_dimensions[1].height = 40
    
    ws_sum_m.append([])
    ws_sum_m.append(headers_sum)
    ws_sum_m.row_dimensions[3].height = 28
    for c_idx in range(1, len(headers_sum) + 1):
        cell = ws_sum_m.cell(row=3, column=c_idx)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border_thin
        
    row_idx = 4
    for rank, r in enumerate(stock_4y_agg, 1):
        ws_sum_m.append([
            f"Rank {rank}", r['sym'], r['name'], r['window'], r['strat'], r['wr'] / 100.0,
            r['pnl_2022'], r['pnl_2023'], r['pnl_2024'], r['pnl_2025'], r['tot_pnl']
        ])
        ws_sum_m.row_dimensions[row_idx].height = 20
        
        ws_sum_m.cell(row=row_idx, column=1).alignment = align_center
        ws_sum_m.cell(row=row_idx, column=1).font = Font(name="Calibri", size=10, bold=True, color="1F4E78")
        ws_sum_m.cell(row=row_idx, column=2).alignment = align_center
        ws_sum_m.cell(row=row_idx, column=3).alignment = align_left
        ws_sum_m.cell(row=row_idx, column=4).alignment = align_center
        ws_sum_m.cell(row=row_idx, column=5).alignment = align_center
        
        cell_wr = ws_sum_m.cell(row=row_idx, column=6)
        cell_wr.number_format = '0.00%'
        cell_wr.alignment = align_right
        cell_wr.font = Font(name="Calibri", size=10, bold=True)
        
        ws_sum_m.cell(row=row_idx, column=7).number_format = '₹#,##0.00;[Red]-₹#,##0.00;₹0.00'
        ws_sum_m.cell(row=row_idx, column=8).number_format = '₹#,##0.00;[Red]-₹#,##0.00;₹0.00'
        ws_sum_m.cell(row=row_idx, column=9).number_format = '₹#,##0.00;[Red]-₹#,##0.00;₹0.00'
        ws_sum_m.cell(row=row_idx, column=10).number_format = '₹#,##0.00;[Red]-₹#,##0.00;₹0.00'
        
        cell_tot = ws_sum_m.cell(row=row_idx, column=11)
        cell_tot.number_format = '₹#,##0.00;[Red]-₹#,##0.00;₹0.00'
        cell_tot.alignment = align_right
        cell_tot.font = Font(name="Calibri", size=10, bold=True)
        
        for c in range(1, 12):
            ws_sum_m.cell(row=row_idx, column=c).border = border_thin
            if row_idx % 2 == 0:
                ws_sum_m.cell(row=row_idx, column=c).fill = zebra_fill
        row_idx += 1

    for col in ws_sum_m.columns:
        col_letter = get_column_letter(col[0].column)
        ws_sum_m.column_dimensions[col_letter].width = 22
    ws_sum_m.column_dimensions['C'].width = 28

    main_multi_file = MAIN_HOLIDAY_DIR / h['filename']
    wb_multi.save(main_multi_file)
    
    print(f"  SUCCESS: Generated Holiday Folder [{h['folder']}] -> 4 Yearly Excel Files + 1 Consolidated Excel File")

print("="*90)
print(f"SUCCESS: Created All 17 Holiday Directories in:\n  -> {REPORTS_DIR.resolve()}")

# Build Master Consolidated Excel File across all 17 holidays
wb_master = openpyxl.Workbook()
wb_master.remove(wb_master.active)

ws_m_sum = wb_master.create_sheet(title="All Holidays Leaderboard")
ws_m_sum.merge_cells("A1:H1")
t_m = ws_m_sum["A1"]
t_m.value = "NIFTY 50 FUTURES ALL 17 HOLIDAYS — BEHAVIORAL ENGINE OPTIMAL LEADERBOARD (2022 - 2025)"
t_m.font = font_title
t_m.fill = title_fill
t_m.alignment = align_center
ws_m_sum.row_dimensions[1].height = 40

headers_m_sum = ["Rank", "Holiday Event", "Optimal Strategy Logic", "Total Trades (4Y)", "Wins", "Losses", "Win Rate (%)", "4-Year Total PnL / Lot (₹)"]
ws_m_sum.append([])
ws_m_sum.append(headers_m_sum)
ws_m_sum.row_dimensions[3].height = 28
for c_idx in range(1, len(headers_m_sum) + 1):
    cell = ws_m_sum.cell(row=3, column=c_idx)
    cell.font = font_header
    cell.fill = header_fill
    cell.alignment = align_center
    cell.border = border_thin

df_master = pd.DataFrame(all_master_trades)

m_summary_rows = []
for h in all_holidays:
    df_h = df_master[df_master['holiday'] == h['sheet']]
    tot = len(df_h)
    wins = len(df_h[df_h['status'] == 'WIN'])
    losses = tot - wins
    wr = (wins / tot * 100.0) if tot > 0 else 0.0
    tot_pnl = df_h['pnl'].sum() if tot > 0 else 0.0
    m_summary_rows.append({
        'holiday': h['sheet'],
        'bias': "BEHAVIORAL OPTIMAL (LONG & SHORT)",
        'tot': tot,
        'wins': wins,
        'losses': losses,
        'wr': wr,
        'pnl': tot_pnl
    })

df_m_summary = pd.DataFrame(m_summary_rows).sort_values(by=['wr', 'pnl'], ascending=[False, False]).reset_index(drop=True)

row_idx = 4
for rank, r in enumerate(df_m_summary.to_dict('records'), 1):
    ws_m_sum.append([
        f"Rank {rank}", r['holiday'], r['bias'], r['tot'], r['wins'], r['losses'], r['wr'] / 100.0, r['pnl']
    ])
    ws_m_sum.row_dimensions[row_idx].height = 20
    
    ws_m_sum.cell(row=row_idx, column=1).font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
    ws_m_sum.cell(row=row_idx, column=1).alignment = align_center
    ws_m_sum.cell(row=row_idx, column=2).alignment = align_left
    ws_m_sum.cell(row=row_idx, column=3).alignment = align_center
    ws_m_sum.cell(row=row_idx, column=4).alignment = align_center
    ws_m_sum.cell(row=row_idx, column=5).alignment = align_center
    ws_m_sum.cell(row=row_idx, column=6).alignment = align_center
    
    cell_wr = ws_m_sum.cell(row=row_idx, column=7)
    cell_wr.number_format = '0.00%'
    cell_wr.alignment = align_right
    cell_wr.font = Font(name="Calibri", size=11, bold=True)
    
    cell_pnl = ws_m_sum.cell(row=row_idx, column=8)
    cell_pnl.number_format = '₹#,##0.00;[Red]-₹#,##0.00;₹0.00'
    cell_pnl.alignment = align_right
    cell_pnl.font = Font(name="Calibri", size=11, bold=True)
    
    for c in range(1, 9):
        ws_m_sum.cell(row=row_idx, column=c).border = border_thin
        if row_idx % 2 == 0:
            ws_m_sum.cell(row=row_idx, column=c).fill = zebra_fill
    row_idx += 1

for col in ws_m_sum.columns:
    col_letter = get_column_letter(col[0].column)
    ws_m_sum.column_dimensions[col_letter].width = 24

try:
    wb_master.save(MASTER_EXCEL_OUT)
    print(f"SUCCESS: Created Master Consolidated Workbook -> {MASTER_EXCEL_OUT.name}")
except Exception as e:
    alt_out = ROOT / 'Nifty50_Futures_All_Holidays_Past_4Years_Master_TradeLog_v5.xlsx'
    wb_master.save(alt_out)
    print(f"SUCCESS: Created Master Consolidated Workbook -> {alt_out.name}")
print("==========================================================================================")
