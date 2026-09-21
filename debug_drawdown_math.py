import openpyxl, pandas as pd, sys

sys.stdout.reconfigure(errors='replace')

master_clean_path = r'D:/behaviour analysis/Nifty50_12_Quarters_Futures_Combined_Turnover_0.5PCT_Cost_Master_CLEAN.xlsx'
wb = openpyxl.load_workbook(master_clean_path, data_only=True)

ws = wb['Q3 2023-24']
print("==========================================================================================")
print("INSPECTING Q3 2023-24 DRAWDOWN COLUMN VALUES vs CALCULATED PEAK")
print("==========================================================================================")

cum_pnl = 0.0
running_max = 0.0

for r in range(8, 28):
    sym     = ws.cell(row=r, column=2).value
    net_pnl = float(ws.cell(row=r, column=18).value or 0)
    c_pnl   = float(ws.cell(row=r, column=19).value or 0)
    dd      = float(ws.cell(row=r, column=20).value or 0)
    
    cum_pnl += net_pnl
    running_max = max(running_max, cum_pnl)
    calc_dd = cum_pnl - running_max
    
    print(f"Row {r:2d} ({sym[:10]:10s}): Net=₹{net_pnl:10.2f} | Cum=₹{c_pnl:10.2f} | Peak=₹{running_max:10.2f} | Cell DD=₹{dd:10.2f} | Calc DD=₹{calc_dd:10.2f} | Diff? {abs(dd - calc_dd) >= 0.01}")

print("==========================================================================================")
