import openpyxl, pandas as pd, sys

sys.stdout.reconfigure(errors='replace')

summary_path = r'D:/behaviour analysis/Nifty50_12_Quarters_Options_1PCT_ITM_Summary.xlsx'
wb = openpyxl.load_workbook(summary_path, data_only=True)
ws = wb['Options_1PCT_ITM_Summary']

print("==========================================================================================")
print("VERIFYING EXACT DAY-BY-DAY PEAK CONCURRENT CAPITAL SUMMARY WORKBOOK")
print("==========================================================================================")
print("Sheet Names:", wb.sheetnames)

print("\nSummary Table Rows 7-21:")
for r in range(7, 22):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 9)]
    print(f"  Row {r:2d}: {vals}")
print("==========================================================================================")
