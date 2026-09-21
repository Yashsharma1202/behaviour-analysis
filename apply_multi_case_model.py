import datetime as dt
import json
import sys
import time
import urllib.request
import warnings
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

warnings.filterwarnings("ignore")
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

ROOT       = Path(__file__).resolve().parent
SRC_PARAMS = ROOT / "UPDATED.xlsx"
SRC_DATES  = ROOT / "Nifty50Stocks_QtyResultDates.xlsx"
OUT_EXCEL  = ROOT / "Nifty50Stocks_MultiCase_Analysis.xlsx"
SYMBOL_MAP = {"TATAMOTORS": "TMPV"}


# ─────────────────────────── helpers ────────────────────────────────────────

def fetch_yahoo_history(sym: str):
    query_sym = SYMBOL_MAP.get(sym, sym)
    t  = query_sym.replace("&", "%26") + ".NS"
    p1 = int(dt.datetime(2024, 6, 1).timestamp())
    p2 = int((dt.datetime.now() + pd.Timedelta(days=2)).timestamp())
    url = (f"https://query1.finance.yahoo.com/v8/finance/chart/{t}"
           f"?period1={p1}&period2={p2}&interval=1d")
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=25) as r:
            j = json.loads(r.read())
        res = j["chart"]["result"][0]
        idx = pd.to_datetime(res["timestamp"], unit="s").normalize()
        q   = res["indicators"]["quote"][0]
        df  = pd.DataFrame(
            {"close": q["close"], "low": q["low"], "high": q["high"]}, index=idx
        ).dropna()
        return df
    except Exception as e:
        print(f"  ! fetch {sym} ({query_sym}): {e}")
        return None


def get_trading_date(cal: pd.DatetimeIndex, result_date, offset: int):
    if cal is None or len(cal) == 0:
        return None
    D = pd.Timestamp(result_date).normalize()
    if D <= cal[-1]:
        pos        = cal.searchsorted(D, side="right") - 1
        if pos < 0: pos = 0
        target_idx = pos + offset
        if 0 <= target_idx < len(cal):
            return cal[target_idx].to_pydatetime()
        elif target_idx < 0:
            return result_date + dt.timedelta(days=offset)
        else:
            rem    = target_idx - (len(cal) - 1)
            future = pd.bdate_range(start=cal[-1], periods=rem + 1)
            return future[-1].to_pydatetime()
    else:
        b_days     = pd.bdate_range(start=cal[-1], end=D)
        n          = len(b_days) - 1
        target_idx = n + offset
        if target_idx >= 0:
            future = pd.bdate_range(start=cal[-1], periods=target_idx + 1)
            return future[-1].to_pydatetime()
        else:
            back = len(cal) - 1 + target_idx
            if back >= 0:
                return cal[back].to_pydatetime()
            return D + dt.timedelta(days=offset)


def get_price_on_date(s: pd.Series, d):
    ts = pd.Timestamp(d).normalize()
    if ts in s.index:
        v = s.loc[ts]
        return float(v) if pd.notna(v) else None
    prior = s.index[s.index <= ts]
    if len(prior) > 0:
        v = s.loc[prior[-1]]
        return float(v) if pd.notna(v) else None
    return None


def check_sl_hit_fixed(df: pd.DataFrame, buy_ts: pd.Timestamp,
                       until_ts: pd.Timestamp, sl_price: float):
    """
    Scan daily lows from day after buy_ts up to until_ts.
    Exit at exactly sl_price when any daily low < sl_price.
    Returns (sl_hit, trigger_date, exit_price).
    """
    mask   = (df.index > buy_ts) & (df.index <= until_ts)
    period = df[mask]
    for date, row in period.iterrows():
        if pd.notna(row["low"]) and row["low"] < sl_price:
            return True, date.to_pydatetime(), sl_price
    return False, None, None


def check_sl_hit(df: pd.DataFrame, buy_ts: pd.Timestamp,
                 until_ts: pd.Timestamp):
    """
    Scan daily lows from day after buy_ts up to until_ts.
    Exit at exactly that day's ema50_low when daily low < ema50_low.
    Returns (sl_hit, trigger_date, exit_price).
    """
    mask   = (df.index > buy_ts) & (df.index <= until_ts)
    period = df[mask]
    for date, row in period.iterrows():
        daily_ema_low = row.get("ema50_low")
        if pd.notna(row["low"]) and pd.notna(daily_ema_low) and row["low"] < daily_ema_low:
            return True, date.to_pydatetime(), float(daily_ema_low)
    return False, None, None


def apply_case(case_num: int, df: pd.DataFrame,
               ema_high: pd.Series, ema_low: pd.Series,
               entry_dt: pd.Timestamp, exit_dt: pd.Timestamp,
               latest_dt: pd.Timestamp) -> dict:
    """
    Returns result dict with keys:
        status, actual_buy_date, buy_price, sl_price,
        actual_exit_date, exit_price, sl_hit, ret
    """
    res = dict(
        status="Filter Out",
        actual_buy_date=None, buy_price=None, sl_price=None,
        actual_exit_date=None, exit_price=None, sl_hit=False, ret=None
    )

    # Upcoming check (before any expensive lookups)
    if entry_dt > latest_dt:
        res["status"] = "Upcoming"
        return res

    close_e = get_price_on_date(df["close"], entry_dt)
    ema_h_e = get_price_on_date(ema_high,    entry_dt)
    ema_l_e = get_price_on_date(ema_low,     entry_dt)

    actual_buy_ts = None
    buy_price     = None
    sl_price      = None

    # ─── Case 1 ─ No filter ──────────────────────────────────────────────────
    if case_num == 1:
        actual_buy_ts = entry_dt
        buy_price     = close_e
        sl_price      = None

    # ─── Case 2 ─ Buy only above EMA50 High band ─────────────────────────────
    elif case_num == 2:
        if (close_e is not None and ema_h_e is not None
                and close_e > ema_h_e):
            actual_buy_ts = entry_dt
            buy_price     = close_e
            sl_price      = ema_l_e

    # ─── Case 3 ─ Buy only WITHIN EMA band ───────────────────────────────────
    elif case_num == 3:
        if (close_e is not None and ema_h_e is not None and ema_l_e is not None
                and ema_l_e <= close_e <= ema_h_e):
            actual_buy_ts = entry_dt
            buy_price     = close_e
            sl_price      = ema_l_e

    # ─── Case 4 ─ Entry below Low band → watch for cross above High ──────────
    elif case_num == 4:
        # Only start watching if entry candle is BELOW EMA Low
        if (close_e is not None and ema_l_e is not None
                and close_e < ema_l_e):
            mask   = (df.index > entry_dt) & (df.index <= exit_dt)
            period = df[mask]
            for date in period.index:
                day_close = period.loc[date, "close"]
                day_ema_h = get_price_on_date(ema_high, date)
                if (day_close is not None and day_ema_h is not None
                        and float(day_close) > day_ema_h):
                    actual_buy_ts = date
                    buy_price     = float(day_close)
                    sl_price      = get_price_on_date(ema_low, date)
                    break

    # ─── Case 5 ─ Within band=Drop; Below=Watch for cross above High ─────────
    elif case_num == 5:
        if (close_e is not None and ema_h_e is not None and ema_l_e is not None):
            if ema_l_e <= close_e <= ema_h_e:
                pass   # Within band → drop (Filter Out)
            elif close_e < ema_l_e:
                # Below Low → scan for close crossing above High
                mask   = (df.index > entry_dt) & (df.index <= exit_dt)
                period = df[mask]
                for date in period.index:
                    day_close = period.loc[date, "close"]
                    day_ema_h = get_price_on_date(ema_high, date)
                    if (day_close is not None and day_ema_h is not None
                            and float(day_close) > day_ema_h):
                        actual_buy_ts = date
                        buy_price     = float(day_close)
                        sl_price      = get_price_on_date(ema_low, date)
                        break
            # else: above High at entry → drop

    # ─── Case 6 ─ No Filter, With Fixed EMA50 Low Stop Loss ────────────────────
    elif case_num == 6:
        actual_buy_ts = entry_dt
        buy_price     = close_e
        sl_price      = ema_l_e

    # ─── Case 7 ─ No Filter, With Trailing EMA50 Low Stop Loss ─────────────────
    elif case_num == 7:
        actual_buy_ts = entry_dt
        buy_price     = close_e
        sl_price      = ema_l_e

    # ─── No trade found ───────────────────────────────────────────────────────
    if buy_price is None or actual_buy_ts is None:
        return res   # status = Filter Out

    buy_ts_norm = pd.Timestamp(actual_buy_ts).normalize()
    if buy_ts_norm > latest_dt:
        res["status"] = "Upcoming"
        return res

    res["actual_buy_date"] = (actual_buy_ts if isinstance(actual_buy_ts, dt.datetime)
                              else actual_buy_ts.to_pydatetime())
    res["buy_price"]       = round(buy_price, 2)
    res["sl_price"]        = round(sl_price, 2) if sl_price is not None else None

    # ─── Determine exit ───────────────────────────────────────────────────────
    if exit_dt <= latest_dt:
        # Planned exit has passed → Realised
        if sl_price is not None:
            if case_num == 6:
                sl_hit, sl_date, sl_ep = check_sl_hit_fixed(df, buy_ts_norm, exit_dt, sl_price)
            else:
                sl_hit, sl_date, sl_ep = check_sl_hit(df, buy_ts_norm, exit_dt)
        else:
            sl_hit, sl_date, sl_ep = False, None, None

        if sl_hit:
            res["sl_hit"]           = True
            res["actual_exit_date"] = sl_date
            res["exit_price"]       = round(sl_ep, 2)
        else:
            ep = get_price_on_date(df["close"], exit_dt)
            res["actual_exit_date"] = exit_dt.to_pydatetime()
            res["exit_price"]       = round(ep, 2) if ep is not None else None

        if res["exit_price"] and res["buy_price"]:
            res["ret"]    = round((res["exit_price"] / res["buy_price"] - 1) * 100, 2)
            res["status"] = "Realised (SL)" if res["sl_hit"] else "Realised"

    else:
        # Exit in future → Open / Till-Date
        if sl_price is not None:
            if case_num == 6:
                sl_hit, sl_date, sl_ep = check_sl_hit_fixed(df, buy_ts_norm, latest_dt, sl_price)
            else:
                sl_hit, sl_date, sl_ep = check_sl_hit(df, buy_ts_norm, latest_dt)
        else:
            sl_hit, sl_date, sl_ep = False, None, None

        if sl_hit:
            res["sl_hit"]           = True
            res["actual_exit_date"] = sl_date
            res["exit_price"]       = round(sl_ep, 2)
            res["ret"]              = round((sl_ep / buy_price - 1) * 100, 2)
            res["status"]           = "Realised (SL)"
        else:
            latest_close            = float(df["close"].iloc[-1])
            res["actual_exit_date"] = df.index[-1].to_pydatetime()
            res["exit_price"]       = round(latest_close, 2)
            res["ret"]              = round((latest_close / buy_price - 1) * 100, 2)
            res["status"]           = "Open"

    return res


def load_parameters() -> dict:
    print("Loading model parameters from UPDATED.xlsx...")
    params = {}
    try:
        wb_p = openpyxl.load_workbook(SRC_PARAMS, data_only=True)
        ws_p = wb_p.active
        # Row 1 = title, Row 2 = column headers, Data starts at Row 3
        # Col 2 = SYMBOL, Col 5 = Buy Before, Col 6 = Sell After, Col 9 = Expected Return
        for r in range(3, ws_p.max_row + 1):
            sym = ws_p.cell(r, 2).value
            if sym:
                sym = str(sym).strip()
                bb  = ws_p.cell(r, 5).value   # Buy Before (Candles)
                sa  = ws_p.cell(r, 6).value   # Sell After (Candles)
                exp = ws_p.cell(r, 9).value   # Expected Return
                if bb is not None and sa is not None:
                    try:
                        params[sym] = (int(bb), int(sa),
                                       round(float(exp), 2) if exp is not None else 0.0)
                    except Exception:
                        pass
        wb_p.close()
    except Exception as e:
        print(f"  ! Error loading UPDATED.xlsx: {e}")

    # Fallback from optimal_rules.csv for historical/missing symbols
    fallback_csv = ROOT / "processed" / "optimal_rules.csv"
    if fallback_csv.exists():
        print("Loading fallback parameters from processed/optimal_rules.csv...")
        try:
            df_rules = pd.read_csv(fallback_csv)
            df_res   = df_rules[
                (df_rules["event_type"] == "RESULTS") &
                (df_rules["mode"] == "descriptive")
            ]
            for _, row in df_res.iterrows():
                s = str(row["symbol"]).strip()
                if s not in params:
                    buy_val = int(row["buy"])
                    bb_val  = -buy_val   # negative offset = candles BEFORE result
                    sa_val  = int(row["sell"])
                    exp_val = round(float(row["train_avg"]), 2)
                    params[s] = (bb_val, sa_val, exp_val)
                    print(f"  Added fallback for {s}: BB={bb_val}, SA={sa_val}, Exp={exp_val}%")
        except Exception as e:
            print(f"  ! Error loading optimal_rules.csv fallbacks: {e}")

    print(f"Loaded parameters for {len(params)} symbols.")
    return params


def make_styles():
    return dict(
        hdr    = Font(name="Segoe UI", size=9, bold=True, color="FFFFFFFF"),
        normal = Font(name="Segoe UI", size=9),
        ital   = Font(name="Segoe UI", size=9, italic=True, color="FF555555"),
        navy   = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid"),
        soft   = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid"),
        gfill  = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        gtxt   = Font(name="Segoe UI", size=9, color="FF006100"),
        rfill  = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        rtxt   = Font(name="Segoe UI", size=9, color="FF9C0006"),
        yfill  = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
        ytxt   = Font(name="Segoe UI", size=9, color="FF9C6500"),
        ofill  = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
        otxt   = Font(name="Segoe UI", size=9, color="FF833C00"),
        border = Border(
            left   = Side(style="thin", color="D9D9D9"),
            right  = Side(style="thin", color="D9D9D9"),
            top    = Side(style="thin", color="D9D9D9"),
            bottom = Side(style="thin", color="D9D9D9"),
        ),
        ac = Alignment(horizontal="center", vertical="center"),
        ar = Alignment(horizontal="right",  vertical="center"),
        al = Alignment(horizontal="left",   vertical="center"),
    )


CASE_HEADERS = [
    "Company", "Symbol", "Quarter", "Result Date",
    "Buy Before", "Sell After", "Expected Return (%)",
    "Planned Entry Date", "Actual Buy Date", "Buy Price",
    "SL Price", "Planned Exit Date", "Actual Exit Date",
    "Exit Price", "SL Hit", "Realised Return (%)", "Status"
]

CASES = [
    (1, "Case1_NoFilter",        "Case 1 — No Filter (Simple Buy/Sell)"),
    (2, "Case2_AboveHigh",       "Case 2 — Buy Only Above EMA50 High Band"),
    (3, "Case3_WithinBand",      "Case 3 — Buy Only Within EMA50 Band"),
    (4, "Case4_BelowWatchCross", "Case 4 — Entry Below Band, Watch Cross Above High"),
    (5, "Case5_NoWithin",        "Case 5 — Within=Drop; Below=Watch Cross Above High"),
    (6, "Case6_NoFilter_WithSL",  "Case 6 — No Filter, With Fixed EMA50 Low Stop Loss"),
    (7, "Case7_NoFilter_TrailSL", "Case 7 — No Filter, With Trailing EMA50 Low Stop Loss"),
]


def write_header_row(ws, headers, S):
    ws.row_dimensions[1].height = 22
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = S["hdr"]; cell.fill = S["navy"]
        cell.alignment = S["ac"]; cell.border = S["border"]


def write_case_row(ws, r, data: dict, res: dict, entry_date, exit_date,
                   case_num: int, S: dict):
    bg = S["soft"] if r % 2 == 0 else PatternFill(fill_type=None)

    def wc(col, val, fnt=None, fll=None, nfmt=None, aln=None):
        cell = ws.cell(row=r, column=col, value=val)
        cell.font      = fnt or S["normal"]
        cell.fill      = fll or bg
        cell.border    = S["border"]
        cell.alignment = aln or S["ac"]
        if nfmt:
            cell.number_format = nfmt
        return cell

    wc(1, data["company"],      aln=S["al"])
    wc(2, data["symbol"])
    wc(3, data["quarter"])
    wc(4, data["result_date"],  nfmt="yyyy-mm-dd")
    wc(5, data["buy_before"])
    wc(6, data["sell_after"])
    wc(7, data["exp"],          nfmt="0.00")
    wc(8, entry_date,           nfmt="yyyy-mm-dd")

    status = res["status"]

    if status in ("Upcoming", "Filter Out"):
        for col in range(9, len(CASE_HEADERS) + 1):
            wc(col, "—")
        if status == "Upcoming":
            wc(17, "Upcoming",   fnt=S["ital"])
        else:
            wc(17, "Filter Out", fnt=S["rtxt"], fll=S["rfill"])
        return

    wc(9,  res["actual_buy_date"],  nfmt="yyyy-mm-dd")
    wc(10, res["buy_price"],        nfmt="#,##0.00", aln=S["ar"])

    sl_val = res["sl_price"] if res["sl_price"] is not None else "N/A"
    wc(11, sl_val, nfmt="#,##0.00" if isinstance(sl_val, (int, float)) else None,
       aln=S["ar"])

    wc(12, exit_date,               nfmt="yyyy-mm-dd")
    wc(13, res["actual_exit_date"], nfmt="yyyy-mm-dd")
    wc(14, res["exit_price"],       nfmt="#,##0.00", aln=S["ar"])

    if case_num == 1:
        wc(15, "N/A")
    elif res["sl_hit"]:
        wc(15, "YES", fnt=S["rtxt"], fll=S["rfill"])
    else:
        wc(15, "No",  fnt=S["gtxt"], fll=S["gfill"])

    ret = res["ret"]
    if ret is not None:
        if res["sl_hit"]:
            wc(16, ret, nfmt="0.00", aln=S["ar"], fnt=S["otxt"], fll=S["ofill"])
        elif ret >= 0:
            wc(16, ret, nfmt="0.00", aln=S["ar"], fnt=S["gtxt"], fll=S["gfill"])
        else:
            wc(16, ret, nfmt="0.00", aln=S["ar"], fnt=S["rtxt"], fll=S["rfill"])
    else:
        wc(16, "—")

    if status == "Realised":
        wc(17, "Realised",      fnt=S["gtxt"], fll=S["gfill"])
    elif status == "Realised (SL)":
        wc(17, "Realised (SL)", fnt=S["otxt"], fll=S["ofill"])
    elif status == "Open":
        wc(17, "Open",          fnt=S["ytxt"], fll=S["yfill"])
    else:
        wc(17, status)


def autofit(ws):
    for col in ws.columns:
        max_len = max(
            (len(str(cell.value)) for cell in col if cell.value is not None),
            default=8
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 3, 30)


def main():
    global OUT_EXCEL

    if not SRC_DATES.exists():
        print(f"Error: {SRC_DATES.name} not found!"); sys.exit(1)

    params = load_parameters()
    if "TATAMOTORS" not in params and "TMPV" in params:
        params["TATAMOTORS"] = params["TMPV"]
        print("  Mapped TATAMOTORS -> TMPV parameters.")

    wb_src = openpyxl.load_workbook(SRC_DATES)

    all_symbols = set()
    for sn in wb_src.sheetnames:
        ws = wb_src[sn]
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, 2).value
            if v:
                all_symbols.add(str(v).strip())

    print("\nPre-fetching price data (close, high, low) for all symbols...")
    price_cache = {}
    for sym in sorted(all_symbols):
        print(f"  Fetching {sym}...", end="", flush=True)
        df = fetch_yahoo_history(sym)
        if df is not None and len(df) > 0:
            df["ema50_high"] = df["high"].ewm(span=50, adjust=False).mean()
            df["ema50_low"]  = df["low"].ewm(span=50, adjust=False).mean()
            price_cache[sym] = df
            print(" -> OK")
        else:
            print(" -> FAILED")
        time.sleep(0.1)

    S = make_styles()

    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)

    case_sheets  = {}
    case_row_ptr = {}
    for case_num, sheet_name, _ in CASES:
        ws = wb_out.create_sheet(title=sheet_name)
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = True
        write_header_row(ws, CASE_HEADERS, S)
        case_sheets[case_num]  = ws
        case_row_ptr[case_num] = 2

    summary_data = []

    # Quarter separator style
    qtr_fill = PatternFill(start_color="1F6B75", end_color="1F6B75", fill_type="solid")
    qtr_font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFFFF")
    qtr_border = Border(
        left=Side(style="medium", color="145260"),
        right=Side(style="medium", color="145260"),
        top=Side(style="medium", color="145260"),
        bottom=Side(style="medium", color="145260"),
    )
    N_COLS = len(CASE_HEADERS)

    for sheet_name in wb_src.sheetnames:
        ws_src  = wb_src[sheet_name]
        quarter = sheet_name
        print(f"\nProcessing quarter: {quarter}...")

        # ── Insert quarter separator row in every case sheet ──────────────────
        for case_num, _, case_title in CASES:
            ws_c = case_sheets[case_num]
            r    = case_row_ptr[case_num]
            ws_c.row_dimensions[r].height = 20
            # Merge A:Q for a full-width banner
            ws_c.merge_cells(start_row=r, start_column=1,
                              end_row=r, end_column=N_COLS)
            cell = ws_c.cell(row=r, column=1,
                             value=f"  ◆  {quarter}  —  {case_title}")
            cell.font      = qtr_font
            cell.fill      = qtr_fill
            cell.border    = qtr_border
            cell.alignment = Alignment(horizontal="left", vertical="center")
            # Fill remaining cells in merged range with same style
            for c in range(2, N_COLS + 1):
                ws_c.cell(row=r, column=c).fill   = qtr_fill
                ws_c.cell(row=r, column=c).border = qtr_border
            case_row_ptr[case_num] += 1

        # ── Process each stock in this quarter ────────────────────────────────
        for r_src in range(2, ws_src.max_row + 1):
            comp    = ws_src.cell(r_src, 1).value
            sym_val = ws_src.cell(r_src, 2).value
            D       = ws_src.cell(r_src, 3).value

            if not sym_val:
                continue
            sym = str(sym_val).strip()

            if sym not in params:
                print(f"  Skip {sym}: no params")
                continue
            if sym not in price_cache:
                print(f"  Skip {sym}: no price data")
                continue
            if not isinstance(D, dt.datetime):
                continue

            bb, sa, exp = params[sym]
            df          = price_cache[sym]
            latest_dt   = df.index[-1]
            ema_high    = df["ema50_high"]
            ema_low     = df["ema50_low"]

            entry_date = get_trading_date(df.index, D, -int(bb))
            exit_date  = get_trading_date(df.index, D,  int(sa))
            if entry_date is None or exit_date is None:
                continue

            entry_dt = pd.Timestamp(entry_date).normalize()
            exit_dt  = pd.Timestamp(exit_date).normalize()

            data_row = dict(company=comp, symbol=sym, quarter=quarter,
                            result_date=D, buy_before=bb, sell_after=sa, exp=exp)
            sum_row  = dict(**data_row, entry_date=entry_date, exit_date=exit_date)

            for case_num, _, _ in CASES:
                res = apply_case(case_num, df, ema_high, ema_low,
                                 entry_dt, exit_dt, latest_dt)
                sum_row[f"c{case_num}_ret"]    = res["ret"]
                sum_row[f"c{case_num}_status"] = res["status"]

                write_case_row(case_sheets[case_num], case_row_ptr[case_num],
                               data_row, res, entry_date, exit_date, case_num, S)
                case_row_ptr[case_num] += 1

            summary_data.append(sum_row)

        # ── Spacer row after each quarter ─────────────────────────────────────
        spacer_fill = PatternFill(start_color="E8F4F6", end_color="E8F4F6", fill_type="solid")
        for case_num, _, _ in CASES:
            ws_c = case_sheets[case_num]
            r    = case_row_ptr[case_num]
            ws_c.row_dimensions[r].height = 8
            for c in range(1, N_COLS + 1):
                ws_c.cell(row=r, column=c).fill = spacer_fill
            case_row_ptr[case_num] += 1

    print("\nApplying column auto-fit...")
    for case_num, _, _ in CASES:
        autofit(case_sheets[case_num])

    # ── Summary Sheet ──────────────────────────────────────────────────────────
    print("Writing Summary sheet...")
    ws_sum = wb_out.create_sheet(title="Summary", index=0)
    ws_sum.freeze_panes = "A2"

    SUM_HEADERS = [
        "Company", "Symbol", "Quarter", "Result Date",
        "Buy Before", "Sell After", "Expected Return (%)",
        "C1 Return", "C1 Status",
        "C2 Return", "C2 Status",
        "C3 Return", "C3 Status",
        "C4 Return", "C4 Status",
        "C5 Return", "C5 Status",
        "C6 Return", "C6 Status",
        "C7 Return", "C7 Status",
    ]
    write_header_row(ws_sum, SUM_HEADERS, S)

    for r_idx, row in enumerate(summary_data, 2):
        bg = S["soft"] if r_idx % 2 == 0 else PatternFill(fill_type=None)

        def sw(col, val, fnt=None, fll=None, nfmt=None):
            cell = ws_sum.cell(row=r_idx, column=col, value=val)
            cell.font = fnt or S["normal"]; cell.fill = fll or bg
            cell.border = S["border"]; cell.alignment = S["ac"]
            if nfmt: cell.number_format = nfmt

        sw(1, row["company"]); sw(2, row["symbol"]); sw(3, row["quarter"])
        sw(4, row["result_date"], nfmt="yyyy-mm-dd")
        sw(5, row["buy_before"]); sw(6, row["sell_after"])
        sw(7, row["exp"], nfmt="0.00")

        for i, cn in enumerate([1, 2, 3, 4, 5, 6, 7]):
            ret    = row.get(f"c{cn}_ret")
            status = row.get(f"c{cn}_status", "")
            rc = 8 + i * 2; sc = 9 + i * 2
            if ret is not None:
                if "SL" in str(status):
                    sw(rc, ret, fnt=S["otxt"], fll=S["ofill"], nfmt="0.00")
                elif ret >= 0:
                    sw(rc, ret, fnt=S["gtxt"], fll=S["gfill"], nfmt="0.00")
                else:
                    sw(rc, ret, fnt=S["rtxt"], fll=S["rfill"], nfmt="0.00")
            else:
                sw(rc, "—")
            sw(sc, status)

    for col in ws_sum.columns:
        ws_sum.column_dimensions[get_column_letter(col[0].column)].width = 16

    # ── Merge Q1 Actual vs Predicted Sheets ─────────────────────────────────
    act_pred_file = ROOT / "Nifty50_Actual_vs_Predicted_Report.xlsx"
    if act_pred_file.exists():
        print("\nMerging Q1 Actual vs Predicted Report sheets...")
        try:
            from copy import copy
            from openpyxl.chart import ScatterChart, Reference, Series
            wb_ap = openpyxl.load_workbook(act_pred_file)
            
            # Helper to copy sheet
            def copy_sheet(src_ws, dest_ws):
                dest_ws.sheet_view.showGridLines = src_ws.sheet_view.showGridLines
                for r in range(1, src_ws.max_row + 1):
                    if r in src_ws.row_dimensions:
                        dest_ws.row_dimensions[r].height = src_ws.row_dimensions[r].height
                    for c in range(1, src_ws.max_column + 1):
                        src_cell = src_ws.cell(row=r, column=c)
                        dest_cell = dest_ws.cell(row=r, column=c, value=src_cell.value)
                        if src_cell.has_style:
                            dest_cell.font = copy(src_cell.font)
                            dest_cell.fill = copy(src_cell.fill)
                            dest_cell.border = copy(src_cell.border)
                            dest_cell.alignment = copy(src_cell.alignment)
                            dest_cell.number_format = src_cell.number_format
                for merged_range in src_ws.merged_cells.ranges:
                    dest_ws.merge_cells(str(merged_range))
                for col_idx, col_dim in src_ws.column_dimensions.items():
                    dest_ws.column_dimensions[col_idx].width = col_dim.width

            # Copy Details sheet
            if "Details" in wb_ap.sheetnames:
                ws_src_det = wb_ap["Details"]
                ws_dest_det = wb_out.create_sheet(title="Q1_Actual_vs_Pred_Details")
                copy_sheet(ws_src_det, ws_dest_det)
                print("  Merged Q1_Actual_vs_Pred_Details sheet.")
                
                # Enrich with multi-case comparison (Cases 2 to 7)
                try:
                    case_data = {2: {}, 3: {}, 4: {}, 5: {}, 6: {}, 7: {}}
                    for c_num in [2, 3, 4, 5, 6, 7]:
                        ws_c = case_sheets[c_num]
                        for r in range(2, ws_c.max_row+1):
                            qtr = ws_c.cell(r, 3).value
                            if qtr == "FY26-27Q1":
                                sym = ws_c.cell(r, 2).value
                                ret = ws_c.cell(r, 16).value
                                status = ws_c.cell(r, 17).value
                                if sym:
                                    case_data[c_num][str(sym).strip()] = (ret, status)
                                    
                    new_headers = [
                        "Case 2 Return (%)", "Case 2 Status",
                        "Case 3 Return (%)", "Case 3 Status",
                        "Case 4 Return (%)", "Case 4 Status",
                        "Case 5 Return (%)", "Case 5 Status",
                        "Case 6 Return (%)", "Case 6 Status",
                        "Case 7 Return (%)", "Case 7 Status"
                    ]
                    for i, h in enumerate(new_headers):
                        col_idx = 16 + i
                        cell = ws_dest_det.cell(row=1, column=col_idx, value=h)
                        cell.font = S["hdr"]
                        cell.fill = S["navy"]
                        cell.alignment = S["ac"]
                        cell.border = S["border"]
                        
                    for r in range(2, ws_dest_det.max_row+1):
                        sym = ws_dest_det.cell(r, 2).value
                        if not sym:
                            continue
                        sym = str(sym).strip()
                        bg = S["soft"] if r % 2 == 0 else PatternFill(fill_type=None)
                        
                        for i, c_num in enumerate([2, 3, 4, 5, 6, 7]):
                            ret_col = 16 + i * 2
                            sta_col = 17 + i * 2
                            ret_val = "—"
                            sta_val = "Filter Out"
                            if sym in case_data[c_num]:
                                ret, status = case_data[c_num][sym]
                                if ret is not None: ret_val = ret
                                if status is not None: sta_val = status
                                
                            c_ret = ws_dest_det.cell(row=r, column=ret_col, value=ret_val)
                            c_ret.font = S["normal"]; c_ret.fill = bg; c_ret.border = S["border"]
                            c_ret.alignment = S["ar"]
                            if isinstance(ret_val, (int, float)):
                                c_ret.number_format = "0.00"
                                if "SL" in str(sta_val):
                                    c_ret.font = S["otxt"]; c_ret.fill = S["ofill"]
                                elif ret_val >= 0:
                                    c_ret.font = S["gtxt"]; c_ret.fill = S["gfill"]
                                else:
                                    c_ret.font = S["rtxt"]; c_ret.fill = S["rfill"]
                                    
                            c_sta = ws_dest_det.cell(row=r, column=sta_col, value=sta_val)
                            c_sta.font = S["normal"]; c_sta.fill = bg; c_sta.border = S["border"]
                            c_sta.alignment = S["ac"]
                            if sta_val == "Realised":
                                c_sta.font = S["gtxt"]; c_sta.fill = S["gfill"]
                            elif sta_val == "Realised (SL)":
                                c_sta.font = S["otxt"]; c_sta.fill = S["ofill"]
                            elif sta_val == "Open":
                                c_sta.font = S["ytxt"]; c_sta.fill = S["yfill"]
                            elif sta_val == "Filter Out":
                                c_sta.font = S["rtxt"]; c_sta.fill = S["rfill"]
                                
                    # Auto-fit new columns
                    for col in range(16, 28):
                        col_letter = get_column_letter(col)
                        max_len = max(len(str(ws_dest_det.cell(r, col).value or "")) for r in range(1, ws_dest_det.max_row+1))
                        ws_dest_det.column_dimensions[col_letter].width = max(max_len + 3, 12)
                    print("  Enriched Q1_Actual_vs_Pred_Details sheet with Case 2-7 returns.")
                except Exception as ee:
                    print(f"  ! Error enriching details: {ee}")
                
            # Copy Summary sheet
            if "Summary" in wb_ap.sheetnames:
                ws_src_sum = wb_ap["Summary"]
                ws_dest_sum = wb_out.create_sheet(title="Q1_Actual_vs_Pred_Summary")
                copy_sheet(ws_src_sum, ws_dest_sum)
                print("  Merged Q1_Actual_vs_Pred_Summary sheet.")
                
                # Regenerate scatter chart inside Q1_Actual_vs_Pred_Summary referencing Q1_Actual_vs_Pred_Details
                if "Details" in wb_ap.sheetnames:
                    try:
                        last_row = ws_dest_det.max_row
                        sc = ScatterChart()
                        sc.title = "Expected vs Actual Returns Scatter"
                        sc.x_axis.title = "Expected Return %"
                        sc.y_axis.title = "Actual Return %"
                        sc.legend = None
                        sc.width = 16
                        sc.height = 11
                        
                        xvalues = Reference(ws_dest_det, min_col=7, min_row=2, max_row=last_row)
                        yvalues = Reference(ws_dest_det, min_col=8, min_row=2, max_row=last_row)
                        ser = Series(yvalues, xvalues, title_from_data=False)
                        ser.marker.symbol = "circle"
                        ser.marker.size = 5
                        ser.graphicalProperties.line.noFill = True
                        sc.series.append(ser)
                        
                        ws_dest_sum.add_chart(sc, "E4")
                        print("  Regenerated Scatter Chart in Q1_Actual_vs_Pred_Summary.")
                    except Exception as ec:
                        print(f"  ! Could not regenerate scatter chart: {ec}")
            wb_ap.close()
        except Exception as em:
            print(f"  ! Error merging actual vs predicted sheets: {em}")
    else:
        print(f"\nWarning: {act_pred_file.name} not found to merge.")

    # ── Save ───────────────────────────────────────────────────────────────────
    print("\nSaving workbook...")
    saved = False
    stem  = "Nifty50Stocks_MultiCase_Analysis"
    for suffix in ["", "_2", "_3", "_4", "_5", "_6"]:
        try:
            out = ROOT / (stem + suffix + ".xlsx")
            wb_out.save(out)
            OUT_EXCEL = out
            saved = True
            if suffix:
                print(f"  (Workbook locked; saved as {out.name})")
            break
        except PermissionError:
            continue
    if not saved:
        out = ROOT / f"{stem}_{int(time.time())}.xlsx"
        wb_out.save(out)
        OUT_EXCEL = out
        print(f"  (All paths locked; saved as {out.name})")

    print(f"\nDone! Saved -> {OUT_EXCEL.name}")
    print("Sheets: Summary | Case1_NoFilter | Case2_AboveHigh | Case3_WithinBand")
    print("        Case4_BelowWatchCross | Case5_NoWithin | Case6_NoFilter_WithSL | Case7_NoFilter_TrailSL")


if __name__ == "__main__":
    main()
