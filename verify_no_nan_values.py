import openpyxl, pandas as pd, sys

sys.stdout.reconfigure(errors='replace')

master_clean_path  = r'D:/behaviour analysis/Nifty50_12_Quarters_Futures_0.5PCT_Total_Turnover_Master_CLEAN.xlsx'
wb = openpyxl.load_workbook(master_clean_path, data_only=True)

print("==========================================================================================")
print("VERIFYING ZERO 'nan' VALUES IN CLEAN MASTER WORKBOOK")
print("==========================================================================================")

total_nan_count = 0
for sheet in wb.sheetnames:
    ws = wb[sheet]
    sheet_nans = 0
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(row=r, column=c).value or '').strip().lower()
            if val in ['nan', 'none', 'null']:
                sheet_nans += 1
                total_nan_count += 1
    print(f"  Sheet '{sheet}': {sheet_nans} 'nan' values found")

print("------------------------------------------------------------------------------------------")
print(f"TOTAL 'nan' VALUES IN WORKBOOK: {total_nan_count}")
print("==========================================================================================")

# Print sample rows from Sheet 'Q3 2023-24' for metadata columns 21, 22, 23
ws_q = wb['Q3 2023-24']
print("\nSample Rows from Q3 2023-24 (Assigned Slot, Re-entry Type, Quarterly Result Date):")
print(f"Header: {[ws_q.cell(row=7, column=c).value for c in range(21, 24)]}")
for r in range(8, 15):
    row_vals = [ws_q.cell(row=r, column=c).value for c in range(21, 24)]
    print(f"  Row {r:2d}: {row_vals}")
print("==========================================================================================")
