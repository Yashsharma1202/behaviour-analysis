"""
1% ITM OPTIONS ONLY — 12-quarter detailed report.
Each quarter sheet: a QUARTER WINDOW header (first entry -> last exit + duration) on top,
then every Nifty-50 stock's 1% ITM option trade (dir, entry/exit dates, strike, premiums,
return %, P&L, outcome). Plus a Rs 20 Lakh/quarter Performance sheet with the real
trade-level drawdown.

Source: GitHub_Dashboard_12_Quarters_1PCT_ITM_Options_Master.xlsx.
ADDITIVE: writes a new workbook.
"""
import openpyxl, statistics as st, math, re
from datetime import date
from collections import defaultdict
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

BASE = r'D:\behaviour analysis'
SRC = BASE + r'\GitHub_Dashboard_12_Quarters_1PCT_ITM_Options_Master.xlsx'
OUT = BASE + r'\GitHub_Dashboard_1PCT_ITM_Options_Only_Report.xlsx'
CAP = 2_000_000
ORDER = ['FY24_Q3', 'FY24_Q4', 'FY25_Q1', 'FY25_Q2', 'FY25_Q3', 'FY25_Q4', 'FY26_Q1', 'FY26_Q2', 'FY26_Q3', 'FY26_Q4', 'FY27_Q1', 'FY27_Q2']
_MONS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']


def num(v):
    try: return float(v)
    except Exception: return None


def pdate(s):
    s = str(s)
    m = re.search(r'(\d{4})-(\d{2})-(\d{2})', s)
    if m: return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.search(r'(\d{1,2})-([A-Za-z]{3})-(\d{4})', s)
    if m: return date(int(m.group(3)), _MONS.index(m.group(2)) + 1, int(m.group(1)))
    return None


def dshort(d): return d.strftime('%d-%b-%Y') if d else '\u2014'


wb0 = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
DATA = {}; premq = defaultdict(float)
for q in ORDER:
    rows = list(wb0[q].iter_rows(values_only=True)); hdr = rows[3]
    idx = {str(h).strip(): j for j, h in enumerate(hdr) if h}
    def C(n):
        for k, j in idx.items():
            if k.lower().startswith(n.lower()): return j
    cS, cN, cB, cW = C('Symbol'), C('Company'), C('Option Bias'), C('Position Window')
    cED, cXD, cK, cEP, cXP = C('Entry Date'), C('Exit Date'), C('Strike'), C('Entry Premium'), C('Exit Premium')
    cR, cP, cPr, cO = C('Option Return'), C('Option P&L'), C('Premium Paid'), C('Outcome')
    recs = []
    for r in rows[4:]:
        if not r or r[cS] is None: continue
        pnl, prem = num(r[cP]), num(r[cPr])
        if pnl is None or prem is None: continue
        recs.append(dict(sym=str(r[cS]), name=str(r[cN] or ''), dir=str(r[cB] or ''), win=str(r[cW] or ''),
                         edt=pdate(r[cED]), xdt=pdate(r[cXD]), strike=r[cK], ep=r[cEP], xp=r[cXP],
                         ret=r[cR], pnl=pnl, prem=prem, out=str(r[cO] or '')))
        premq[q] += prem
    DATA[q] = recs

# ---- trade-level drawdown (Rs 20L/quarter) ----
alltr = []
for q in ORDER:
    sc = CAP / premq[q] if premq[q] else 0
    for r in DATA[q]:
        alltr.append(dict(q=q, xdt=r['xdt'], spnl=r['pnl'] * sc))
alltr.sort(key=lambda t: (t['xdt'] or date(1900, 1, 1)))
cum = peak = mdd = 0.0
for t in alltr:
    cum += t['spnl']; peak = max(peak, cum); mdd = min(mdd, cum - peak)
byq = defaultdict(float)
for t in alltr: byq[t['q']] += t['spnl']
rets = [byq[q] / CAP * 100 for q in ORDER]
trade_win = 100 * sum(1 for t in alltr if t['spnl'] > 0) / len(alltr)
std = st.pstdev(rets); pos = [r for r in rets if r > 0]; neg = [r for r in rets if r < 0]

# ---- styles ----
NAVY = PatternFill('solid', fgColor='1A202C'); WHITE = Font(bold=True, color='FFFFFF', size=9)
TITLE = Font(bold=True, color='1A202C', size=13); SUB = Font(color='4A5568', size=10, italic=True)
WINH = Font(bold=True, color='2B6CB0', size=11)
GF = PatternFill('solid', fgColor='E6F4EA'); RF = PatternFill('solid', fgColor='FCE8E6')
AM = PatternFill('solid', fgColor='FEF3C7'); LBL = PatternFill('solid', fgColor='F1F5F9')
GOOD = Font(color='137333', bold=True); BAD = Font(color='C5221F', bold=True); BOLD = Font(bold=True)
THIN = Border(*[Side(style='thin', color='D9E2EC')] * 4)
CEN = Alignment(horizontal='center'); RIGHT = Alignment(horizontal='right'); WRAP = Alignment(wrap_text=True, vertical='top')

wb = openpyxl.Workbook()
HEAD = ['#', 'Symbol', 'Company', 'Dir', 'Entry Date', 'Exit Date', 'Window', 'Strike (1% ITM)',
        'Entry Prem', 'Exit Prem', 'Option Ret %', 'Option P&L (1 Lot)', 'Premium Paid', 'Outcome']
WID = [5, 12, 24, 9, 13, 13, 12, 14, 11, 11, 12, 17, 13, 9]
DAY = {5, 6}; PNL = {9, 10, 12, 13}

for q in ORDER:
    recs = DATA[q]
    en = [r['edt'] for r in recs if r['edt']]; ex = [r['xdt'] for r in recs if r['xdt']]
    fe = min(en) if en else None; lx = max(ex) if ex else None
    dur = (lx - fe).days if (fe and lx) else None
    ws = wb.create_sheet(q)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEAD))
    ws.cell(1, 1, '%s  |  1%% ITM OPTIONS  |  Nifty-50 detailed trade sheet (1 lot)' % q).font = TITLE
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEAD))
    ws.cell(2, 1, 'QUARTER WINDOW:  First Entry %s   \u2192   Last Exit (quarter close) %s   |   Duration: %s calendar days'
            % (dshort(fe), dshort(lx), (dur if dur is not None else '\u2014'))).font = WINH
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(HEAD))
    for j, h in enumerate(HEAD, 1):
        c = ws.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
    tot = tprem = 0.0; wcount = 0; row = 5
    for i, r in enumerate(recs, 1):
        win = (r['pnl'] > 0); wcount += win
        tot += r['pnl']; tprem += r['prem']
        vals = [i, r['sym'], r['name'][:22], r['dir'].replace('BUY ', ''), dshort(r['edt']), dshort(r['xdt']),
                r['win'], round(r['strike'], 2) if isinstance(r['strike'], (int, float)) else r['strike'],
                round(r['ep'], 2) if isinstance(r['ep'], (int, float)) else r['ep'],
                round(r['xp'], 2) if isinstance(r['xp'], (int, float)) else r['xp'],
                round(r['ret'], 1) if isinstance(r['ret'], (int, float)) else r['ret'],
                round(r['pnl']), round(r['prem']), 'WIN' if win else 'LOSS']
        fill = GF if win else RF
        for j, v in enumerate(vals, 1):
            c = ws.cell(row, j, v); c.border = THIN; c.fill = fill
            if j in (1, 4, 7, 14) or j in DAY: c.alignment = CEN
            if j in (8,) or j in PNL: c.alignment = RIGHT
            if j == 12: c.font = GOOD if win else BAD
        row += 1
    n = len(recs)
    tr = ['', 'TOTAL', '%d stocks' % n, '', dshort(fe), dshort(lx), '', '', '', '', '', round(tot), round(tprem), '%dW/%dL' % (wcount, n - wcount)]
    for j, v in enumerate(tr, 1):
        c = ws.cell(row, j, v); c.fill = NAVY; c.font = WHITE; c.border = THIN
        if j in PNL or j in DAY: c.alignment = RIGHT if j in PNL else CEN
    ws.cell(3, 1, 'Net option P&L: Rs %s  |  Premium deployed: Rs %s  |  Win %d/%d  |  Return on premium: %s%%'
            % (format(round(tot), ','), format(round(tprem), ','), wcount, n, round(tot / tprem * 100, 1) if tprem else 0)).font = SUB
    for j, w in enumerate(WID, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.freeze_panes = 'A5'

# ---- performance sheet (front) ----
ps = wb.create_sheet('Performance_20L', 0)
ps.merge_cells('A1:D1'); ps.cell(1, 1, '1% ITM OPTIONS ONLY \u2014 Rs 20 Lakh / Quarter Performance (12 quarters)').font = TITLE
ps.merge_cells('A2:D2'); ps.cell(2, 1, 'Fresh Rs 20L each quarter. Return %% = quarter P&L / premium. Drawdown is trade-level (real).').font = SUB
hd = ['Quarter', 'First Entry', 'Last Exit (close)', 'Return %', 'Rs P&L on 20L']
for j, h in enumerate(hd, 1):
    c = ps.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
for i, q in enumerate(ORDER, 5):
    recs = DATA[q]
    en = [r['edt'] for r in recs if r['edt']]; ex = [r['xdt'] for r in recs if r['xdt']]
    r = byq[q] / CAP * 100
    for j, v in enumerate([q, dshort(min(en) if en else None), dshort(max(ex) if ex else None), round(r, 1), round(byq[q])], 1):
        c = ps.cell(i, j, v); c.border = THIN; c.fill = GF if r >= 0 else RF
        if j in (2, 3): c.alignment = CEN
        if j in (4, 5): c.alignment = RIGHT
        if j == 4: c.font = GOOD if r >= 0 else BAD
r0 = 5 + 12 + 1
ps.cell(r0, 1, 'PERFORMANCE MEASURES').font = BOLD
meas = [
    ('Total P&L (12 x Rs 20L)', 'Rs {:+,.0f}'.format(sum(byq.values()))),
    ('Overall return (Rs 2.4 Cr deployed)', '{:+.1f}%'.format(sum(byq.values()) / (CAP * 12) * 100)),
    ('Quarter win rate', '{:.0f}%'.format(100 * sum(1 for r in rets if r > 0) / 12)),
    ('Trade-level win rate', '{:.1f}%'.format(trade_win)),
    ('Avg / Median return per quarter', '{:+.1f}% / {:+.1f}%'.format(st.mean(rets), st.median(rets))),
    ('Best / Worst quarter', '{:+.1f}% / {:+.1f}%'.format(max(rets), min(rets))),
    ('Sharpe (annualised)', '{:.2f}'.format(st.mean(rets) / std * math.sqrt(4) if std else 0)),
    ('Profit factor (quarterly)', ('{:.2f}'.format(sum(pos) / abs(sum(neg))) if neg else 'n/a')),
    ('>> MAX DRAWDOWN (real, trade-level)', 'Rs {:,.0f}  ({:.1f}% of a 20L book)'.format(mdd, mdd / CAP * 100)),
]
for k, (lbl, val) in enumerate(meas):
    rr = r0 + 1 + k
    a = ps.cell(rr, 1, lbl); a.fill = (AM if lbl.startswith('>>') else LBL); a.font = BOLD; a.border = THIN
    b = ps.cell(rr, 2, val); b.border = THIN; ps.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=3)
note = ps.cell(r0 + 1 + len(meas) + 1, 1, 'NOTE: 1% synthetic option model + 2023-26 bull sample where every quarter won = optimistic ceiling. The -{:.0f}% drawdown and {:.0f}% trade win rate are the honest risk.'.format(abs(mdd / CAP * 100), trade_win))
note.font = Font(color='7C2D12', size=9, italic=True); note.alignment = WRAP
ps.merge_cells(start_row=r0 + 1 + len(meas) + 1, start_column=1, end_row=r0 + 3 + len(meas), end_column=4)
for j, w in enumerate([36, 16, 18, 16], 1):
    ps.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w

if 'Sheet' in wb.sheetnames:
    wb.remove(wb['Sheet'])
wb.save(OUT)
print('Saved:', OUT)
print('ITM-only report: 12 quarter sheets + performance. Rs20L total Rs {:+,.0f}, trade-win {:.0f}%, maxDD {:.1f}%'
      .format(sum(byq.values()), trade_win, mdd / CAP * 100))
