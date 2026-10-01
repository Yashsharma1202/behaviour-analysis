"""
DUSSEHRA DETAILED TRADE SHEET (Nifty-50 + BSE, 51 stocks) \u2014 Excel.
Every real walk-forward out-of-sample trade (fix_dussehra_position_window_walkforward.py),
reconstructed with its actual historical entry/exit DATE and PRICE (read
directly off the real trading-day price array, not a synthetic calendar),
margin, rupee PnL, and running cumulative PnL \u2014 plus per-stock CAGR, Max
Drawdown, Sharpe-like ratio and a portfolio-level cumulative PnL curve.

ADDITIVE: reads the walk-forward pickle + price data; writes a new report.
"""
import json, pickle, os, subprocess, statistics as st
from datetime import date
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

BASE = r'D:\behaviour analysis'
import os as _os
OUT_XLSX = _os.environ.get('DUSSTS_OUT_XLSX', BASE + r'\Dussehra_Detailed_Trade_Sheet.xlsx')

results = pickle.load(open(BASE + r'\scraped_parquet\_tmp_dussehra_walkforward.pkl', 'rb'))
fo = json.load(open(BASE + r'\dashboard_data\fo_stocks_211.json', encoding='utf-8'))
meta = {f['symbol']: f for f in fo}
d = json.load(open(BASE + r'\dashboard_data\event_dashboard_data.json', encoding='utf-8'))
dus = next(h for h in d['holidays'] if h['id'] == 'dussehra_2026')
NIFTY_SET = set(s['symbol'] for s in dus['stocks'])  # the live 51 (50 Nifty-50 + BSE)
live_meta = {s['symbol']: s for s in dus['stocks']}

# DATA QUALITY FIX: the cash-equity file is missing 98.3% of all Fridays
# across its entire 26yr history (1,371 of 1,395, 2000-2026) -- not a few
# gaps, it is effectively a 4-day trading week for its whole life, silently
# corrupting every trading-day countback and every return compounded on it.
# Using the futures file for ALL stocks instead: shorter history (2019-2026)
# but a complete, trustworthy trading calendar (~6.5% Friday gaps, consistent
# with genuine NSE holidays landing on a Friday).
fut = pd.read_parquet(BASE + r'\scraped_parquet\fo_futures_near_month_continuous.parquet', columns=['Date', 'Instrument', 'Close'])
fut['Date'] = pd.to_datetime(fut['Date']).dt.date
fut_closes = {s: g.sort_values('Date')[['Date', 'Close']].values.tolist() for s, g in fut.groupby('Instrument')}

beh = json.load(open(BASE + r'\dashboard_data\nifty_futures_holiday_behaviour.json', encoding='utf-8'))
dh = [h for h in beh['holidays'] if 'Dussehra' in h['name']][0]
dussehra_dates = sorted(date.fromisoformat(t['holiday_date']) for t in dh['trades'])


def closes_for(sym):
    if sym in fut_closes:
        return fut_closes[sym], [dt for dt in dussehra_dates if dt.year >= 2019]
    return None, []


def lot_for(sym):
    m = meta.get(sym, {})
    if m.get('lot_size'):
        return m['lot_size']
    lm = live_meta.get(sym, {})
    return lm.get('lot') or 500


# ---- Reconstruct every trade with real date + price + margin + rupee PnL ----
all_trades = []   # one row per (symbol, year) trade
for sym in NIFTY_SET:
    r = results.get(sym)
    if not r or r['oos_n'] < 1:
        continue
    a, dates_to_use = closes_for(sym)
    if a is None:
        continue
    lot = lot_for(sym)
    date_to_idx = {}
    ptr = -1
    for hd in dates_to_use:
        while ptr + 1 < len(a) and a[ptr + 1][0] <= hd:
            ptr += 1
        date_to_idx[hd] = ptr

    cum_pnl = 0.0
    for t in sorted(r['oos_trades'], key=lambda x: x['year']):
        hd = next((dt for dt in dates_to_use if dt.year == t['year']), None)
        if hd is None:
            continue
        ib = date_to_idx.get(hd)
        if ib is None or ib < 0:
            continue
        en_idx, ex_idx = ib - t['bo'], ib + t['so']
        if en_idx < 0 or ex_idx >= len(a):
            continue
        en_date, en_px = a[en_idx]
        ex_date, ex_px = a[ex_idx]
        margin = round(en_px * lot * 0.20, 2)
        pnl = round((t['ret'] / 100) * en_px * lot, 2)
        cum_pnl += pnl
        all_trades.append(dict(symbol=sym, year=t['year'], direction=t['side'],
                                window=f"T-{t['bo']} to T+{t['so']}",
                                entry_date=en_date, exit_date=ex_date,
                                entry_px=round(en_px, 2), exit_px=round(ex_px, 2),
                                lot=lot, margin=margin, ret_pct=t['ret'], pnl=pnl,
                                cum_pnl=round(cum_pnl, 2)))

print(f"Reconstructed {len(all_trades)} real walk-forward trades across {len(NIFTY_SET)} stocks.")

# ---- Per-stock performance measures (CAGR, Max DD, Sharpe) on the yearly return series ----
perf = {}
for sym in NIFTY_SET:
    trades = sorted([t for t in all_trades if t['symbol'] == sym], key=lambda x: x['year'])
    if not trades:
        continue
    rets = [t['ret_pct'] for t in trades]
    n = len(rets)
    cum = 1.0
    for v in rets: cum *= (1 + v / 100)
    cagr = (cum ** (1 / n) - 1) * 100 if cum > 0 else None
    eq = [100.0]
    for v in rets: eq.append(eq[-1] * (1 + v / 100))
    peak = eq[0]; maxdd = 0.0
    for v in eq[1:]:
        peak = max(peak, v)
        dd = (v - peak) / peak * 100
        maxdd = min(maxdd, dd)
    sharpe = (st.mean(rets) / st.pstdev(rets)) if n > 1 and st.pstdev(rets) > 0 else None
    wins = sum(1 for v in rets if v > 0)
    wr = round(100 * wins / n, 1)
    total_pnl = round(sum(t['pnl'] for t in trades), 2)
    perf[sym] = dict(n=n, cagr=cagr, maxdd=maxdd, sharpe=sharpe, wr=wr,
                      avg_ret=round(sum(rets) / n, 2), total_pnl=total_pnl,
                      first_year=trades[0]['year'], last_year=trades[-1]['year'])

# ---- Portfolio-level cumulative PnL (sum of all stocks' PnL per year, 1 lot each) ----
by_year = {}
for t in all_trades:
    by_year.setdefault(t['year'], []).append(t['pnl'])
years_sorted = sorted(by_year)
port_cum = 0.0
portfolio_curve = []
for y in years_sorted:
    year_pnl = round(sum(by_year[y]), 2)
    port_cum += year_pnl
    portfolio_curve.append((y, len(by_year[y]), year_pnl, round(port_cum, 2)))

NAVY = PatternFill('solid', fgColor='1A202C'); WHITE = Font(bold=True, color='FFFFFF', size=9)
TITLE = Font(bold=True, color='1A202C', size=13); SUB = Font(color='4A5568', size=10, italic=True)
GF = PatternFill('solid', fgColor='E6F4EA'); RF = PatternFill('solid', fgColor='FCE8E6')
AM = PatternFill('solid', fgColor='FEF3C7'); LBL = PatternFill('solid', fgColor='F1F5F9')
GOOD = Font(color='137333', bold=True); BAD = Font(color='C5221F', bold=True); BOLD = Font(bold=True)
THIN = Border(*[Side(style='thin', color='D9E2EC')] * 4)
CEN = Alignment(horizontal='center'); RIGHT = Alignment(horizontal='right'); WRAP = Alignment(wrap_text=True, vertical='top')

wb = Workbook()

# ============ Sheet 1: All Trades Log (every year, every stock) ============
tl = wb.active; tl.title = 'All_Trades_Log'
tl.merge_cells('A1:M1')
tl.cell(1, 1, 'DUSSEHRA \u2014 EVERY REAL WALK-FORWARD TRADE (Nifty-50 + BSE, 1 lot each)').font = TITLE
tl.merge_cells('A2:M2')
tl.cell(2, 1, f'{len(all_trades)} genuine out-of-sample trades. Margin = 20% x entry price x lot. PnL = return% x entry price x lot. Cumulative PnL runs per stock, year ascending.').font = SUB
hdr = ['Symbol', 'Year', 'Direction', 'Window', 'Entry Date', 'Entry Price', 'Exit Date', 'Exit Price', 'Lot', 'Margin (\u20b9)', 'Return %', 'PnL (\u20b9)', 'Cumulative PnL (\u20b9)']
for j, h in enumerate(hdr, 1):
    c = tl.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
row = 5
for sym in sorted(NIFTY_SET):
    trades = sorted([t for t in all_trades if t['symbol'] == sym], key=lambda x: x['year'])
    for t in trades:
        vals = [t['symbol'], t['year'], t['direction'], t['window'], t['entry_date'].strftime('%d-%b-%Y'),
                t['entry_px'], t['exit_date'].strftime('%d-%b-%Y'), t['exit_px'], t['lot'], t['margin'],
                t['ret_pct'], t['pnl'], t['cum_pnl']]
        fill = GF if t['pnl'] >= 0 else RF
        for j, v in enumerate(vals, 1):
            c = tl.cell(row, j, v); c.border = THIN; c.fill = fill
            if j in (2, 3, 9): c.alignment = CEN
            if j in (6, 8, 10, 11, 12, 13): c.alignment = RIGHT
            if j in (11, 12, 13) and isinstance(v, (int, float)): c.font = GOOD if v >= 0 else BAD
        row += 1
for j, w in enumerate([12, 7, 10, 14, 13, 11, 13, 11, 7, 12, 10, 12, 15], 1):
    tl.column_dimensions[get_column_letter(j)].width = w
tl.freeze_panes = 'A5'

# ============ Sheet 2: Last Available Year's Trades ============
ly = wb.create_sheet('Last_Year_Trades')
ly.merge_cells('A1:K1')
ly.cell(1, 1, "MOST RECENT TRADE PER STOCK (each stock's latest available Dussehra, usually 2025)").font = TITLE
ly.merge_cells('A2:K2')
ly.cell(2, 1, 'Not all stocks have a 2025 trade (some have shorter history) \u2014 each row is that stock\u2019s own most recent walk-forward result.').font = SUB
hdr = ['Symbol', 'Year', 'Direction', 'Window', 'Entry Date', 'Entry Price', 'Exit Date', 'Exit Price', 'Margin (\u20b9)', 'Return %', 'PnL (\u20b9)']
for j, h in enumerate(hdr, 1):
    c = ly.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
row = 5
last_rows = []
for sym in NIFTY_SET:
    trades = sorted([t for t in all_trades if t['symbol'] == sym], key=lambda x: x['year'])
    if trades:
        last_rows.append(trades[-1])
last_rows.sort(key=lambda t: (-t['year'], t['symbol']))
for t in last_rows:
    vals = [t['symbol'], t['year'], t['direction'], t['window'], t['entry_date'].strftime('%d-%b-%Y'),
            t['entry_px'], t['exit_date'].strftime('%d-%b-%Y'), t['exit_px'], t['margin'], t['ret_pct'], t['pnl']]
    fill = GF if t['pnl'] >= 0 else RF
    for j, v in enumerate(vals, 1):
        c = ly.cell(row, j, v); c.border = THIN; c.fill = fill
        if j in (2, 3): c.alignment = CEN
        if j in (6, 8, 9, 10, 11): c.alignment = RIGHT
        if j in (10, 11) and isinstance(v, (int, float)): c.font = GOOD if v >= 0 else BAD
    row += 1
for j, w in enumerate([12, 7, 10, 14, 13, 11, 13, 11, 12, 10, 12], 1):
    ly.column_dimensions[get_column_letter(j)].width = w
ly.freeze_panes = 'A5'

# ============ Sheet 3: Performance Measures ============
pm = wb.create_sheet('Performance_Measures')
pm.merge_cells('A1:K1')
pm.cell(1, 1, 'PER-STOCK PERFORMANCE MEASURES (walk-forward, genuine OOS trades)').font = TITLE
pm.merge_cells('A2:K2')
pm.cell(2, 1, 'CAGR and Max DD computed on each stock\u2019s own yearly trade equity curve (one trade/year, starting at 100). Sharpe-like = mean / population-stdev of yearly returns.').font = SUB
hdr = ['Symbol', 'Company', 'Trades (n)', 'Years Covered', 'Win Rate %', 'Avg Return %', 'CAGR %', 'Max DD %', 'Sharpe-like', 'Total PnL (\u20b9, 1 lot)']
for j, h in enumerate(hdr, 1):
    c = pm.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
row = 5
perf_ranked = sorted(perf.items(), key=lambda kv: -(kv[1]['cagr'] if kv[1]['cagr'] is not None else -999))
for sym, p in perf_ranked:
    m = meta.get(sym, {})
    vals = [sym, m.get('name', sym), p['n'], f"{p['first_year']}-{p['last_year']}", p['wr'], p['avg_ret'],
            (round(p['cagr'], 2) if p['cagr'] is not None else '\u2014'), round(p['maxdd'], 2),
            (round(p['sharpe'], 2) if p['sharpe'] is not None else '\u2014'), p['total_pnl']]
    fill = AM if p['n'] < 3 else (GF if p['total_pnl'] >= 0 else RF)
    for j, v in enumerate(vals, 1):
        c = pm.cell(row, j, v); c.border = THIN; c.fill = fill
        if j in (3, 5): c.alignment = CEN
        if j in (6, 7, 8, 9, 10): c.alignment = RIGHT
        if j == 7 and isinstance(v, (int, float)): c.font = GOOD if v >= 0 else BAD
        if j == 8: c.font = BAD if p['maxdd'] < 0 else GOOD
        if j == 10: c.font = GOOD if p['total_pnl'] >= 0 else BAD
    row += 1
for j, w in enumerate([12, 26, 10, 14, 11, 12, 10, 10, 11, 16], 1):
    pm.column_dimensions[get_column_letter(j)].width = w
pm.freeze_panes = 'A5'

# ============ Sheet 4: Portfolio Cumulative PnL ============
pc = wb.create_sheet('Portfolio_Cumulative_PnL')
pc.merge_cells('A1:D1')
pc.cell(1, 1, 'PORTFOLIO VIEW \u2014 all stocks\u2019 Dussehra trades combined, 1 lot each, year by year').font = TITLE
hdr = ['Year', 'Stocks Traded', "Year's Total PnL (\u20b9)", 'Cumulative PnL (\u20b9)']
for j, h in enumerate(hdr, 1):
    c = pc.cell(3, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
row = 4
for y, n_stocks, year_pnl, cum in portfolio_curve:
    vals = [y, n_stocks, year_pnl, cum]
    fill = GF if year_pnl >= 0 else RF
    for j, v in enumerate(vals, 1):
        c = pc.cell(row, j, v); c.border = THIN; c.fill = fill
        if j in (1, 2): c.alignment = CEN
        if j in (3, 4): c.alignment = RIGHT; c.font = GOOD if v >= 0 else BAD
    row += 1
for j, w in enumerate([8, 14, 20, 20], 1):
    pc.column_dimensions[get_column_letter(j)].width = w
pc.freeze_panes = 'A4'

# ============ Sheet 5: Overview ============
ov = wb.create_sheet('Overview', 0)
ov.merge_cells('A1:B1')
ov.cell(1, 1, 'DUSSEHRA DETAILED TRADE SHEET \u2014 OVERVIEW').font = TITLE
n_stocks_with_trades = len(perf)
total_portfolio_pnl = portfolio_curve[-1][3] if portfolio_curve else 0
best = max(perf.items(), key=lambda kv: kv[1]['total_pnl'])
worst = min(perf.items(), key=lambda kv: kv[1]['total_pnl'])
notes = [
    ('Universe', f'{n_stocks_with_trades} of {len(NIFTY_SET)} Nifty-50 + BSE stocks have at least 1 genuine walk-forward trade (BSE and JIOFIN excluded \u2014 insufficient history).'),
    ('Trade source', 'Every trade is a REAL out-of-sample (OOS) result from the walk-forward 8x8 grid search \u2014 window and direction for each year were chosen using ONLY years strictly before it, never future data. See All_Trades_Log for every single trade.'),
    ('Margin', 'Estimated as 20% of (entry price \u00d7 lot size) \u2014 a standard SPAN-margin proxy, not the exact historical broker margin on that date.'),
    ('PnL calculation', 'PnL (\u20b9) = return% \u00d7 entry price \u00d7 lot size, for 1 lot. Scale linearly for more lots/capital.'),
    ('CAGR & Max DD', "Computed on each stock's own yearly trade equity curve (starting at 100, compounding one trade per year) \u2014 not a calendar-year price CAGR."),
    ('Last_Year_Trades sheet', 'Shows each stock\u2019s own most recent available trade (mostly 2025, some earlier if a stock\u2019s history is shorter).'),
    ('Portfolio view', f'If 1 lot of every stock\u2019s Dussehra trade had been taken every year since each stock\u2019s own first qualifying year, cumulative PnL across the whole period is \u20b9{total_portfolio_pnl:,.0f} (see Portfolio_Cumulative_PnL). Best single stock: {best[0]} (\u20b9{best[1]["total_pnl"]:,.0f} total). Worst: {worst[0]} (\u20b9{worst[1]["total_pnl"]:,.0f} total).'),
    ('CAVEAT', 'This is a backtest of 1 lot per stock, no compounding of capital across stocks, no position sizing, and no transaction costs/slippage. Real deployment needs a capital allocation and cost model on top of these numbers.'),
]
r0 = 3
for lbl, txt in notes:
    a = ov.cell(r0, 1, lbl); a.font = BOLD; a.fill = AM; a.border = THIN; a.alignment = WRAP
    b = ov.cell(r0, 2, txt); b.alignment = WRAP; b.border = THIN
    ov.row_dimensions[r0].height = 70; r0 += 1
ov.column_dimensions['A'].width = 22; ov.column_dimensions['B'].width = 105

wb.save(OUT_XLSX)
print('Saved Excel:', OUT_XLSX)
print(f'Stocks with trades: {n_stocks_with_trades} | Total trades logged: {len(all_trades)} | Portfolio cumulative PnL: Rs {total_portfolio_pnl:,.0f}')
