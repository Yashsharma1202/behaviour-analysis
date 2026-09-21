import datetime as dt
import json
import sys
import time
import urllib.request
import warnings
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import numpy as np

warnings.filterwarnings("ignore")
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

ROOT       = Path(__file__).resolve().parent
SRC_PARAMS = ROOT / "UPDATED.xlsx"
SRC_DATES  = ROOT / "Nifty50Stocks_QtyResultDates.xlsx"
OUT_EXCEL  = ROOT / "Nifty50Stocks_QtyResultDates_Filtered_EMA.xlsx"
SYMBOL_MAP = {"TATAMOTORS": "TMPV"}


# ─────────────────────────── helpers ────────────────────────────────────────

def fetch_yahoo_history(sym: str):
    query_sym = SYMBOL_MAP.get(sym, sym)
    t  = query_sym.replace("&", "%26") + ".NS"
    p1 = int(dt.datetime(2024, 6, 1).timestamp())
    p2 = int((dt.datetime.now() + pd.Timedelta(days=2)).timestamp())
    url = (f"https://query1.finance.yahoo.com/v8/finance/chart/{t}"
           f"?period1={p1}&period2={p2}&interval=1d")
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=25) as r:
            j = json.loads(r.read())
        res = j["chart"]["result"][0]
        idx = pd.to_datetime(res["timestamp"], unit="s").normalize()
        q   = res["indicators"]["quote"][0]
        df  = pd.DataFrame(
            {"close": q["close"], "low": q["low"], "high": q["high"]}, index=idx
        ).dropna()
        return df
    except Exception as e:
        print(f"  ! fetch {sym} ({query_sym}): {e}")
        return None


def get_trading_date(cal: pd.DatetimeIndex, result_date, offset: int):
    if cal is None or len(cal) == 0:
        return None
    D = pd.Timestamp(result_date).normalize()
    if D <= cal[-1]:
        pos        = cal.searchsorted(D, side="right") - 1
        if pos < 0: pos = 0
        target_idx = pos + offset
        if 0 <= target_idx < len(cal):
            return cal[target_idx].to_pydatetime()
        elif target_idx < 0:
            return result_date + dt.timedelta(days=offset)
        else:
            rem    = target_idx - (len(cal) - 1)
            future = pd.bdate_range(start=cal[-1], periods=rem + 1)
            return future[-1].to_pydatetime()
    else:
        b_days     = pd.bdate_range(start=cal[-1], end=D)
        n          = len(b_days) - 1
        target_idx = n + offset
        if target_idx >= 0:
            future = pd.bdate_range(start=cal[-1], periods=target_idx + 1)
            return future[-1].to_pydatetime()
        else:
            back = len(cal) - 1 + target_idx
            if back >= 0:
                return cal[back].to_pydatetime()
            return D + dt.timedelta(days=offset)


def get_price_on_date(s: pd.Series, d):
    ts = pd.Timestamp(d).normalize()
    if ts in s.index:
        v = s.loc[ts]
        return float(v) if pd.notna(v) else None
    prior = s.index[s.index <= ts]
    if len(prior) > 0:
        v = s.loc[prior[-1]]
        return float(v) if pd.notna(v) else None
    return None


def check_sl_hit_trailing(df: pd.DataFrame, buy_ts: pd.Timestamp, until_ts: pd.Timestamp):
    """
    Scan daily lows from day after buy_ts up to until_ts.
    Exit at exactly that day's ema50_low when daily low < ema50_low.
    """
    mask   = (df.index > buy_ts) & (df.index <= until_ts)
    period = df[mask]
    for date, row in period.iterrows():
        daily_ema_low = row.get("ema50_low")
        if pd.notna(row["low"]) and pd.notna(daily_ema_low) and row["low"] < daily_ema_low:
            return True, date.to_pydatetime(), float(daily_ema_low)
    return False, None, None


def apply_strategy(df: pd.DataFrame, entry_dt: pd.Timestamp, exit_dt: pd.Timestamp,
                   latest_dt: pd.Timestamp) -> dict:
    """
    Returns result dict with strategy calculation keys:
        status, actual_buy_date, buy_price, sl_price,
        actual_exit_date, exit_price, sl_hit, ret
    """
    res = dict(
        status="Filter Out",
        actual_buy_date=None, buy_price=None, sl_price=None,
        actual_exit_date=None, exit_price=None, sl_hit=False, ret=None
    )

    if entry_dt > latest_dt:
        res["status"] = "Upcoming"
        return res

    close_e = get_price_on_date(df["close"], entry_dt)
    ema_l_e = get_price_on_date(df["ema50_low"], entry_dt)

    # Entry Rule: Buy only if close >= ema50_low on entry day. Else Drop.
    if (close_e is None or ema_l_e is None or close_e < ema_l_e):
        return res # status = Filter Out

    res["actual_buy_date"] = entry_dt.to_pydatetime()
    res["buy_price"]       = round(close_e, 2)
    res["sl_price"]        = round(ema_l_e, 2) # starting SL on entry day

    # Check exit
    if exit_dt <= latest_dt:
        # Planned exit has passed
        sl_hit, sl_date, sl_ep = check_sl_hit_trailing(df, entry_dt, exit_dt)
        if sl_hit:
            res["sl_hit"]           = True
            res["actual_exit_date"] = sl_date
            res["exit_price"]       = round(sl_ep, 2)
            res["ret"]              = round((sl_ep / close_e - 1) * 100, 2)
            res["status"]           = "Realised (SL)"
        else:
            ep = get_price_on_date(df["close"], exit_dt)
            res["actual_exit_date"] = exit_dt.to_pydatetime()
            res["exit_price"]       = round(ep, 2) if ep is not None else None
            if ep is not None:
                res["ret"]          = round((ep / close_e - 1) * 100, 2)
            res["status"]           = "Realised"
    else:
        # Exit in future -> Open / Till-Date
        sl_hit, sl_date, sl_ep = check_sl_hit_trailing(df, entry_dt, latest_dt)
        if sl_hit:
            res["sl_hit"]           = True
            res["actual_exit_date"] = sl_date
            res["exit_price"]       = round(sl_ep, 2)
            res["ret"]              = round((sl_ep / close_e - 1) * 100, 2)
            res["status"]           = "Realised (SL)"
        else:
            latest_close            = float(df["close"].iloc[-1])
            res["actual_exit_date"] = df.index[-1].to_pydatetime()
            res["exit_price"]       = round(latest_close, 2)
            res["ret"]              = round((latest_close / close_e - 1) * 100, 2)
            res["status"]           = "Open"

    return res


def load_parameters() -> dict:
    print("Loading model parameters from UPDATED.xlsx...")
    params = {}
    try:
        wb_p = openpyxl.load_workbook(SRC_PARAMS, data_only=True)
        ws_p = wb_p.active
        for r in range(3, ws_p.max_row + 1):
            sym = ws_p.cell(r, 2).value
            if sym:
                sym = str(sym).strip()
                bb  = ws_p.cell(r, 5).value
                sa  = ws_p.cell(r, 6).value
                exp = ws_p.cell(r, 9).value
                if bb is not None and sa is not None:
                    try:
                        params[sym] = (int(bb), int(sa),
                                       round(float(exp), 2) if exp is not None else 0.0)
                    except Exception:
                        pass
        wb_p.close()
    except Exception as e:
        print(f"  ! Error loading UPDATED.xlsx: {e}")

    fallback_csv = ROOT / "processed" / "optimal_rules.csv"
    if fallback_csv.exists():
        print("Loading fallback parameters from processed/optimal_rules.csv...")
        try:
            df_rules = pd.read_csv(fallback_csv)
            df_res   = df_rules[
                (df_rules["event_type"] == "RESULTS") &
                (df_rules["mode"] == "descriptive")
            ]
            for _, row in df_res.iterrows():
                s = str(row["symbol"]).strip()
                if s not in params:
                    buy_val = int(row["buy"])
                    bb_val  = -buy_val
                    sa_val  = int(row["sell"])
                    exp_val = round(float(row["train_avg"]), 2)
                    params[s] = (bb_val, sa_val, exp_val)
                    print(f"  Added fallback for {s}: BB={bb_val}, SA={sa_val}, Exp={exp_val}%")
        except Exception as e:
            print(f"  ! Error loading optimal_rules.csv fallbacks: {e}")

    print(f"Loaded parameters for {len(params)} symbols.")
    return params


def make_styles():
    return dict(
        hdr    = Font(name="Segoe UI", size=9, bold=True, color="FFFFFFFF"),
        normal = Font(name="Segoe UI", size=9),
        ital   = Font(name="Segoe UI", size=9, italic=True, color="FF555555"),
        navy   = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid"),
        soft   = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid"),
        gfill  = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        gtxt   = Font(name="Segoe UI", size=9, color="FF006100"),
        rfill  = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        rtxt   = Font(name="Segoe UI", size=9, color="FF9C0006"),
        yfill  = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
        ytxt   = Font(name="Segoe UI", size=9, color="FF9C6500"),
        ofill  = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
        otxt   = Font(name="Segoe UI", size=9, color="FF833C00"),
        border = Border(
            left   = Side(style="thin", color="D9D9D9"),
            right  = Side(style="thin", color="D9D9D9"),
            top    = Side(style="thin", color="D9D9D9"),
            bottom = Side(style="thin", color="D9D9D9"),
        ),
        ac = Alignment(horizontal="center", vertical="center"),
        ar = Alignment(horizontal="right",  vertical="center"),
        al = Alignment(horizontal="left",   vertical="center"),
    )


HEADERS = [
    "Company Name", "SYMBOL", "Quarter", "Result Date",
    "Buy Before (Candles)", "Sell After (Candles)", "Expected Return (%)",
    "Planned Entry Date", "Actual Buy Date", "Buy Price",
    "SL Price (EMA Low)", "Planned Exit Date", "Actual Exit Date",
    "Exit Price", "SL Hit", "Realised Return (%)", "Status"
]


def write_header_row(ws, headers, S):
    ws.row_dimensions[1].height = 22
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = S["hdr"]; cell.fill = S["navy"]
        cell.alignment = S["ac"]; cell.border = S["border"]


def write_row(ws, r, data: dict, res: dict, entry_date, exit_date, S: dict):
    bg = S["soft"] if r % 2 == 0 else PatternFill(fill_type=None)

    def wc(col, val, fnt=None, fll=None, nfmt=None, aln=None):
        cell = ws.cell(row=r, column=col, value=val)
        cell.font      = fnt or S["normal"]
        cell.fill      = fll or bg
        cell.border    = S["border"]
        cell.alignment = aln or S["ac"]
        if nfmt:
            cell.number_format = nfmt
        return cell

    wc(1, data["company"],      aln=S["al"])
    wc(2, data["symbol"])
    wc(3, data["quarter"])
    wc(4, data["result_date"],  nfmt="yyyy-mm-dd")
    wc(5, data["buy_before"])
    wc(6, data["sell_after"])
    wc(7, data["exp"],          nfmt="0.00")
    wc(8, entry_date,           nfmt="yyyy-mm-dd")

    status = res["status"]

    if status in ("Upcoming", "Filter Out"):
        for col in range(9, len(HEADERS) + 1):
            wc(col, "—")
        if status == "Upcoming":
            wc(17, "Upcoming",   fnt=S["ital"])
        else:
            wc(17, "Filter Out", fnt=S["rtxt"], fll=S["rfill"])
        return

    wc(9,  res["actual_buy_date"],  nfmt="yyyy-mm-dd")
    wc(10, res["buy_price"],        nfmt="#,##0.00", aln=S["ar"])

    sl_val = res["sl_price"] if res["sl_price"] is not None else "N/A"
    wc(11, sl_val, nfmt="#,##0.00" if isinstance(sl_val, (int, float)) else None,
       aln=S["ar"])

    wc(12, exit_date,               nfmt="yyyy-mm-dd")
    wc(13, res["actual_exit_date"], nfmt="yyyy-mm-dd")
    wc(14, res["exit_price"],       nfmt="#,##0.00", aln=S["ar"])

    if res["sl_hit"]:
        wc(15, "YES", fnt=S["rtxt"], fll=S["rfill"])
    else:
        wc(15, "No",  fnt=S["gtxt"], fll=S["gfill"])

    ret = res["ret"]
    if ret is not None:
        if res["sl_hit"]:
            wc(16, ret, nfmt="0.00", aln=S["ar"], fnt=S["otxt"], fll=S["ofill"])
        elif ret >= 0:
            wc(16, ret, nfmt="0.00", aln=S["ar"], fnt=S["gtxt"], fll=S["gfill"])
        else:
            wc(16, ret, nfmt="0.00", aln=S["ar"], fnt=S["rtxt"], fll=S["rfill"])
    else:
        wc(16, "—")

    if status == "Realised":
        wc(17, "Realised",      fnt=S["gtxt"], fll=S["gfill"])
    elif status == "Realised (SL)":
        wc(17, "Realised (SL)", fnt=S["otxt"], fll=S["ofill"])
    elif status == "Open":
        wc(17, "Open",          fnt=S["ytxt"], fll=S["yfill"])
    else:
        wc(17, status)


def autofit(ws):
    for col in ws.columns:
        max_len = max(
            (len(str(cell.value)) for cell in col if cell.value is not None),
            default=8
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 3, 30)


def main():
    global OUT_EXCEL

    if not SRC_DATES.exists():
        print(f"Error: {SRC_DATES.name} not found!"); sys.exit(1)

    params = load_parameters()
    if "TATAMOTORS" not in params and "TMPV" in params:
        params["TATAMOTORS"] = params["TMPV"]

    wb_src = openpyxl.load_workbook(SRC_DATES)

    all_symbols = set()
    for sn in wb_src.sheetnames:
        ws = wb_src[sn]
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, 2).value
            if v:
                all_symbols.add(str(v).strip())

    print("\nPre-fetching price data (close, high, low) for all symbols...")
    price_cache = {}
    for sym in sorted(all_symbols):
        print(f"  Fetching {sym}...", end="", flush=True)
        df = fetch_yahoo_history(sym)
        if df is not None and len(df) > 0:
            df["ema50_low"]  = df["low"].ewm(span=50, adjust=False).mean()
            price_cache[sym] = df
            print(" -> OK")
        else:
            print(" -> FAILED")
        time.sleep(0.1)

    S = make_styles()

    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)

    # 1. Create Sheets for each Quarter
    quarter_sheets = {}
    quarter_row_ptr = {}
    
    # Process from newest to oldest sheets
    qtr_names = wb_src.sheetnames
    
    for qtr in qtr_names:
        ws = wb_out.create_sheet(title=qtr)
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = True
        write_header_row(ws, HEADERS, S)
        quarter_sheets[qtr]  = ws
        quarter_row_ptr[qtr] = 2

    # Quarter separator style
    qtr_fill = PatternFill(start_color="1F6B75", end_color="1F6B75", fill_type="solid")
    qtr_font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFFFF")
    qtr_border = Border(
        left=Side(style="medium", color="145260"),
        right=Side(style="medium", color="145260"),
        top=Side(style="medium", color="145260"),
        bottom=Side(style="medium", color="145260"),
    )
    N_COLS = len(HEADERS)

    # For Summary statistics
    summary_stats = []

    for quarter in qtr_names:
        ws_src  = wb_src[quarter]
        ws_dest = quarter_sheets[quarter]
        print(f"\nProcessing quarter: {quarter}...")

        # Merged Quarter Banner Header row
        ws_dest.merge_cells(start_row=2, start_column=1, end_row=2, end_column=N_COLS)
        ws_dest.row_dimensions[2].height = 20
        cell = ws_dest.cell(row=2, column=1, value=f"  ◆  {quarter}  —  EMA Low Support Filtered Strategy")
        cell.font = qtr_font
        cell.fill = qtr_fill
        cell.border = qtr_border
        cell.alignment = Alignment(horizontal="left", vertical="center")
        for c in range(2, N_COLS + 1):
            ws_dest.cell(row=2, column=c).fill   = qtr_fill
            ws_dest.cell(row=2, column=c).border = qtr_border
        
        quarter_row_ptr[quarter] += 1
        
        # Quarter stats tracking variables
        total_stocks = 0
        trades_taken = 0
        profits = 0
        losses = 0
        sl_hits = 0
        total_ret = 0.0

        for r_src in range(2, ws_src.max_row + 1):
            comp    = ws_src.cell(r_src, 1).value
            sym_val = ws_src.cell(r_src, 2).value
            D       = ws_src.cell(r_src, 3).value

            if not sym_val:
                continue
            sym = str(sym_val).strip()

            if sym not in params or sym not in price_cache:
                continue
            if not isinstance(D, dt.datetime):
                continue

            bb, sa, exp = params[sym]
            df          = price_cache[sym]
            latest_dt   = df.index[-1]

            entry_date = get_trading_date(df.index, D, -int(bb))
            exit_date  = get_trading_date(df.index, D,  int(sa))
            if entry_date is None or exit_date is None:
                continue

            entry_dt = pd.Timestamp(entry_date).normalize()
            exit_dt  = pd.Timestamp(exit_date).normalize()

            data_row = dict(company=comp, symbol=sym, quarter=quarter,
                            result_date=D, buy_before=bb, sell_after=sa, exp=exp)

            res = apply_strategy(df, entry_dt, exit_dt, latest_dt)
            
            # Statistics tracking
            total_stocks += 1
            if res["status"] in ("Realised", "Realised (SL)", "Open"):
                trades_taken += 1
                ret = res["ret"]
                if ret is not None:
                    total_ret += ret
                    if ret >= 0: profits += 1
                    else: losses += 1
                if res["sl_hit"]:
                    sl_hits += 1

            r_dest = quarter_row_ptr[quarter]
            write_row(ws_dest, r_dest, data_row, res, entry_date, exit_date, S)
            quarter_row_ptr[quarter] += 1

        # Save stats
        avg_ret = total_ret / trades_taken if trades_taken > 0 else 0.0
        win_rate = (profits / trades_taken * 100) if trades_taken > 0 else 0.0
        summary_stats.append({
            "quarter": quarter,
            "total": total_stocks,
            "traded": trades_taken,
            "win_rate": win_rate,
            "avg_ret": avg_ret,
            "sl_hits": sl_hits
        })

    # Auto-fit columns for all quarter sheets
    for qtr in qtr_names:
        autofit(quarter_sheets[qtr])

    # ── Summary Sheet ──────────────────────────────────────────────────────────
    print("\nWriting Summary sheet...")
    ws_sum = wb_out.create_sheet(title="Summary", index=0)
    ws_sum.freeze_panes = "A2"
    ws_sum.sheet_view.showGridLines = True

    SUM_HEADERS = [
        "Quarter", "Total Stock Events", "Trades Taken", "Win Rate (%)",
        "Average Return (%)", "SL Hits", "Status"
    ]
    write_header_row(ws_sum, SUM_HEADERS, S)

    for r_idx, row in enumerate(summary_stats, 2):
        bg = S["soft"] if r_idx % 2 == 0 else PatternFill(fill_type=None)

        def sw(col, val, fnt=None, fll=None, nfmt=None):
            cell = ws_sum.cell(row=r_idx, column=col, value=val)
            cell.font = fnt or S["normal"]; cell.fill = fll or bg
            cell.border = S["border"]; cell.alignment = S["ac"]
            if nfmt: cell.number_format = nfmt

        sw(1, row["quarter"])
        sw(2, row["total"])
        sw(3, row["traded"])
        
        wr = row["win_rate"]
        sw(4, wr, fnt=S["gtxt"] if wr >= 50 else S["rtxt"],
           fll=S["gfill"] if wr >= 50 else S["rfill"], nfmt="0.00")
           
        avg = row["avg_ret"]
        sw(5, avg, fnt=S["gtxt"] if avg >= 0 else S["rtxt"],
           fll=S["gfill"] if avg >= 0 else S["rfill"], nfmt="0.00")
           
        sw(6, row["sl_hits"], fnt=S["otxt"] if row["sl_hits"] > 0 else None,
           fll=S["ofill"] if row["sl_hits"] > 0 else None)
        sw(7, "Completed")

    for col in ws_sum.columns:
        ws_sum.column_dimensions[get_column_letter(col[0].column)].width = 18

    # ── Save ───────────────────────────────────────────────────────────────────
    print("\nSaving workbook...")
    saved = False
    stem  = "Nifty50Stocks_QtyResultDates_Filtered_EMA"
    for suffix in ["", "_2", "_3", "_4", "_5", "_6"]:
        try:
            out = ROOT / (stem + suffix + ".xlsx")
            wb_out.save(out)
            OUT_EXCEL = out
            saved = True
            if suffix:
                print(f"  (Workbook locked; saved as {out.name})")
            break
        except PermissionError:
            continue
    if not saved:
        out = ROOT / f"{stem}_{int(time.time())}.xlsx"
        wb_out.save(out)
        OUT_EXCEL = out
        print(f"  (All paths locked; saved as {out.name})")

    print(f"\nDone! Saved -> {OUT_EXCEL.name}")


if __name__ == "__main__":
    main()
