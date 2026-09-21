import openpyxl, sys

sys.stdout.reconfigure(errors='replace')

master_path = r'D:/behaviour analysis/Nifty50_12_Quarters_Options_1PCT_ITM_Detailed_Master.xlsx'
wb = openpyxl.load_workbook(master_path, data_only=True)

print("==========================================================================================")
print("VERIFYING DETAILED MASTER 13-SHEET WORKBOOK")
print("==========================================================================================")
print("Total Sheets:", len(wb.sheetnames))
print("Sheet Names:", wb.sheetnames)

for sheet in wb.sheetnames:
    ws = wb[sheet]
    print(f"\nSheet '{sheet}': Max Row={ws.max_row}, Max Col={ws.max_column}")
    title_val = ws.cell(row=1, column=1).value
    print(f"  Title Banner: '{title_val}'")

print("==========================================================================================")
