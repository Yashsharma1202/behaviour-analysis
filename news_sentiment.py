"""
news_sentiment.py
===============================================================================
ONLINE sentiment for the Nifty-50 behaviour project — a SEPARATE, additive module
(it never touches the offline dashboard core). For a stock it pulls recent news
headlines from Google News RSS (free, no API key) and scores their tone with
VADER + a small finance lexicon, returning:

    { symbol, name, buzz, avg_score, pos, neg, neu, label, headlines[...] }

It also flags headlines that look like a board-meeting / quarterly-results date,
so it doubles as a finder for the upcoming result dates missing from the sheet.

    python news_sentiment.py                 # demo on a few stocks
    python news_sentiment.py ICICIBANK SBIN  # specific symbols
    python news_sentiment.py --missing       # the 12 stocks with no Q1 date yet

Requires internet (this is the ONLINE feature). No API key, no paid service.
===============================================================================
"""
from __future__ import annotations

import re
import sys
import time
import urllib.parse
import urllib.request
from html import unescape

sys.stdout.reconfigure(encoding="utf-8")

from vaderSentiment.vaderSentiment import SentimentIntensityAnalyzer

_VADER = SentimentIntensityAnalyzer()

# Small finance-specific booster so VADER understands market words it doesn't
# ship with. Values are added to VADER's lexicon (scale roughly -4..+4).
_FINANCE_LEX = {
    "beat": 2.2, "beats": 2.2, "outperform": 2.5, "upgrade": 2.6, "upgraded": 2.6,
    "buyback": 1.8, "bullish": 2.6, "surge": 2.4, "surges": 2.4, "jumps": 2.2,
    "rally": 2.2, "record": 1.6, "profit": 1.4, "bonus": 1.6, "dividend": 1.2,
    "orderwin": 2.0, "multibagger": 2.6, "upside": 1.8, "strong": 1.5,
    "miss": -2.2, "misses": -2.2, "downgrade": -2.6, "downgraded": -2.6,
    "bearish": -2.6, "plunge": -2.6, "plunges": -2.6, "slump": -2.4, "slumps": -2.4,
    "falls": -1.8, "drop": -1.6, "drops": -1.6, "cut": -1.4, "cuts": -1.4,
    "loss": -2.0, "losses": -2.0, "fraud": -3.2, "probe": -2.0, "raid": -2.4,
    "resign": -1.8, "resigns": -1.8, "weak": -1.6, "decline": -1.6, "warning": -1.8,
    "downside": -1.8, "lawsuit": -2.0, "default": -2.8,
}
_VADER.lexicon.update(_FINANCE_LEX)

# headline patterns that reveal a results / board-meeting date
_DATE_HINT = re.compile(
    r"(board meeting|q1|q2|q3|q4|quarter|results?|earnings)", re.I)


def _fetch_rss(query: str, timeout: int = 20) -> bytes:
    url = ("https://news.google.com/rss/search?q="
           + urllib.parse.quote(query)
           + "&hl=en-IN&gl=IN&ceid=IN:en")
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    return urllib.request.urlopen(req, timeout=timeout).read()


def _parse_items(xml: bytes) -> list[dict]:
    text = xml.decode("utf-8", "replace")
    items = []
    for block in re.findall(r"<item>(.*?)</item>", text, re.DOTALL):
        def grab(tag):
            m = re.search(rf"<{tag}>(.*?)</{tag}>", block, re.DOTALL)
            return unescape(re.sub(r"<.*?>", "", m.group(1)).strip()) if m else ""
        title = grab("title")
        if not title:
            continue
        # Google News appends " - <source>" to titles; keep it, it's harmless
        items.append({"title": title, "date": grab("pubDate"),
                      "source": grab("source"), "link": grab("link")})
    return items


def _label(score: float) -> str:
    return "positive" if score >= 0.15 else ("negative" if score <= -0.15 else "neutral")


def stock_sentiment(symbol: str, name: str, max_items: int = 25) -> dict:
    """Fetch + score recent news for one stock."""
    query = f'"{name}" share OR stock OR results'
    try:
        items = _parse_items(_fetch_rss(query))[:max_items]
    except Exception as e:                                    # noqa: BLE001
        return {"symbol": symbol, "name": name, "error": f"{type(e).__name__}: {e}",
                "buzz": 0, "headlines": []}
    scored = []
    for it in items:
        s = _VADER.polarity_scores(it["title"])["compound"]
        it = {**it, "score": round(s, 3), "tone": _label(s),
              "date_hint": bool(_DATE_HINT.search(it["title"]))}
        scored.append(it)
    n = len(scored)
    avg = round(sum(i["score"] for i in scored) / n, 3) if n else 0.0
    pos = sum(1 for i in scored if i["tone"] == "positive")
    neg = sum(1 for i in scored if i["tone"] == "negative")
    return {
        "symbol": symbol, "name": name, "buzz": n, "avg_score": avg,
        "pos": pos, "neg": neg, "neu": n - pos - neg, "label": _label(avg),
        "headlines": scored,
    }


# ── symbol -> company name (read from the sheet if present, else a small map) ──
def _names() -> dict:
    try:
        import openpyxl
        wb = openpyxl.load_workbook("Nifty50_Q1Results_BehaviourEngine.xlsx")
        ws = wb.active
        return {(ws.cell(r, 2).value or "").strip(): (ws.cell(r, 1).value or "").strip()
                for r in range(3, ws.max_row + 1) if ws.cell(r, 2).value}
    except Exception:                                         # noqa: BLE001
        return {}


MISSING = ["APOLLOHOSP", "BHARTIARTL", "GRASIM", "HINDALCO", "ICICIBANK",
           "MAXHEALTH", "ONGC", "POWERGRID", "SBIN", "TITAN", "TMPV", "TRENT"]


def main() -> None:
    names = _names()
    args = [a for a in sys.argv[1:] if not a.startswith("-")]
    if "--missing" in sys.argv:
        syms = MISSING
    elif args:
        syms = args
    else:
        syms = ["RELIANCE", "ICICIBANK", "TITAN"]

    emo = {"positive": "🟢", "negative": "🔴", "neutral": "⚪"}
    for sym in syms:
        name = names.get(sym, sym)
        r = stock_sentiment(sym, name)
        if r.get("error"):
            print(f"{sym:11} ! {r['error']}")
            continue
        print(f"\n{emo[r['label']]} {sym:11} {name}")
        print(f"   buzz={r['buzz']:>2}  avg={r['avg_score']:+.2f}  "
              f"(+{r['pos']}/-{r['neg']}/·{r['neu']})  → {r['label'].upper()}")
        for h in r["headlines"][:5]:
            flag = "📅" if h["date_hint"] else "  "
            print(f"   {emo[h['tone']]}{flag} {h['score']:+.2f}  {h['title'][:95]}")
        time.sleep(0.4)                                       # be polite to Google


if __name__ == "__main__":
    main()
