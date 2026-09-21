import os
import pathlib
import sys
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(errors='replace')

ROOT = pathlib.Path('D:/behaviour analysis')
OI_DIR = ROOT / 'OI_DATA'
PROC = ROOT / 'processed'
PRICE_CACHE = PROC / 'price_cache'

SYMBOLS = sorted([e.name.strip().upper() for e in os.scandir(OI_DIR) if e.is_dir()])

print("==========================================================================================")
print("BUILDING BEST PERFORMING STOCKS BY QUARTER MASTER WORKBOOK (5 CLEAN SHEETS)")
print("==========================================================================================")
print(f"Total Target Universe: {len(SYMBOLS)} Stocks")

LOT_SIZES = {
    "RELIANCE": 250, "TCS": 175, "INFY": 400, "HDFCBANK": 550, "ICICIBANK": 700,
    "BHARTIARTL": 950, "ITC": 1600, "SBIN": 1500, "LTIM": 150, "LT": 300,
    "HINDUNILVR": 300, "AXISBANK": 625, "KOTAKBANK": 400, "BAJFINANCE": 125,
    "M&M": 350, "MARUTI": 100, "SUNPHARMA": 350, "TATASTEEL": 5500,
    "NTPC": 1500, "POWERGRID": 1800, "TITAN": 175, "ADANIENT": 300,
    "ADANIPORTS": 625, "ULTRACEMCO": 100, "ASIANPAINT": 200, "COALINDIA": 2100,
    "BAJAJ-AUTO": 125, "JSWSTEEL": 675, "TATAMOTORS": 1425, "HCLTECH": 350,
    "GRASIM": 250, "HEROMOTOCO": 150, "EICHERMOT": 175, "CIPLA": 650,
    "HDFCLIFE": 1100, "SBILIFE": 375, "DRREDDY": 125, "BRITANNIA": 200,
    "APOLLOHOSP": 125, "TATACONSUM": 450, "HINDALCO": 1400, "BPCL": 1800,
    "INDUSINDBK": 500, "DIVISLAB": 200, "BAJAJFINSV": 500, "NESTLEIND": 200,
    "WIPRO": 1500, "ONGC": 3750, "TECHM": 600, "ASHOKLEY": 5000, "ADANIGREEN": 500
}

def get_stock_window_params(sym):
    h = sum(ord(c) for c in sym)
    pre_options = [8, 7, 5, 4, 3, 2, 1]
    post_options = [1, 2, 3, 4, 5, 8]
    pre = pre_options[h % len(pre_options)]
    post = post_options[(h * 3) % len(post_options)]
    return pre, post

# Container for all stock data across the 4 quarter types
quarter_data_store = {
    "Q1": [], # April, May, June
    "Q2": [], # July, August, September
    "Q3": [], # October, November, December
    "Q4": []  # January, February, March
}

stock_best_quarter_map = {}

for idx, sym in enumerate(SYMBOLS, 1):
    stock_folder = ROOT / sym
    fr_path = stock_folder / 'financial_results.csv'
    price_path = PRICE_CACHE / f"{sym}.csv"
    
    df_prices = pd.DataFrame()
    if price_path.exists():
        try:
            df_prices = pd.read_csv(price_path)
            df_prices['date'] = pd.to_datetime(df_prices['date'])
            df_prices = df_prices.sort_values('date').reset_index(drop=True)
        except Exception:
            pass
            
    if df_prices.empty:
        continue
        
    n_quarters_total = 0
    q_dates_history = []
    
    if fr_path.exists() and fr_path.stat().st_size > 10:
        try:
            df_fr = pd.read_csv(fr_path, dtype=str).fillna('')
            if not df_fr.empty and 'broadCastDate' in df_fr.columns:
                dts = pd.to_datetime(df_fr['broadCastDate'], errors='coerce').dropna().sort_values(ascending=False)
                n_quarters_total = len(dts)
                q_dates_history = dts.tolist()
        except Exception:
            pass
            
    trading_status = "QUALIFIED" if n_quarters_total >= 12 else "AVOID BUT MONITOR IT"
    pre_days, post_days = get_stock_window_params(sym)
    window_str = f"T-{pre_days} to T+{post_days}"
    valid_dates = df_prices['date'].tolist()
    if not valid_dates:
        continue
        
    lot_size = LOT_SIZES.get(sym, max(100, int(1000000 / float(df_prices['adj'].iloc[-1]))))
    last_price = float(df_prices['adj'].iloc[-1])
    margin_for_stock = round(0.20 * lot_size * last_price, 2)
    
    # Categorize historical broadcast dates strictly into Q1, Q2, Q3, Q4
    q_dates_by_type = {"Q1": [], "Q2": [], "Q3": [], "Q4": []}
    for dt in q_dates_history:
        m = dt.month
        if m in (4, 5, 6): q_t = "Q1"
        elif m in (7, 8, 9): q_t = "Q2"
        elif m in (10, 11, 12): q_t = "Q3"
        else: q_t = "Q4"
        q_dates_by_type[q_t].append(dt)

    stock_q_metrics = {}

    for q_t in ["Q1", "Q2", "Q3", "Q4"]:
        dts = q_dates_by_type[q_t]
        tot_q_count = len(dts)
        
        if tot_q_count == 0:
            continue

        emp_long_wins = 0; emp_short_wins = 0
        rets_long = []; rets_short = []
        pnl_long_total = 0.0; pnl_short_total = 0.0

        for dt in dts:
            r_i = min(range(len(valid_dates)), key=lambda i: abs((valid_dates[i] - dt).days))
            e_i = max(0, r_i - pre_days)
            x_i = min(len(valid_dates) - 1, r_i + post_days)
            pe = float(df_prices.loc[e_i, 'adj']); px = float(df_prices.loc[x_i, 'adj'])
            if pe > 0:
                rl = (px - pe) / pe * 100
                rs = (pe - px) / pe * 100
                rets_long.append(rl); rets_short.append(rs)
                
                b_to_l = round(lot_size * pe, 2); s_to_l = round(lot_size * px, 2)
                cost_l = round((b_to_l + s_to_l) * 0.0005, 2)
                net_pnl_l = round(s_to_l - b_to_l - cost_l, 2)
                pnl_long_total += net_pnl_l
                
                b_to_s = round(lot_size * px, 2); s_to_s = round(lot_size * pe, 2)
                cost_s = round((b_to_s + s_to_s) * 0.0005, 2)
                net_pnl_s = round(s_to_s - b_to_s - cost_s, 2)
                pnl_short_total += net_pnl_s
                
                if net_pnl_l > 0: emp_long_wins += 1
                if net_pnl_s > 0: emp_short_wins += 1

        wr_long = (emp_long_wins / tot_q_count) * 100
        wr_short = (emp_short_wins / tot_q_count) * 100
        
        if wr_long >= wr_short:
            opt_strat = "FUTURE LONG"
            opt_wr = round(wr_long, 2)
            opt_wins = emp_long_wins
            opt_losses = tot_q_count - emp_long_wins
            opt_avg_ret = round(np.mean(rets_long), 2) if rets_long else 0.0
            opt_pnl = round(pnl_long_total, 2)
        else:
            opt_strat = "FUTURE SHORT"
            opt_wr = round(wr_short, 2)
            opt_wins = emp_short_wins
            opt_losses = tot_q_count - emp_short_wins
            opt_avg_ret = round(np.mean(rets_short), 2) if rets_short else 0.0
            opt_pnl = round(pnl_short_total, 2)

        # STRICT FILTER FOR GOOD PERFORMING STOCKS: Win Rate >= 60% AND Avg Return > 0%
        if opt_wr >= 60.0 and opt_avg_ret > 0.0:
            item = {
                "sym": sym, "trading_status": trading_status, "n_quarters_total": n_quarters_total,
                "n_quarters_qtype": tot_q_count, "window": window_str, "strategy": opt_strat,
                "win_rate": opt_wr, "wins": opt_wins, "losses": opt_losses, "avg_ret": opt_avg_ret,
                "margin": margin_for_stock, "net_pnl": opt_pnl, "is_available": True
            }
            quarter_data_store[q_t].append(item)
            stock_q_metrics[q_t] = item

    if stock_q_metrics:
        # Determine the SINGLE BEST QUARTER for this stock (highest Avg Return)
        best_q_type = max(stock_q_metrics.keys(), key=lambda k: (stock_q_metrics[k]['avg_ret'], stock_q_metrics[k]['win_rate']))
        best_m = stock_q_metrics[best_q_type]
        stock_best_quarter_map[sym] = {
            "best_quarter": best_q_type,
            "strategy": best_m['strategy'],
            "trading_status": trading_status,
            "n_quarters_total": n_quarters_total,
            "win_rate": best_m['win_rate'],
            "avg_ret": best_m['avg_ret'],
            "net_pnl": best_m['net_pnl'],
            "margin": margin_for_stock
        }

out_path = ROOT / "Nifty211_Best_Performing_Stocks_By_Quarter_Master.xlsx"
wb = openpyxl.Workbook()

title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
title_fill = PatternFill("solid", fgColor="1F4E79")
card_val_font = Font(name="Calibri", size=15, bold=True, color="1F4E79")
hdr_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1F4E79")
subhdr_fill = PatternFill("solid", fgColor="2F5597")
border_thin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                     top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
card_fill = PatternFill("solid", fgColor="F2F4F7")
pos_fill = PatternFill("solid", fgColor="E2EFDA")
neg_fill = PatternFill("solid", fgColor="FCE4D6")
amb_fill = PatternFill("solid", fgColor="FFF2CC")
pos_font = Font(name="Calibri", size=11, bold=True, color="276A3C")
neg_font = Font(name="Calibri", size=11, bold=True, color="9C0006")
amb_font = Font(name="Calibri", size=11, bold=True, color="B25900")

quarter_names = {
    "Q1": "Q1 Best Stocks (April)",
    "Q2": "Q2 Best Stocks (July)",
    "Q3": "Q3 Best Stocks (October)",
    "Q4": "Q4 Best Stocks (January)"
}

# WRITE SHEETS 1 to 4: Q1, Q2, Q3, Q4 BEST STOCKS
for q_idx, q_t in enumerate(["Q1", "Q2", "Q3", "Q4"]):
    sheet_title = quarter_names[q_t]
    if q_idx == 0:
        ws = wb.active
        ws.title = sheet_title
    else:
        ws = wb.create_sheet(title=sheet_title)
        
    ws.views.sheetView[0].showGridLines = True
    
    raw_list = quarter_data_store[q_t]
    df_qtype = pd.DataFrame(raw_list)
    if not df_qtype.empty:
        df_qtype = df_qtype.sort_values(['avg_ret', 'win_rate', 'net_pnl'], ascending=[False, False, False]).reset_index(drop=True)
    
    # Title Banner
    ws.merge_cells("A1:L1")
    ws["A1"] = f"NIFTY 211 STOCKS - {sheet_title.upper()} (WIN RATE >= 60% & AVG RETURN > 0%)"
    ws["A1"].font = title_font; ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 38

    # KPI Summary Cards
    n_good = len(df_qtype)
    top_stock = df_qtype.iloc[0]['sym'] if not df_qtype.empty else "N/A"
    top_ret = df_qtype.iloc[0]['avg_ret'] if not df_qtype.empty else 0.0
    avg_ret_q = round(df_qtype['avg_ret'].mean(), 2) if not df_qtype.empty else 0.0
    avg_wr_q = round(df_qtype['win_rate'].mean(), 2) if not df_qtype.empty else 0.0
    tot_pnl_q = round(df_qtype['net_pnl'].sum(), 2) if not df_qtype.empty else 0.0
    
    cards = [
        ("GOOD PERFORMING STOCKS", f"{n_good} Stocks", "A3:B4", "A3"),
        ("TOP PERFORMING STOCK", f"{top_stock} (+{top_ret}%)", "C3:D4", "C3"),
        ("AVERAGE QUARTER RETURN %", f"{avg_ret_q}%", "E3:F4", "E3"),
        ("AVERAGE QUARTER WIN RATE %", f"{avg_wr_q}%", "G3:H4", "G3"),
        ("TOTAL QUARTER NET P&L", f"₹{tot_pnl_q:,.2f}", "I3:J4", "I3")
    ]
    for title, val, merge_range, top_left in cards:
        ws.merge_cells(merge_range)
        ws[top_left] = f"{title}\n{val}"
        ws[top_left].font = card_val_font
        ws[top_left].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws[top_left].fill = card_fill; ws[top_left].border = border_thin

    # Main Ranking Table
    rank_start_row = 6
    ws.cell(row=rank_start_row, column=1, value=f"{n_good} GOOD PERFORMING STOCKS IN {q_t} (WIN RATE >= 60% & AVG RETURN > 0%)").font = hdr_font
    ws.cell(row=rank_start_row, column=1).fill = hdr_fill
    ws.merge_cells(start_row=rank_start_row, start_column=1, end_row=rank_start_row, end_column=12)

    headers = [
        "Rank", "Stock Symbol", "Optimal Quarter Strategy", "Action Status", "Historical Quarter Cycles",
        "Position Window", "Quarter Win Rate %", "Winning Quarters", "Losing Quarters",
        "Quarter Average Return %", "Margin Capital (₹)", "Quarter Cumulative Net P&L (₹)"
    ]
    ws.row_dimensions[rank_start_row+1].height = 26
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=rank_start_row+1, column=c_idx, value=h)
        cell.font = hdr_font; cell.fill = subhdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, r in enumerate(df_qtype.to_dict('records'), start=rank_start_row+2):
        ws.cell(row=r_idx, column=1, value=r_idx - rank_start_row - 1).border = border_thin
        ws.cell(row=r_idx, column=1).alignment = Alignment(horizontal="center")
        
        ws.cell(row=r_idx, column=2, value=r['sym']).border = border_thin
        ws.cell(row=r_idx, column=2).font = Font(bold=True)
        
        c_st = ws.cell(row=r_idx, column=3, value=r['strategy'])
        c_st.border = border_thin; c_st.alignment = Alignment(horizontal="center")
        if r['strategy'] == 'FUTURE LONG': c_st.fill = pos_fill; c_st.font = pos_font
        else: c_st.fill = neg_fill; c_st.font = neg_font

        c_act = ws.cell(row=r_idx, column=4, value=r['trading_status'])
        c_act.border = border_thin; c_act.alignment = Alignment(horizontal="center")
        if r['trading_status'] == 'QUALIFIED': c_act.fill = pos_fill; c_act.font = pos_font
        else: c_act.fill = amb_fill; c_act.font = amb_font
        
        ws.cell(row=r_idx, column=5, value=f"{r['n_quarters_qtype']} Quarters").border = border_thin
        ws.cell(row=r_idx, column=5).alignment = Alignment(horizontal="center")
        
        ws.cell(row=r_idx, column=6, value=r['window']).border = border_thin
        ws.cell(row=r_idx, column=6).alignment = Alignment(horizontal="center")
        
        c_wr = ws.cell(row=r_idx, column=7, value=f"{r['win_rate']}%")
        c_wr.border = border_thin; c_wr.alignment = Alignment(horizontal="right")
        c_wr.font = Font(bold=True); c_wr.fill = pos_fill; c_wr.font = pos_font
        
        ws.cell(row=r_idx, column=8, value=r['wins']).border = border_thin
        ws.cell(row=r_idx, column=8).alignment = Alignment(horizontal="center")
        
        ws.cell(row=r_idx, column=9, value=r['losses']).border = border_thin
        ws.cell(row=r_idx, column=9).alignment = Alignment(horizontal="center")
        
        c_ret = ws.cell(row=r_idx, column=10, value=f"{r['avg_ret']}%")
        c_ret.border = border_thin; c_ret.alignment = Alignment(horizontal="right")
        c_ret.font = Font(bold=True); c_ret.fill = pos_fill; c_ret.font = pos_font
        
        ws.cell(row=r_idx, column=11, value=f"₹{r['margin']:,.2f}").border = border_thin
        ws.cell(row=r_idx, column=11).alignment = Alignment(horizontal="right")
        
        c_p = ws.cell(row=r_idx, column=12, value=f"₹{r['net_pnl']:,.2f}")
        c_p.border = border_thin; c_p.alignment = Alignment(horizontal="right")
        c_p.fill = pos_fill; c_p.font = pos_font

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = 24

# WRITE SHEET 5: BEST QUARTER MATRIX
ws_matrix = wb.create_sheet(title="Best Quarter Matrix")
ws_matrix.views.sheetView[0].showGridLines = True

df_matrix = pd.DataFrame([
    {
        "sym": k,
        "best_quarter": v['best_quarter'],
        "strategy": v['strategy'],
        "trading_status": v['trading_status'],
        "n_quarters_total": v['n_quarters_total'],
        "win_rate": v['win_rate'],
        "avg_ret": v['avg_ret'],
        "net_pnl": v['net_pnl'],
        "margin": v['margin']
    }
    for k, v in stock_best_quarter_map.items()
])

df_matrix = df_matrix.sort_values(['best_quarter', 'avg_ret'], ascending=[True, False]).reset_index(drop=True)

ws_matrix.merge_cells("A1:I1")
ws_matrix["A1"] = "NIFTY 211 STOCKS - SINGLE BEST PERFORMING QUARTER MATCHING MATRIX"
ws_matrix["A1"].font = title_font; ws_matrix["A1"].fill = title_fill
ws_matrix["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws_matrix.row_dimensions[1].height = 38

headers_m = [
    "Sr. No.", "Stock Symbol", "Single Best Quarter", "Optimal Strategy", "Action Status",
    "Total History Quarters", "Best Quarter Win Rate %", "Best Quarter Avg Return %", "Best Quarter Net P&L (₹)"
]
ws_matrix.row_dimensions[3].height = 26
for c_idx, h in enumerate(headers_m, 1):
    cell = ws_matrix.cell(row=3, column=c_idx, value=h)
    cell.font = hdr_font; cell.fill = subhdr_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for r_idx, r in enumerate(df_matrix.to_dict('records'), start=4):
    ws_matrix.cell(row=r_idx, column=1, value=r_idx - 3).border = border_thin
    ws_matrix.cell(row=r_idx, column=1).alignment = Alignment(horizontal="center")
    
    ws_matrix.cell(row=r_idx, column=2, value=r['sym']).border = border_thin
    ws_matrix.cell(row=r_idx, column=2).font = Font(bold=True)
    
    c_bq = ws_matrix.cell(row=r_idx, column=3, value=r['best_quarter'])
    c_bq.border = border_thin; c_bq.alignment = Alignment(horizontal="center")
    c_bq.font = Font(bold=True)
    
    c_st = ws_matrix.cell(row=r_idx, column=4, value=r['strategy'])
    c_st.border = border_thin; c_st.alignment = Alignment(horizontal="center")
    if r['strategy'] == 'FUTURE LONG': c_st.fill = pos_fill; c_st.font = pos_font
    else: c_st.fill = neg_fill; c_st.font = neg_font

    c_act = ws_matrix.cell(row=r_idx, column=5, value=r['trading_status'])
    c_act.border = border_thin; c_act.alignment = Alignment(horizontal="center")
    if r['trading_status'] == 'QUALIFIED': c_act.fill = pos_fill; c_act.font = pos_font
    else: c_act.fill = amb_fill; c_act.font = amb_font
    
    ws_matrix.cell(row=r_idx, column=6, value=f"{r['n_quarters_total']} Quarters").border = border_thin
    ws_matrix.cell(row=r_idx, column=6).alignment = Alignment(horizontal="center")
    
    c_wr = ws_matrix.cell(row=r_idx, column=7, value=f"{r['win_rate']}%")
    c_wr.border = border_thin; c_wr.alignment = Alignment(horizontal="right")
    c_wr.font = Font(bold=True); c_wr.fill = pos_fill; c_wr.font = pos_font
    
    c_ret = ws_matrix.cell(row=r_idx, column=8, value=f"{r['avg_ret']}%")
    c_ret.border = border_thin; c_ret.alignment = Alignment(horizontal="right")
    c_ret.font = Font(bold=True); c_ret.fill = pos_fill; c_ret.font = pos_font
    
    c_p = ws_matrix.cell(row=r_idx, column=9, value=f"₹{r['net_pnl']:,.2f}")
    c_p.border = border_thin; c_p.alignment = Alignment(horizontal="right")
    c_p.fill = pos_fill; c_p.font = pos_font

for col in ws_matrix.columns:
    col_letter = get_column_letter(col[0].column)
    ws_matrix.column_dimensions[col_letter].width = 24

wb.save(out_path)
print(f"  • Successfully generated: Nifty211_Best_Performing_Stocks_By_Quarter_Master.xlsx")
print("==========================================================================================")
print("BEST PERFORMING STOCKS BY QUARTER MASTER WORKBOOK CREATED SUCCESSFULLY!")
print("==========================================================================================")
