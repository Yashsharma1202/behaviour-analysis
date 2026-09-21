import json
import os
import pandas as pd
import pathlib
import sys
import warnings
import concurrent.futures
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore')
sys.stdout.reconfigure(errors='replace')

ROOT = pathlib.Path('D:/behaviour analysis')
OI_DIR = ROOT / 'OI_DATA'
PROC = ROOT / 'processed'
PROC.mkdir(parents=True, exist_ok=True)

SYMBOLS = sorted([e.name.strip().upper() for e in os.scandir(OI_DIR) if e.is_dir()])

NIFTY50_SYMBOLS = set([
    "ADANIENT", "ADANIPORTS", "APOLLOHOSP", "ASIANPAINT", "AXISBANK",
    "BAJAJ-AUTO", "BAJAJFINSV", "BAJFINANCE", "BEL", "BHARTIARTL",
    "CIPLA", "COALINDIA", "DRREDDY", "EICHERMOT", "ETERNAL",
    "GRASIM", "HCLTECH", "HDFCBANK", "HDFCLIFE", "HINDALCO",
    "HINDUNILVR", "ICICIBANK", "INDIGO", "INFY", "ITC", "JIOFIN",
    "JSWSTEEL", "KOTAKBANK", "LT", "M&M", "MARUTI", "MAXHEALTH", "NESTLEIND",
    "NTPC", "ONGC", "POWERGRID", "RELIANCE", "SBILIFE", "SBIN", "SHRIRAMFIN",
    "SUNPHARMA", "TATACONSUM", "TATASTEEL", "TCS", "TECHM",
    "TITAN", "TMPV", "TRENT", "ULTRACEMCO", "WIPRO"
])

print("==========================================================================================")
print("AUDIT & EXTRACT: QUARTERLY RESULT DATES & BOARD MEETING INTIMATION ANNOUNCEMENT DATES")
print("==========================================================================================")
print(f"Target Universe: {len(SYMBOLS)} Stocks")
print(f"Current System Time: {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M:%S')}")
print("------------------------------------------------------------------------------------------", flush=True)

matched_records = []

def process_stock(sym):
    stock_folder = ROOT / sym
    fr_path = stock_folder / 'financial_results.csv'
    bm_path = stock_folder / 'board_meetings.csv'
    ann_path = stock_folder / 'announcements.csv'
    
    if not fr_path.exists() or fr_path.stat().st_size < 10:
        return []
        
    stock_records = []
    try:
        df_fr = pd.read_csv(fr_path, dtype=str).fillna('')
        df_bm = pd.read_csv(bm_path, dtype=str).fillna('') if bm_path.exists() and bm_path.stat().st_size > 10 else pd.DataFrame()
        df_ann = pd.read_csv(ann_path, dtype=str).fillna('') if ann_path.exists() and ann_path.stat().st_size > 10 else pd.DataFrame()
        
        intimations = []
        if not df_bm.empty:
            for _, r in df_bm.iterrows():
                dt_str = r.get('bm_date', '') or r.get('an_dt', '') or r.get('sort_date', '')
                dt = pd.to_datetime(dt_str, errors='coerce')
                purpose = str(r.get('bm_purpose', '')) + " " + str(r.get('bm_desc', '')) + " " + str(r.get('purpose', ''))
                if pd.notna(dt):
                    intimations.append({'dt': dt, 'dt_str': dt_str, 'purpose': purpose})
                    
        if not df_ann.empty:
            for _, r in df_ann.iterrows():
                dt_str = r.get('an_dt', '') or r.get('sort_date', '')
                dt = pd.to_datetime(dt_str, errors='coerce')
                desc = str(r.get('desc', '')) + " " + str(r.get('attchmntText', ''))
                if pd.notna(dt) and any(k in desc.lower() for k in ['intimation', 'board meeting', 'financial result', 'consider results']):
                    intimations.append({'dt': dt, 'dt_str': dt_str, 'purpose': desc})
                    
        df_int = pd.DataFrame(intimations) if intimations else pd.DataFrame()
        
        for _, r in df_fr.iterrows():
            res_dt_raw = r.get('broadCastDate', '')
            res_dt = pd.to_datetime(res_dt_raw, errors='coerce')
            rel_to = r.get('relatingTo', '')
            to_date = r.get('toDate', '')
            
            if pd.notna(res_dt):
                intimation_date_str = "N/A"
                days_notice = "N/A"
                intimation_details = "N/A"
                
                if not df_int.empty:
                    priors = df_int[(df_int['dt'] <= res_dt) & (df_int['dt'] >= res_dt - pd.Timedelta(days=35))]
                    if not priors.empty:
                        best_int = priors.sort_values('dt', ascending=True).iloc[0]
                        intimation_date_str = best_int['dt'].strftime('%Y-%m-%d')
                        days_diff = (res_dt.normalize() - best_int['dt'].normalize()).days
                        days_notice = f"{days_diff} Days Prior Notice" if days_diff >= 0 else "Same Day"
                        intimation_details = str(best_int['purpose'])[:120]
                        
                stock_records.append({
                    'symbol': sym,
                    'is_nifty50': 'YES ✅' if sym in NIFTY50_SYMBOLS else 'NO',
                    'financial_period': rel_to if rel_to else f"Qtr ended {to_date}",
                    'result_date_announced_to_nse': intimation_date_str,
                    'actual_result_declaration_date': res_dt.strftime('%Y-%m-%d'),
                    'prior_notice_days': days_notice,
                    'intimation_details': intimation_details,
                    'raw_broadcast_time': res_dt_raw
                })
    except Exception:
        pass
        
    return stock_records

print("Processing universe stocks in parallel pool...", flush=True)

with concurrent.futures.ThreadPoolExecutor(max_workers=12) as executor:
    futures = {executor.submit(process_stock, sym): sym for sym in SYMBOLS}
    count = 0
    for future in concurrent.futures.as_completed(futures):
        count += 1
        recs = future.result()
        if recs:
            matched_records.extend(recs)
        if count % 50 == 0 or count == len(SYMBOLS):
            print(f"  Progress: {count}/{len(SYMBOLS)} stocks completed...", flush=True)

df_all = pd.DataFrame(matched_records)
if not df_all.empty:
    df_all['_dt'] = pd.to_datetime(df_all['actual_result_declaration_date'], errors='coerce')
    df_all = df_all.sort_values(['_dt', 'symbol'], ascending=[False, True]).drop(columns=['_dt'])

print("------------------------------------------------------------------------------------------")
print("RECENT QUARTERLY RESULTS: INTIMATION ANNOUNCEMENT DATE vs ACTUAL RESULT DATE:")
print("------------------------------------------------------------------------------------------")

if not df_all.empty:
    print(f"Total Matched Result Records: {len(df_all)}")
    print("\nMost Recent 25 Quarterly Results across Universe:")
    print(df_all[['symbol', 'financial_period', 'result_date_announced_to_nse', 'actual_result_declaration_date', 'prior_notice_days']].head(25).to_string(index=False))

# Generate Master Excel Workbook
out_excel = ROOT / 'Recent_Quarterly_Result_Intimation_vs_Declaration_Dates_Master.xlsx'
wb = openpyxl.Workbook()

hdr_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1F4E79")
border_thin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                     top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))

# Sheet 1: Recent Result Intimation vs Declaration (All 211 Stocks)
ws1 = wb.active
ws1.title = "Result Intimation vs Declaration"
ws1.views.sheetView[0].showGridLines = True

headers1 = [
    "Sr. No.", "Stock Symbol", "Nifty 50 Status", "Financial Period / Description",
    "Date Result Date Announced to NSE (Intimation Date)", "Actual Result Declaration Date",
    "Prior Notice Period", "Intimation Details"
]

ws1.row_dimensions[1].height = 26
for c_idx, h in enumerate(headers1, 1):
    cell = ws1.cell(row=1, column=c_idx, value=h)
    cell.font = hdr_font
    cell.fill = hdr_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")

if not df_all.empty:
    for r_idx, r in enumerate(df_all.to_dict('records'), start=2):
        ws1.cell(row=r_idx, column=1, value=r_idx - 1).alignment = Alignment(horizontal="center")
        ws1.cell(row=r_idx, column=2, value=r['symbol'])
        ws1.cell(row=r_idx, column=3, value=r['is_nifty50']).alignment = Alignment(horizontal="center")
        ws1.cell(row=r_idx, column=4, value=r['financial_period'])
        ws1.cell(row=r_idx, column=5, value=r['result_date_announced_to_nse']).alignment = Alignment(horizontal="center")
        ws1.cell(row=r_idx, column=6, value=r['actual_result_declaration_date']).alignment = Alignment(horizontal="center")
        ws1.cell(row=r_idx, column=7, value=r['prior_notice_days']).alignment = Alignment(horizontal="center")
        ws1.cell(row=r_idx, column=8, value=r['intimation_details'])
        for c in range(1, 9):
            ws1.cell(row=r_idx, column=c).border = border_thin

for col in ws1.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

# Sheet 2: Nifty 50 Result Intimation vs Declaration
ws2 = wb.create_sheet(title="Nifty 50 Intimation vs Declaration")
ws2.views.sheetView[0].showGridLines = True

headers2 = [
    "Sr. No.", "Nifty 50 Symbol", "Financial Period / Description",
    "Date Result Date Announced to NSE (Intimation Date)", "Actual Result Declaration Date",
    "Prior Notice Period", "Intimation Details"
]
ws2.row_dimensions[1].height = 26
for c_idx, h in enumerate(headers2, 1):
    cell = ws2.cell(row=1, column=c_idx, value=h)
    cell.font = hdr_font
    cell.fill = PatternFill("solid", fgColor="2F5597")
    cell.alignment = Alignment(horizontal="center", vertical="center")

if not df_all.empty:
    df_nifty_all = df_all[df_all['is_nifty50'] == 'YES ✅'].copy()
    for r_idx, r in enumerate(df_nifty_all.to_dict('records'), start=2):
        ws2.cell(row=r_idx, column=1, value=r_idx - 1).alignment = Alignment(horizontal="center")
        ws2.cell(row=r_idx, column=2, value=r['symbol'])
        ws2.cell(row=r_idx, column=3, value=r['financial_period'])
        ws2.cell(row=r_idx, column=4, value=r['result_date_announced_to_nse']).alignment = Alignment(horizontal="center")
        ws2.cell(row=r_idx, column=5, value=r['actual_result_declaration_date']).alignment = Alignment(horizontal="center")
        ws2.cell(row=r_idx, column=6, value=r['prior_notice_days']).alignment = Alignment(horizontal="center")
        ws2.cell(row=r_idx, column=7, value=r['intimation_details'])
        for c in range(1, 8):
            ws2.cell(row=r_idx, column=c).border = border_thin

for col in ws2.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws2.column_dimensions[col_letter].width = max(max_len + 3, 12)

wb.save(out_excel)
print(f"\nSaved Master Excel: {out_excel}", flush=True)

# Save CSV
if not df_all.empty:
    df_all.to_csv(PROC / 'recent_result_intimation_vs_declaration_dates.csv', index=False)

print("==========================================================================================")
print("SUCCESS: Result Intimation Announcement Date vs Result Declaration Date Extracted!")
print("==========================================================================================")
