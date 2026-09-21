import openpyxl, sys

sys.stdout.reconfigure(errors='replace')

path = r'D:/behaviour analysis/Nifty50_12_Quarters_Drawdown_Charts_Only.xlsx'

wb = openpyxl.load_workbook(path, data_only=True)
print("==========================================")
print("VERIFYING CHARTS-ONLY EXCEL WORKBOOK")
print("==========================================")
print("Total Visible Sheets:", [s for s in wb.sheetnames if wb[s].sheet_state != 'hidden'])
print("Hidden Sheets       :", [s for s in wb.sheetnames if wb[s].sheet_state == 'hidden'])

total_charts = 0
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    num_charts = len(ws._charts)
    if num_charts > 0:
        total_charts += num_charts
        print(f"  Sheet '{sheet_name}': {num_charts} chart(s)")

print(f"\nTotal Charts across Workbook: {total_charts}")
print("==========================================")
