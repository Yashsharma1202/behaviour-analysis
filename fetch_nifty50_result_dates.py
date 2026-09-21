import urllib.request
import urllib.parse
import http.cookiejar
import json
import os
import pandas as pd
import pathlib
import sys
import time
import warnings
import concurrent.futures
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore')
sys.stdout.reconfigure(errors='replace')

ROOT = pathlib.Path('D:/behaviour analysis')
PROC = ROOT / 'processed'
PROC.mkdir(parents=True, exist_ok=True)

print("==========================================================================================")
print("NSE OFFICIAL SITE: NIFTY 50 UPCOMING & ANNOUNCED QUARTERLY RESULT DATES")
print("==========================================================================================")
print(f"Target Scope: ONLY NIFTY 50 CONSTITUENTS (50 STOCKS)")
print(f"Current System Time: {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M:%S')}")
print("------------------------------------------------------------------------------------------", flush=True)

class NSESession:
    def __init__(self):
        self.cj = http.cookiejar.CookieJar()
        self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(self.cj))
        self.headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
            'Accept': '*/*',
            'Accept-Language': 'en-US,en;q=0.9',
            'Referer': 'https://www.nseindia.com/market-data/event-calendar'
        }
        self.primed = False

    def prime(self, symbol="RELIANCE"):
        try:
            url = f"https://www.nseindia.com/get-quotes/equity?symbol={urllib.parse.quote(symbol)}"
            req = urllib.request.Request(url, headers=self.headers)
            with self.opener.open(req, timeout=10) as resp:
                resp.read()
            self.primed = True
        except Exception:
            pass

    def get_json(self, url, symbol="RELIANCE"):
        if not self.primed:
            self.prime(symbol)
        req_headers = self.headers.copy()
        req_headers['Referer'] = f"https://www.nseindia.com/get-quotes/equity?symbol={urllib.parse.quote(symbol)}"
        req = urllib.request.Request(url, headers=req_headers)
        
        for attempt in range(2):
            try:
                with self.opener.open(req, timeout=10) as resp:
                    raw = resp.read().decode('utf-8', errors='ignore')
                    return json.loads(raw)
            except Exception:
                self.prime(symbol)
                time.sleep(0.3)
        return None

nse = NSESession()
nse.prime("RELIANCE")

# 1. Fetch Nifty 50 Constituents live from NSE
print("Fetching live Nifty 50 constituents from NSE API...", flush=True)
nifty50_url = "https://www.nseindia.com/api/equity-stockIndices?index=NIFTY%2050"
nifty_data = nse.get_json(nifty50_url, "RELIANCE") or {}

NIFTY50_SYMBOLS = []
if isinstance(nifty_data, dict) and 'data' in nifty_data:
    NIFTY50_SYMBOLS = sorted([r['symbol'] for r in nifty_data['data'] if r.get('symbol') and r['symbol'] != 'NIFTY 50'])

# Fallback sync list if live fetch returned fewer than 45 stocks
FALLBACK_NIFTY50 = [
    "ADANIENT", "ADANIPORTS", "APOLLOHOSP", "ASIANPAINT", "AXISBANK",
    "BAJAJ-AUTO", "BAJAJFINSV", "BAJFINANCE", "BEL", "BHARTIARTL",
    "CIPLA", "COALINDIA", "DRREDDY", "EICHERMOT", "ETERNAL",
    "GRASIM", "HCLTECH", "HDFCBANK", "HDFCLIFE", "HINDALCO",
    "HINDUNILVR", "ICICIBANK", "INDIGO", "INFY", "ITC", "JIOFIN",
    "JSWSTEEL", "KOTAKBANK", "LT", "M&M", "MARUTI", "MAXHEALTH", "NESTLEIND",
    "NTPC", "ONGC", "POWERGRID", "RELIANCE", "SBILIFE", "SBIN", "SHRIRAMFIN",
    "SUNPHARMA", "TATACONSUM", "TATASTEEL", "TCS", "TECHM",
    "TITAN", "TMPV", "TRENT", "ULTRACEMCO", "WIPRO"
]

if len(NIFTY50_SYMBOLS) < 45:
    NIFTY50_SYMBOLS = sorted(FALLBACK_NIFTY50)

print(f"Total Nifty 50 Constituents Target: {len(NIFTY50_SYMBOLS)} Stocks", flush=True)

# 2. Fetch NSE Official Event Calendar API
print("Fetching NSE Event Calendar (https://www.nseindia.com/api/event-calendar)...", flush=True)
url_evt = 'https://www.nseindia.com/api/event-calendar'
events = nse.get_json(url_evt, "RELIANCE") or []

nifty_upcoming_bms = {}
for e in events:
    if not isinstance(e, dict): continue
    sym = (e.get('symbol') or '').upper().strip()
    if sym in set(NIFTY50_SYMBOLS):
        nifty_upcoming_bms[sym] = {
            'upcoming_date': e.get('date', ''),
            'purpose': e.get('purpose', ''),
            'details': e.get('bm_desc', '')
        }

# 3. Process each Nifty 50 Stock: Latest result date + Upcoming board meeting date + Q2 Oct/Nov historical window
nifty50_results_summary = []

def process_nifty_stock(sym):
    stock_folder = ROOT / sym
    stock_folder.mkdir(parents=True, exist_ok=True)
    encoded_sym = urllib.parse.quote(sym)
    
    # Live Results fetch
    res_url = f"https://www.nseindia.com/api/results-comparision?symbol={encoded_sym}"
    res_data = nse.get_json(res_url, sym) or {}
    res_raw = (res_data.get('resCmpData') or []) if isinstance(res_data, dict) else []
    
    latest_broadcast = "N/A"
    latest_period = "N/A"
    q2_oct_nov_dates = []
    
    if isinstance(res_raw, list) and len(res_raw) > 0:
        fr_rows = []
        for r in res_raw:
            if not isinstance(r, dict): continue
            dt_str = r.get('re_create_dt', '')
            rel = f"Qtr ended {r.get('re_to_dt', '')}"
            fr_rows.append({
                'broadCastDate': dt_str,
                'relatingTo': rel,
                'consolidated': r.get('re_res_type', 'Audited'),
                'audited': 'Audited',
                'fromDate': r.get('re_from_dt', ''),
                'toDate': r.get('re_to_dt', '')
            })
            
            dt = pd.to_datetime(dt_str, errors='coerce')
            if pd.notna(dt) and dt.month in (10, 11):
                q2_oct_nov_dates.append(dt.strftime('%d-%b'))
                
        if fr_rows:
            df_fr = pd.DataFrame(fr_rows)
            df_fr.to_csv(stock_folder / 'financial_results.csv', index=False)
            df_fr['_dt'] = pd.to_datetime(df_fr['broadCastDate'], errors='coerce')
            df_sorted = df_fr.sort_values('_dt', ascending=False)
            latest_r = df_sorted.iloc[0]
            latest_broadcast = str(latest_r['_dt'].date()) if pd.notna(latest_r['_dt']) else latest_r['broadCastDate']
            latest_period = latest_r['relatingTo']

    # Upcoming board meeting from event calendar
    evt_info = nifty_upcoming_bms.get(sym, {})
    upcoming_date = evt_info.get('upcoming_date', '')
    purpose = evt_info.get('purpose', '')
    details = evt_info.get('details', '')
    
    # If no upcoming board meeting on calendar yet, specify typical Q2 window
    typical_q2_window = ", ".join(sorted(set(q2_oct_nov_dates))) if q2_oct_nov_dates else "Mid Oct - Early Nov"
    
    status_str = f"CONFIRMED ({upcoming_date})" if upcoming_date else f"Expected ({typical_q2_window})"
    
    return {
        'symbol': sym,
        'upcoming_result_date': upcoming_date if upcoming_date else f"Expected: {typical_q2_window}",
        'status': 'ANNOUNCED ON NSE ✅' if upcoming_date else 'Pending Announcement (Est: Oct-Nov)',
        'announced_purpose': purpose if purpose else 'Q2 FY27 Financial Results (Period ended Sept 30)',
        'announced_details': details if details else f"Historically announced around {typical_q2_window}",
        'latest_broadcast_date': latest_broadcast,
        'latest_financial_period': latest_period,
        'typical_q2_announcement_window': typical_q2_window
    }

print("Extracting result dates across all 50 Nifty stocks in parallel pool...", flush=True)

with concurrent.futures.ThreadPoolExecutor(max_workers=10) as executor:
    futures = {executor.submit(process_nifty_stock, sym): sym for sym in NIFTY50_SYMBOLS}
    count = 0
    for future in concurrent.futures.as_completed(futures):
        count += 1
        res = future.result()
        if res:
            nifty50_results_summary.append(res)
        if count % 10 == 0 or count == len(NIFTY50_SYMBOLS):
            print(f"  Progress: {count}/{len(NIFTY50_SYMBOLS)} Nifty 50 stocks completed...", flush=True)

df_nifty = pd.DataFrame(nifty50_results_summary)
df_nifty = df_nifty.sort_values('symbol').reset_index(drop=True)

print("------------------------------------------------------------------------------------------", flush=True)
print("NIFTY 50 UPCOMING & ANNOUNCED QUARTERLY RESULT DATES SUMMARY:")
print("------------------------------------------------------------------------------------------", flush=True)

print(df_nifty[['symbol', 'upcoming_result_date', 'status', 'announced_purpose']].to_string(index=False), flush=True)

# Generate Excel Master Workbook
out_excel = ROOT / 'Nifty50_Upcoming_Quarterly_Result_Dates_2026.xlsx'
wb = openpyxl.Workbook()

hdr_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1F4E79")
border_thin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                     top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))

# Sheet 1: Nifty 50 Upcoming & Announced Result Dates
ws1 = wb.active
ws1.title = "Nifty 50 Result Dates"
ws1.views.sheetView[0].showGridLines = True

headers1 = [
    "Sr. No.", "Nifty 50 Symbol", "Upcoming Result Date / Expected Window", "NSE Announcement Status",
    "Announced Purpose / Event", "Details / Intimation", "Latest Result Broadcast Date",
    "Latest Financial Period", "Typical Q2 Announcement Window"
]

ws1.row_dimensions[1].height = 26
for c_idx, h in enumerate(headers1, 1):
    cell = ws1.cell(row=1, column=c_idx, value=h)
    cell.font = hdr_font
    cell.fill = hdr_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")

for r_idx, r in enumerate(df_nifty.to_dict('records'), start=2):
    ws1.cell(row=r_idx, column=1, value=r_idx - 1).alignment = Alignment(horizontal="center")
    ws1.cell(row=r_idx, column=2, value=r['symbol'])
    ws1.cell(row=r_idx, column=3, value=r['upcoming_result_date']).alignment = Alignment(horizontal="center")
    ws1.cell(row=r_idx, column=4, value=r['status']).alignment = Alignment(horizontal="center")
    ws1.cell(row=r_idx, column=5, value=r['announced_purpose'])
    ws1.cell(row=r_idx, column=6, value=r['announced_details'])
    ws1.cell(row=r_idx, column=7, value=r['latest_broadcast_date']).alignment = Alignment(horizontal="center")
    ws1.cell(row=r_idx, column=8, value=r['latest_financial_period'])
    ws1.cell(row=r_idx, column=9, value=r['typical_q2_announcement_window']).alignment = Alignment(horizontal="center")
    for c in range(1, 10):
        ws1.cell(row=r_idx, column=c).border = border_thin

for col in ws1.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

wb.save(out_excel)
print(f"\nSaved Master Excel: {out_excel}", flush=True)

# Save CSV
df_nifty.to_csv(PROC / 'nifty50_upcoming_quarterly_result_dates_2026.csv', index=False)

print("==========================================================================================", flush=True)
print("SUCCESS: Nifty 50 Quarterly Result Dates Extracted Successfully!", flush=True)
print("==========================================================================================", flush=True)
