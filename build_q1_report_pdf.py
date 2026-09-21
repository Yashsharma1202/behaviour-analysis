"""
build_q1_report_pdf.py
===============================================================================
Render the Nifty-50 Q1-Result Behaviour Engine sheet into a clean 2-page A4
(landscape) PDF — same look as the Excel print (header band, coloured returns,
italic = still-open, upcoming-date highlight, legend).

    python build_q1_report_pdf.py [source.xlsx]

Reads the workbook (default: the corrected UPDATED file), evaluates the sheet's
date/holding formulas the way Excel does, and writes an HTML + a PDF next to it.
Re-run after a trading day settles to refresh the returns.
===============================================================================
"""
from __future__ import annotations

import datetime as dt
import re
import subprocess
import sys
import time
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else \
    ROOT / "Nifty50_Q1Results_BehaviourEngine_UPDATED_06Aug_fixed.xlsx"
CUTOFF = dt.datetime(2026, 7, 31)          # dates after this are "upcoming"
ROWS_PER_PAGE = 25
CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"


def num(v, dp=2):
    """Format a number with thousands separators; '' for None."""
    if isinstance(v, (int, float)):
        return f"{v:,.{dp}f}".rstrip("0").rstrip(".") if dp else f"{v:,.0f}"
    return ""


def fdate(v):
    return v.strftime("%d-%b-%Y") if isinstance(v, dt.datetime) else ""


def eval_cell(val, D, E, F):
    """Evaluate the sheet's simple formulas the way Excel would.
       entry '=D-E' -> D-E calendar days · exit '=D+F' -> D+F · holding '=SUM(E:F)+n'."""
    if not isinstance(val, str) or not val.startswith("="):
        return val
    if val.upper().startswith("=SUM"):
        m = re.search(r"\+(\d+)", val)
        add = int(m.group(1)) if m else 1
        return (int(E) + int(F) + add) if (E is not None and F is not None) else None
    if D is None:
        return None
    if re.match(r"=D\d+-E\d+", val):
        return D - dt.timedelta(days=int(E))
    if re.match(r"=D\d+\+F\d+", val):
        return D + dt.timedelta(days=int(F))
    return None


def load_rows():
    ws = openpyxl.load_workbook(SRC).active
    out = []
    for r in range(3, ws.max_row + 1):
        g = lambda c: ws.cell(r, c).value
        D, E, F = g(4), g(5), g(6)
        out.append({
            "name": g(1) or "", "sym": (g(2) or "").strip(), "ltp": g(3),
            "result": D, "before": E, "after": F, "win": g(7),
            "holding": eval_cell(g(8), D, E, F), "exp": g(9),
            "entry": eval_cell(g(11), D, E, F), "buy": g(12),
            "exit": eval_cell(g(13), D, E, F), "sell": g(14), "ret": g(15),
        })
    return out


def ret_cell(v):
    """(text, css-class) for the Realised-Return cell."""
    if v is None or v == "-" or v == "":
        return "–", "dash"
    if isinstance(v, str):                         # "+3.43% (unrealised)"
        if "Filter Out" in v:
            return "Filter Out", "dash filter-out"
        m = re.match(r"\s*([+-]?\d+(?:\.\d+)?)", v)
        sign = float(m.group(1)) if m else 0.0
        cls = "open pos" if sign >= 0 else "open neg"
        return v.replace("(unrealised)", "(open)"), cls
    return (f"{v:+.2f}", "pos" if v >= 0 else "neg")


def upcoming(v):
    return isinstance(v, dt.datetime) and v > CUTOFF


def date_td(v):
    cls = "dt up" if upcoming(v) else "dt"
    return f'<td class="{cls}">{fdate(v) or "–"}</td>'


def row_html(x):
    rt, rc = ret_cell(x["ret"])
    exp = x["exp"]
    expc = "neg" if isinstance(exp, (int, float)) and exp < 0 else "exp"
    cells = [
        f'<td class="name">{x["name"]}</td>',
        f'<td class="sym">{x["sym"]}</td>',
        f'<td class="n">{num(x["ltp"])}</td>',
        date_td(x["result"]),
        f'<td class="c">{x["before"] if x["before"] is not None else "–"}</td>',
        f'<td class="c">{x["after"] if x["after"] is not None else "–"}</td>',
        f'<td class="c">{x["win"] if x["win"] is not None else "–"}</td>',
        f'<td class="c">{x["holding"] if x["holding"] is not None else "–"}</td>',
        f'<td class="n {expc}">{num(exp) if exp is not None else "–"}</td>',
        date_td(x["entry"]),
        f'<td class="n">{num(x["buy"]) or "–"}</td>',
        date_td(x["exit"]),
        f'<td class="n">{num(x["sell"]) or "–"}</td>',
        f'<td class="ret {rc}">{rt}</td>',
    ]
    return "<tr>" + "".join(cells) + "</tr>"


HEAD = ["Company Name", "SYMBOL", "LTP", "FY26-27 Q1<br>Result Date",
        "Buy Before<br>(Candles)", "Sell After<br>(Candles)", "Win<br>Rate",
        "Holding<br>Candles", "Expected<br>Return", "Entry Date", "Buy Price",
        "Exit Date", "Sell Price", "Realised Return"]


def table(rows):
    thead = (
        '<tr class="band"><th colspan="9">NIFTY 50 Stocks&nbsp; |&nbsp; Quarterly '
        'Results&nbsp; |&nbsp; Behaviour Analysis Engine</th>'
        '<th colspan="5">FY26-27 Q1 Result&nbsp; |&nbsp; Performance Review</th></tr>'
        '<tr class="cols">' + "".join(f"<th>{h}</th>" for h in HEAD) + "</tr>")
    body = "".join(row_html(x) for x in rows)
    return f'<table>{thead}{body}</table>'


def build_html(rows):
    pages = [rows[i:i + ROWS_PER_PAGE] for i in range(0, len(rows), ROWS_PER_PAGE)]
    gen = dt.date.today().strftime("%d-%b-%Y")
    legend = (
        '<div class="legend">'
        '<span><i class="sw pos"></i>Positive return</span>'
        '<span><i class="sw neg"></i>Negative return</span>'
        '<span><i class="em">Italic</i> = position still open (unrealised)</span>'
        '<span><i class="sw up"></i>Upcoming date (after 31-Jul-2026)</span>'
        f'<span class="gen">Generated {gen} · NIFTY 50 Q1 Result Behaviour Analysis Engine</span>'
        '</div>')
    body = ""
    for i, pg in enumerate(pages):
        last = i == len(pages) - 1
        body += (f'<section class="page">{table(pg)}'
                 + (legend if last else "")
                 + f'<div class="pnum">Page {i+1} of {len(pages)}</div></section>')
    return PAGE_CSS + body


PAGE_CSS = """<style>
@page { size: A4 landscape; margin: 8mm; }
*{ box-sizing:border-box; }
body{ margin:0; font-family:'Segoe UI',Arial,sans-serif; color:#1a1a1a; -webkit-print-color-adjust:exact; print-color-adjust:exact; }
.page{ page-break-after:always; }
.page:last-child{ page-break-after:auto; }
table{ width:100%; border-collapse:collapse; table-layout:fixed; }
th,td{ border:1px solid #c9d2dd; padding:3px 5px; font-size:8.4px; overflow:hidden; white-space:nowrap; text-overflow:ellipsis; }
tr.band th{ background:#1f4e79; color:#fff; font-size:11px; font-weight:700; text-align:left; padding:7px 9px; border-color:#1f4e79; }
tr.band th:last-child{ text-align:left; }
tr.cols th{ background:#d9e1f2; color:#1f3b5b; font-weight:700; text-align:center; font-size:8px; line-height:1.15; vertical-align:middle; }
td.name{ text-align:left; font-weight:600; }
td.sym{ text-align:center; font-weight:700; color:#1f4e79; }
td.n{ text-align:right; }
td.c,td.dt{ text-align:center; }
td.dash{ text-align:center; color:#9aa4b0; }
td.exp{ color:#0f6f8b; font-weight:600; }
td.pos{ color:#1a7f37; font-weight:700; }
td.neg{ color:#c0392b; font-weight:700; }
td.ret{ text-align:right; }
td.ret.open{ font-style:italic; }
td.dt.up{ background:#dbe9f7; color:#1f4e79; font-weight:700; }
/* column widths */
colgroup,col{}
td.name,th:nth-child(1){ }
table tr td:nth-child(1),table tr.cols th:nth-child(1){ width:16%; }
table tr td:nth-child(2){ width:6.5%; }
table tr td:nth-child(14){ width:9%; }
.legend{ margin-top:8px; display:flex; gap:18px; align-items:center; font-size:8px; color:#333; flex-wrap:wrap; }
.legend .sw{ display:inline-block; width:10px; height:10px; border-radius:2px; margin-right:4px; vertical-align:-1px; }
.legend .sw.pos{ background:#1a7f37; } .legend .sw.neg{ background:#c0392b; } .legend .sw.up{ background:#dbe9f7; border:1px solid #1f4e79; }
.legend .em{ font-style:italic; } .legend .gen{ margin-left:auto; color:#7a8699; }
.pnum{ text-align:right; font-size:8px; color:#7a8699; margin-top:3px; }
</style>"""


def main():
    rows = load_rows()
    html = build_html(rows)
    hpath = ROOT / "Nifty50_Q1_Result_Progress.html"
    ppath = ROOT / "Nifty50_Q1_Result_Progress.pdf"
    hpath.write_text(html, encoding="utf-8")
    print("wrote", hpath.name, f"({len(rows)} stocks)")
    out = ppath
    for cand in (ppath, ROOT / "Nifty50_Q1_Result_Progress_v2.pdf"):
        r = subprocess.run([CHROME, "--headless", "--disable-gpu",
                            "--no-pdf-header-footer", f"--print-to-pdf={cand}",
                            hpath.as_uri()], capture_output=True, text=True)
        if cand.exists() and cand.stat().st_size > 0:
            out = cand
            break
        time.sleep(1)
    print("wrote", out.name, f"({out.stat().st_size:,} bytes)" if out.exists() else "(FAILED)")


if __name__ == "__main__":
    main()
