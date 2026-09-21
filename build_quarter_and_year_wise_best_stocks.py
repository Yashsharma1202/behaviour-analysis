import pandas as pd
import numpy as np
import pathlib
import sys
import time
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

sys.stdout.reconfigure(errors='replace')

t0 = time.time()

BASE_DIR = pathlib.Path('D:/behaviour analysis')
FUT_V13 = BASE_DIR / 'Nifty50_12_Quarters_Futures_OI_Master_v13.xlsx'
EQ_V2   = BASE_DIR / 'Nifty50_12_Quarters_Equity_Master_v2.xlsx'
OUTPUT_EXCEL = BASE_DIR / 'Nifty50_QuarterWise_And_YearWise_Best_Stock_Performance.xlsx'

# Executive Color Palette Tokens
COLOR_TITLE_BG     = '1B365D'
COLOR_SECTION_BG   = '2B4C7E'
COLOR_HEADER_BG    = '1A202C'
COLOR_SUBHEADER_BG = '2C5282'
COLOR_ZEBRA_BG     = 'F8FAFC'
COLOR_CARD_BG      = 'F7FAFC'

COLOR_WIN_BG       = 'E6F4EA'
COLOR_WIN_TEXT     = '137333'
COLOR_LOSS_BG      = 'FCE8E6'
COLOR_LOSS_TEXT    = 'C5221F'
COLOR_GOLD_BG      = 'FEF3C7'
COLOR_GOLD_TEXT    = '92400E'

COLOR_WHITE        = 'FFFFFF'
COLOR_TEXT_MAIN    = '2D3748'

font_title     = Font(name='Segoe UI', size=14, bold=True, color=COLOR_WHITE)
font_section   = Font(name='Segoe UI', size=11, bold=True, color=COLOR_WHITE)
font_header    = Font(name='Segoe UI', size=10, bold=True, color=COLOR_WHITE)
font_card_lbl  = Font(name='Segoe UI', size=9, bold=True, color=COLOR_WHITE)
font_data_bold = Font(name='Segoe UI', size=10, bold=True, color=COLOR_TEXT_MAIN)

font_win       = Font(name='Segoe UI', size=10, bold=True, color=COLOR_WIN_TEXT)
font_loss      = Font(name='Segoe UI', size=10, bold=True, color=COLOR_LOSS_TEXT)
font_gold      = Font(name='Segoe UI', size=10, bold=True, color=COLOR_GOLD_TEXT)

fill_title     = PatternFill(start_color=COLOR_TITLE_BG, end_color=COLOR_TITLE_BG, fill_type='solid')
fill_section   = PatternFill(start_color=COLOR_SECTION_BG, end_color=COLOR_SECTION_BG, fill_type='solid')
fill_header    = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type='solid')
fill_subheader = PatternFill(start_color=COLOR_SUBHEADER_BG, end_color=COLOR_SUBHEADER_BG, fill_type='solid')
fill_zebra     = PatternFill(start_color=COLOR_ZEBRA_BG, end_color=COLOR_ZEBRA_BG, fill_type='solid')
fill_card      = PatternFill(start_color=COLOR_CARD_BG, end_color=COLOR_CARD_BG, fill_type='solid')

fill_win       = PatternFill(start_color=COLOR_WIN_BG, end_color=COLOR_WIN_BG, fill_type='solid')
fill_loss       = PatternFill(start_color=COLOR_LOSS_BG, end_color=COLOR_LOSS_BG, fill_type='solid')
fill_gold      = PatternFill(start_color=COLOR_GOLD_BG, end_color=COLOR_GOLD_BG, fill_type='solid')

border_thin = Border(
    left=Side(style='thin', color='CBD5E0'),
    right=Side(style='thin', color='CBD5E0'),
    top=Side(style='thin', color='CBD5E0'),
    bottom=Side(style='thin', color='CBD5E0')
)

border_double_bottom = Border(
    left=Side(style='thin', color='CBD5E0'),
    right=Side(style='thin', color='CBD5E0'),
    top=Side(style='thin', color='1A202C'),
    bottom=Side(style='double', color='1A202C')
)

print("Loading Master Trade Data for Equity & Futures...", flush=True)
xl_eq = pd.ExcelFile(EQ_V2, engine='openpyxl')
df_eq = xl_eq.parse('All_12Q_Equity_Trades')

xl_fut = pd.ExcelFile(FUT_V13, engine='openpyxl')
df_fut = xl_fut.parse('All_12Q_Futures_Trades')

# Extract Quarter Season (Q1, Q2, Q3, Q4) and Financial Year (FY24, FY25, FY26, FY27)
def add_quarter_season_and_fy(df):
    df['Season'] = df['Quarter'].apply(lambda x: str(x).split()[0]) # Q1, Q2, Q3, Q4
    df['FY_Code'] = df['Quarter'].apply(lambda x: str(x).split()[1] if len(str(x).split()) > 1 else 'FY25')
    return df

df_eq  = add_quarter_season_and_fy(df_eq)
df_fut = add_quarter_season_and_fy(df_fut)

eq_pnl_col = next(c for c in df_eq.columns if 'P&L' in c or 'Profit' in c)
eq_ret_col = next(c for c in df_eq.columns if 'Return' in c and 'Expected' not in c)

fut_pnl_col = next(c for c in df_fut.columns if 'P&L' in c or 'Profit' in c)
fut_ret_col = next(c for c in df_fut.columns if 'Margin' in c or 'Return' in c)

# Group Stock Performance by Season (Q1, Q2, Q3, Q4)
def compute_season_performance(df_eq_sub, df_fut_sub, season_name):
    # Equity stock metrics
    eq_grp = df_eq_sub.groupby(['Symbol', 'Company Name', 'Sector']).agg(
        Eq_Trades=(eq_pnl_col, 'count'),
        Eq_Wins=(eq_pnl_col, lambda x: (x > 0).sum()),
        Eq_Total_PnL=(eq_pnl_col, 'sum'),
        Eq_Avg_Return=(eq_ret_col, 'mean')
    ).reset_index()
    eq_grp['Eq_Win_Rate'] = eq_grp['Eq_Wins'] / eq_grp['Eq_Trades']

    # Futures stock metrics
    fut_grp = df_fut_sub.groupby('Symbol').agg(
        Fut_Trades=(fut_pnl_col, 'count'),
        Fut_Wins=(fut_pnl_col, lambda x: (x > 0).sum()),
        Fut_Total_PnL=(fut_pnl_col, 'sum'),
        Fut_Avg_Return=(fut_ret_col, 'mean')
    ).reset_index()
    fut_grp['Fut_Win_Rate'] = fut_grp['Fut_Wins'] / fut_grp['Fut_Trades']

    merged = pd.merge(eq_grp, fut_grp, on='Symbol', how='outer')
    merged['Season'] = season_name
    merged = merged.sort_values(by='Fut_Total_PnL', ascending=False).reset_index(drop=True)
    merged['Rank'] = range(1, len(merged) + 1)
    return merged

seasons = ['Q1', 'Q2', 'Q3', 'Q4']
season_df_dict = {}

for s in seasons:
    df_eq_s  = df_eq[df_eq['Season'] == s]
    df_fut_s = df_fut[df_fut['Season'] == s]
    season_df_dict[s] = compute_season_performance(df_eq_s, df_fut_s, s)

fys = ['2023-24', '2024-25', '2025-26', '2026-27']
fy_df_dict = {}
for fy_code in fys:
    df_eq_fy  = df_eq[df_eq['Quarter'].str.contains(fy_code)]
    df_fut_fy = df_fut[df_fut['Quarter'].str.contains(fy_code)]
    
    if not df_eq_fy.empty and not df_fut_fy.empty:
        fy_df_dict[fy_code] = compute_season_performance(df_eq_fy, df_fut_fy, fy_code)

# ------------------- 3. Build Full Excel Workbook -------------------
print("Creating Master Quarter-Wise & Year-Wise Best Stock Workbook...", flush=True)
wb = Workbook()
wb.remove(wb.active)

def style_table_headers(ws, start_row, headers, fill_style=fill_header):
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=c_idx, value=h)
        cell.font = font_header; cell.fill = fill_style
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border_thin

def autofit_columns(ws, max_cols=15):
    ws.views.sheetView[0].showGridLines = True
    for col_idx in range(1, max_cols + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in range(1, min(ws.max_row + 1, 100)):
            val = str(ws.cell(row=row, column=col_idx).value or '')
            if len(val) > max_len: max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

# --- SHEET 1: Executive_Q1_Q4_Summary ---
ws_exec = wb.create_sheet(title='Executive_Q1_Q4_Summary')

ws_exec.merge_cells('A1:K1')
c_t = ws_exec['A1']
c_t.value = "NIFTY 50 STRATEGY — QUARTER-WISE (Q1, Q2, Q3, Q4) & YEAR-WISE BEST STOCK LEADERBOARD"
c_t.font = font_title; c_t.fill = fill_title; c_t.alignment = Alignment(horizontal='center', vertical='center')
ws_exec.row_dimensions[1].height = 40

curr_row = 3

for s in seasons:
    df_s = season_df_dict[s]
    top_5_fut = df_s.head(5)
    
    ws_exec.merge_cells(f'A{curr_row}:K{curr_row}')
    c_s = ws_exec[f'A{curr_row}']
    c_s.value = f"TOP 5 BEST PERFORMING STOCKS IN {s} EARNINGS SEASON (MULTI-YEAR AGGREGATE)"
    c_s.font = font_section; c_s.fill = fill_section; c_s.alignment = Alignment(horizontal='left', vertical='center')
    ws_exec.row_dimensions[curr_row].height = 24
    curr_row += 1
    
    headers_exec = ['Rank', 'Symbol', 'Company Name', 'Sector', 'Qtr Trades', 'Equity Net P&L (₹)', 'Equity Win Rate (%)', 'Futures Net P&L (₹)', 'Futures Return on Margin (%)', 'Futures Win Rate (%)', 'Overall Rating']
    style_table_headers(ws_exec, curr_row, headers_exec, fill_header)
    ws_exec.row_dimensions[curr_row].height = 24
    curr_row += 1
    
    for r_idx, r_data in top_5_fut.iterrows():
        ws_exec.row_dimensions[curr_row].height = 21
        ws_exec.cell(row=curr_row, column=1, value=r_idx + 1).alignment = Alignment(horizontal='center')
        ws_exec.cell(row=curr_row, column=2, value=r_data['Symbol']).alignment = Alignment(horizontal='center')
        ws_exec.cell(row=curr_row, column=3, value=r_data['Company Name'])
        ws_exec.cell(row=curr_row, column=4, value=r_data['Sector'])
        ws_exec.cell(row=curr_row, column=5, value=int(r_data['Eq_Trades'])).alignment = Alignment(horizontal='center')
        
        c_eqp = ws_exec.cell(row=curr_row, column=6, value=r_data['Eq_Total_PnL'])
        c_eqp.number_format = '₹#,##0.00'; c_eqp.font = font_win if r_data['Eq_Total_PnL'] >= 0 else font_loss
        
        c_eqw = ws_exec.cell(row=curr_row, column=7, value=r_data['Eq_Win_Rate'])
        c_eqw.number_format = '0.00%'; c_eqw.alignment = Alignment(horizontal='center')
        
        c_futp = ws_exec.cell(row=curr_row, column=8, value=r_data['Fut_Total_PnL'])
        c_futp.number_format = '₹#,##0.00'; c_futp.font = font_win if r_data['Fut_Total_PnL'] >= 0 else font_loss; c_futp.fill = fill_win if r_data['Fut_Total_PnL'] >= 0 else fill_loss
        
        c_futr = ws_exec.cell(row=curr_row, column=9, value=r_data['Fut_Avg_Return'])
        c_futr.number_format = '+0.00%'
        
        c_futw = ws_exec.cell(row=curr_row, column=10, value=r_data['Fut_Win_Rate'])
        c_futw.number_format = '0.00%'; c_futw.alignment = Alignment(horizontal='center')
        
        ws_exec.cell(row=curr_row, column=11, value="⭐⭐⭐⭐⭐ TOP PERFORMER").alignment = Alignment(horizontal='center')
        
        for c in range(1, 12): ws_exec.cell(row=curr_row, column=c).border = border_thin
        curr_row += 1
        
    curr_row += 1

autofit_columns(ws_exec, max_cols=11)

# --- SHEETS 2-5: Q1, Q2, Q3, Q4 DETAILED LEADERBOARDS ---
headers_detail = ['Rank', 'Symbol', 'Company Name', 'Sector', 'Total Qtr Trades', 'Equity Total P&L (₹)', 'Equity Avg Return (%)', 'Equity Win Rate (%)', 'Futures Total P&L (₹)', 'Futures Return on Margin (%)', 'Futures Win Rate (%)']

for s in seasons:
    sheet_title = f"{s}_Best_Stocks_Equity_&_Fut"
    ws_s = wb.create_sheet(title=sheet_title)
    df_s = season_df_dict[s]
    
    ws_s.merge_cells('A1:K1')
    c_t = ws_s['A1']
    c_t.value = f"NIFTY 50 STOCKS PERFORMANCE LEADERBOARD — {s} EARNINGS SEASON (EQUITY vs FUTURES)"
    c_t.font = font_title; c_t.fill = fill_title; c_t.alignment = Alignment(horizontal='center', vertical='center')
    ws_s.row_dimensions[1].height = 36
    
    style_table_headers(ws_s, 3, headers_detail, fill_header)
    ws_s.row_dimensions[3].height = 26
    
    for r_idx, r_data in df_s.iterrows():
        row_num = 4 + r_idx
        row_fill = fill_zebra if r_idx % 2 == 1 else PatternFill(fill_type=None)
        ws_s.row_dimensions[row_num].height = 21
        
        ws_s.cell(row=row_num, column=1, value=r_idx + 1).alignment = Alignment(horizontal='center')
        ws_s.cell(row=row_num, column=2, value=r_data['Symbol']).alignment = Alignment(horizontal='center')
        ws_s.cell(row=row_num, column=3, value=r_data['Company Name'])
        ws_s.cell(row=row_num, column=4, value=r_data['Sector'])
        ws_s.cell(row=row_num, column=5, value=int(r_data['Eq_Trades'])).alignment = Alignment(horizontal='center')
        
        c_eqp = ws_s.cell(row=row_num, column=6, value=r_data['Eq_Total_PnL'])
        c_eqp.number_format = '₹#,##0.00'; c_eqp.font = font_win if r_data['Eq_Total_PnL'] >= 0 else font_loss
        
        c_eqr = ws_s.cell(row=row_num, column=7, value=r_data['Eq_Avg_Return'])
        c_eqr.number_format = '0.00%'
        
        c_eqw = ws_s.cell(row=row_num, column=8, value=r_data['Eq_Win_Rate'])
        c_eqw.number_format = '0.00%'; c_eqw.alignment = Alignment(horizontal='center')
        
        c_futp = ws_s.cell(row=row_num, column=9, value=r_data['Fut_Total_PnL'])
        c_futp.number_format = '₹#,##0.00'; c_futp.font = font_win if r_data['Fut_Total_PnL'] >= 0 else font_loss
        if r_data['Fut_Total_PnL'] >= 0: c_futp.fill = fill_win
        
        c_futr = ws_s.cell(row=row_num, column=10, value=r_data['Fut_Avg_Return'])
        c_futr.number_format = '+0.00%'
        
        c_futw = ws_s.cell(row=row_num, column=11, value=r_data['Fut_Win_Rate'])
        c_futw.number_format = '0.00%'; c_futw.alignment = Alignment(horizontal='center')
        
        for c in range(1, 12):
            cell = ws_s.cell(row=row_num, column=c)
            cell.border = border_thin
            if c not in [6, 9] and row_fill.fill_type: cell.fill = row_fill

    autofit_columns(ws_s, max_cols=11)
    ws_s.freeze_panes = 'A4'

# --- SHEETS 6-9: YEAR-WISE BREAKDOWNS (FY 2023-24, FY 2024-25, FY 2025-26, FY 2026-27) ---
for fy_code in fys:
    if fy_code in fy_df_dict:
        df_fy = fy_df_dict[fy_code]
        ws_fy = wb.create_sheet(title=f"FY_{fy_code}_YearWise_Perf")
        
        ws_fy.merge_cells('A1:K1')
        c_t = ws_fy['A1']
        c_t.value = f"NIFTY 50 STOCKS ANNUAL PERFORMANCE LEADERBOARD — FINANCIAL YEAR {fy_code}"
        c_t.font = font_title; c_t.fill = fill_title; c_t.alignment = Alignment(horizontal='center', vertical='center')
        ws_fy.row_dimensions[1].height = 36
        
        style_table_headers(ws_fy, 3, headers_detail, fill_header)
        ws_fy.row_dimensions[3].height = 26
        
        for r_idx, r_data in df_fy.iterrows():
            row_num = 4 + r_idx
            row_fill = fill_zebra if r_idx % 2 == 1 else PatternFill(fill_type=None)
            ws_fy.row_dimensions[row_num].height = 21
            
            ws_fy.cell(row=row_num, column=1, value=r_idx + 1).alignment = Alignment(horizontal='center')
            ws_fy.cell(row=row_num, column=2, value=r_data['Symbol']).alignment = Alignment(horizontal='center')
            ws_fy.cell(row=row_num, column=3, value=r_data['Company Name'])
            ws_fy.cell(row=row_num, column=4, value=r_data['Sector'])
            ws_fy.cell(row=row_num, column=5, value=int(r_data['Eq_Trades'])).alignment = Alignment(horizontal='center')
            
            c_eqp = ws_fy.cell(row=row_num, column=6, value=r_data['Eq_Total_PnL'])
            c_eqp.number_format = '₹#,##0.00'; c_eqp.font = font_win if r_data['Eq_Total_PnL'] >= 0 else font_loss
            
            c_eqr = ws_fy.cell(row=row_num, column=7, value=r_data['Eq_Avg_Return'])
            c_eqr.number_format = '0.00%'
            
            c_eqw = ws_fy.cell(row=row_num, column=8, value=r_data['Eq_Win_Rate'])
            c_eqw.number_format = '0.00%'; c_eqw.alignment = Alignment(horizontal='center')
            
            c_futp = ws_fy.cell(row=row_num, column=9, value=r_data['Fut_Total_PnL'])
            c_futp.number_format = '₹#,##0.00'; c_futp.font = font_win if r_data['Fut_Total_PnL'] >= 0 else font_loss
            if r_data['Fut_Total_PnL'] >= 0: c_futp.fill = fill_win
            
            c_futr = ws_fy.cell(row=row_num, column=10, value=r_data['Fut_Avg_Return'])
            c_futr.number_format = '+0.00%'
            
            c_futw = ws_fy.cell(row=row_num, column=11, value=r_data['Fut_Win_Rate'])
            c_futw.number_format = '0.00%'; c_futw.alignment = Alignment(horizontal='center')
            
            for c in range(1, 12):
                cell = ws_fy.cell(row=row_num, column=c)
                cell.border = border_thin
                if c not in [6, 9] and row_fill.fill_type: cell.fill = row_fill

        autofit_columns(ws_fy, max_cols=11)
        ws_fy.freeze_panes = 'A4'

# Save Workbook
wb.save(OUTPUT_EXCEL)
print(f"\n🎉 MASTER QUARTER-WISE & YEAR-WISE BEST STOCK WORKBOOK CREATED IN {time.time() - t0:.1f}s!", flush=True)
