import os
import pathlib
import sys
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(errors='replace')

ROOT = pathlib.Path('D:/behaviour analysis')
OI_DIR = ROOT / 'OI_DATA'
PROC = ROOT / 'processed'
PRICE_CACHE = PROC / 'price_cache'
PORTFOLIO_MASTER = ROOT / 'Nifty211_All_Stocks_Portfolio_Master.xlsx'
OUTPUT_FILE = ROOT / 'Nifty211_Last_8_Quarters_Detailed_Futures_Master.xlsx'

def clean_pct(val):
    if pd.isna(val): return 0.0
    val_str = str(val).replace('%', '').replace('₹', '').strip()
    try:
        v = float(val_str)
        return v / 100.0 if v > 1.0 else v
    except Exception:
        return 0.0

print("="*90)
print(f"Loading stock strategy & predicted benchmarks from: {PORTFOLIO_MASTER.name} ...")
df_pm = pd.read_excel(PORTFOLIO_MASTER, sheet_name='ALL QUARTERS COMBINED', skiprows=6)
df_pm['Stock Symbol'] = df_pm['Stock Symbol'].astype(str).str.strip().str.upper()

STOCK_MAP = {}
for _, row in df_pm.iterrows():
    sym = row['Stock Symbol']
    if pd.isna(sym) or not sym or sym == 'NAN':
        continue
    strat = str(row['Strategy']).strip().upper()
    window_raw = str(row['Position Window']).strip()
    status = str(row['Action Status']).strip()
    hist_q_str = str(row['Total History Quarters']).strip()
    sector = str(row['Industry Sector']).strip() if 'Industry Sector' in row and pd.notna(row['Industry Sector']) else "N/A"
    
    pred_win_rate = clean_pct(row.get('Win Rate %', 0))
    pred_avg_ret = clean_pct(row.get('Average Return %', 0))
    
    import re
    match = re.search(r'(\d+)', hist_q_str)
    hist_q_num = int(match.group(1)) if match else 5
    
    try:
        parts = window_raw.replace('T-', '').replace('T+', '').split(' to ')
        pre_days = int(parts[0])
        post_days = int(parts[1])
    except Exception:
        pre_days, post_days = 5, 5
        
    STOCK_MAP[sym] = {
        "sector": sector,
        "strategy": "LONG" if "LONG" in strat else "SHORT",
        "raw_strategy": strat,
        "window_str": window_raw,
        "pre_days": pre_days,
        "post_days": post_days,
        "status": status,
        "hist_q_str": hist_q_str,
        "hist_q_num": hist_q_num,
        "pred_win_rate": pred_win_rate,
        "pred_avg_ret": pred_avg_ret
    }

SYMBOLS = sorted(list(STOCK_MAP.keys()))

LOT_SIZES = {
    "RELIANCE": 250, "TCS": 175, "INFY": 400, "HDFCBANK": 550, "ICICIBANK": 700,
    "BHARTIARTL": 950, "ITC": 1600, "SBIN": 1500, "LTIM": 150, "LT": 300,
    "HINDUNILVR": 300, "AXISBANK": 625, "KOTAKBANK": 400, "BAJFINANCE": 125,
    "M&M": 350, "MARUTI": 100, "SUNPHARMA": 350, "TATASTEEL": 5500,
    "NTPC": 1500, "POWERGRID": 1800, "TITAN": 175, "ADANIENT": 300,
    "ADANIPORTS": 625, "ULTRACEMCO": 100, "ASIANPAINT": 200, "COALINDIA": 2100,
    "BAJAJ-AUTO": 125, "JSWSTEEL": 675, "TATAMOTORS": 1425, "HCLTECH": 350,
    "GRASIM": 250, "HEROMOTOCO": 150, "EICHERMOT": 175, "CIPLA": 650,
    "HDFCLIFE": 1100, "SBILIFE": 375, "DRREDDY": 125, "BRITANNIA": 200,
    "APOLLOHOSP": 125, "TATACONSUM": 450, "HINDALCO": 1400, "BPCL": 1800,
    "INDUSINDBK": 500, "DIVISLAB": 200, "BAJAJFINSV": 500, "NESTLEIND": 200,
    "WIPRO": 1500, "ONGC": 3750, "TECHM": 600, "ASHOKLEY": 5000, "ADANIGREEN": 500
}

TARGET_QUARTERS_8 = [
    {"q_idx": 1, "code": "Q3_FY23-24", "label": "Q3 FY2023-24", "q_type": "Q3", "target_year": 2024, "default_dt": "2024-01-22"},
    {"q_idx": 2, "code": "Q4_FY23-24", "label": "Q4 FY2023-24", "q_type": "Q4", "target_year": 2024, "default_dt": "2024-04-22"},
    {"q_idx": 3, "code": "Q1_FY24-25", "label": "Q1 FY2024-25", "q_type": "Q1", "target_year": 2024, "default_dt": "2024-07-22"},
    {"q_idx": 4, "code": "Q2_FY24-25", "label": "Q2 FY2024-25", "q_type": "Q2", "target_year": 2024, "default_dt": "2024-10-22"},
    {"q_idx": 5, "code": "Q3_FY24-25", "label": "Q3 FY2024-25", "q_type": "Q3", "target_year": 2025, "default_dt": "2025-01-22"},
    {"q_idx": 6, "code": "Q4_FY24-25", "label": "Q4 FY2024-25", "q_type": "Q4", "target_year": 2025, "default_dt": "2025-04-22"},
    {"q_idx": 7, "code": "Q1_FY25-26", "label": "Q1 FY2025-26", "q_type": "Q1", "target_year": 2025, "default_dt": "2025-07-22"},
    {"q_idx": 8, "code": "Q2_FY25-26", "label": "Q2 FY2025-26", "q_type": "Q2", "target_year": 2025, "default_dt": "2025-10-22"},
]

print("Target Universe          : 211 Nifty F&O / Broad Market Stocks")
print("Target Financial Quarters : LAST 8 QUARTERS (Q3 FY23-24 to Q2 FY25-26)")
print("-" * 90)

quarter_trades = {q['label']: [] for q in TARGET_QUARTERS_8}
stock_aggregated_trades = {sym: [] for sym in SYMBOLS}

for sym in SYMBOLS:
    stock_info = STOCK_MAP[sym]
    stock_folder = ROOT / sym
    fr_path = stock_folder / 'financial_results.csv'
    price_path = PRICE_CACHE / f"{sym}.csv"
    
    df_prices = pd.DataFrame()
    if price_path.exists():
        try:
            df_prices = pd.read_csv(price_path)
            df_prices['date'] = pd.to_datetime(df_prices['date'])
            df_prices = df_prices.sort_values('date').reset_index(drop=True)
        except Exception:
            pass
            
    if df_prices.empty:
        for q_info in TARGET_QUARTERS_8:
            quarter_trades[q_info['label']].append({
                "symbol": sym, "sector": stock_info['sector'], "status": stock_info['status'],
                "raw_strat": stock_info['raw_strategy'], "window_str": stock_info['window_str'],
                "hist_q_str": stock_info['hist_q_str'], "pred_win_rate": stock_info['pred_win_rate'],
                "pred_avg_ret": stock_info['pred_avg_ret'], "executed": False
            })
        continue
        
    valid_dates = df_prices['date'].tolist()
    last_price = float(df_prices['adj'].iloc[-1])
    lot_size = LOT_SIZES.get(sym, max(100, int(1000000 / last_price)))
    
    df_fr_clean = pd.DataFrame()
    if fr_path.exists() and fr_path.stat().st_size > 10:
        try:
            df_fr = pd.read_csv(fr_path, dtype=str).fillna('')
            if not df_fr.empty and 'broadCastDate' in df_fr.columns:
                df_fr['bDate'] = pd.to_datetime(df_fr['broadCastDate'], errors='coerce')
                df_fr = df_fr.dropna(subset=['bDate'])
                if 'toDate' in df_fr.columns:
                    df_fr['tDate'] = pd.to_datetime(df_fr['toDate'], errors='coerce')
                else:
                    df_fr['tDate'] = pd.NaT
                df_fr_clean = df_fr.sort_values('bDate', ascending=False)
        except Exception:
            pass

    stock_status = stock_info['status']
    pre_days = stock_info['pre_days']
    post_days = stock_info['post_days']
    window_str = stock_info['window_str']
    trade_dir = stock_info['strategy']
    raw_strat = stock_info['raw_strategy']
    hist_q_num = stock_info['hist_q_num']
    sector = stock_info['sector']
    pred_win_rate = stock_info['pred_win_rate']
    pred_avg_ret = stock_info['pred_avg_ret']
    
    hist_q_dates = {"Q1": [], "Q2": [], "Q3": [], "Q4": []}
    if not df_fr_clean.empty:
        for _, row in df_fr_clean.iterrows():
            b_dt = row['bDate']
            t_dt = row['tDate']
            ref_dt = t_dt if pd.notnull(t_dt) else b_dt
            m = ref_dt.month
            if m in [4, 5, 6]: hist_q_dates["Q1"].append(b_dt)
            elif m in [7, 8, 9]: hist_q_dates["Q2"].append(b_dt)
            elif m in [10, 11, 12]: hist_q_dates["Q3"].append(b_dt)
            elif m in [1, 2, 3]: hist_q_dates["Q4"].append(b_dt)

    for q_info in TARGET_QUARTERS_8:
        q_label = q_info['label']
        q_idx = q_info['q_idx']
        q_type = q_info['q_type']
        target_year = q_info['target_year']
        
        if q_idx > hist_q_num:
            t_record = {
                "symbol": sym, "sector": sector, "status": stock_status,
                "raw_strat": raw_strat, "window_str": window_str,
                "hist_q_str": stock_info['hist_q_str'], "pred_win_rate": pred_win_rate,
                "pred_avg_ret": pred_avg_ret, "executed": False
            }
            quarter_trades[q_label].append(t_record)
            continue

        b_dt = None
        if not df_fr_clean.empty:
            for _, row in df_fr_clean.iterrows():
                cand_b = row['bDate']
                cand_t = row['tDate']
                ref_dt = cand_t if pd.notnull(cand_t) else cand_b
                m = ref_dt.month
                y = ref_dt.year
                
                match_q = False
                if q_type == "Q1" and m in [4, 5, 6] and y == target_year: match_q = True
                elif q_type == "Q2" and m in [7, 8, 9] and y == target_year: match_q = True
                elif q_type == "Q3" and m in [10, 11, 12] and y == (target_year - 1): match_q = True
                elif q_type == "Q4" and m in [1, 2, 3] and y == target_year: match_q = True
                
                if match_q:
                    b_dt = cand_b
                    break

        if b_dt is None:
            candidates = hist_q_dates[q_type]
            if candidates:
                avg_day = int(np.mean([d.day for d in candidates]))
            else:
                avg_day = 22
                
            if q_type == "Q3": month_num = 1
            elif q_type == "Q4": month_num = 4
            elif q_type == "Q1": month_num = 7
            elif q_type == "Q2": month_num = 10
            
            try:
                b_dt = pd.Timestamp(year=target_year, month=month_num, day=min(avg_day, 28))
            except Exception:
                b_dt = pd.to_datetime(q_info['default_dt'])

        matching_indices = [i for i, d in enumerate(valid_dates) if d >= b_dt]
        if not matching_indices:
            idx_event = len(valid_dates) - 1
        else:
            idx_event = matching_indices[0]
            
        entry_idx = max(0, idx_event - pre_days)
        exit_idx = min(len(valid_dates) - 1, idx_event + post_days)
        
        entry_date = valid_dates[entry_idx]
        exit_date = valid_dates[exit_idx]
        
        row_entry = df_prices.iloc[entry_idx]
        row_exit = df_prices.iloc[exit_idx]
        
        entry_price = float(row_entry['adj'])
        exit_price = float(row_exit['adj'])
        
        pos_val = entry_price * lot_size
        margin_req = pos_val * 0.20
        
        if trade_dir == "LONG":
            gross_pnl = (exit_price - entry_price) * lot_size
        else:
            gross_pnl = (entry_price - exit_price) * lot_size
            
        fric_cost = pos_val * 0.0005
        net_pnl = gross_pnl - fric_cost
        act_ret_pct = (net_pnl / margin_req) * 100.0 if margin_req > 0 else 0.0
        win_flag = "WIN" if net_pnl > 0 else "LOSS"
        
        t_record = {
            "symbol": sym, "sector": sector, "status": stock_status,
            "raw_strat": raw_strat, "window_str": window_str,
            "hist_q_str": stock_info['hist_q_str'], "pred_win_rate": pred_win_rate,
            "pred_avg_ret": pred_avg_ret, "executed": True,
            "result_date": b_dt.strftime('%Y-%m-%d'), "entry_date": entry_date.strftime('%Y-%m-%d'),
            "exit_date": exit_date.strftime('%Y-%m-%d'), "entry_price": entry_price,
            "exit_price": exit_price, "lot_size": lot_size, "pos_val": pos_val,
            "margin_req": margin_req, "net_pnl": net_pnl, "act_ret_pct": act_ret_pct, "win_flag": win_flag
        }
        
        quarter_trades[q_label].append(t_record)
        stock_aggregated_trades[sym].append(t_record)

# Calculate Stock Performance Metrics for Ranking over 8 Quarters
stock_rank_list = []
for sym in SYMBOLS:
    trades = stock_aggregated_trades[sym]
    executed = [t for t in trades if t.get('executed', False)]
    tot_trades = len(executed)
    
    if tot_trades > 0:
        wins = sum(1 for t in executed if t['win_flag'] == 'WIN')
        losses = tot_trades - wins
        win_rate = (wins / tot_trades) * 100.0
        total_pnl = sum(t['net_pnl'] for t in executed)
        avg_act_ret = np.mean([t['act_ret_pct'] for t in executed])
        avg_margin = np.mean([t['margin_req'] for t in executed])
    else:
        wins, losses, win_rate, total_pnl, avg_act_ret, avg_margin = 0, 0, 0.0, 0.0, 0.0, 0.0
        
    s_info = STOCK_MAP[sym]
    stock_rank_list.append({
        "symbol": sym,
        "sector": s_info['sector'],
        "status": s_info['status'],
        "strategy": s_info['raw_strategy'],
        "window_str": s_info['window_str'],
        "hist_q_str": s_info['hist_q_str'],
        "pred_win_rate": s_info['pred_win_rate'],
        "pred_avg_ret": s_info['pred_avg_ret'],
        "tot_trades": tot_trades,
        "wins": wins,
        "losses": losses,
        "win_rate": win_rate,
        "avg_act_ret": avg_act_ret,
        "avg_margin": avg_margin,
        "total_pnl": total_pnl
    })

# Sort all stocks primarily by Actual Win Rate desc, then Total PnL desc
stock_rank_list = sorted(stock_rank_list, key=lambda x: (x['win_rate'], x['total_pnl']), reverse=True)

# Assign Overall Rank numbers
SYM_RANK_MAP = {}
for r_idx, item in enumerate(stock_rank_list, start=1):
    item['rank'] = r_idx
    SYM_RANK_MAP[item['symbol']] = r_idx

# Filter & Rank QUALIFIED Stocks
qual_rank_list = [dict(item) for item in stock_rank_list if item['status'] == 'QUALIFIED']
qual_rank_list = sorted(qual_rank_list, key=lambda x: (x['win_rate'], x['total_pnl']), reverse=True)

for r_idx, item in enumerate(qual_rank_list, start=1):
    item['qual_rank'] = r_idx

# Calculate Executive Performance Summary KPIs for 8 Quarters
total_portfolio_pnl = sum(item['total_pnl'] for item in stock_rank_list)
total_portfolio_trades = sum(item['tot_trades'] for item in stock_rank_list)
total_portfolio_wins = sum(item['wins'] for item in stock_rank_list)
overall_win_rate = (total_portfolio_wins / max(1, total_portfolio_trades)) * 100.0

qual_pnl = sum(item['total_pnl'] for item in qual_rank_list)
qual_trades = sum(item['tot_trades'] for item in qual_rank_list)
qual_wins = sum(item['wins'] for item in qual_rank_list)
qual_win_rate = (qual_wins / max(1, qual_trades)) * 100.0

top_stock = stock_rank_list[0]['symbol']
top_stock_pnl = stock_rank_list[0]['total_pnl']
top_stock_wr = stock_rank_list[0]['win_rate']

top_qual_stock = qual_rank_list[0]['symbol']
top_qual_pnl = qual_rank_list[0]['total_pnl']
top_qual_wr = qual_rank_list[0]['win_rate']

print("="*105)
print(f"LAST 8 QUARTERS EXECUTIVE PERFORMANCE SUMMARY:")
print(f"Total Portfolio Net Realised P&L : ₹{total_portfolio_pnl:,.2f}")
print(f"Overall Portfolio Win Rate %     : {overall_win_rate:.2f}% ({total_portfolio_wins} Wins / {total_portfolio_trades} Executed Trades)")
print(f"QUALIFIED Stocks Total Net P&L   : ₹{qual_pnl:,.2f}")
print(f"QUALIFIED Stocks Win Rate %       : {qual_win_rate:.2f}% ({qual_wins} Wins / {qual_trades} Executed Trades)")
print(f"Top Ranked Stock                 : {top_stock} ({top_stock_wr:.2f}% Win Rate, ₹{top_stock_pnl:,.2f} P&L)")
print("="*105)

# Create Workbook
wb = openpyxl.Workbook()
wb.remove(wb.active)

# Styles
header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

title_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")

kpi_header_fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
kpi_label_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
na_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

# 1. MASTER SUMMARY & RANKINGS SHEET (ALL 211 STOCKS)
ws_sum = wb.create_sheet(title="ALL 211 STOCKS RANKINGS")
ws_sum.views.sheetView[0].showGridLines = True

ws_sum.merge_cells("A1:P1")
ws_sum["A1"] = "NIFTY 211 - LAST 8 QUARTERS FUTURES PERFORMANCE & WIN RATE RANKINGS"
ws_sum["A1"].font = title_font
ws_sum["A1"].fill = title_fill
ws_sum["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws_sum.row_dimensions[1].height = 35

ws_sum.merge_cells("A3:P3")
ws_sum["A3"] = "EXECUTIVE PORTFOLIO PERFORMANCE SUMMARY (LAST 8 FINANCIAL QUARTERS)"
ws_sum["A3"].font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
ws_sum["A3"].fill = kpi_header_fill
ws_sum["A3"].alignment = Alignment(horizontal="center", vertical="center")
ws_sum.row_dimensions[3].height = 24

ws_sum.cell(row=4, column=1, value="Total Portfolio Net P&L (₹):").font = Font(bold=True)
ws_sum.cell(row=4, column=2, value=total_portfolio_pnl).number_format = '₹#,##0.00'
ws_sum.cell(row=4, column=2).font = Font(bold=True, color="006100")

ws_sum.cell(row=4, column=4, value="Overall Portfolio Win Rate %:").font = Font(bold=True)
ws_sum.cell(row=4, column=5, value=overall_win_rate / 100.0).number_format = '0.00%'
ws_sum.cell(row=4, column=5).font = Font(bold=True)

ws_sum.cell(row=4, column=7, value="Total Executed Trades (8Q):").font = Font(bold=True)
ws_sum.cell(row=4, column=8, value=total_portfolio_trades).number_format = '#,##0'

ws_sum.cell(row=4, column=10, value="Total N/A Rows:").font = Font(bold=True)
ws_sum.cell(row=4, column=11, value=(211 * 8) - total_portfolio_trades).number_format = '#,##0'

ws_sum.cell(row=5, column=1, value="QUALIFIED Stocks Net P&L (₹):").font = Font(bold=True)
ws_sum.cell(row=5, column=2, value=qual_pnl).number_format = '₹#,##0.00'
ws_sum.cell(row=5, column=2).font = Font(bold=True, color="006100")

ws_sum.cell(row=5, column=4, value="QUALIFIED Stocks Win Rate %:").font = Font(bold=True)
ws_sum.cell(row=5, column=5, value=qual_win_rate / 100.0).number_format = '0.00%'
ws_sum.cell(row=5, column=5).font = Font(bold=True)

ws_sum.cell(row=5, column=7, value="QUALIFIED Executed Trades:").font = Font(bold=True)
ws_sum.cell(row=5, column=8, value=qual_trades).number_format = '#,##0'

ws_sum.cell(row=5, column=10, value="Top Performing Stock:").font = Font(bold=True)
ws_sum.cell(row=5, column=11, value=f"{top_stock} ({top_stock_wr:.1f}% WR)").font = Font(bold=True, color="002060")

headers_rank = [
    "Rank", "Stock Symbol", "Industry Sector", "Action Status", "Strategy",
    "Position Window", "Total History Quarters", "Pred Historical Win Rate %", "Pred Avg Return %",
    "Total Trades Executed", "Winning Trades", "Losing Trades", "Actual Win Rate %",
    "Actual Avg Return %", "Margin Capital (20%) (₹)", "Total Net Realised P&L (₹)"
]

ws_sum.append([])
ws_sum.append(headers_rank)
ws_sum.row_dimensions[7].height = 25

for col_num, h in enumerate(headers_rank, 1):
    cell = ws_sum.cell(row=7, column=col_num)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for idx, item in enumerate(stock_rank_list, start=1):
    r_idx = idx + 7
    ws_sum.append([
        item['rank'], item['symbol'], item['sector'], item['status'], item['strategy'],
        item['window_str'], item['hist_q_str'], item['pred_win_rate'], item['pred_avg_ret'],
        item['tot_trades'], item['wins'], item['losses'], item['win_rate'] / 100.0,
        item['avg_act_ret'] / 100.0, item['avg_margin'], item['total_pnl']
    ])
    ws_sum.row_dimensions[r_idx].height = 20
    
    ws_sum.cell(row=r_idx, column=1).alignment = Alignment(horizontal="center")
    ws_sum.cell(row=r_idx, column=2).alignment = Alignment(horizontal="left")
    ws_sum.cell(row=r_idx, column=2).font = Font(bold=True)
    ws_sum.cell(row=r_idx, column=4).alignment = Alignment(horizontal="center")
    ws_sum.cell(row=r_idx, column=5).alignment = Alignment(horizontal="center")
    ws_sum.cell(row=r_idx, column=6).alignment = Alignment(horizontal="center")
    ws_sum.cell(row=r_idx, column=7).alignment = Alignment(horizontal="center")
    
    ws_sum.cell(row=r_idx, column=8).number_format = '0.00%'
    ws_sum.cell(row=r_idx, column=9).number_format = '0.00%'
    ws_sum.cell(row=r_idx, column=10).number_format = '#,##0'
    ws_sum.cell(row=r_idx, column=11).number_format = '#,##0'
    ws_sum.cell(row=r_idx, column=12).number_format = '#,##0'
    ws_sum.cell(row=r_idx, column=13).number_format = '0.00%'
    ws_sum.cell(row=r_idx, column=14).number_format = '0.00%'
    ws_sum.cell(row=r_idx, column=15).number_format = '₹#,##0.00'
    ws_sum.cell(row=r_idx, column=16).number_format = '₹#,##0.00'
    
    if item['total_pnl'] > 0:
        ws_sum.cell(row=r_idx, column=16).fill = green_fill
    elif item['total_pnl'] < 0:
        ws_sum.cell(row=r_idx, column=16).fill = red_fill

for col in ws_sum.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_sum.column_dimensions[col_letter].width = max(max_len + 3, 15)

# 2. QUALIFIED STOCKS RANKINGS SHEET (67 STOCKS)
ws_q = wb.create_sheet(title="QUALIFIED STOCKS RANKINGS")
ws_q.views.sheetView[0].showGridLines = True

ws_q.merge_cells("A1:P1")
ws_q["A1"] = "QUALIFIED STOCKS RANKINGS & PERFORMANCE SUMMARY (LAST 8 QUARTERS)"
ws_q["A1"].font = title_font
ws_q["A1"].fill = PatternFill(start_color="004B50", end_color="004B50", fill_type="solid")
ws_q["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws_q.row_dimensions[1].height = 35

ws_q.merge_cells("A3:P3")
ws_q["A3"] = "QUALIFIED STOCKS PERFORMANCE HIGHLIGHTS (LAST 8 QUARTERS)"
ws_q["A3"].font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
ws_q["A3"].fill = PatternFill(start_color="006666", end_color="006666", fill_type="solid")
ws_q["A3"].alignment = Alignment(horizontal="center", vertical="center")
ws_q.row_dimensions[3].height = 24

ws_q.cell(row=4, column=1, value="QUALIFIED Stocks Count:").font = Font(bold=True)
ws_q.cell(row=4, column=2, value=len(qual_rank_list)).number_format = '#,##0'

ws_q.cell(row=4, column=4, value="Total Executed Trades (8Q):").font = Font(bold=True)
ws_q.cell(row=4, column=5, value=qual_trades).number_format = '#,##0'

ws_q.cell(row=4, column=7, value="Total Winning Trades:").font = Font(bold=True)
ws_q.cell(row=4, column=8, value=qual_wins).number_format = '#,##0'

ws_q.cell(row=4, column=10, value="QUALIFIED Portfolio Win Rate %:").font = Font(bold=True)
ws_q.cell(row=4, column=11, value=qual_win_rate / 100.0).number_format = '0.00%'
ws_q.cell(row=4, column=11).font = Font(bold=True)

ws_q.cell(row=5, column=1, value="QUALIFIED Stocks Net P&L (₹):").font = Font(bold=True)
ws_q.cell(row=5, column=2, value=qual_pnl).number_format = '₹#,##0.00'
ws_q.cell(row=5, column=2).font = Font(bold=True, color="006100")

ws_q.cell(row=5, column=4, value="Top QUALIFIED Stock:").font = Font(bold=True)
ws_q.cell(row=5, column=5, value=f"{top_qual_stock} ({top_qual_wr:.1f}% WR)").font = Font(bold=True, color="002060")

ws_q.cell(row=5, column=7, value="Top QUALIFIED Net P&L:").font = Font(bold=True)
ws_q.cell(row=5, column=8, value=top_qual_pnl).number_format = '₹#,##0.00'
ws_q.cell(row=5, column=8).font = Font(bold=True, color="006100")

headers_qual_rank = [
    "Qual Rank", "Overall Rank", "Stock Symbol", "Industry Sector", "Strategy",
    "Position Window", "Total History Quarters", "Pred Historical Win Rate %", "Pred Avg Return %",
    "Total Trades Executed", "Winning Trades", "Losing Trades", "Actual Win Rate %",
    "Actual Avg Return %", "Margin Capital (20%) (₹)", "Net Realised P&L (₹)"
]

ws_q.append([])
ws_q.append(headers_qual_rank)
ws_q.row_dimensions[7].height = 25

for col_num, h in enumerate(headers_qual_rank, 1):
    cell = ws_q.cell(row=7, column=col_num)
    cell.font = header_font
    cell.fill = PatternFill(start_color="004B50", end_color="004B50", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for idx, item in enumerate(qual_rank_list, start=1):
    r_idx = idx + 7
    ws_q.append([
        item['qual_rank'], item['rank'], item['symbol'], item['sector'], item['strategy'],
        item['window_str'], item['hist_q_str'], item['pred_win_rate'], item['pred_avg_ret'],
        item['tot_trades'], item['wins'], item['losses'], item['win_rate'] / 100.0,
        item['avg_act_ret'] / 100.0, item['avg_margin'], item['total_pnl']
    ])
    ws_q.row_dimensions[r_idx].height = 20
    
    ws_q.cell(row=r_idx, column=1).alignment = Alignment(horizontal="center")
    ws_q.cell(row=r_idx, column=1).font = Font(bold=True)
    ws_q.cell(row=r_idx, column=2).alignment = Alignment(horizontal="center")
    ws_q.cell(row=r_idx, column=3).alignment = Alignment(horizontal="left")
    ws_q.cell(row=r_idx, column=3).font = Font(bold=True)
    ws_q.cell(row=r_idx, column=5).alignment = Alignment(horizontal="center")
    ws_q.cell(row=r_idx, column=6).alignment = Alignment(horizontal="center")
    ws_q.cell(row=r_idx, column=7).alignment = Alignment(horizontal="center")
    
    ws_q.cell(row=r_idx, column=8).number_format = '0.00%'
    ws_q.cell(row=r_idx, column=9).number_format = '0.00%'
    ws_q.cell(row=r_idx, column=10).number_format = '#,##0'
    ws_q.cell(row=r_idx, column=11).number_format = '#,##0'
    ws_q.cell(row=r_idx, column=12).number_format = '#,##0'
    ws_q.cell(row=r_idx, column=13).number_format = '0.00%'
    ws_q.cell(row=r_idx, column=14).number_format = '0.00%'
    ws_q.cell(row=r_idx, column=15).number_format = '₹#,##0.00'
    ws_q.cell(row=r_idx, column=16).number_format = '₹#,##0.00'
    
    if item['total_pnl'] > 0:
        ws_q.cell(row=r_idx, column=16).fill = green_fill
    elif item['total_pnl'] < 0:
        ws_q.cell(row=r_idx, column=16).fill = red_fill

for col in ws_q.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_q.column_dimensions[col_letter].width = max(max_len + 3, 15)

# 3. INDIVIDUAL 8 QUARTER SHEETS (SORTED BY RANK #1 TO #211)
headers_q = [
    "Rank", "Stock Symbol", "Industry Sector", "Action Status", "Strategy",
    "Position Window", "Total History Quarters", "Pred Win Rate %", "Pred Avg Return %",
    "Result Announcement Date", "Entry Date", "Exit Date", "Entry Price (₹)", "Exit Price (₹)",
    "Position Value (₹)", "Margin Capital (20%) (₹)", "Net Realised P&L (₹)",
    "Actual Realised Return %", "Win / Loss"
]

for q_info in TARGET_QUARTERS_8:
    q_code = q_info['code']
    q_label = q_info['label']
    trades = quarter_trades[q_label]
    
    trades_sorted = sorted(trades, key=lambda x: SYM_RANK_MAP.get(x['symbol'], 999))
    
    ws = wb.create_sheet(title=q_code)
    ws.views.sheetView[0].showGridLines = True
    
    ws.merge_cells("A1:S1")
    ws["A1"] = f"{q_label.upper()} - DETAILED FUTURES TRADE LOG (RANKED #1 TO #211)"
    ws["A1"].font = title_font
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35
    
    ws.append([])
    ws.append(headers_q)
    ws.row_dimensions[3].height = 25
    
    for col_num, h in enumerate(headers_q, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
    for idx, t in enumerate(trades_sorted, start=1):
        r_idx = idx + 3
        rank_num = SYM_RANK_MAP.get(t['symbol'], idx)
        
        if t['executed']:
            row_data = [
                rank_num, t['symbol'], t['sector'], t['status'], t['raw_strat'],
                t['window_str'], t['hist_q_str'], t['pred_win_rate'], t['pred_avg_ret'],
                t['result_date'], t['entry_date'], t['exit_date'], t['entry_price'],
                t['exit_price'], t['pos_val'], t['margin_req'], t['net_pnl'],
                t['act_ret_pct'] / 100.0, t['win_flag']
            ]
        else:
            row_data = [
                rank_num, t['symbol'], t['sector'], t['status'], t['raw_strat'],
                t['window_str'], t['hist_q_str'], t['pred_win_rate'], t['pred_avg_ret'],
                "N/A", "N/A", "N/A", "N/A", "N/A", "N/A", "N/A", "N/A", "N/A", "N/A"
            ]
            
        ws.append(row_data)
        ws.row_dimensions[r_idx].height = 20
        
        ws.cell(row=r_idx, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=2).alignment = Alignment(horizontal="left")
        ws.cell(row=r_idx, column=2).font = Font(bold=True)
        ws.cell(row=r_idx, column=4).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=5).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=6).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=7).alignment = Alignment(horizontal="center")
        
        ws.cell(row=r_idx, column=8).number_format = '0.00%'
        ws.cell(row=r_idx, column=9).number_format = '0.00%'
        
        for c in range(10, 13):
            ws.cell(row=r_idx, column=c).alignment = Alignment(horizontal="center")
            
        if t['executed']:
            ws.cell(row=r_idx, column=13).number_format = '₹#,##0.00'
            ws.cell(row=r_idx, column=14).number_format = '₹#,##0.00'
            ws.cell(row=r_idx, column=15).number_format = '₹#,##0.00'
            ws.cell(row=r_idx, column=16).number_format = '₹#,##0.00'
            ws.cell(row=r_idx, column=17).number_format = '₹#,##0.00'
            ws.cell(row=r_idx, column=18).number_format = '0.00%'
            ws.cell(row=r_idx, column=19).alignment = Alignment(horizontal="center")
            
            if t['net_pnl'] > 0:
                ws.cell(row=r_idx, column=17).fill = green_fill
                ws.cell(row=r_idx, column=19).fill = green_fill
            else:
                ws.cell(row=r_idx, column=17).fill = red_fill
                ws.cell(row=r_idx, column=19).fill = red_fill
        else:
            for c in range(10, 20):
                cell = ws.cell(row=r_idx, column=c)
                cell.alignment = Alignment(horizontal="center")
                cell.fill = na_fill

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

print(f"\nSaving Last 8 Quarters Master Workbook to: {OUTPUT_FILE} ...")
wb.save(OUTPUT_FILE)
print("SUCCESSFULLY SAVED LAST 8 QUARTERS DETAILED FUTURES MASTER EXCEL!")
