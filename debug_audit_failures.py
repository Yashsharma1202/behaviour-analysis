import openpyxl, sys

sys.stdout.reconfigure(errors='replace')

master_clean_path = r'D:/behaviour analysis/Nifty50_12_Quarters_Futures_Combined_Turnover_0.5PCT_Cost_Master_CLEAN.xlsx'
wb = openpyxl.load_workbook(master_clean_path, data_only=True)
ws = wb['Q3 2023-24']

print("Header Row 7:")
for c in range(1, 24):
    print(f"  Col {c:2d} ({openpyxl.utils.get_column_letter(c)}): {ws.cell(row=7, column=c).value}")

print("\nRow 8 Data:")
for c in range(1, 24):
    print(f"  Col {c:2d} ({openpyxl.utils.get_column_letter(c)}): {ws.cell(row=8, column=c).value}")
