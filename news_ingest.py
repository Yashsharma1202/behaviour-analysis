"""
news_ingest.py  —  timestamped news ingestion (foundation for the Expectation-Bar engine)
===============================================================================
Fetches recent news headlines per Nifty-50 stock from Google News RSS (free, no
API key) and stores each article with the moment it entered OUR data
(`ingested_at`, UTC). That timestamp is what makes the causality rule
(ingested_at < decision_ts) honest — so this MUST run continuously going
forward; it cannot be reconstructed for the past.

    python news_ingest.py                 # ingest all 50
    python news_ingest.py APOLLOHOSP TMPV # ingest specific symbols

Store: news.db (SQLite, stdlib). De-duped on (symbol, link).
Separate & additive — touches nothing in the main dashboard.
===============================================================================
"""
from __future__ import annotations

import datetime as dt
import re
import sqlite3
import sys
import time
import urllib.parse
import urllib.request
from html import unescape
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parent
DB = ROOT / "news.db"


def _conn():
    c = sqlite3.connect(DB)
    c.execute("""CREATE TABLE IF NOT EXISTS articles(
        symbol TEXT, title TEXT, source TEXT, link TEXT,
        published_at TEXT, ingested_at TEXT,
        UNIQUE(symbol, link))""")
    return c


def names() -> dict:
    """symbol -> company name (from the Q1 sheet if present)."""
    try:
        import openpyxl
        ws = openpyxl.load_workbook(ROOT / "Nifty50_Q1Results_BehaviourEngine.xlsx").active
        return {(ws.cell(r, 2).value or "").strip(): (ws.cell(r, 1).value or "").strip()
                for r in range(3, ws.max_row + 1) if ws.cell(r, 2).value}
    except Exception:                                        # noqa: BLE001
        return {}


def _fetch(query: str) -> list[dict]:
    url = ("https://news.google.com/rss/search?q="
           + urllib.parse.quote(query) + "&hl=en-IN&gl=IN&ceid=IN:en")
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    xml = urllib.request.urlopen(req, timeout=25).read().decode("utf-8", "replace")
    out = []
    for block in re.findall(r"<item>(.*?)</item>", xml, re.DOTALL):
        def g(tag):
            m = re.search(rf"<{tag}>(.*?)</{tag}>", block, re.DOTALL)
            return unescape(re.sub(r"<.*?>", "", m.group(1)).strip()) if m else ""
        title = g("title")
        if title:
            out.append({"title": title, "source": g("source"),
                        "link": g("link"), "published_at": g("pubDate")})
    return out


def ingest(symbols: list[str]) -> None:
    nm = names()
    c = _conn()
    now = dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")
    total_new = 0
    for sym in symbols:
        q = f'"{nm.get(sym, sym)}" stock OR results OR shares'
        try:
            items = _fetch(q)
        except Exception as e:                              # noqa: BLE001
            print(f"  ! {sym}: {e}")
            continue
        new = 0
        for it in items:
            cur = c.execute(
                "INSERT OR IGNORE INTO articles VALUES(?,?,?,?,?,?)",
                (sym, it["title"], it["source"], it["link"],
                 it["published_at"], now))
            new += cur.rowcount
        c.commit()
        total_new += new
        print(f"  {sym:12} fetched {len(items):>3}  new {new:>3}")
        time.sleep(0.25)
    c.close()
    print(f"\ningested_at={now}  |  {total_new} new article(s) stored in {DB.name}")


def articles_for(symbol: str) -> list[dict]:
    c = _conn()
    rows = c.execute(
        "SELECT title,source,link,published_at,ingested_at FROM articles "
        "WHERE symbol=? ORDER BY ingested_at", (symbol,)).fetchall()
    c.close()
    keys = ("title", "source", "link", "published_at", "ingested_at")
    return [dict(zip(keys, r)) for r in rows]


if __name__ == "__main__":
    from download_feeds import NIFTY50_FALLBACK
    args = [a for a in sys.argv[1:] if not a.startswith("-")]
    syms = args or sorted(set(NIFTY50_FALLBACK))
    print(f"Ingesting news for {len(syms)} symbol(s)…")
    ingest(syms)
