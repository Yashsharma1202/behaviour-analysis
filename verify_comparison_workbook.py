import openpyxl, pandas as pd, sys

sys.stdout.reconfigure(errors='replace')

comp_path = r'D:/behaviour analysis/Nifty50_12_Quarters_Futures_vs_Options_Comparison.xlsx'
wb = openpyxl.load_workbook(comp_path, data_only=True)

print("==========================================================================================")
print("VERIFYING FUTURES VS OPTIONS COMPARISON MASTER WORKBOOK")
print("==========================================================================================")
print("Sheet Names:", wb.sheetnames)

ws1 = wb['Summary_&_Win_Rate_Compare']
print("\nSheet 1 Summary Rows 7-20:")
for r in range(7, 21):
    vals = [ws1.cell(row=r, column=c).value for c in range(1, 11)]
    print(f"  Row {r:2d}: {vals}")

ws2 = wb['Futures_Win_Options_Loss_Audit']
print(f"\nSheet 2 Audit Journal Total Rows: {ws2.max_row}")
print("First 5 Audit Trades:")
for r in range(4, 9):
    vals = [ws2.cell(row=r, column=c).value for c in [1, 2, 3, 6, 9, 10, 11, 13, 17, 19]]
    print(f"  Trade {r-3}: {vals}")

print("==========================================================================================")
