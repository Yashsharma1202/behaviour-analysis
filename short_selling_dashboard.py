import streamlit as st
import pandas as pd
import openpyxl
from pathlib import Path
import plotly.express as px
import plotly.graph_objects as go
import urllib.request
import json
import datetime as dt

# --- PAGE SETUP ---
st.set_page_config(
    page_title="Short Selling Dashboard",
    page_icon="📉",
    layout="wide",
    initial_sidebar_state="expanded"
)

# --- CUSTOM CSS FOR EXACT MOCKUP AESTHETICS ---
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Outfit:wght@300;400;600;700&display=swap');
    
    /* Global Styles */
    html, body, [class*="css"], .stApp {
        font-family: 'Outfit', sans-serif !important;
        background-color: #0f0505 !important;
        color: #f9fafb !important;
    }
    
    /* Sidebar Overrides */
    section[data-testid="stSidebar"] {
        background-color: #180808 !important;
        border-right: 1px solid #3f1010 !important;
    }
    
    /* Hide Streamlit Default UI Elements */
    .stDeployButton {
        display: none !important;
    }
    header[data-testid="stHeader"] {
        display: none !important;
    }
    #MainMenu {
        visibility: hidden;
    }
    footer {
        visibility: hidden;
    }
    
    /* Card Styles */
    .kpi-container {
        display: grid;
        grid-template-columns: repeat(5, 1fr);
        gap: 16px;
        margin-bottom: 24px;
    }
    .card-maroon-crimson {
        background: linear-gradient(135deg, #7f1d1d 0%, #dc2626 100%);
        border-radius: 16px;
        padding: 20px;
        box-shadow: 0 8px 16px rgba(220, 38, 38, 0.15);
        color: white;
    }
    .card-red-orange {
        background: linear-gradient(135deg, #991b1b 0%, #ea580c 100%);
        border-radius: 16px;
        padding: 20px;
        box-shadow: 0 8px 16px rgba(234, 88, 12, 0.15);
        color: white;
    }
    .card-dark {
        background-color: #1a0a0a;
        border: 1px solid #331414;
        border-radius: 16px;
        padding: 20px;
        position: relative;
        box-shadow: 0 4px 10px rgba(0,0,0,0.3);
    }
    .card-title {
        font-size: 14px;
        font-weight: 300;
        opacity: 0.8;
    }
    .card-title-dark {
        font-size: 14px;
        font-weight: 400;
        color: #9ca3af;
    }
    .card-value {
        font-size: 26px;
        font-weight: 700;
        margin-top: 8px;
    }
    .badge {
        position: absolute;
        top: 15px;
        right: 15px;
        padding: 3px 8px;
        border-radius: 20px;
        font-size: 11px;
        font-weight: 600;
    }
    .badge-green {
        background-color: rgba(63, 185, 80, 0.15);
        color: #3fb950;
    }
    .badge-red {
        background-color: rgba(248, 81, 73, 0.15);
        color: #f85149;
    }
    
    /* Tabs Overrides */
    .stTabs [data-baseweb="tab-list"] {
        gap: 10px;
        background-color: #180808;
        padding: 6px;
        border-radius: 12px;
        border: 1px solid #3f1010;
    }
    .stTabs [data-baseweb="tab"] {
        padding: 10px 16px !important;
        border-radius: 8px !important;
        color: #9ca3af !important;
        background-color: transparent !important;
        border: none !important;
        font-weight: 600 !important;
    }
    .stTabs [aria-selected="true"] {
        color: white !important;
        background-color: #2b0e0e !important;
        box-shadow: 0 4px 6px rgba(0,0,0,0.4) !important;
    }
</style>
""", unsafe_allow_html=True)

# --- FILE PATHS ---
ROOT = Path("D:/behaviour analysis")
QTY_DATES_EXCEL = ROOT / "Nifty50Stocks_QtyResultDates.xlsx"
UPDATED_SHORT_EXCEL = ROOT / "UPDATED_Short.xlsx"

# --- LOAD CUSTOM SHORT PARAMETERS ---
@st.cache_data(ttl=3600)
def load_custom_params():
    params = {}
    if not UPDATED_SHORT_EXCEL.exists():
        return params
    wb = openpyxl.load_workbook(UPDATED_SHORT_EXCEL, read_only=True)
    ws = wb["Q2 2026-27_Result Behaviour"]
    for r in range(3, ws.max_row + 1):
        sym = ws.cell(r, 2).value
        sb = ws.cell(r, 5).value
        ca = ws.cell(r, 6).value
        if sym and sb is not None and ca is not None:
            params[str(sym).strip().upper()] = (int(sb), int(ca))
    wb.close()
    return params

# --- LOAD QUARTER DATES ---
@st.cache_data(ttl=3600)
def load_quarter_dates():
    if not QTY_DATES_EXCEL.exists():
        return {}
    wb = openpyxl.load_workbook(QTY_DATES_EXCEL, read_only=True)
    quarter_data = {}
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = []
        for r in range(2, ws.max_row + 1):
            comp = ws.cell(r, 1).value
            sym = ws.cell(r, 2).value
            r_dt = ws.cell(r, 3).value
            if sym and r_dt:
                rows.append({
                    "company": str(comp or ""),
                    "symbol": str(sym).strip().upper(),
                    "result_date": pd.Timestamp(r_dt).normalize()
                })
        quarter_data[sn] = rows
    wb.close()
    return quarter_data

# --- FETCH YAHOO PRICES (CACHED) ---
@st.cache_data(ttl=3600, show_spinner="Fetching stock prices from Yahoo Finance...")
def fetch_prices(symbols):
    p1 = int(dt.datetime(2023, 1, 1).timestamp())
    p2 = int((dt.datetime.now() + pd.Timedelta(days=2)).timestamp())
    headers = {"User-Agent": "Mozilla/5.0"}
    cache = {}
    for sym in symbols:
        ticker = sym.replace("&", "%26") + ".NS"
        url = f"https://query1.finance.yahoo.com/v8/finance/chart/{ticker}?period1={p1}&period2={p2}&interval=1d"
        try:
            req = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(req) as resp:
                data = json.loads(resp.read().decode())
                result = data["chart"]["result"][0]
                timestamps = result["timestamp"]
                quotes = result["indicators"]["quote"][0]
                df = pd.DataFrame({
                    "date": pd.to_datetime(timestamps, unit="s").date,
                    "close": quotes["close"]
                }).dropna()
                df["date"] = pd.to_datetime(df["date"])
                df = df.sort_values("date").reset_index(drop=True)
                cache[sym] = df
        except Exception:
            pass
    return cache

def get_trading_date(cal, result_date, offset):
    D = pd.Timestamp(result_date).normalize()
    if len(cal) == 0:
        return None
    pos = cal.searchsorted(D, side="right") - 1
    pos = max(0, pos)
    target = pos + offset
    if 0 <= target < len(cal):
        return cal[target]
    return None

# --- LOAD ALL DATA ---
params = load_custom_params()
quarter_data = load_quarter_dates()
all_syms = set(params.keys())
price_cache = fetch_prices(all_syms)

# --- SIDEBAR FILTERS ---
st.sidebar.image("https://img.icons8.com/color/120/000000/empty-filter.png", width=65)
st.sidebar.markdown("<h2 style='color:#ef4444; font-weight:700; font-size:20px; margin-top:10px;'>Shorting Dashboard</h2>", unsafe_allow_html=True)
st.sidebar.write("---")

available_quarters = sorted(list(quarter_data.keys()), reverse=True)
selected_q = st.sidebar.selectbox("Select Quarter", available_quarters)

# Sizing Model Selector
model_type = st.sidebar.radio(
    "Sizing Model",
    ["Model A: 1-Share Booked P&L", "Model B: Compounded Slots"]
)

# --- BUILD SHORT TRADES FOR SELECTED QUARTER ---
q_rows = quarter_data.get(selected_q, [])
trades_list = []
for row in q_rows:
    sym = row["symbol"]
    res_date = row["result_date"]
    if sym not in price_cache or price_cache[sym].empty:
        continue
    df = price_cache[sym]
    cal = pd.DatetimeIndex(df["date"])
    bo, so = params.get(sym, (2, 5))
    en_dt = get_trading_date(cal, res_date, -bo)
    ex_dt = get_trading_date(cal, res_date, so)
    if en_dt is None or ex_dt is None:
        continue
    p_en_rows = df[df["date"] == en_dt]
    p_ex_rows = df[df["date"] == ex_dt]
    if p_en_rows.empty or p_ex_rows.empty:
        continue
    p_en = float(p_en_rows.iloc[0]["close"])
    p_ex = float(p_ex_rows.iloc[0]["close"])
    ret_pct = ((p_en - p_ex) / p_en) * 100
    trades_list.append({
        "symbol": sym,
        "company": row["company"],
        "sector": "",
        "short_before": bo,
        "cover_after": so,
        "entry_date": en_dt,
        "exit_date": ex_dt,
        "short_entry": p_en,
        "short_cover": p_ex,
        "short_return": ret_pct,
        "short_profit": p_en - p_ex
    })

trades_list.sort(key=lambda x: x["entry_date"])

# Sector filter (load from existing data)
sectors_set = set()
for t in trades_list:
    if t["sector"]:
        sectors_set.add(t["sector"])
sectors = ["All"] + sorted(list(sectors_set))
selected_sector = st.sidebar.selectbox("Filter by Sector", sectors)

# Symbol Search
search_symbol = st.sidebar.text_input("Search Stock Symbol", "").strip().upper()

# Apply Filters
filtered_trades = trades_list.copy()
if selected_sector != "All":
    filtered_trades = [t for t in filtered_trades if t["sector"] == selected_sector]
if search_symbol:
    filtered_trades = [t for t in filtered_trades if search_symbol in t["symbol"]]

# Simulate slots
slots = []
for t in filtered_trades:
    assigned = False
    for s in slots:
        if t["entry_date"] >= s["last_exit"]:
            s["trades"].append(t)
            s["last_exit"] = t["exit_date"]
            assigned = True
            break
    if not assigned:
        slots.append({
            "id": len(slots) + 1,
            "trades": [t],
            "last_exit": t["exit_date"]
        })

# Calculate KPIs
if not slots:
    margin_utilised = 0.0
    final_value = 0.0
    net_profit = 0.0
    overall_return = 0.0
    total_reentries = 0
    win_rate = 0.0
else:
    total_reentries = sum(len(s["trades"]) - 1 for s in slots)
    wins = sum(1 for t in filtered_trades if t["short_profit"] > 0)
    win_rate = (wins / len(filtered_trades)) * 100 if filtered_trades else 0.0
    
    if model_type == "Model A: 1-Share Booked P&L":
        margin_utilised = sum(max(t["short_entry"] for t in s["trades"]) for s in slots)
        net_profit = sum(sum(t["short_profit"] for t in s["trades"]) for s in slots)
        final_value = margin_utilised + net_profit
        overall_return = (net_profit / margin_utilised) * 100 if margin_utilised > 0 else 0.0
    else:
        # Model B: Compounded Slots
        margin_utilised = 0.0
        final_value = 0.0
        for s in slots:
            cap_start = s["trades"][0]["short_entry"]
            cap_end = cap_start
            for t in s["trades"]:
                cap_end = cap_end * (1 + t["short_return"] / 100)
            margin_utilised += cap_start
            final_value += cap_end
        net_profit = final_value - margin_utilised
        overall_return = (net_profit / margin_utilised) * 100 if margin_utilised > 0 else 0.0

# --- MAIN PAGE HEADER ---
st.title("📉 Short-Selling Capital Utilisation Dashboard")
st.markdown(f"**Part 2 Strategy: Sell Before → Event → Buy After (Cover)** &middot; Quarter: `{selected_q}` &middot; Sizing: `{model_type.split(':')[0]}`")
st.write("---")

# --- KPI DISPLAY CARDS (CUSTOM HTML) ---
profit_sign = "+" if net_profit >= 0 else ""
profit_color_class = "badge-green" if net_profit >= 0 else "badge-red"
comp_sign = "+" if overall_return >= 0 else ""

html_kpi_block = f"""
<div class="kpi-container">
    <div class="card-maroon-crimson">
        <div class="card-title">Margin Utilised</div>
        <div class="card-value">₹{margin_utilised:,.2f}</div>
    </div>
    <div class="card-red-orange">
        <div class="card-title">Final Value</div>
        <div class="card-value">₹{final_value:,.2f}</div>
    </div>
    <div class="card-dark">
        <div class="card-title-dark">Net Short Profit</div>
        <div class="card-value" style="color: {'#3fb950' if net_profit >= 0 else '#f85149'};">₹{profit_sign}{net_profit:,.2f}</div>
        <span class="badge {profit_color_class}">{profit_sign}{overall_return:.2f}%</span>
    </div>
    <div class="card-dark">
        <div class="card-title-dark">Short Win Rate</div>
        <div class="card-value" style="color: #f59e0b;">{win_rate:.1f}%</div>
        <span class="badge badge-green" style="background-color:rgba(245,158,11,0.15); color:#f59e0b;">{len(filtered_trades)} Trades</span>
    </div>
    <div class="card-dark">
        <div class="card-title-dark">Re-entries</div>
        <div class="card-value">{total_reentries}</div>
        <span class="badge badge-green" style="background-color:rgba(124,58,237,0.15); color:#a78bfa;">{len(slots)} Slots</span>
    </div>
</div>
"""
st.markdown(html_kpi_block, unsafe_allow_html=True)

# --- VISUAL TABS ---
tab1, tab2, tab3, tab4, tab5 = st.tabs([
    "📂 Slots Allocation Timeline", 
    "📋 Detailed Trade Log", 
    "📈 Cumulative Short Returns",
    "🍩 Sector Short Distribution",
    "⚙️ Per-Stock Short Parameters"
])

with tab1:
    st.subheader("Capital Slots Re-entry Timeline")
    
    if not slots:
        st.info("No slots found for the selected filters.")
    else:
        # Create Gantt Chart
        gantt_data = []
        for s in slots:
            for t in s["trades"]:
                gantt_data.append({
                    "Slot ID": f"Slot {s['id']}",
                    "Stock": t["symbol"],
                    "Start": t["entry_date"],
                    "End": t["exit_date"],
                    "Profit": t["short_profit"]
                })
        df_gantt = pd.DataFrame(gantt_data)
        
        fig = px.timeline(
            df_gantt, 
            x_start="Start", 
            x_end="End", 
            y="Slot ID", 
            color="Profit",
            color_continuous_scale=["#f85149", "#3fb950"],
            title="Short Recycled Margins Slot Gantt Chart",
            hover_data=["Stock"]
        )
        fig.update_yaxes(categoryorder="category descending")
        fig.update_layout(
            paper_bgcolor="#150909",
            plot_bgcolor="#150909",
            font_color="white",
            margin=dict(l=20, r=20, t=40, b=20),
            height=400
        )
        st.plotly_chart(fig, use_container_width=True)
        
        # Expandable Cards for slots
        st.write("### Detailed Trade Flow Path per Slot")
        for s in slots:
            flow_str = " → ".join([f"**{t['symbol']}** (T-{t['short_before']}→T+{t['cover_after']}, {t['entry_date'].strftime('%d-%b')} to {t['exit_date'].strftime('%d-%b')})" for t in s["trades"]])
            slot_prof = sum(t["short_profit"] for t in s["trades"])
            prof_str = f"₹{slot_prof:+,.2f}"
            prof_color = "green" if slot_prof >= 0 else "red"
            
            with st.expander(f"Slot {s['id']} &middot; Net Short Profit: :{prof_color}[{prof_str}] &middot; Re-entries: {len(s['trades'])-1}"):
                st.markdown(flow_str)

with tab2:
    st.subheader("Chronological Trade Ledger")
    if not filtered_trades:
        st.info("No trades matched the filters.")
    else:
        df_trades_display = pd.DataFrame([{
            "Symbol": t["symbol"],
            "Company": t["company"],
            "Short Before (Days)": f"T-{t['short_before']}",
            "Cover After (Days)": f"T+{t['cover_after']}",
            "Entry Date": t["entry_date"].strftime("%Y-%m-%d"),
            "Exit Date": t["exit_date"].strftime("%Y-%m-%d"),
            "Entry Price (Short Sell)": t["short_entry"],
            "Cover Price (Buy Back)": t["short_cover"],
            "Return (%)": t["short_return"],
            "Short Profit/Loss (Rs.)": t["short_profit"]
        } for t in filtered_trades])
        
        st.dataframe(
            df_trades_display.style.format({
                "Entry Price (Short Sell)": "₹{:,.2f}",
                "Cover Price (Buy Back)": "₹{:,.2f}",
                "Return (%)": "{:+.2f}%",
                "Short Profit/Loss (Rs.)": "₹{:+,.2f}"
            }),
            use_container_width=True
        )

with tab3:
    st.subheader("Cumulative Short Returns Growth")
    if not filtered_trades:
        st.info("No trades to display return curve.")
    else:
        # Accumulate returns chronologically
        df_cum = pd.DataFrame(filtered_trades)
        df_cum = df_cum.sort_values("exit_date").reset_index(drop=True)
        df_cum["cumulative_profit"] = df_cum["short_profit"].cumsum()
        
        fig_curve = go.Figure()
        fig_curve.add_trace(go.Scatter(
            x=df_cum["exit_date"],
            y=df_cum["cumulative_profit"],
            mode='lines+markers',
            name='Cumulative Short Profit',
            line=dict(color='#ef4444', width=3, shape='spline'),
            fill='tozeroy',
            fillcolor='rgba(239, 68, 68, 0.1)'
        ))
        fig_curve.update_layout(
            paper_bgcolor="#150909",
            plot_bgcolor="#150909",
            font_color="white",
            xaxis_title="Exit Date",
            yaxis_title="Cumulative Realised Short Profit (Rs.)",
            margin=dict(l=20, r=20, t=30, b=20),
            height=400
        )
        st.plotly_chart(fig_curve, use_container_width=True)

with tab4:
    st.subheader("Short Exposure by Sector")
    if not filtered_trades:
        st.info("No trades to display sector counts.")
    else:
        df_sec = pd.DataFrame(filtered_trades)
        sector_counts = df_sec["sector"].value_counts().reset_index()
        sector_counts.columns = ["Sector", "Trades Count"]
        
        fig_pie = px.pie(
            sector_counts, 
            values="Trades Count", 
            names="Sector", 
            hole=0.62,
            color_discrete_sequence=px.colors.sequential.Reds_r
        )
        fig_pie.update_layout(
            paper_bgcolor="#150909",
            font_color="white",
            margin=dict(l=20, r=20, t=30, b=20),
            height=400
        )
        st.plotly_chart(fig_pie, use_container_width=True)

with tab5:
    st.subheader("⚙️ Per-Stock Optimized Short Parameters (from UPDATED_Short.xlsx)")
    st.markdown("Each stock has its **own unique optimal Short Before and Cover After** offset, discovered via historical grid search across all past events.")
    
    if not params:
        st.warning("UPDATED_Short.xlsx not found or empty.")
    else:
        param_rows = []
        for sym, (sb, ca) in sorted(params.items()):
            # Find this stock's performance in current quarter
            match = [t for t in filtered_trades if t["symbol"] == sym]
            ret_str = f"{match[0]['short_return']:+.2f}%" if match else "N/A"
            prof_str = f"₹{match[0]['short_profit']:+,.2f}" if match else "N/A"
            param_rows.append({
                "Symbol": sym,
                "Short Before (Candles)": f"T-{sb}",
                "Cover After (Candles)": f"T+{ca}",
                "Holding Period (Days)": sb + ca,
                f"Return in {selected_q} (%)": ret_str,
                f"Profit in {selected_q} (Rs.)": prof_str
            })
        
        df_params = pd.DataFrame(param_rows)
        st.dataframe(df_params, use_container_width=True, height=600)
