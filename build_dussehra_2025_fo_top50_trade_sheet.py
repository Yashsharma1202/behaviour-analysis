"""
TOP 50 F&O STOCKS (full 211 universe) \u2014 LAST DUSSEHRA (2025) RESULT, DETAILED
TRADE SHEET. Reconstructs each F&O stock's real 2025 Dussehra walk-forward
trade (point-in-time window/direction, chosen using only pre-2025 years --
same method as the live dashboard), with real entry/exit date + price
(Open/Close per the project's 09:20/15:15 execution convention), margin and
PnL, ranked by that single occurrence's actual result, top 50 of 211.

ADDITIVE: reads the walk-forward pickle + futures price data; writes a new report.
"""
import json, pickle, os, subprocess
from datetime import date
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

BASE = r'D:\behaviour analysis'
OUT_XLSX = os.environ.get('DUSS25TOP50_OUT_XLSX', BASE + r'\Dussehra_2025_FO_Top50_Detailed_Trade_Sheet.xlsx')
OUT_HTML = BASE + r'\Dussehra_2025_FO_Top50_Detailed_Trade_Sheet.html'
OUT_PDF = BASE + r'\Dussehra_2025_FO_Top50_Detailed_Trade_Sheet.pdf'

results = pickle.load(open(BASE + r'\scraped_parquet\_tmp_dussehra_walkforward.pkl', 'rb'))
fo = json.load(open(BASE + r'\dashboard_data\fo_stocks_211.json', encoding='utf-8'))
meta = {f['symbol']: f for f in fo}

fut = pd.read_parquet(BASE + r'\scraped_parquet\fo_futures_near_month_continuous.parquet', columns=['Date', 'Instrument', 'Open', 'Close'])
fut['Date'] = pd.to_datetime(fut['Date']).dt.date
fut_closes = {s: g.sort_values('Date')[['Date', 'Close', 'Open']].values.tolist() for s, g in fut.groupby('Instrument')}

beh = json.load(open(BASE + r'\dashboard_data\nifty_futures_holiday_behaviour.json', encoding='utf-8'))
dh = [h for h in beh['holidays'] if 'Dussehra' in h['name']][0]
dussehra_dates = sorted(date.fromisoformat(t['holiday_date']) for t in dh['trades'])
anchor_2025 = next(dt for dt in dussehra_dates if dt.year == 2025)


def lot_for(sym):
    m = meta.get(sym, {})
    return m.get('lot_size') or 500


trades = []
for sym, r in results.items():
    if r.get('oos_n', 0) < 1:
        continue
    t = next((x for x in r['oos_trades'] if x['year'] == 2025), None)
    if not t:
        continue
    a = fut_closes.get(sym)
    if not a:
        continue
    ib = -1
    for i, row in enumerate(a):
        if row[0] <= anchor_2025: ib = i
        else: break
    if ib < 0:
        continue
    bo, so, side = t['bo'], t['so'], t['side']
    en_idx, ex_idx = ib - bo, ib + so
    if en_idx < 0 or ex_idx >= len(a):
        continue
    lot = lot_for(sym)
    sign = 1 if side == 'LONG' else -1
    en_date, en_close, en_px = a[en_idx]
    ex_date, ex_close, _ = a[ex_idx]

    g = 1.0
    if en_px:
        dr0 = (en_close - en_px) / en_px
        if abs(dr0) <= 0.20: g *= (1 + dr0)
    for k in range(en_idx + 1, ex_idx + 1):
        p0, p1 = a[k - 1][1], a[k][1]
        if p0:
            dr = (p1 - p0) / p0
            if abs(dr) <= 0.20: g *= (1 + dr)
    ret_pct = (g - 1) * 100 * sign
    pnl = round(ret_pct / 100 * en_px * lot, 2)
    margin = round(en_px * lot * 0.20, 2)
    m = meta.get(sym, {})
    trades.append(dict(symbol=sym, name=m.get('name', sym), sector=m.get('sector', ''), direction=side,
                        window=f"T-{bo} to T+{so}", entry_date=en_date, entry_px=round(en_px, 2),
                        exit_date=ex_date, exit_px=round(ex_close, 2), lot=lot, margin=margin,
                        ret_pct=round(ret_pct, 2), pnl=pnl, low_confidence=r.get('low_confidence', False),
                        oos_n=r['oos_n'], oos_wr=r.get('oos_wr'), source=r.get('source')))

print(f"Reconstructed real 2025 Dussehra trades for {len(trades)} of 211 F&O stocks.")
ranked = sorted(trades, key=lambda t: -t['pnl'])
top50 = ranked[:50]

NAVY = PatternFill('solid', fgColor='1A202C'); WHITE = Font(bold=True, color='FFFFFF', size=9)
TITLE = Font(bold=True, color='1A202C', size=13); SUB = Font(color='4A5568', size=10, italic=True)
GF = PatternFill('solid', fgColor='E6F4EA'); RF = PatternFill('solid', fgColor='FCE8E6')
AM = PatternFill('solid', fgColor='FEF3C7'); LBL = PatternFill('solid', fgColor='F1F5F9')
GOOD = Font(color='137333', bold=True); BAD = Font(color='C5221F', bold=True); BOLD = Font(bold=True)
THIN = Border(*[Side(style='thin', color='D9E2EC')] * 4)
CEN = Alignment(horizontal='center'); RIGHT = Alignment(horizontal='right'); WRAP = Alignment(wrap_text=True, vertical='top')

wb = Workbook()

# ============ Sheet 1: Trade Log (Top 50) ============
tl = wb.active; tl.title = 'Top50_Trade_Log_2025'
tl.merge_cells('A1:O1')
tl.cell(1, 1, 'TOP 50 F&O STOCKS \u2014 LAST DUSSEHRA (2-Oct-2025) RESULT, DETAILED TRADE LOG').font = TITLE
tl.merge_cells('A2:O2')
tl.cell(2, 1, f'Real 2025 walk-forward trade (point-in-time, chosen using only pre-2025 years) for each of {len(trades)} F&O stocks with usable data; ranked by PnL, top 50 shown. Futures-sourced prices (no Friday data-quality defect).').font = SUB
hdr = ['Rank', 'Symbol', 'Company', 'Sector', 'Direction', 'Window', 'Entry Date', 'Entry Price (\u20b9)',
       'Exit Date', 'Exit Price (\u20b9)', 'Lot', 'Margin (\u20b9)', 'Return %', 'PnL (\u20b9)', 'Confidence']
for j, h in enumerate(hdr, 1):
    c = tl.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
row = 5
for i, t in enumerate(top50, 1):
    vals = [i, t['symbol'], t['name'], t['sector'], t['direction'], t['window'],
            t['entry_date'].strftime('%d-%b-%Y'), t['entry_px'], t['exit_date'].strftime('%d-%b-%Y'), t['exit_px'],
            t['lot'], t['margin'], t['ret_pct'], t['pnl'], ('LOW (n<3)' if t['low_confidence'] else 'OK')]
    fill = AM if t['low_confidence'] else (GF if t['pnl'] >= 0 else RF)
    for j, v in enumerate(vals, 1):
        c = tl.cell(row, j, v); c.border = THIN; c.fill = fill
        if j in (1, 5, 7, 9, 15): c.alignment = CEN
        if j in (8, 10, 12, 13, 14): c.alignment = RIGHT
        if j == 13: c.font = GOOD if v >= 0 else BAD
        if j == 14: c.font = GOOD if v >= 0 else BAD
        if j == 15 and t['low_confidence']: c.font = Font(color='92400E', bold=True, size=9)
    row += 1
for j, w in enumerate([6, 12, 24, 20, 10, 13, 15, 13, 15, 13, 7, 13, 10, 13, 12], 1):
    tl.column_dimensions[get_column_letter(j)].width = w
tl.freeze_panes = 'A5'

# ============ Sheet 2: Overview ============
ov = wb.create_sheet('Overview', 0)
ov.merge_cells('A1:B1')
ov.cell(1, 1, 'TOP 50 F&O STOCKS \u2014 LAST DUSSEHRA RESULT \u2014 OVERVIEW').font = TITLE
wins = sum(1 for t in top50 if t['pnl'] > 0)
win_rate = round(100 * wins / len(top50), 1)
total_pnl = round(sum(t['pnl'] for t in top50), 2)
total_margin = round(sum(t['margin'] for t in top50), 2)
n_long = sum(1 for t in top50 if t['direction'] == 'LONG'); n_short = len(top50) - n_long
n_low = sum(1 for t in top50 if t['low_confidence'])
best = top50[0]; worst = top50[-1]
notes = [
    ('Universe & ranking', f'Full 211 F&O universe; {len(trades)} stocks had usable 2025 Dussehra data. Top 50 shown, ranked by that single occurrence\u2019s actual PnL (1 lot).'),
    ('What "last Dussehra" means here', 'The real 2-Oct-2025 Dussehra trade for each stock \u2014 window and direction were picked using ONLY years strictly before 2025 (walk-forward, no hindsight), exactly as the live dashboard does for the upcoming 2026 trade, just one year earlier.'),
    ('Data source', 'Futures prices (2019-2026) \u2014 the cash-equity file was found missing 98.3% of all Fridays across its 26yr history and is not used anywhere in this report.'),
    ('This occurrence\u2019s result (top 50)', f'{wins}/50 profitable ({win_rate}%). Total PnL \u20b9{total_pnl:,.0f} on \u20b9{total_margin:,.0f} margin (1 lot each) \u2014 return on margin {(total_pnl/total_margin*100 if total_margin else 0):+.2f}%. {n_long} LONG / {n_short} SHORT.'),
    ('Confidence', f'{n_low} of the top 50 are LOW confidence (fewer than 3 prior years available to pick their window) \u2014 amber highlighted. A single profitable occurrence is one data point, not a validated edge.'),
    ('Best / Worst (within top 50)', f'Best: {best["symbol"]} ({best["pnl"]:+,.0f} \u20b9, {best["direction"]}). Worst shown: {worst["symbol"]} ({worst["pnl"]:+,.0f} \u20b9, {worst["direction"]}).'),
    ('CAVEAT', 'Ranking by a single year\u2019s PnL favours whichever stocks happened to move most in 2025 \u2014 it is not the same as a multi-year validated edge (see the companion FO211_Dussehra_WalkForward_Ranked_Report for the across-years view). Treat this as "what happened last time", not "what to expect next time".'),
]
r0 = 3
for lbl, txt in notes:
    a = ov.cell(r0, 1, lbl); a.font = BOLD; a.fill = AM; a.border = THIN; a.alignment = WRAP
    b = ov.cell(r0, 2, txt); b.alignment = WRAP; b.border = THIN
    ov.row_dimensions[r0].height = 65; r0 += 1
ov.column_dimensions['A'].width = 24; ov.column_dimensions['B'].width = 105

wb.save(OUT_XLSX)
print('Saved Excel:', OUT_XLSX)
print(f"Top50: win rate {win_rate}% ({wins}/50) | Total PnL Rs{total_pnl:,.0f} | Margin Rs{total_margin:,.0f} | {n_long} LONG / {n_short} SHORT | {n_low} low-confidence")

# ============ PDF ============
def esc(s): return str(s).replace('&', '&amp;').replace('<', '&lt;')
def pdf_row(rank, t):
    low = t['low_confidence']
    rowstyle = ' style="background:#fffbeb;"' if low else ''
    conf = '<span style="color:#b45309;font-weight:700;">LOW</span>' if low else 'OK'
    return (f'<tr{rowstyle}><td>{rank}</td><td><b>{esc(t["symbol"])}</b></td><td>{esc(t["name"])[:22]}</td>'
            f'<td class="c">{t["direction"]}</td><td class="c">{esc(t["window"])}</td>'
            f'<td class="r {"pos" if t["ret_pct"]>=0 else "neg"}">{t["ret_pct"]:+.2f}%</td>'
            f'<td class="r {"pos" if t["pnl"]>=0 else "neg"}">{t["pnl"]:+,.0f}</td><td class="c">{conf}</td></tr>')
rows_all = ''.join(pdf_row(i + 1, t) for i, t in enumerate(top50))

html = f"""<!doctype html><html><head><meta charset="utf-8">
<title>Top 50 F&amp;O Last Dussehra Result</title>
<style>
 *{{box-sizing:border-box}} body{{font-family:'Segoe UI',Arial,sans-serif;margin:0;color:#0f172a;background:#fff}}
 .wrap{{max-width:1050px;margin:0 auto;padding:28px 36px}}
 h1{{font-size:22px;margin:0 0 3px}} .sub{{color:#64748b;font-size:12px;margin-bottom:16px}}
 h2{{font-size:15px;margin:20px 0 8px;border-left:4px solid #0f766e;padding-left:9px}}
 table{{width:100%;border-collapse:collapse;font-size:11px}}
 th,td{{padding:5px 7px;text-align:right;border-bottom:1px solid #eef2f7}} th{{background:#f1f5f9;font-size:9.5px;text-transform:uppercase;color:#475569}}
 td.c,th.c{{text-align:center}} .pos{{color:#059669;font-weight:700}} .neg{{color:#dc2626;font-weight:700}}
 .cards{{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;margin:10px 0}}
 .mc{{background:#fff;border:1px solid #e2e8f0;border-radius:8px;padding:10px 12px}}
 .mcl{{font-size:9.5px;text-transform:uppercase;color:#64748b;font-weight:700}} .mcv{{font-size:19px;font-weight:800}}
 .note{{background:#fff7ed;border:1px solid #fed7aa;border-radius:8px;padding:11px 14px;font-size:11px;color:#7c2d12;margin-top:10px;line-height:1.5}}
 @page{{size:A4;margin:12mm}}
</style></head><body><div class="wrap">
 <h1>Top 50 F&amp;O Stocks &mdash; Last Dussehra (2-Oct-2025) Result</h1>
 <div class="sub">Real point-in-time trade per stock, walk-forward (no hindsight) &bull; {len(trades)} of 211 F&amp;O stocks had usable data &bull; top 50 by PnL shown</div>
 <div class="cards">
  <div class="mc"><div class="mcl">Win Rate</div><div class="mcv pos">{win_rate}%</div></div>
  <div class="mc"><div class="mcl">Total PnL</div><div class="mcv {"pos" if total_pnl>=0 else "neg"}">\u20b9{total_pnl:,.0f}</div></div>
  <div class="mc"><div class="mcl">Margin Deployed</div><div class="mcv">\u20b9{total_margin:,.0f}</div></div>
  <div class="mc"><div class="mcl">LONG / SHORT</div><div class="mcv">{n_long} / {n_short}</div></div>
 </div>
 <h2>Top 50 \u2014 Ranked by PnL</h2>
 <table><thead><tr><th>Rank</th><th class="c">Symbol</th><th class="c">Company</th><th class="c">Dir</th><th class="c">Window</th><th>Return</th><th>PnL (\u20b9)</th><th class="c">Conf.</th></tr></thead>
 <tbody>{rows_all}</tbody></table>
 <div class="note"><b>Caveat:</b> ranked by this ONE occurrence's PnL \u2014 that favours whichever stocks happened to move most in 2025, not necessarily a validated multi-year edge. See the companion FO211_Dussehra_WalkForward_Ranked_Report for the across-years ranking.</div>
</div></body></html>"""
open(OUT_HTML, 'w', encoding='utf-8').write(html)
print('Saved HTML:', OUT_HTML)

for exe in [r'C:\Program Files\Google\Chrome\Application\chrome.exe', r'C:\Program Files (x86)\Google\Chrome\Application\chrome.exe']:
    if os.path.exists(exe):
        try:
            subprocess.run([exe, '--headless=new', '--disable-gpu', '--no-pdf-header-footer',
                            '--print-to-pdf=' + OUT_PDF, 'file:///' + OUT_HTML.replace('\\', '/')], timeout=60, capture_output=True)
            print('Saved PDF:', OUT_PDF, 'exists:', os.path.exists(OUT_PDF))
        except Exception as e:
            print('PDF render failed:', e)
        break
