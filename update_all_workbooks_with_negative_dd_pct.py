import openpyxl, pandas as pd, numpy as np, pathlib, sys

sys.stdout.reconfigure(errors='replace')

BASE_DIR = pathlib.Path('D:/behaviour analysis')
FUT_V13 = BASE_DIR / 'Nifty50_12_Quarters_Futures_OI_Master_v13.xlsx'

wb_fut = openpyxl.load_workbook(FUT_V13)

# Update Exec Summary Sheet
ws_exec = wb_fut['Exec_12Q_Combined_Summary']
for row in range(8, 20):
    val_ddp = ws_exec.cell(row=row, column=11).value
    if val_ddp is not None and type(val_ddp) in (int, float) and val_ddp > 0:
        ws_exec.cell(row=row, column=11, value=-val_ddp)
        
ws_conc = wb_fut['Concurrent_Trades_Analysis']
for row in range(5, 17):
    v8 = ws_conc.cell(row=row, column=8).value
    if v8 is not None and type(v8) in (int, float) and v8 > 0:
        ws_conc.cell(row=row, column=8, value=-v8)
        
    v10 = ws_conc.cell(row=row, column=10).value
    if v10 is not None and type(v10) in (int, float) and v10 > 0:
        ws_conc.cell(row=row, column=10, value=-v10)

# Update Individual Quarter Sheets Card Row 4
quarters_order = [
    'Q3 2023-24', 'Q4 2023-24', 'Q1 2024-25', 'Q2 2024-25',
    'Q3 2024-25', 'Q4 2024-25', 'Q1 2025-26', 'Q2 2025-26',
    'Q3 2025-26', 'Q4 2025-26', 'Q1 2026-27', 'Q2 2026-27'
]

for qtr in quarters_order:
    if qtr in wb_fut.sheetnames:
        ws_q = wb_fut[qtr]
        v_ddp = ws_q.cell(row=4, column=11).value
        if v_ddp is not None and type(v_ddp) in (int, float) and v_ddp > 0:
            ws_q.cell(row=4, column=11, value=-v_ddp)

wb_fut.save(FUT_V13)
print("Updated Futures Master v13 with Negative Drawdown Percentages!")
