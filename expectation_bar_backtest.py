"""
expectation_bar_backtest.py  —  PROXY study: does the bar relate to the move?
===============================================================================
STRICT causality needs ingested_at, which we can't reconstruct for the past.
This is a PROXY: it uses each article's published_at < decision_ts as the
pre-event filter, scores a bar with the same lexicons, then correlates it with
the REALISED Q1 return already in the tracker.

Treat the numbers as INDICATIVE, not proof — published_at can be re-dated by
syndication, coverage is recency-biased, and the sample is one quarter. The
real test is forward (score now via ingested_at, measure the move after).

    python expectation_bar_backtest.py
===============================================================================
"""
from __future__ import annotations
import datetime as dt
import re
import sys
from email.utils import parsedate_to_datetime
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
import expectation_bar as EB
import news_ingest

ROOT = Path(__file__).resolve().parent
XLSX = ROOT / "Nifty50_Q1Results_BehaviourEngine_UPDATED_06Aug_fixed.xlsx"


def realised_events():
    ws = openpyxl.load_workbook(XLSX).active
    out = []
    for r in range(3, ws.max_row + 1):
        sym = (ws.cell(r, 2).value or "").strip()
        D = ws.cell(r, 4).value
        o = ws.cell(r, 15).value
        if isinstance(D, dt.datetime) and isinstance(o, (int, float)):
            out.append((sym, D, float(o)))
    return out


def proxy_bar(symbol, decision_ts):
    """Same rules as the engine but causality via published_at (PROXY)."""
    nm = EB.names(); name = nm.get(symbol, symbol)
    num = den = specific = 0.0
    for a in news_ingest.articles_for(symbol):
        try:
            pub = parsedate_to_datetime(a["published_at"]).replace(tzinfo=None)
        except Exception:                                  # noqa: BLE001
            continue
        if pub >= decision_ts:                             # not pre-event
            continue
        if EB.REACTIVE.search(a["title"]):                 # exclude reactive
            continue
        if EB._attribution(symbol, name, a["title"]) != "high":
            continue
        specific += 1
        s = (1 if EB.POS.search(a["title"]) else 0) - (1 if EB.NEG.search(a["title"]) else 0)
        if s:
            days = max(0.0, (decision_ts - pub).total_seconds() / 86400)
            w = EB._src_weight(a["source"]) * max(0.3, 1 - days / 30.0)
            num += s * w; den += w
    if specific < 2 or den == 0:
        return None, int(specific)
    return round(max(-1, min(1, num / den)), 3), int(specific)


def compute(fetch_missing: bool = True) -> dict:
    """Run the proxy backtest and return structured results (for host/CLI)."""
    events = realised_events()
    if fetch_missing:
        missing = sorted({s for s, _, _ in events if not news_ingest.articles_for(s)})
        if missing:
            news_ingest.ingest(missing)

    rows = []
    for sym, D, ret in events:
        bar, n = proxy_bar(sym, D)
        rows.append({"symbol": sym, "result": D.date().isoformat(),
                     "bar": bar, "arts": n, "realised": round(ret, 2)})
    rows.sort(key=lambda x: (x["bar"] is None, -(x["bar"] or 0)))

    def avg(xs): return round(sum(xs) / len(xs), 2) if xs else None
    hi = [r["realised"] for r in rows if r["bar"] is not None and r["bar"] >= 0.34]
    md = [r["realised"] for r in rows if r["bar"] is not None and -0.34 < r["bar"] < 0.34]
    lo = [r["realised"] for r in rows if r["bar"] is not None and r["bar"] <= -0.34]
    nul = [r["realised"] for r in rows if r["bar"] is None]

    pairs = [(r["bar"], r["realised"]) for r in rows if r["bar"] is not None]
    corr = None
    if len(pairs) >= 3:
        import statistics as st
        bs = [p[0] for p in pairs]; rs = [p[1] for p in pairs]
        mb, mr = st.mean(bs), st.mean(rs)
        cov = sum((b - mb) * (r - mr) for b, r in pairs)
        db = sum((b - mb) ** 2 for b in bs) ** .5
        dr = sum((r - mr) ** 2 for r in rs) ** .5
        corr = round(cov / (db * dr), 3) if db and dr else None

    return {
        "rows": rows,
        "buckets": {"high": {"n": len(hi), "avg": avg(hi)},
                    "neutral": {"n": len(md), "avg": avg(md)},
                    "low": {"n": len(lo), "avg": avg(lo)},
                    "null": {"n": len(nul), "avg": avg(nul)}},
        "correlation": corr, "n": len(pairs),
        "note": ("PROXY (published_at, not ingested_at) · one quarter · headline-based bar "
                 "saturates → treat as indicative, not proof."),
    }


def main():
    d = compute()
    print(f"\n{'SYM':11}{'result':>11}{'bar':>7}{'arts':>5}{'realised%':>10}")
    for r in d["rows"]:
        bs = "null" if r["bar"] is None else f"{r['bar']:+.2f}"
        print(f"{r['symbol']:11}{r['result']:>11}{bs:>7}{r['arts']:>5}{r['realised']:>+10.2f}")
    b = d["buckets"]
    print("\n── average realised move by bar bucket (PROXY) ──")
    for k, lab in [("high", "HIGH bar (>= +0.34)"), ("neutral", "NEUTRAL (-0.34..+0.34)"),
                   ("low", "LOW bar (<= -0.34)"), ("null", "NO COVERAGE (null)")]:
        print(f"  {lab:24} n={b[k]['n']:>2}  avg move {b[k]['avg']}")
    print(f"\n  correlation(bar, realised move) = {d['correlation']}  over n={d['n']}")
    print("  (hypothesis: HIGH bar → weaker move; a NEGATIVE corr would support 'sell the news')")


if __name__ == "__main__":
    main()
