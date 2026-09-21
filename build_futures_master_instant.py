import pandas as pd
import pathlib
import sys
import time
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(errors='replace')

t0 = time.time()

# ------------------- Paths -------------------
BASE_DIR = pathlib.Path('D:/behaviour analysis')
OPTIONS_MASTER_PATH = BASE_DIR / 'Nifty50_12_Quarters_Options_OI_Master_v4.xlsx'
RESULT_DATES_PATH = BASE_DIR / 'Nifty50Stocks_QtyResultDates.xlsx'
FUTURES_MASTER_EXISTING = BASE_DIR / 'Futures_Master_Only.xlsx'

OUTPUT_FUTURES_ONLY = BASE_DIR / 'Futures_Master_Only.xlsx'
OUTPUT_FUTURES_MASTER = BASE_DIR / 'Nifty50_12_Quarters_Futures_OI_Master.xlsx'

# ------------------- 1. Load Existing Daily Futures Data -------------------
print("Loading existing daily futures master data...", flush=True)
try:
    df_daily = pd.read_excel(FUTURES_MASTER_EXISTING, sheet_name='Futures_Master', engine='openpyxl')
    print(f"Loaded {len(df_daily)} daily futures rows in {time.time() - t0:.1f}s", flush=True)
except Exception as e:
    print(f"Error loading existing futures master: {e}", flush=True)
    sys.exit(1)

# Ensure Date column is string YYYY-MM-DD
df_daily['Date'] = pd.to_datetime(df_daily['Date']).dt.strftime('%Y-%m-%d')
df_daily['Symbol'] = df_daily['Symbol'].astype(str).str.strip()

# ------------------- 2. Load Result Dates -------------------
print("Loading Quarterly Result Dates...", flush=True)
result_date_map = {} # (symbol, quarter_str) -> date_str

if RESULT_DATES_PATH.exists():
    try:
        xl_rd = pd.ExcelFile(RESULT_DATES_PATH, engine='openpyxl')
        for sheet in xl_rd.sheet_names:
            df_rd = xl_rd.parse(sheet)
            sym_col = next((c for c in df_rd.columns if 'SYM' in str(c).upper()), None)
            date_col = next((c for c in df_rd.columns if 'DATE' in str(c).upper() or 'REPORTING' in str(c).upper()), None)
            if sym_col and date_col:
                for _, r in df_rd.iterrows():
                    sym = str(r[sym_col]).strip()
                    dt_val = r[date_col]
                    if pd.notnull(dt_val):
                        dt_str = pd.to_datetime(dt_val).strftime('%Y-%m-%d')
                        result_date_map[(sym, sheet)] = dt_str
    except Exception as e:
        print(f"Error loading result dates: {e}", flush=True)

print(f"Loaded {len(result_date_map)} quarterly result date mappings.", flush=True)

# ------------------- 3. Load Trades from Options Master -------------------
print("Loading Trades from Options Master...", flush=True)
try:
    df_trades = pd.read_excel(OPTIONS_MASTER_PATH, sheet_name='All_12Q_Options_Trades', engine='openpyxl')
except Exception as e:
    print(f"Error reading trades: {e}", flush=True)
    sys.exit(1)

print(f"Loaded {len(df_trades)} trade records.", flush=True)

# Quick lookup for price from df_daily
price_lookup = {}
for _, r in df_daily.iterrows():
    price_lookup[(r['Symbol'], str(r['Date']))] = r.iloc[2] # Close price col

futures_trades = []
trade_lookup = {} # (Symbol, Date_str) -> trade dict

for idx, row in df_trades.iterrows():
    sym = str(row['Symbol']).strip()
    entry_dt = pd.to_datetime(row['Entry Date']).strftime('%Y-%m-%d')
    exit_dt = pd.to_datetime(row['Exit Date']).strftime('%Y-%m-%d')
    orig_strat = str(row['Strategy']).strip()
    strat = 'FUTURE LONG' if 'LONG' in orig_strat.upper() else 'FUTURE SHORT'
    direction = 1 if 'LONG' in strat else -1
    lot_size = row.get('Lot Size (Qty)', 1)
    if pd.isnull(lot_size) or lot_size == 0:
        lot_size = 1

    entry_price = price_lookup.get((sym, entry_dt), 0.0)
    exit_price = price_lookup.get((sym, exit_dt), 0.0)
    
    booked_pnl = direction * (exit_price - entry_price) * lot_size
    qtr = str(row['Quarter']).strip()
    result_dt = result_date_map.get((sym, qtr), '')
    
    trade_info = {
        'Quarter': qtr,
        'FY': row.get('FY', ''),
        'Trade No': row.get('Trade No', idx + 1),
        'Symbol': sym,
        'Company Name': row.get('Company Name', ''),
        'Sector': row.get('Sector', ''),
        'Strategy': strat,
        'Futures Expiry Date': '',
        'Entry Date': entry_dt,
        'Exit Date': exit_dt,
        'Entry Futures Price (₹)': entry_price,
        'Exit Futures Price (₹)': exit_price,
        'Lot Size (Qty)': lot_size,
        'Booked Futures P&L (₹)': booked_pnl,
        'Expected Return (%)': row.get('Expected Return (%)', 0.0),
        'Return We Get (%)': row.get('Return We Get (%)', 0.0),
        'Quarterly Result Date': result_dt,
        'Re-entry Type': row.get('Re-entry Type', 'First Entry'),
        'Data Query Status': row.get('Data Query Status', 'Success')
    }
    futures_trades.append(trade_info)
    trade_lookup[(sym, entry_dt)] = trade_info
    trade_lookup[(sym, exit_dt)] = trade_info

df_fut_trades = pd.DataFrame(futures_trades)
print(f"Processed {len(df_fut_trades)} Futures trades.", flush=True)

# ------------------- 4. Populate Daily Trades Columns -------------------
print("Populating daily trade information...", flush=True)
positions = []
entry_dates = []
exit_dates = []
exp_returns = []
ret_we_gets = []
booked_pnls = []
result_dates = []

for _, r in df_daily.iterrows():
    sym = r['Symbol']
    dt = str(r['Date'])
    t_match = trade_lookup.get((sym, dt))
    if t_match:
        positions.append(t_match['Strategy'])
        entry_dates.append(t_match['Entry Date'])
        exit_dates.append(t_match['Exit Date'])
        exp_returns.append(t_match['Expected Return (%)'])
        ret_we_gets.append(t_match['Return We Get (%)'])
        booked_pnls.append(t_match['Booked Futures P&L (₹)'])
        result_dates.append(t_match['Quarterly Result Date'])
    else:
        positions.append('')
        entry_dates.append('')
        exit_dates.append('')
        exp_returns.append('')
        ret_we_gets.append('')
        booked_pnls.append('')
        result_dates.append('')

df_daily['Position'] = positions
df_daily['Entry Date'] = entry_dates
df_daily['Exit Date'] = exit_dates
df_daily['Expected Return (%)'] = exp_returns
df_daily['Return We Get (%)'] = ret_we_gets
df_daily['Booked P&L (₹)'] = booked_pnls
df_daily['Quarterly Result Date'] = result_dates

# ------------------- 5. Safe Fallback Saver -------------------
def save_fast_with_fallback(out_path):
    t_s = time.time()
    target_file = out_path
    
    # Try primary target, fall back if locked
    for attempt in range(1, 5):
        try:
            print(f"Writing {target_file.name}...", flush=True)
            with pd.ExcelWriter(target_file, engine='openpyxl') as writer:
                df_daily.to_excel(writer, sheet_name='Futures_Master', index=False)
                df_fut_trades.to_excel(writer, sheet_name='All_12Q_Futures_Trades', index=False)

            # Styling header
            wb = load_workbook(target_file)
            header_font = Font(bold=True, color='FFFFFFFF')
            fill_teal = PatternFill(start_color='FF006688', end_color='FF006688', fill_type='solid')
            fill_navy = PatternFill(start_color='FF1F497D', end_color='FF1F497D', fill_type='solid')

            ws_daily = wb['Futures_Master']
            ws_daily.freeze_panes = 'A2'
            for cell in ws_daily[1]:
                cell.font = header_font
                cell.fill = fill_teal
                cell.alignment = Alignment(horizontal='center')
            for c_idx in range(1, len(df_daily.columns) + 1):
                col_name = str(df_daily.columns[c_idx - 1])
                ws_daily.column_dimensions[get_column_letter(c_idx)].width = max(len(col_name) + 4, 14)

            ws_trades = wb['All_12Q_Futures_Trades']
            ws_trades.freeze_panes = 'A2'
            for cell in ws_trades[1]:
                cell.font = header_font
                cell.fill = fill_navy
                cell.alignment = Alignment(horizontal='center')
            for c_idx in range(1, len(df_fut_trades.columns) + 1):
                col_name = str(df_fut_trades.columns[c_idx - 1])
                ws_trades.column_dimensions[get_column_letter(c_idx)].width = max(len(col_name) + 4, 14)

            wb.save(target_file)
            print(f"✅ Successfully saved {target_file.name} in {time.time() - t_s:.1f}s", flush=True)
            return
        except PermissionError:
            target_file = out_path.with_name(f"{out_path.stem}_v{attempt+1}{out_path.suffix}")
            print(f"⚠️ {out_path.name} is locked/open in Excel. Saving as {target_file.name}...", flush=True)

save_fast_with_fallback(OUTPUT_FUTURES_MASTER)
save_fast_with_fallback(OUTPUT_FUTURES_ONLY)

print(f"🎉 INSTANT SUCCESS in total {time.time() - t0:.1f}s!", flush=True)
