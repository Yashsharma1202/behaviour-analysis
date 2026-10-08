import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

a = pd.read_parquet('all_stocks_daily_ohlc.parquet')
last = a.sort_values('Date').groupby('Symbol').tail(1)[['Symbol', 'Date', 'Close']]
last.columns = ['Symbol', 'LastDate', 'LastClose']
lots = pd.read_csv('lot_sizes.csv')
df = lots.merge(last, on='Symbol', how='left')
df['NotionalPerLot'] = (df.LotSize * df.LastClose).round(0)
df['IndicativeMargin_20pct'] = (0.20 * df.NotionalPerLot).round(0)
df['LotStatus'] = df.LotSize.apply(lambda x: 'Local master (unverified vs NSE)' if pd.notna(x) else 'MISSING')
df = df.sort_values('Symbol').reset_index(drop=True)

wb = Workbook(); ws = wb.active; ws.title = 'Lot_Margin'
hdr = ['Symbol', 'LotSize', 'LotStatus', 'LastDate', 'LastClose', 'NotionalPerLot', 'IndicativeMargin_20pct']
ws.append(hdr)
for r in df[hdr].itertuples(index=False):
    ws.append([None if pd.isna(v) else (v.date() if hasattr(v, 'date') else v) for v in r])
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78')
    c.alignment = Alignment(horizontal='center', wrap_text=True)
for col, w in zip('ABCDEFG', [14, 10, 34, 12, 12, 16, 22]):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    row[4].number_format = '#,##0.00'; row[5].number_format = '#,##0'; row[6].number_format = '#,##0'
ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions

notes = wb.create_sheet('Notes')
for line in [
    'Lot size & margin report - 213 stocks (Stock_Data_213)',
    '',
    'LotSize: from dashboard_data/fo_stocks_211.json. NOT verified against NSE fo_mktlots.csv (NSE unreachable).',
    'LastClose: latest Yahoo Finance daily Close (unadjusted for dividends).',
    'NotionalPerLot = LotSize x LastClose.',
    'IndicativeMargin_20pct = 20% of NotionalPerLot. This is NOT the exchange SPAN+exposure margin; it is a rough proxy.',
    'Stocks with LotStatus MISSING have no lot size and no margin.',
    'The old margin_20pct in fo_stocks_211.json is not used: 161 of 211 rows use a placeholder price of 1000.',
]:
    notes.append([line])
notes.column_dimensions['A'].width = 120
wb.save('Lot_Margin_Report_213.xlsx')
print(df[['Symbol', 'LotSize', 'LastClose', 'NotionalPerLot', 'IndicativeMargin_20pct']].head(5).to_string())
print('missing lots:', df.LotSize.isna().sum(), '| rows:', len(df))
