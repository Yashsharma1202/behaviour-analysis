"""
F&O OCTOBER SEASONALITY REPORT (Phase 1: single month) \u2014 split-robust.
Builds the Excel deliverable from the validated split-robust results computed by
fo_october_seasonality.py (run that first to refresh the intermediate parquet/csv).

Concept: for every F&O stock, test entering on each trading-day-of-October and
holding to the October month-end close; find which entry day has the strongest,
most ROBUST (risk-adjusted, not just highest raw mean) historical return, with a
full year-by-year comparison. Corporate-action (split/bonus) days are excluded
from return compounding \u2014 see Overview_Assumptions sheet.
ADDITIVE: reads the intermediate results; writes a new workbook.
"""
import pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

import os, sys
BASE = r'D:\behaviour analysis'
OUT = os.environ.get('FO_OCT_OUT', BASE + r'\FO_October_Seasonality_Report.xlsx')

res = pd.read_parquet(BASE + r'\scraped_parquet\_tmp_oct_seasonality_results.parquet')

# ---- per-stock best day (computed here, not a separate script, so the
# pipeline is just: run fo_october_seasonality.py, then this file) ----
_rows = []
for sym, g in res.groupby('Instrument'):
    piv = g.pivot_table(index='DayIdx', columns='Year', values='Ret', aggfunc='mean')
    nyears = piv.shape[1]
    if nyears < 4:
        continue
    mean_, std_ = piv.mean(axis=1), piv.std(axis=1)
    wr = (piv > 0).mean(axis=1) * 100
    sharpe = mean_ / std_
    bd = sharpe.idxmax()
    _rows.append(dict(sym=sym, nyears=nyears, best_day=bd, best_mean=mean_[bd],
                      best_median=piv.loc[bd].median(), best_wr=wr[bd],
                      best_sharpe=sharpe[bd], best_std=std_[bd]))
perstock = pd.DataFrame(_rows)

basket_yr = res.groupby(['Year', 'DayIdx'])['Ret'].mean().reset_index()
pivot = basket_yr.pivot(index='DayIdx', columns='Year', values='Ret').sort_index()
avg_hold = res.groupby('DayIdx')['HoldDays'].mean()
stats = pivot.agg(['mean', 'median', 'std'], axis=1)
stats['win_rate'] = (pivot > 0).mean(axis=1) * 100
ex2019 = pivot.drop(columns=2019, errors='ignore')
stats['mean_ex2019'] = ex2019.mean(axis=1)
stats['win_rate_ex2019'] = (ex2019 > 0).mean(axis=1) * 100
stats['avg_hold_days'] = avg_hold.reindex(stats.index)
stats['sharpe_like'] = stats['mean'] / stats['std']
YEARS = sorted(pivot.columns)

# calendar reference: for each year, map DayIdx -> date using RELIANCE (long, liquid history)
d = pd.read_parquet(BASE + r'\scraped_parquet\fo_futures_near_month_continuous.parquet')
d['Date'] = pd.to_datetime(d['Date'])
rel = d[(d['Instrument'] == 'RELIANCE') & (d['Date'].dt.month == 10)].copy()
rel['Year'] = rel['Date'].dt.year
rel = rel.sort_values(['Year', 'Date'])
rel['DayIdx'] = rel.groupby('Year').cumcount() + 1
cal = rel.pivot(index='DayIdx', columns='Year', values='Date')

# ---- styles ----
NAVY = PatternFill('solid', fgColor='1A202C'); WHITE = Font(bold=True, color='FFFFFF', size=10)
TITLE = Font(bold=True, color='1A202C', size=13); SUB = Font(color='4A5568', size=10, italic=True)
GF = PatternFill('solid', fgColor='E6F4EA'); RF = PatternFill('solid', fgColor='FCE8E6')
AM = PatternFill('solid', fgColor='FEF3C7'); LBL = PatternFill('solid', fgColor='F1F5F9')
GOLD = PatternFill('solid', fgColor='FEF9C3')
GOOD = Font(color='137333', bold=True); BAD = Font(color='C5221F', bold=True); BOLD = Font(bold=True)
THIN = Border(*[Side(style='thin', color='CBD5E0')] * 4)
CEN = Alignment(horizontal='center'); RIGHT = Alignment(horizontal='right'); WRAP = Alignment(wrap_text=True, vertical='top')

wb = openpyxl.Workbook()

# ============ Sheet 1: Overview & Assumptions ============
ov = wb.active; ov.title = 'Overview_Assumptions'
ov.merge_cells('A1:B1'); ov.cell(1, 1, 'F&O OCTOBER SEASONALITY \u2014 PHASE 1 (single month) \u2014 Methodology & Assumptions').font = TITLE
notes = [
    ('Concept', 'For every F&O stock, test entering a LONG position on each trading-day-of-October and holding to the October month-end close. Find which entry day has historically produced the strongest, most consistent forward return.'),
    ('Universe & source', '211 NSE F&O stocks, near-month continuous FUTURES prices (scraped_parquet/fo_futures_near_month_continuous.parquet). Trading days are exactly what appears in the data \u2014 weekends and NSE holidays (e.g. 2-Oct Gandhi Jayanti, Dussehra) are naturally absent, verified directly, no synthetic calendar needed.'),
    ('Years covered', '7 Octobers: %s (varies per stock \u2014 newer F&O listings have fewer years; per-stock results require >=4 years).' % ', '.join(str(y) for y in YEARS)),
    ('Entry / Exit', 'Entry = daily Close on the trading day being tested (proxy for the trade price; no intraday data available). Exit = Close on the LAST trading day of that October (month-end close).'),
    ('CRITICAL FIX \u2014 split/bonus adjustment', 'Raw futures closes are NOT corporate-action adjusted. A stock split/bonus shows as a fake single-day move of 50-90% (found: DRREDDY 5:1 split crashed the close from Rs 6512 to Rs 1310 on 28-Oct-2024; RELIANCE had one the same day; 135+ such days found across the full dataset). FIX APPLIED: every return is built by compounding DAILY %% changes from entry to exit, SKIPPING any single day where |daily change| > 20%% (treated as a corporate action, not a real price move). Verified: DRREDDY Oct-2024 went from a fake -80%% to a real -3%% to -7%% range.'),
    ('Why NOT just "highest average return"', 'The naive highest-mean entry day (early October, ~Day-Index 4) is a MIRAGE \u2014 it only looks best because of the 2019 outlier (a corporate-tax-cut rally). Excluding 2019, that day flips NEGATIVE and its win rate drops from 57%% to 50%% with high volatility. The BEST_ENTRY_DAYS sheet ranks by a risk-adjusted (Sharpe-like = mean/std across years) score instead, and shows both the raw and ex-2019 figures so you can see what is robust vs what is a single-year artifact.'),
    ('Holding-period caveat', 'Later entry days (e.g. Day-Index 16, ~24-Oct) mechanically hold for FEWER days to month-end than early entries (Day-Index 1 holds ~19-21 days; Day-Index 16 holds ~4 days) \u2014 shorter holds have less time to accumulate risk, which partly explains their lower volatility. The Sharpe-like ranking and the avg_hold_days column make this comparable and visible; it is not hidden.'),
    ('Scope', 'PHASE 1: October only. Once reviewed, the same method extends to every month for a full 12-month seasonality calendar.'),
]
r = 3
for lbl, txt in notes:
    a = ov.cell(r, 1, lbl); a.font = BOLD; a.fill = AM; a.border = THIN; a.alignment = WRAP
    b = ov.cell(r, 2, txt); b.alignment = WRAP; b.border = THIN
    ov.row_dimensions[r].height = 85; r += 1
ov.column_dimensions['A'].width = 26; ov.column_dimensions['B'].width = 108

# ============ Sheet 2: Basket Day x Year matrix ============
bk = wb.create_sheet('Basket_Day_x_Year')
bk.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(YEARS) + 8)
bk.cell(1, 1, 'BASKET (all F&O stocks, equal-weight) \u2014 avg return %% by Entry Trading-Day-of-October, by YEAR \u2014 split-robust').font = TITLE
bk.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(YEARS) + 8)
bk.cell(2, 1, 'Each cell = basket avg %% return from entering on that trading-day-of-October to the October month-end close, for that year.').font = SUB
hdr = ['DayIdx'] + [str(y) for y in YEARS] + ['Mean', 'Mean(ex-2019)', 'Median', 'Std', 'Win%', 'Win%(ex-2019)']
for j, h in enumerate(hdr, 1):
    c = bk.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
best_idx = stats['sharpe_like'].idxmax()
row = 5
for di in pivot.index:
    vals = [di] + [round(pivot.loc[di, y], 2) if pd.notna(pivot.loc[di, y]) else None for y in YEARS]
    vals += [round(stats.loc[di, 'mean'], 2), round(stats.loc[di, 'mean_ex2019'], 2), round(stats.loc[di, 'median'], 2),
             round(stats.loc[di, 'std'], 2), round(stats.loc[di, 'win_rate'], 0), round(stats.loc[di, 'win_rate_ex2019'], 0)]
    rowfill = GOLD if di == best_idx else None
    for j, v in enumerate(vals, 1):
        c = bk.cell(row, j, v); c.border = THIN
        if j == 1: c.alignment = CEN
        else:
            c.alignment = RIGHT
            if isinstance(v, (int, float)) and v is not None and j <= len(YEARS) + 1:
                c.font = GOOD if v >= 0 else BAD
        if rowfill: c.fill = rowfill
    row += 1
bk.cell(row + 1, 1, '(Gold row = the risk-adjusted best entry day \u2014 see Best_Entry_Days_Ranked sheet)').font = Font(italic=True, color='92400E', size=9)
for j, w in enumerate([8] + [9] * len(YEARS) + [9, 14, 9, 7, 7, 13], 1):
    bk.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
bk.freeze_panes = 'B5'

# ============ Sheet 3: Best Entry Days Ranked ============
rk = wb.create_sheet('Best_Entry_Days_Ranked')
rk.merge_cells('A1:J1'); rk.cell(1, 1, 'BEST ENTRY TRADING-DAY-OF-OCTOBER \u2014 ranked by risk-adjusted (Sharpe-like) score').font = TITLE
rk.merge_cells('A2:J2'); rk.cell(2, 1, 'Sharpe-like = mean/std of the basket\u2019s yearly avg return for that entry day, across the 7 years.').font = SUB
hdr = ['Rank', 'Entry Day-Index', '~Calendar (2025 ref)', 'Avg Hold Days', 'Mean Ret %', 'Mean ex-2019 %', 'Median %', 'Std %', 'Win Rate %', 'Sharpe-like']
for j, h in enumerate(hdr, 1):
    c = rk.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
ranked = stats.sort_values('sharpe_like', ascending=False)
row = 5
for i, (di, s) in enumerate(ranked.iterrows(), 1):
    cal_dates = cal.loc[di].dropna() if di in cal.index else pd.Series(dtype='object')
    calref = cal_dates[2025].strftime('%d-%b') if 2025 in cal_dates.index else '\u2014'
    vals = [i, int(di), calref, round(s['avg_hold_days'], 1), round(s['mean'], 2), round(s['mean_ex2019'], 2),
            round(s['median'], 2), round(s['std'], 2), round(s['win_rate'], 0), round(s['sharpe_like'], 2)]
    fill = GOLD if i == 1 else (GF if s['sharpe_like'] > 0 else RF)
    for j, v in enumerate(vals, 1):
        c = rk.cell(row, j, v); c.border = THIN; c.fill = fill
        if j in (1, 2, 3): c.alignment = CEN
        else: c.alignment = RIGHT
    row += 1
for j, w in enumerate([7, 15, 17, 14, 12, 14, 11, 9, 11, 12], 1):
    rk.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
rk.freeze_panes = 'A5'

# ============ Sheet 4: Calendar Reference ============
cr = wb.create_sheet('Calendar_Reference')
cr.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(YEARS) + 1)
cr.cell(1, 1, 'Trading-Day-Index -> actual calendar date, by year (reference: RELIANCE, liquid full-history stock)').font = TITLE
hdr = ['DayIdx'] + [str(y) for y in YEARS]
for j, h in enumerate(hdr, 1):
    c = cr.cell(3, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
row = 4
for di in cal.index:
    vals = [di] + [(cal.loc[di, y].strftime('%d-%b') if y in cal.columns and pd.notna(cal.loc[di, y]) else '\u2014') for y in YEARS]
    fill = GOLD if di == best_idx else None
    for j, v in enumerate(vals, 1):
        c = cr.cell(row, j, v); c.border = THIN; c.alignment = CEN
        if fill: c.fill = fill
    row += 1
for j, w in enumerate([8] + [10] * len(YEARS), 1):
    cr.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w

# ============ Sheet 5: Per-Stock Trade Log (ALL Octobers, real dates+prices) ============
tl = wb.create_sheet('Per_Stock_Trade_Log')
tl.merge_cells('A1:K1')
tl.cell(1, 1, 'PER-STOCK TRADE LOG — every stock’s OWN best entry day, every October year, real dates & prices').font = TITLE
tl.merge_cells('A2:K2')
tl.cell(2, 1, 'One row per (stock, year): entry on that stock’s historically best trading-day-of-October (own Sharpe-like rank), held to month-end close. Split-robust return.').font = SUB
hdr = ['Symbol', 'Best Day-Index', 'Year', 'Entry Date', 'Entry Close (Rs)', 'Exit Date (month-end)',
       'Exit Close (Rs)', 'Return %', 'Had Split?', 'Outcome', 'Stock Overall Sharpe-like']
for j, h in enumerate(hdr, 1):
    c = tl.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
row = 5
perstock_map = perstock.set_index('sym')
for sym in perstock.sort_values('best_sharpe', ascending=False)['sym']:
    bd = int(perstock_map.loc[sym, 'best_day'])
    sh = perstock_map.loc[sym, 'best_sharpe']
    yr_rows = res[(res['Instrument'] == sym) & (res['DayIdx'] == bd)].sort_values('Year')
    for _, r in yr_rows.iterrows():
        win = r['Ret'] > 0
        vals = [sym, bd, int(r['Year']), r['EntryDate'].strftime('%d-%b-%Y'), round(r['EntryClose'], 2),
                r['ExitDate'].strftime('%d-%b-%Y'), round(r['ExitClose'], 2), round(r['Ret'], 2),
                ('YES — raw prices would mislead' if r['HadSplit'] else 'no'),
                'WIN' if win else 'LOSS', round(sh, 2)]
        fill = GF if win else RF
        for j, v in enumerate(vals, 1):
            c = tl.cell(row, j, v); c.border = THIN; c.fill = fill
            if j in (2, 3, 9, 10): c.alignment = CEN
            elif j == 1: c.alignment = Alignment(horizontal='left')
            else: c.alignment = RIGHT
            if j == 9 and r['HadSplit']: c.font = Font(color='B45309', bold=True, size=8)
            if j == 8: c.font = GOOD if win else BAD
        row += 1
for j, w in enumerate([13, 13, 7, 13, 15, 18, 14, 10, 24, 8, 18], 1):
    tl.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
tl.freeze_panes = 'A5'

# ============ Sheet 6: Per-Stock Best Day (summary) ============
ps = wb.create_sheet('Per_Stock_Best_Day')
ps.merge_cells('A1:H1'); ps.cell(1, 1, 'PER-STOCK best entry day \u2014 F&O stocks with >=4 years October history (split-robust)').font = TITLE
ps.merge_cells('A2:H2'); ps.cell(2, 1, '%d of 211 F&O stocks have enough history; ranked by risk-adjusted (Sharpe-like) score of their own best day.' % len(perstock)).font = SUB
hdr = ['Rank', 'Symbol', 'Years', 'Best Entry Day-Index', 'Mean Ret %', 'Median Ret %', 'Win Rate %', 'Sharpe-like']
for j, h in enumerate(hdr, 1):
    c = ps.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
perstock_sorted = perstock.sort_values('best_sharpe', ascending=False).reset_index(drop=True)
row = 5
for i, r_ in perstock_sorted.iterrows():
    vals = [i + 1, r_['sym'], int(r_['nyears']), int(r_['best_day']), round(r_['best_mean'], 2),
            round(r_['best_median'], 2), round(r_['best_wr'], 0), round(r_['best_sharpe'], 2)]
    fill = GF if r_['best_sharpe'] > 0 else RF
    for j, v in enumerate(vals, 1):
        c = ps.cell(row, j, v); c.border = THIN; c.fill = fill
        if j in (1, 3, 4, 7): c.alignment = CEN
        elif j == 2: c.alignment = Alignment(horizontal='left')
        else: c.alignment = RIGHT
    row += 1
for j, w in enumerate([7, 13, 8, 16, 12, 14, 12, 12], 1):
    ps.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
ps.freeze_panes = 'A5'

wb.move_sheet('Overview_Assumptions', -(len(wb.sheetnames) - 1))
wb.save(OUT)
print('Saved:', OUT)
print('Best basket entry day (Sharpe-like ranked): Day-Index %d (%s)' % (best_idx, cal.loc[best_idx, 2025].strftime('%d-%b-%Y') if best_idx in cal.index and 2025 in cal.columns else '?'))
print('Per-stock table: %d stocks with >=4yr history' % len(perstock))
