import openpyxl, pandas as pd, sys

sys.stdout.reconfigure(errors='replace')

path = r'D:/behaviour analysis/Nifty50_All_12_Quarters_And_YearWise_Best_Stock_Rankings.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

print("==========================================================================================")
print("VERIFYING ALL 12-QUARTER & YEAR-WISE STOCK RANKINGS WORKBOOK")
print("==========================================================================================")
print("Total Sheets:", len(wb.sheetnames))
print("Sheet Names :", wb.sheetnames)

quarters_sample = ['Q3 2023-24', 'Q1 2024-25', 'Q4 2025-26', 'Q2 2026-27']

for qtr in quarters_sample:
    df_q = pd.read_excel(path, sheet_name=qtr, engine='openpyxl')
    print(f"\n--- TOP 3 STOCKS IN {qtr} ---")
    print(df_q.iloc[2:5, [0, 1, 2, 5, 7, 8, 10]].to_string())

print("==========================================================================================")
