"""
Detailed portfolio-performance Excel for the November Top-N walk-forward strategy: CAGR, fund utilisation, Sharpe and ~80 other measures.
Reads the simulation results produced by the engine run (portfolio_engine.py). Output: November_Portfolio_Performance.xlsx (new file).
"""
import sys, pickle
sys.path.insert(0, 'D:/behaviour analysis'); sys.stdout.reconfigure(encoding='utf-8')
import numpy as np, pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import portfolio_engine as pe

SC = 'D:/behaviour analysis/portfolio_cache/'
picks = pickle.load(open(SC + 'picks.pkl', 'rb'))
sims, scaled, sens, rows, rows_sc, rows_se, G = pickle.load(open(SC + 'sims.pkl', 'rb'))
OUT = pe.ROOT + 'November_Portfolio_Performance.xlsx'
ORDER = [(p, n, sy) for n in pe.SIZES for (p, sy) in (('ALL', 2013), ('HINDSIGHT', 2013), ('EXPANDING', 2018), ('ALL', 2018), ('HINDSIGHT', 2018))]
HONEST = {'ALL': 'Yes - no date selection', 'EXPANDING': 'Yes - dates chosen from prior years only', 'HINDSIGHT': 'NO - dates chosen with full-period hindsight'}
POLNAME = {'ALL': 'All 30 dates', 'HINDSIGHT': 'Verdict (TRADE) dates - hindsight', 'EXPANDING': 'Verdict dates - expanding window'}
key = lambda r: (r['Policy'], r['N'], int(r['Window'][:4]))
R = {key(r): r for r in rows}; RS = {key(r): r for r in rows_sc}
cols = [(p, n, sy) for (p, n, sy) in ORDER]
colname = lambda k: f'{k[0]}-Top{k[1]} ({k[2]}-25)'

GROUPS = [
 ('GROWTH', [('Window start', 'Start', 'd'), ('Window end', 'End', 'd'), ('Years', 'Years', '0.00'), ('Starting capital (Rs)', 'StartCapital', '#,##0'), ('Ending capital - idle cash earns 0% (Rs)', 'EndCapital', '#,##0'),
             ('Ending capital - idle cash earns risk-free (Rs)', 'RfCash_EndCapital', '#,##0'), ('Total return % (idle cash 0%)', 'TotalReturnPct', '0.00'),
             ('CAGR % (idle cash 0%)', 'CAGRPct', '0.00'), ('CAGR % (idle cash earns 6.5%)', 'RfCash_CAGRPct', '0.00'), ('Nifty 50 CAGR % (same window)', 'NiftyCAGRPct', '0.00'),
             ('Excess CAGR vs Nifty % (idle cash 0%)', 'ExcessCAGRvsNiftyPct', '0.00'), ('Excess CAGR vs Nifty % (idle cash 6.5%)', 'RfCash_ExcessCAGRvsNiftyPct', '0.00'),
             ('Best year %', 'BestYearPct', '0.00'), ('Worst year %', 'WorstYearPct', '0.00'), ('Positive years %', 'PositiveYearsPct', '0.0')]),
 ('FUND UTILISATION', [('Average fund utilised % (of equity, all days)', 'AvgFundUtilisedPct', '0.00'), ('Peak fund utilised %', 'PeakFundUtilisedPct', '0.0'), ('Average utilised % on days invested', 'AvgUtilisedWhenInvestedPct', '0.0'),
             ('Days in market %', 'DaysInMarketPct', '0.0'), ('Idle cash % (average)', 'IdleCashPct', '0.0'), ('Average rupees deployed', 'AvgRupeesDeployed', '#,##0'),
             ('Return on deployed capital % per year', 'ReturnOnDeployedCapitalAnnPct', '0.00'), ('Return on margin % per year - futures scenario at 20% margin', 'ReturnOnMargin20pct_Scenario', '0.0'),
             ('Trades per year', 'TradesPerYear', '0.0'), ('Total trades', 'TotalTrades', '#,##0'), ('Average hold (trading days)', 'AvgHoldDays', '0.0')]),
 ('RISK-ADJUSTED (idle cash earns 0%)', [('Annualised volatility %', 'VolatilityAnnPct', '0.00'), ('Sharpe ratio (annualised, rf 6.5%)', 'SharpeAnn', '0.00'), ('Sortino ratio', 'SortinoAnn', '0.00'),
             ('Calmar ratio (CAGR / max DD)', 'CalmarRatio', '0.00'), ('Maximum drawdown %', 'MaxDrawdownPct', '0.00'), ('Longest time below a peak (days)', 'LongestDrawdownDays', '#,##0'), ('Ulcer index', 'UlcerIndex', '0.00')]),
 ('RISK-ADJUSTED (idle cash earns 6.5% in a liquid fund)', [('Annualised volatility %', 'RfCash_VolatilityAnnPct', '0.00'), ('Sharpe ratio (annualised)', 'RfCash_SharpeAnn', '0.00'), ('Sortino ratio', 'RfCash_SortinoAnn', '0.00'),
             ('Calmar ratio', 'RfCash_CalmarRatio', '0.00'), ('Maximum drawdown %', 'RfCash_MaxDrawdownPct', '0.00')]),
 ('TAIL RISK AND EXTREMES', [('Daily VaR 95% (days invested) %', 'VaR95_DailyWhenInvestedPct', '0.00'), ('Daily CVaR 95% (days invested) %', 'CVaR95_DailyWhenInvestedPct', '0.00'), ('Best day %', 'BestDayPct', '0.00'), ('Worst day %', 'WorstDayPct', '0.00'),
             ('Best month %', 'BestMonthPct', '0.00'), ('Worst month %', 'WorstMonthPct', '0.00'), ('Positive months % (months with a position)', 'PositiveMonthsInvestedPct', '0.0')]),
 ('VERSUS NIFTY 50', [('Beta (daily)', 'Beta', '0.000'), ('Alpha % per year (daily regression)', 'AlphaAnnPct', '0.00'), ('Correlation', 'CorrelationNifty', '0.000'), ('Up-market capture', 'UpCapture', '0.000'), ('Down-market capture', 'DownCapture', '0.000'),
             ('Information ratio', 'InformationRatio', '0.00'), ('Nifty volatility % (same window)', 'NiftyVolPct', '0.00'), ('Nifty max drawdown % (same window)', 'NiftyMaxDDPct', '0.00'), ('Nifty Sharpe (same window)', 'NiftySharpe', '0.00')]),
 ('TRADE QUALITY', [('Trade win %', 'TradeWinPct', '0.00'), ('Average trade net %', 'AvgTradeNetPct', '0.00'), ('Median trade net %', 'MedianTradeNetPct', '0.00'), ('Average win %', 'AvgWinPct', '0.00'), ('Average loss %', 'AvgLossPct', '0.00'),
             ('Payoff ratio', 'PayoffRatio', '0.00'), ('Profit factor', 'ProfitFactor', '0.00'), ('Expectancy per Rs 100 deployed', 'ExpectancyPer100', '0.00'), ('Kelly fraction', 'KellyFraction', '0.000'),
             ('Longest winning streak (trades)', 'LongestWinStreakTrades', '0'), ('Longest losing streak (trades)', 'LongestLossStreakTrades', '0'), ('Top-3 trades share of total profit %', 'Top3TradesShareOfProfitPct', '0.0'),
             ('Total cost paid (Rs)', 'TotalCostRs', '#,##0'), ('Cost drag % of gross profit', 'CostDragPctOfGrossProfit', '0.0')]),
 ('STATISTICAL CONFIDENCE', [('t-stat of yearly returns', 'YearlyReturnTstat', '0.00'), ('Bootstrap CAGR 5th percentile %', 'CAGR_Boot5th_Pct', '0.00'), ('Bootstrap CAGR 95th percentile %', 'CAGR_Boot95th_Pct', '0.00'),
             ('Probability CAGR > 0 (bootstrap over years) %', 'ProbCAGRpositive_Pct', '0.0'), ('Probabilistic Sharpe: P(true Sharpe > 0) % (indicative)', 'ProbSharpePositive_PSR_Pct', '0.0'),
             ('Deflated Sharpe probability % (20 strategy trials, indicative)', 'DeflatedSharpeProb_Pct', '0.0')]),
]
hdr = PatternFill('solid', fgColor='1F3864'); hf = Font(bold=True, color='FFFFFF'); gfill = PatternFill('solid', fgColor='D9E1F2')
GREEN = PatternFill('solid', fgColor='E3F1E4'); AMB = PatternFill('solid', fgColor='FBF1DC'); RED = PatternFill('solid', fgColor='F8E1E1')
POLFILL = {'ALL': GREEN, 'EXPANDING': GREEN, 'HINDSIGHT': RED}
def hrow(ws, r, n, fill=hdr):
    for c in range(1, n + 1):
        x = ws.cell(r, c); x.fill = fill; x.font = hf; x.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')

README = [
 'NOVEMBER TOP-N STRATEGY - PORTFOLIO PERFORMANCE WORKBOOK (CAGR, fund utilisation, Sharpe and more)',
 'What it simulates: a rupee account trading the walk-forward Top-N stocks of the November calendar-date study. Data: Stock_Data_213_Adjusted (cleaned Yahoo Finance), 213 F&O stocks; Nifty 50 index as benchmark. Test years 2013-2025.',
 '',
 'ASSUMPTIONS (all changeable - tell me and I rerun)',
 '1. Starting capital Rs 10,00,000. Cash-settled P&L (LONG and SHORT positions both modelled as stake x return; SHORT would need futures in practice). No leverage.',
 '2. Stock selection: for each year Y and each November date d (1..30) the Top-N stocks and their (buy N days before, sell M days after) rule are chosen from years BEFORE Y only (>= 10 prior events, positive average, >= 15 eligible stocks, ranked by win % then average).',
 '3. Sizing: each date is a "sleeve". At the start of each year the account equity is split equally across the K sleeves in use that year, and each sleeve equally across its N stocks. Positions are not rebalanced during the year; equity compounds from year to year.',
 '4. Execution: buy/sell at the close of T-n / T+m (trading days around the event day T). Cost 0.15% of stake per position (round trip), charged on exit. Daily mark-to-market on the combined Nifty + stock trading calendar.',
 '5. Idle cash: shown two ways - earning 0% (conservative) and earning the 6.5% risk-free rate (liquid-fund view). Risk-free rate 6.5% a year for Sharpe/Sortino.',
 '6. Date policies: ALL = all 30 November dates (no selection). HINDSIGHT = dates that pass the TRADE test (avg >= +1.0%, winning years >= 62%) over the WHOLE 2013-2025 period - this uses future knowledge and is BIASED (shown for reference only). EXPANDING = a date is used in year Y only if its walk-forward record over earlier years (at least 5) passes the TRADE test - honest; its test window is 2018-2025.',
 '7. Windows: 2013-2025 (13 years) for ALL/HINDSIGHT; 2018-2025 for EXPANDING, with ALL and HINDSIGHT also shown over 2018-2025 so all three can be compared like for like.',
 '8. "Scaled" sheet: same strategies with positions enlarged so that peak deployment is about 98% of equity (no leverage), to show what low fund utilisation costs. Cost sensitivity: same strategies at 0.30% and 0.50% cost per position.',
 '',
 'KEY DEFINITIONS',
 'CAGR = (ending capital / starting capital)^(1/years) - 1. Fund utilised % = capital in open positions / equity (daily average over ALL days, so idle periods count). Days in market % = share of days with at least one open position.',
 'Return on deployed capital = total P&L / (average rupees deployed x years): what the money earned while it was actually working. Return on margin (scenario) = that figure / 0.20, i.e. if the same trades were done in futures with 20% margin.',
 'Sharpe = mean(daily return - rf/252) / std(daily return) x sqrt(252). Sortino divides by downside deviation below the risk-free rate. Calmar = CAGR / |max drawdown|. Ulcer index = root-mean-square of daily drawdown %.',
 'Max drawdown = largest fall from a previous equity peak (daily mark-to-market). VaR/CVaR 95% are measured on days with a position open (otherwise the many idle zero-return days hide the risk).',
 'Beta/alpha come from a daily regression on the Nifty 50. Up/down capture = average strategy daily return divided by average Nifty return on Nifty up / down days. Information ratio = mean(strategy - Nifty daily) / std x sqrt(252).',
 'Profit factor = gains / losses in rupees. Payoff = average win % / average loss %. Kelly = win% - (1 - win%)/payoff. Expectancy per Rs 100 = average net trade %.',
 'Bootstrap: yearly returns are resampled 5,000 times to give the 5th-95th percentile range of CAGR and the probability that CAGR is above zero. Probabilistic / Deflated Sharpe (Bailey & Lopez de Prado) use daily returns and are INDICATIVE only: daily returns of held positions are autocorrelated, which flatters them; the deflation assumes 20 strategy trials.',
 '',
 'HOW TO READ THE RESULTS HONESTLY',
 'The strategy is invested only about 12-14% of days, so CAGR on the whole account is small even when the trades earn well while open - read CAGR together with fund utilisation and return on deployed capital.',
 'Only ALL and EXPANDING are free of selection bias. HINDSIGHT rows look better because the dates were chosen after seeing the results. Compare EXPANDING with ALL over the same 2018-2025 window.',
 'Limits: 13 test years; 30 dates overlap in time; survivorship bias (today\'s F&O list); flat 0.15% cost, no slippage, margin interest or tax; SHORT trades assumed available at the same cost; not investment advice.']

with pd.ExcelWriter(OUT, engine='openpyxl') as xw:
    pd.DataFrame({'README': README}).to_excel(xw, sheet_name='ReadMe', index=False); xw.sheets['ReadMe'].column_dimensions['A'].width = 210

    # ---------- Dashboard (key measures) ----------
    KEY = [('Date policy', None), ('Selection bias-free?', None), ('Window', 'Window'), ('Ending capital - idle cash 0% (Rs)', 'EndCapital'), ('CAGR % (idle cash 0%)', 'CAGRPct'), ('CAGR % (idle cash 6.5%)', 'RfCash_CAGRPct'),
           ('Nifty 50 CAGR %', 'NiftyCAGRPct'), ('Average fund utilised %', 'AvgFundUtilisedPct'), ('Peak fund utilised %', 'PeakFundUtilisedPct'), ('Days in market %', 'DaysInMarketPct'),
           ('Return on deployed capital % / yr', 'ReturnOnDeployedCapitalAnnPct'), ('Sharpe (idle cash 0%)', 'SharpeAnn'), ('Sharpe (idle cash 6.5%)', 'RfCash_SharpeAnn'), ('Sortino (idle cash 6.5%)', 'RfCash_SortinoAnn'),
           ('Max drawdown % (idle cash 0%)', 'MaxDrawdownPct'), ('Calmar (idle cash 6.5%)', 'RfCash_CalmarRatio'), ('Volatility % (idle cash 0%)', 'VolatilityAnnPct'), ('Beta vs Nifty', 'Beta'), ('Alpha % / yr', 'AlphaAnnPct'),
           ('Trades', 'TotalTrades'), ('Trade win %', 'TradeWinPct'), ('Profit factor', 'ProfitFactor'), ('Average trade net %', 'AvgTradeNetPct'), ('Worst year %', 'WorstYearPct'), ('Positive years %', 'PositiveYearsPct'),
           ('Bootstrap CAGR 5th-95th %', None), ('Prob. CAGR > 0 %', 'ProbCAGRpositive_Pct')]
    D = pd.DataFrame(index=[k[0] for k in KEY], columns=[colname(c) for c in cols])
    for c in cols:
        r = R[c]; nm = colname(c)
        for lab, kk in KEY:
            if lab == 'Date policy': D.loc[lab, nm] = POLNAME[c[0]]
            elif lab == 'Selection bias-free?': D.loc[lab, nm] = HONEST[c[0]].split(' - ')[0]
            elif lab.startswith('Bootstrap'): D.loc[lab, nm] = f"{r['CAGR_Boot5th_Pct']:.2f} to {r['CAGR_Boot95th_Pct']:.2f}"
            else:
                v = r[kk]; D.loc[lab, nm] = v if isinstance(v, str) else (round(float(v), 2) if v == v else None)
    D.index.name = 'Measure'; D.reset_index().to_excel(xw, sheet_name='Dashboard', index=False)
    ws = xw.sheets['Dashboard']; hrow(ws, 1, len(cols) + 1); ws.freeze_panes = 'B2'; ws.row_dimensions[1].height = 46
    ws.column_dimensions['A'].width = 38
    for i, c in enumerate(cols, 2):
        ws.column_dimensions[get_column_letter(i)].width = 17
        ws.cell(2, i).fill = POLFILL[c[0]]; ws.cell(3, i).fill = POLFILL[c[0]]
    for r_ in range(4, ws.max_row + 1):
        ws.cell(r_, 1).font = Font(bold=True)
        for i in range(2, len(cols) + 2):
            v = ws.cell(r_, i).value
            if isinstance(v, (int, float)): ws.cell(r_, i).number_format = '#,##0' if abs(v) > 5000 else '0.00'
    for r_ in range(2, ws.max_row + 1):
        for i in range(2, len(cols) + 2): ws.cell(r_, i).alignment = Alignment(horizontal='center', wrap_text=True)

    # ---------- All_Measures ----------
    rowsout = []
    for g, items in GROUPS:
        rowsout.append((g, None, None))
        for lab, kk, fm in items: rowsout.append((lab, kk, fm))
    ws = xw.book.create_sheet('All_Measures'); xw.sheets['All_Measures'] = ws
    ws.cell(1, 1, 'Measure'); ws.cell(1, 2, 'Unit / note')
    for i, c in enumerate(cols, 3): ws.cell(1, i, colname(c))
    hrow(ws, 1, len(cols) + 2); ws.row_dimensions[1].height = 50
    ws.cell(2, 1, 'Date policy'); ws.cell(3, 1, 'Selection bias-free?')
    for i, c in enumerate(cols, 3): ws.cell(2, i, POLNAME[c[0]]); ws.cell(3, i, HONEST[c[0]]); ws.cell(2, i).fill = POLFILL[c[0]]; ws.cell(3, i).fill = POLFILL[c[0]]
    rr = 4
    for lab, kk, fm in rowsout:
        if kk is None:
            ws.cell(rr, 1, lab).font = Font(bold=True)
            for q in range(1, len(cols) + 3): ws.cell(rr, q).fill = gfill
        else:
            ws.cell(rr, 1, lab)
            for i, c in enumerate(cols, 3):
                v = R[c][kk]
                if isinstance(v, (np.floating, float)): v = None if np.isnan(v) else float(v)
                elif isinstance(v, (np.integer,)): v = int(v)
                x = ws.cell(rr, i, v)
                if fm == 'd': x.number_format = 'yyyy-mm-dd'
                elif isinstance(v, (int, float)): x.number_format = fm
        rr += 1
    ws.column_dimensions['A'].width = 56; ws.column_dimensions['B'].width = 6; ws.freeze_panes = 'C4'
    for i in range(3, len(cols) + 3): ws.column_dimensions[get_column_letter(i)].width = 17

    # ---------- Scaled ----------
    SCOLS = [('Strategy', 'Strategy'), ('Window', 'Window'), ('Position size multiplier', 'scale'), ('Peak fund utilised %', 'PeakFundUtilisedPct'), ('Average fund utilised %', 'AvgFundUtilisedPct'),
             ('CAGR % (idle cash 0%)', 'CAGRPct'), ('CAGR % (idle cash 6.5%)', 'RfCash_CAGRPct'), ('Sharpe (idle cash 6.5%)', 'RfCash_SharpeAnn'), ('Max drawdown % (idle cash 0%)', 'MaxDrawdownPct'),
             ('Worst year %', 'WorstYearPct'), ('Worst day %', 'WorstDayPct'), ('Ending capital (Rs)', 'EndCapital'), ('Prob. CAGR > 0 %', 'ProbCAGRpositive_Pct')]
    S2 = []
    for c in cols:
        r = RS[c]; sc = scaled[c]['scale']
        S2.append({lab: (sc if k == 'scale' else r[k]) for lab, k in SCOLS})
    S2 = pd.DataFrame(S2); S2.round(3).to_excel(xw, sheet_name='Scaled_Positions', index=False)
    ws = xw.sheets['Scaled_Positions']; hrow(ws, 1, S2.shape[1]); ws.freeze_panes = 'C2'
    for i in range(1, S2.shape[1] + 1): ws.column_dimensions[get_column_letter(i)].width = 17
    ws.cell(len(S2) + 3, 1, 'Positions enlarged by the multiplier so peak deployment is ~98% of equity (no leverage). Shows the effect of low fund utilisation; risk (drawdown, worst day) scales up with it.')

    # ---------- Cost sensitivity ----------
    cs = []
    for r in rows_se:
        cs.append(dict(Strategy=r['Strategy'], Window=r['Window'], CostPerPositionPct=r['Cost'], CAGR_IdleCash0_Pct=r['CAGRPct'], CAGR_IdleCashRf_Pct=r['RfCash_CAGRPct'], SharpeIdleCashRf=r['RfCash_SharpeAnn'],
                       MaxDrawdownPct=r['MaxDrawdownPct'], ProfitFactor=r['ProfitFactor'], AvgTradeNetPct=r['AvgTradeNetPct']))
    for c in cols:
        if (c[0], c[1], c[2]) in [(p, n, sy) for p, n, sy in cols if (p, sy) in (('ALL', 2013), ('EXPANDING', 2018), ('HINDSIGHT', 2013))]:
            r = R[c]; cs.append(dict(Strategy=r['Strategy'], Window=r['Window'], CostPerPositionPct=0.15, CAGR_IdleCash0_Pct=r['CAGRPct'], CAGR_IdleCashRf_Pct=r['RfCash_CAGRPct'], SharpeIdleCashRf=r['RfCash_SharpeAnn'],
                                     MaxDrawdownPct=r['MaxDrawdownPct'], ProfitFactor=r['ProfitFactor'], AvgTradeNetPct=r['AvgTradeNetPct']))
    CS = pd.DataFrame(cs).sort_values(['Strategy', 'Window', 'CostPerPositionPct']).round(3); CS.to_excel(xw, sheet_name='Cost_Sensitivity', index=False)
    ws = xw.sheets['Cost_Sensitivity']; hrow(ws, 1, CS.shape[1]); ws.freeze_panes = 'A2'
    for i in range(1, CS.shape[1] + 1): ws.column_dimensions[get_column_letter(i)].width = 20

    # ---------- Equity daily ----------
    idx = G.index[G.index >= '2013-01-01']
    E = pd.DataFrame(index=idx); E['Nifty 50 (rebased to Rs 10,00,000)'] = (G.reindex(idx) / G.reindex(idx).iloc[0] * pe.CAP0)
    for c in cols:
        e = sims[c]['equity']; E[colname(c) + ' equity'] = e.reindex(idx)
    for c in cols:
        d_ = (sims[c]['deployed'] / sims[c]['equity']) * 100; E[colname(c) + ' utilised %'] = d_.reindex(idx)
    E.index.name = 'Date'; E = E.reset_index(); E['Date'] = E['Date'].dt.date
    E.round(2).to_excel(xw, sheet_name='Equity_Daily', index=False); ws = xw.sheets['Equity_Daily']; hrow(ws, 1, E.shape[1]); ws.freeze_panes = 'B2'; ws.row_dimensions[1].height = 60
    for i in range(1, E.shape[1] + 1): ws.column_dimensions[get_column_letter(i)].width = 16
    print('equity rows', len(E))

    # ---------- Yearly & monthly ----------
    def yearly(eq):
        ye = eq.groupby(eq.index.year).last(); y = ye.pct_change(); y.iloc[0] = ye.iloc[0] / pe.CAP0 - 1; return y * 100
    Y = pd.DataFrame({colname(c): yearly(sims[c]['equity']) for c in cols})
    nf = G.reindex(idx); ny = nf.groupby(nf.index.year).last().pct_change() * 100; ny.iloc[0] = (nf.groupby(nf.index.year).last().iloc[0] / nf.iloc[0] - 1) * 100
    Y.insert(0, 'Nifty 50 %', ny); Y.index.name = 'Year'; Y.round(2).reset_index().to_excel(xw, sheet_name='Yearly_Returns', index=False)
    ws = xw.sheets['Yearly_Returns']; hrow(ws, 1, Y.shape[1] + 1); ws.freeze_panes = 'B2'; ws.row_dimensions[1].height = 50
    for i in range(1, Y.shape[1] + 2): ws.column_dimensions[get_column_letter(i)].width = 16
    def monthly(eq):
        me = eq.resample('ME').last(); m = me.pct_change(); m.iloc[0] = me.iloc[0] / pe.CAP0 - 1; return m * 100
    M = pd.DataFrame({colname(c): monthly(sims[c]['equity']) for c in cols}); M.index = M.index.strftime('%Y-%m'); M.index.name = 'Month'
    M.round(3).reset_index().to_excel(xw, sheet_name='Monthly_Returns', index=False); ws = xw.sheets['Monthly_Returns']; hrow(ws, 1, M.shape[1] + 1); ws.freeze_panes = 'B2'; ws.row_dimensions[1].height = 50
    for i in range(1, M.shape[1] + 2): ws.column_dimensions[get_column_letter(i)].width = 16

    # ---------- Utilisation by month of year ----------
    U = pd.DataFrame({colname(c): ((sims[c]['deployed'] / sims[c]['equity']) * 100).groupby(sims[c]['equity'].index.month).mean() for c in cols})
    U.index = [pd.Timestamp(2000, m, 1).strftime('%b') for m in U.index]; U.index.name = 'Calendar month'
    U.round(2).reset_index().to_excel(xw, sheet_name='Utilisation_By_Month', index=False); ws = xw.sheets['Utilisation_By_Month']; hrow(ws, 1, U.shape[1] + 1); ws.row_dimensions[1].height = 50
    for i in range(1, U.shape[1] + 2): ws.column_dimensions[get_column_letter(i)].width = 16
    ws.cell(U.shape[0] + 3, 1, 'Average % of equity deployed, by calendar month - the strategy is only invested from late October to mid-December.')

    # ---------- Drawdowns ----------
    dd = []
    for c in cols:
        for k, (s_, t_, e_, depth) in enumerate(pe.dd_episodes(sims[c]['equity'], 5), 1):
            dd.append(dict(Strategy=colname(c), Rank=k, PeakBefore=s_.date(), Trough=t_.date(), Recovered=e_.date() if e_ is not None else 'not yet', DepthPct=round(depth * 100, 2),
                           DaysPeakToTrough=(t_ - s_).days, DaysToRecover=(e_ - t_).days if e_ is not None else None))
    DD = pd.DataFrame(dd); DD.to_excel(xw, sheet_name='Drawdowns', index=False); ws = xw.sheets['Drawdowns']; hrow(ws, 1, DD.shape[1]); ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for i in range(1, DD.shape[1] + 1): ws.column_dimensions[get_column_letter(i)].width = 22

    # ---------- Date selection ----------
    ds = []
    for N in pe.SIZES:
        r_ = pe.wf_returns(picks, N)
        for d in range(1, 31):
            v = np.array([r_[(d, y)] for y in range(2013, 2026) if (d, y) in r_])
            used = [y for y in range(2018, 2026) if d in pe.select_dates(picks, N, 'EXPANDING', y)]
            ds.append(dict(Portfolio=f'Top {N}', Date=f'{d:02d}-Nov', WF_AvgNetPct=round(v.mean(), 2), WF_WinYearsPct=round((v > 0).mean() * 100, 1), WorstYearPct=round(v.min(), 2),
                           VerdictHindsight='TRADE' if (v.mean() >= 1.0 and (v > 0).mean() * 100 >= 62) else 'not selected', YearsSelectedByExpandingRule=len(used), YearsList=', '.join(map(str, used))))
    DS = pd.DataFrame(ds); DS.to_excel(xw, sheet_name='Date_Selection', index=False); ws = xw.sheets['Date_Selection']; hrow(ws, 1, DS.shape[1]); ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for i in range(1, DS.shape[1] + 1): ws.column_dimensions[get_column_letter(i)].width = 22
    ws.column_dimensions['H'].width = 40

    # ---------- Sleeves per year ----------
    sl = []
    for c in cols:
        for (Y_, K, dts) in sims[c]['sleeves']: sl.append(dict(Strategy=colname(c), Year=Y_, SleevesUsed=K, Dates=', '.join(f'{d:02d}' for d in dts)))
    SL = pd.DataFrame(sl); SL.to_excel(xw, sheet_name='Sleeves_By_Year', index=False); ws = xw.sheets['Sleeves_By_Year']; hrow(ws, 1, SL.shape[1]); ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate([34, 8, 14, 110], 1): ws.column_dimensions[get_column_letter(i)].width = w

    # ---------- Trades ----------
    T = pd.concat([sims[c]['trades'].assign(Window=f'{c[2]}-2025') for c in cols], ignore_index=True)
    T['Entry'] = T.Entry.dt.date; T['Exit'] = T.Exit.dt.date
    T = T[['Strategy', 'Window', 'Year', 'Sleeve', 'Rank', 'Stock', 'Side', 'Rule', 'Entry', 'Exit', 'HoldDays', 'Stake', 'GrossPct', 'NetPct', 'PnL']].round({'Stake': 0, 'GrossPct': 2, 'NetPct': 2, 'PnL': 0})
    T.columns = ['Strategy', 'Window', 'Year', 'Date sleeve', 'Rank in sleeve', 'Stock', 'Side', 'Rule', 'Entry', 'Exit', 'Hold days', 'Stake (Rs)', 'Gross %', 'Net % (after 0.15%)', 'Net P&L (Rs)']
    T.to_excel(xw, sheet_name='Trades', index=False); ws = xw.sheets['Trades']; hrow(ws, 1, T.shape[1]); ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for i in range(1, T.shape[1] + 1): ws.column_dimensions[get_column_letter(i)].width = 15
    print('trade rows', len(T))

    # ---------- Picks ----------
    pk = []
    for (d, Y_), l in sorted(picks.items()):
        for rank, p in enumerate(l, 1):
            pk.append(dict(Date=f'{d:02d}-Nov', Year=Y_, Rank=rank, Stock=p['sym'], Side='LONG' if p['side'] == 1 else 'SHORT', Rule=f'T-{p["n"]} to T+{p["m"]}', Entry=p['entry'].date(), Exit=p['exit'].date(),
                           TrainWinPct=round(p['train_win'], 1), TrainAvgPct=round(p['train_avg'], 2), GrossReturnPct=round(p['gross'], 2)))
    PK = pd.DataFrame(pk); PK.to_excel(xw, sheet_name='WF_Picks_All', index=False); ws = xw.sheets['WF_Picks_All']; hrow(ws, 1, PK.shape[1]); ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for i in range(1, PK.shape[1] + 1): ws.column_dimensions[get_column_letter(i)].width = 16
print('saved', OUT)
