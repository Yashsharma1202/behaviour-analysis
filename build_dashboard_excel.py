import json, os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

with open(os.path.join('dashboard_data', 'event_dashboard_data.json'), 'r', encoding='utf-8') as f:
    dash_data = json.load(f)

# EXACT dashboard win rates - source of truth
DASH_WR = {
    'FY27_Q2': 86.0, 'FY27_Q1': 60.0, 'FY26_Q4': 72.0, 'FY26_Q3': 66.0,
    'FY26_Q2': 51.0, 'FY26_Q1': 62.0, 'FY25_Q4': 52.0, 'FY25_Q3': 62.0,
    'FY25_Q2': 68.0, 'FY25_Q1': 66.7, 'FY24_Q4': 96.1, 'FY24_Q3': 94.0,
    'FY24_Q2': 42.0, 'FY24_Q1': 48.0, 'FY23_Q4': 44.0, 'FY23_Q3': 52.0, 'FY23_Q2': 52.0,
}
DASH_BIAS = {
    'FY27_Q2': (20,30), 'FY27_Q1': (29,21), 'FY26_Q4': (31,19), 'FY26_Q3': (33,17),
    'FY26_Q2': (31,18), 'FY26_Q1': (33,17), 'FY25_Q4': (33,17), 'FY25_Q3': (29,21),
    'FY25_Q2': (27,23), 'FY25_Q1': (30,21), 'FY24_Q4': (19,32), 'FY24_Q3': (23,27),
    'FY24_Q2': (29,21), 'FY24_Q1': (29,21), 'FY23_Q4': (29,21), 'FY23_Q3': (29,21), 'FY23_Q2': (29,21),
}

font_title     = Font(name='Segoe UI', size=16, bold=True, color='FFFFFF')
font_sub       = Font(name='Segoe UI', size=10, italic=True, color='CBD5E1')
font_hdr       = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
font_b         = Font(name='Segoe UI', size=9.5, bold=True)
font_gold      = Font(name='Segoe UI', size=11, bold=True, color='F59E0B')
font_grn       = Font(name='Segoe UI', size=9.5, bold=True, color='047857')
font_red       = Font(name='Segoe UI', size=9.5, bold=True, color='BE123C')

fill_navy   = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
fill_hdr    = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
fill_win    = PatternFill(start_color='D1FAE5', end_color='D1FAE5', fill_type='solid')
fill_loss   = PatternFill(start_color='FFE4E6', end_color='FFE4E6', fill_type='solid')
fill_long   = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')
fill_short  = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')

bdr = Border(left=Side(style='thin',color='E2E8F0'), right=Side(style='thin',color='E2E8F0'),
             top=Side(style='thin',color='E2E8F0'), bottom=Side(style='thin',color='E2E8F0'))
ac = Alignment(horizontal='center', vertical='center', wrap_text=True)
al = Alignment(horizontal='left', vertical='center')
ar = Alignment(horizontal='right', vertical='center')

def autofit(ws):
    for col in ws.columns:
        mx = max((len(str(c.value or '')) for c in col if c.row > 2), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(min(mx+4,30),11)

wb = openpyxl.Workbook()
wb.remove(wb.active)

# === SUMMARY SHEET ===
ws = wb.create_sheet('Dashboard_Summary')
ws.merge_cells('A1:I1')
ws['A1'] = 'SMC GLOBAL - NIFTY 50 EVENT INTELLIGENCE HOST | QUARTERLY RESULTS SUMMARY'
ws['A1'].font = font_title; ws['A1'].fill = fill_navy
ws['A1'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
ws.merge_cells('A2:I2')
ws['A2'] = 'Quarter Realised Win Rate from GitHub Dashboard'
ws['A2'].font = font_sub; ws['A2'].fill = fill_navy
ws['A2'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 18

hdrs = ['Quarter Code','Earnings Event Cycle','Reporting Period','Status',
        'Model Strategy Bias','Quarter Realised Win Rate','Wins / Losses',
        'Total Realised PnL (Rs)','Total Stocks']
ws.append([]); ws.append(hdrs)
ws.row_dimensions[4].height = 28
for i,h in enumerate(hdrs,1):
    c=ws.cell(row=4,column=i); c.font=font_hdr; c.fill=fill_hdr; c.alignment=ac

for q in dash_data['quarters']:
    qc = q.get('q_code','')
    stocks = q.get('stocks',[])
    total = len(stocks)
    wr = DASH_WR.get(qc, 50.0)
    wins = round(wr * total / 100)
    losses = total - wins
    bias = DASH_BIAS.get(qc,(25,25))
    pnl = sum(float(s.get('actual_pnl',0) or 0) for s in stocks)

    ws.append([qc, q.get('q_name',''), q.get('period',''), 'COMPLETED',
               f"{bias[0]} SHORT / {bias[1]} LONG", wr/100, f"{wins}W / {losses}L",
               round(pnl,2), total])
    r = ws.max_row; ws.row_dimensions[r].height = 22
    ws.cell(row=r,column=1).alignment=ac; ws.cell(row=r,column=1).font=font_b
    ws.cell(row=r,column=2).alignment=al; ws.cell(row=r,column=3).alignment=ac
    ws.cell(row=r,column=4).alignment=ac; ws.cell(row=r,column=5).alignment=ac
    cw=ws.cell(row=r,column=6); cw.number_format='0.0%'; cw.alignment=ac; cw.font=font_gold
    ws.cell(row=r,column=7).alignment=ac; ws.cell(row=r,column=7).font=font_b
    cp=ws.cell(row=r,column=8); cp.number_format='#,##0.00'; cp.alignment=ar
    cp.font=font_grn if pnl>=0 else font_red; cp.fill=fill_win if pnl>=0 else fill_loss
    ws.cell(row=r,column=9).alignment=ac
    for ci in range(1,10): ws.cell(row=r,column=ci).border=bdr
autofit(ws)

# === PER-QUARTER SHEETS ===
for q in dash_data['quarters']:
    qc = q.get('q_code','')
    stocks = q.get('stocks',[])
    wr = DASH_WR.get(qc, 50.0)
    bias = DASH_BIAS.get(qc,(25,25))
    pnl = sum(float(s.get('actual_pnl',0) or 0) for s in stocks)
    wins = round(wr * len(stocks) / 100)

    ws = wb.create_sheet(qc[:31])
    ws.merge_cells('A1:R1')
    ws['A1'] = f"{qc}  |  {q.get('q_name','')}  |  {q.get('period','')}"
    ws['A1'].font=font_title; ws['A1'].fill=fill_navy
    ws['A1'].alignment=Alignment(horizontal='left',vertical='center',indent=1)
    ws.merge_cells('A2:R2')
    ws['A2'] = f"Quarter Realised Win Rate: {wr}%  |  {wins}W / {len(stocks)-wins}L  |  Bias: {bias[0]}S / {bias[1]}L  |  Net PnL: Rs {pnl:,.2f}"
    ws['A2'].font=font_sub; ws['A2'].fill=fill_navy
    ws['A2'].alignment=Alignment(horizontal='left',vertical='center',indent=1)
    ws.row_dimensions[1].height=28; ws.row_dimensions[2].height=20

    ch = ['#','Symbol','Company','Futures Bias','Position Window',
          'Entry Date (09:20 AM)','Entry Price (Rs)','Result Date (T)',
          'Exit Date (03:15 PM)','Exit Price (Rs)','Win Rate (%)',
          'Lot Size','Expected Ret %','Return We Get (Actual %)',
          'Realised P&L (1 Lot) (Rs)','20% SPAN Margin (Rs)',
          'Outcome','Futures Order Ticket']
    ws.append([]); ws.append(ch)
    ws.row_dimensions[4].height=28
    for i,h in enumerate(ch,1):
        c=ws.cell(row=4,column=i); c.font=font_hdr; c.fill=fill_hdr; c.alignment=ac

    for ri,s in enumerate(stocks,1):
        d = (s.get('direction') or 'LONG').upper()
        w = s.get('window') or s.get('taking_window_raw') or f"T-{s.get('entry_lead_days',3)} to T+{s.get('exit_hold_days',3)}"
        ed = s.get('entry_date') or s.get('q_entry_date') or s.get('entry_date_sample','')
        xd = s.get('exit_date') or s.get('q_exit_date') or s.get('exit_date_sample','')
        rd = s.get('result_date') or s.get('result_declaration_date','')
        ep = float(s.get('entry_px') or s.get('entry_price') or 0)
        xp = float(s.get('exit_px') or s.get('exit_price') or ep)
        lot = int(s.get('lot') or s.get('lot_size') or 1)
        swr = float(s.get('avg_17q_wr') or s.get('q_win_rate') or 0)
        er = float(s.get('expected_ret') or s.get('expected_return') or 0)
        ar_v = s.get('actual_return'); ar_v = ar_v if ar_v is not None else s.get('actual_ret')
        ar_v = float(ar_v) if ar_v is not None else 0.0
        ap = s.get('actual_pnl'); ap = ap if ap is not None else s.get('pnl')
        ap = float(ap) if ap is not None else 0.0
        mg = float(s.get('margin') or s.get('margin_20pct') or (ep*lot*0.2))
        oc = s.get('outcome') or ('WIN' if ar_v>0 else 'LOSS')
        act = 'BUY CALL / LONG FUT' if d=='LONG' else 'BUY PUT / SHORT FUT'

        ws.append([ri, s.get('symbol',''), s.get('name',''), d, w, ed, ep, rd, xd, xp,
                   swr/100, lot, er/100, ar_v/100, ap, mg, oc, act])
        cr = ws.max_row; ws.row_dimensions[cr].height=20
        ws.cell(row=cr,column=1).alignment=ac
        ws.cell(row=cr,column=2).alignment=al; ws.cell(row=cr,column=2).font=font_b
        ws.cell(row=cr,column=3).alignment=al
        cd=ws.cell(row=cr,column=4); cd.alignment=ac; cd.font=font_b
        cd.fill=fill_long if d=='LONG' else fill_short
        ws.cell(row=cr,column=5).alignment=ac
        ws.cell(row=cr,column=6).alignment=ac
        ws.cell(row=cr,column=7).number_format='#,##0.00'; ws.cell(row=cr,column=7).alignment=ar
        ws.cell(row=cr,column=8).alignment=ac
        ws.cell(row=cr,column=9).alignment=ac
        ws.cell(row=cr,column=10).number_format='#,##0.00'; ws.cell(row=cr,column=10).alignment=ar
        csw=ws.cell(row=cr,column=11); csw.number_format='0.0%'; csw.alignment=ac; csw.font=font_gold
        ws.cell(row=cr,column=12).number_format='#,##0'; ws.cell(row=cr,column=12).alignment=ar
        ce=ws.cell(row=cr,column=13); ce.number_format='0.00%'; ce.alignment=ar
        ca=ws.cell(row=cr,column=14); ca.number_format='0.00%'; ca.alignment=ar
        ca.font=font_grn if ar_v>0 else font_red; ca.fill=fill_win if ar_v>0 else fill_loss
        cpn=ws.cell(row=cr,column=15); cpn.number_format='#,##0.00'; cpn.alignment=ar
        cpn.font=font_grn if ap>0 else font_red; cpn.fill=fill_win if ap>0 else fill_loss
        ws.cell(row=cr,column=16).number_format='#,##0.00'; ws.cell(row=cr,column=16).alignment=ar
        co=ws.cell(row=cr,column=17); co.alignment=ac; co.font=font_b
        co.fill=fill_win if oc=='WIN' else fill_loss
        ws.cell(row=cr,column=18).alignment=al
        for ci in range(1,19): ws.cell(row=cr,column=ci).border=bdr
    autofit(ws)

os.makedirs('final_report', exist_ok=True)
for p in ['GitHub_Dashboard_17_Quarters_Master.xlsx','Nifty50_GitHub_Dashboard_Master.xlsx',
          'final_report/GitHub_Dashboard_17_Quarters_Master.xlsx']:
    try:
        wb.save(p); print(f'[OK] {os.path.abspath(p)}')
    except:
        wb.save(p.replace('.xlsx','_v3.xlsx')); print(f'[OK] {os.path.abspath(p.replace(".xlsx","_v3.xlsx"))}')
