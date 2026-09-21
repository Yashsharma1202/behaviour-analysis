import openpyxl
import pathlib
import sys

sys.stdout.reconfigure(errors='replace')

ROOT = pathlib.Path('D:/behaviour analysis')
fpath = ROOT / 'Nifty211_Best_Performing_Stocks_By_Quarter_Master.xlsx'

wb = openpyxl.load_workbook(fpath, data_only=True)

for q_sheet in ["Q1 Best Stocks (April)", "Q2 Best Stocks (July)", "Q3 Best Stocks (October)", "Q4 Best Stocks (January)"]:
    ws = wb[q_sheet]
    print(f"\n==========================================================================================")
    print(f"TOP 10 GOOD PERFORMING STOCKS IN {q_sheet.upper()} (WIN RATE >= 60% & AVG RET > 0%)")
    print(f"==========================================================================================")
    r = 8
    while True:
        val = ws.cell(row=r, column=1).value
        if val is None or r >= 18:
            break
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 13)]
        rank, sym, strat, status, n_q, win_str, wr, w_c, l_c, avg_r, margin, pnl = row_vals
        print(f"Rank #{rank:<2} | {sym:<12} | {strat:<13} | {status:<20} | Cycles: {n_q:<10} | Win Rate: {wr:<7} | Avg Ret: {avg_r:<7} | Net P&L: {pnl}")
        r += 1

print("\n==========================================================================================")
print("TOP 15 STOCKS IN SINGLE BEST QUARTER MATCHING MATRIX")
print("==========================================================================================")
ws_m = wb['Best Quarter Matrix']
for r in range(4, 19):
    row_vals = [ws_m.cell(row=r, column=c).value for c in range(1, 10)]
    sr, sym, bq, strat, status, n_q, wr, avg_r, pnl = row_vals
    print(f"#{sr:<2} | {sym:<12} | Best Qtr: {bq:<4} | {strat:<13} | Win Rate: {wr:<7} | Avg Ret: {avg_r:<7} | Net P&L: {pnl}")

print("==========================================================================================")
