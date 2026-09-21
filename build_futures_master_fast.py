import pandas as pd
import pathlib
import sys
import time
import numpy as np
from concurrent.futures import ThreadPoolExecutor
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(errors='replace')

t0 = time.time()

# ------------------- Paths -------------------
BASE_DIR = pathlib.Path('D:/behaviour analysis')
OI_ROOT = BASE_DIR / 'OI_DATA'
OPTIONS_MASTER_PATH = BASE_DIR / 'Nifty50_12_Quarters_Options_OI_Master_v4.xlsx'
RESULT_DATES_PATH = BASE_DIR / 'Nifty50Stocks_QtyResultDates.xlsx'

OUTPUT_FUTURES_ONLY = BASE_DIR / 'Futures_Master_Only.xlsx'
OUTPUT_FUTURES_MASTER = BASE_DIR / 'Nifty50_12_Quarters_Futures_OI_Master.xlsx'

# ------------------- 1. Load Result Dates -------------------
print("Loading Quarterly Result Dates...", flush=True)
result_date_map = {} # (symbol, quarter_str) -> date_str

if RESULT_DATES_PATH.exists():
    try:
        xl_rd = pd.ExcelFile(RESULT_DATES_PATH, engine='openpyxl')
        for sheet in xl_rd.sheet_names:
            df_rd = xl_rd.parse(sheet)
            sym_col = next((c for c in df_rd.columns if 'SYM' in str(c).upper()), None)
            date_col = next((c for c in df_rd.columns if 'DATE' in str(c).upper() or 'REPORTING' in str(c).upper()), None)
            if sym_col and date_col:
                for _, r in df_rd.iterrows():
                    sym = str(r[sym_col]).strip()
                    dt_val = r[date_col]
                    if pd.notnull(dt_val):
                        dt_str = pd.to_datetime(dt_val).strftime('%Y-%m-%d')
                        result_date_map[(sym, sheet)] = dt_str
    except Exception as e:
        print(f"Error loading result dates: {e}", flush=True)

print(f"Loaded {len(result_date_map)} quarterly result date mappings.", flush=True)

# ------------------- 2. Load Trades from Options Master -------------------
print("Loading Trades from Options Master...", flush=True)
try:
    df_trades = pd.read_excel(OPTIONS_MASTER_PATH, sheet_name='All_12Q_Options_Trades', engine='openpyxl')
except Exception as e:
    print(f"Error reading trades: {e}", flush=True)
    sys.exit(1)

print(f"Loaded {len(df_trades)} trade records.", flush=True)

# Helper to read parquet file quickly
def parse_parquet_file(args):
    stock, p_file = args
    date_str = p_file.stem
    try:
        df_p = pd.read_parquet(p_file)
        price = None
        if 'Time' in df_p.columns:
            row = df_p[df_p['Time'] == '15:29:59']
            if not row.empty:
                price = float(row.iloc[0]['Close'])
            else:
                price = float(df_p.iloc[-1]['Close'])
        else:
            price = float(df_p.iloc[-1]['Close'])
        
        exp_date = ''
        if 'ExpiryDate' in df_p.columns:
            u = df_p['ExpiryDate'].dropna().unique()
            if len(u) > 0:
                exp_date = str(u[0])
        return (stock, date_str, price, exp_date)
    except Exception:
        return None

# Single file price/expiry lookup for trade mapping
def load_futures_data(stock: str, date_str: str):
    stock_dir = OI_ROOT / stock
    fut_file = stock_dir / 'futures_parquet' / f"{date_str}.parquet"
    price, expiry = None, ''
    if fut_file.exists():
        res = parse_parquet_file((stock, fut_file))
        if res:
            _, _, price, expiry = res

    # Spot fallback if futures missing
    if price is None:
        spot_file = stock_dir / 'spot_parquet' / f"{date_str}.parquet"
        if spot_file.exists():
            try:
                df_s = pd.read_parquet(spot_file)
                if 'Time' in df_s.columns:
                    row = df_s[df_s['Time'] == '15:29:59']
                    if not row.empty:
                        price = float(row.iloc[0]['Close'])
                    else:
                        price = float(df_s.iloc[-1]['Close'])
                else:
                    price = float(df_s.iloc[-1]['Close'])
            except Exception:
                pass
    return price, expiry

# Process Trades
futures_trades = []
trade_lookup = {}

for idx, row in df_trades.iterrows():
    sym = str(row['Symbol']).strip()
    entry_dt = pd.to_datetime(row['Entry Date']).strftime('%Y-%m-%d')
    exit_dt = pd.to_datetime(row['Exit Date']).strftime('%Y-%m-%d')
    orig_strat = str(row['Strategy']).strip()
    strat = 'FUTURE LONG' if 'LONG' in orig_strat.upper() else 'FUTURE SHORT'
    direction = 1 if 'LONG' in strat else -1
    lot_size = row.get('Lot Size (Qty)', 1)
    if pd.isnull(lot_size) or lot_size == 0:
        lot_size = 1

    entry_price, entry_exp = load_futures_data(sym, entry_dt)
    exit_price, exit_exp = load_futures_data(sym, exit_dt)
    
    if entry_price is None:
        entry_price = 0.0
    if exit_price is None:
        exit_price = 0.0
    
    expiry_dt = entry_exp or exit_exp
    booked_pnl = direction * (exit_price - entry_price) * lot_size
    qtr = str(row['Quarter']).strip()
    
    result_dt = result_date_map.get((sym, qtr), '')
    
    trade_info = {
        'Quarter': qtr,
        'FY': row.get('FY', ''),
        'Trade No': row.get('Trade No', idx + 1),
        'Symbol': sym,
        'Company Name': row.get('Company Name', ''),
        'Sector': row.get('Sector', ''),
        'Strategy': strat,
        'Futures Expiry Date': expiry_dt,
        'Entry Date': entry_dt,
        'Exit Date': exit_dt,
        'Entry Futures Price (₹)': entry_price,
        'Exit Futures Price (₹)': exit_price,
        'Lot Size (Qty)': lot_size,
        'Booked Futures P&L (₹)': booked_pnl,
        'Expected Return (%)': row.get('Expected Return (%)', 0.0),
        'Return We Get (%)': row.get('Return We Get (%)', 0.0),
        'Quarterly Result Date': result_dt,
        'Re-entry Type': row.get('Re-entry Type', 'First Entry'),
        'Data Query Status': row.get('Data Query Status', 'Success')
    }
    futures_trades.append(trade_info)
    trade_lookup[(sym, entry_dt)] = trade_info
    trade_lookup[(sym, exit_dt)] = trade_info

df_fut_trades = pd.DataFrame(futures_trades)
print(f"Processed {len(df_fut_trades)} Futures trades.", flush=True)

# ------------------- 3. Parallel Scanning of OI_DATA -------------------
print("Gathering parquet files...", flush=True)
file_tasks = []
for stock_dir in sorted(OI_ROOT.iterdir()):
    if not stock_dir.is_dir():
        continue
    stock = stock_dir.name
    fut_folder = stock_dir / 'futures_parquet'
    if not fut_folder.is_dir():
        continue
    for p_file in sorted(fut_folder.glob('*.parquet')):
        file_tasks.append((stock, p_file))

print(f"Parsing {len(file_tasks)} files with 32 workers...", flush=True)
daily_records = []

with ThreadPoolExecutor(max_workers=32) as executor:
    results = executor.map(parse_parquet_file, file_tasks)
    for res in results:
        if res is None:
            continue
        stock, date_str, price, exp_date = res
        t_match = trade_lookup.get((stock, date_str))
        
        daily_records.append({
            'Symbol': stock,
            'Date': date_str,
            'Futures Close Price (₹)': price,
            'Futures Expiry Date': exp_date,
            'Position': t_match['Strategy'] if t_match else '',
            'Entry Date': t_match['Entry Date'] if t_match else '',
            'Exit Date': t_match['Exit Date'] if t_match else '',
            'Expected Return (%)': t_match['Expected Return (%)'] if t_match else '',
            'Return We Get (%)': t_match['Return We Get (%)'] if t_match else '',
            'Booked P&L (₹)': t_match['Booked Futures P&L (₹)'] if t_match else '',
            'Quarterly Result Date': t_match['Quarterly Result Date'] if t_match else ''
        })

df_daily_master = pd.DataFrame(daily_records)
df_daily_master = df_daily_master.sort_values(by=['Symbol', 'Date']).reset_index(drop=True)
print(f"Total daily records built: {len(df_daily_master)} in {time.time() - t0:.1f}s", flush=True)

# ------------------- 4. Fast Export using pandas ExcelWriter -------------------
def save_fast_excel(out_path):
    t_start = time.time()
    print(f"Exporting to {out_path.name}...", flush=True)
    with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
        df_daily_master.to_excel(writer, sheet_name='Futures_Master', index=False)
        df_fut_trades.to_excel(writer, sheet_name='All_12Q_Futures_Trades', index=False)
    
    # Format header styles cleanly
    wb = load_workbook(out_path)
    header_font = Font(bold=True, color='FFFFFFFF')
    fill_teal = PatternFill(start_color='FF006688', end_color='FF006688', fill_type='solid')
    fill_navy = PatternFill(start_color='FF1F497D', end_color='FF1F497D', fill_type='solid')

    ws_daily = wb['Futures_Master']
    ws_daily.freeze_panes = 'A2'
    for cell in ws_daily[1]:
        cell.font = header_font
        cell.fill = fill_teal
        cell.alignment = Alignment(horizontal='center')
    for c_idx in range(1, len(df_daily_master.columns) + 1):
        col_name = str(df_daily_master.columns[c_idx - 1])
        ws_daily.column_dimensions[get_column_letter(c_idx)].width = max(len(col_name) + 4, 14)

    ws_trades = wb['All_12Q_Futures_Trades']
    ws_trades.freeze_panes = 'A2'
    for cell in ws_trades[1]:
        cell.font = header_font
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal='center')
    for c_idx in range(1, len(df_fut_trades.columns) + 1):
        col_name = str(df_fut_trades.columns[c_idx - 1])
        ws_trades.column_dimensions[get_column_letter(c_idx)].width = max(len(col_name) + 4, 14)

    wb.save(out_path)
    print(f"✅ Finished {out_path.name} in {time.time() - t_start:.1f}s", flush=True)

save_fast_excel(OUTPUT_FUTURES_ONLY)
save_fast_excel(OUTPUT_FUTURES_MASTER)

print(f"🎉 ALL DONE in total {time.time() - t0:.1f}s!", flush=True)
