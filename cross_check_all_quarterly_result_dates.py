import os
import pathlib
import sys
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(errors='replace')

ROOT = pathlib.Path('D:/behaviour analysis')
OI_DIR = ROOT / 'OI_DATA'
PROC = ROOT / 'processed'
PRICE_CACHE = PROC / 'price_cache'

# Discover all 211 symbols
SYMBOLS = sorted([e.name.strip().upper() for e in os.scandir(OI_DIR) if e.is_dir()])

print("==========================================================================================")
print("CROSS-CHECKING QUARTERLY RESULT DATES ACROSS ALL QUARTERS FOR ALL 211 STOCKS")
print("==========================================================================================")
print(f"Total Target Universe: {len(SYMBOLS)} Stocks")

result_audit_rows = []

for idx, sym in enumerate(SYMBOLS, 1):
    stock_folder = ROOT / sym
    fr_path = stock_folder / 'financial_results.csv'
    price_path = PRICE_CACHE / f"{sym}.csv"
    
    # Load prices to verify trading dates
    valid_dates = set()
    if price_path.exists():
        try:
            df_p = pd.read_csv(price_path)
            valid_dates = set(pd.to_datetime(df_p['date']).dt.strftime('%Y-%m-%d').tolist())
        except Exception:
            pass
            
    n_results = 0
    earliest_date = "N/A"
    latest_date = "N/A"
    matching_trading_days = 0
    
    if fr_path.exists() and fr_path.stat().st_size > 10:
        try:
            df_fr = pd.read_csv(fr_path, dtype=str).fillna('')
            if not df_fr.empty and 'broadCastDate' in df_fr.columns:
                dts = pd.to_datetime(df_fr['broadCastDate'], errors='coerce').dropna().sort_values()
                if not dts.empty:
                    n_results = len(dts)
                    earliest_date = dts.iloc[0].strftime('%Y-%m-%d')
                    latest_date = dts.iloc[-1].strftime('%Y-%m-%d')
                    
                    for dt in dts:
                        d_str = dt.strftime('%Y-%m-%d')
                        if d_str in valid_dates:
                            matching_trading_days += 1
                            
                    # Detailed row output per quarter date
                    for _, r in df_fr.iterrows():
                        raw_dt = pd.to_datetime(r.get('broadCastDate', ''), errors='coerce')
                        if pd.notna(raw_dt):
                            dt_str = raw_dt.strftime('%Y-%m-%d')
                            is_valid_trading = dt_str in valid_dates
                            result_audit_rows.append({
                                "sym": sym,
                                "result_date": dt_str,
                                "relating_to": r.get('relatingTo', 'Financial Results'),
                                "audited_status": r.get('audited', 'Audited'),
                                "from_date": r.get('fromDate', ''),
                                "to_date": r.get('toDate', ''),
                                "is_trading_day": "YES ✅" if is_valid_trading else "WEEKEND / HOLIDAY ⚠️"
                            })
        except Exception:
            pass

print(f"Total Quarterly Result Broadcast Dates Analyzed: {len(result_audit_rows)}")

# Export Audit Workbook to Excel
out_excel = ROOT / 'Verified_All_Quarters_Result_Dates_211_Stocks.xlsx'
wb = openpyxl.Workbook()

# Sheet 1: Detailed Quarterly Dates Log
ws1 = wb.active
ws1.title = "All Quarters Result Dates"
ws1.views.sheetView[0].showGridLines = True

hdr_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1F4E79")
border_thin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                     top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))

headers1 = [
    "Sr. No.", "Stock Symbol", "Result Announcement Date", "Financial Period / Description",
    "Audit Status", "Period From Date", "Period To Date", "Trading Day Verified"
]

ws1.row_dimensions[1].height = 28
for c_idx, h in enumerate(headers1, 1):
    cell = ws1.cell(row=1, column=c_idx, value=h)
    cell.font = hdr_font
    cell.fill = hdr_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")

for r_idx, r in enumerate(result_audit_rows, start=2):
    ws1.cell(row=r_idx, column=1, value=r_idx - 1).alignment = Alignment(horizontal="center")
    ws1.cell(row=r_idx, column=2, value=r['sym'])
    ws1.cell(row=r_idx, column=3, value=r['result_date']).alignment = Alignment(horizontal="center")
    ws1.cell(row=r_idx, column=4, value=r['relating_to'])
    ws1.cell(row=r_idx, column=5, value=r['audited_status']).alignment = Alignment(horizontal="center")
    ws1.cell(row=r_idx, column=6, value=r['from_date']).alignment = Alignment(horizontal="center")
    ws1.cell(row=r_idx, column=7, value=r['to_date']).alignment = Alignment(horizontal="center")
    
    c_tr = ws1.cell(row=r_idx, column=8, value=r['is_trading_day'])
    c_tr.alignment = Alignment(horizontal="center")
    
    for c in range(1, 9):
        ws1.cell(row=r_idx, column=c).border = border_thin

for col in ws1.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

wb.save(out_excel)

print("==========================================================================================")
print(f"✅ SUCCESSFULLY SAVED AUDIT EXCEL: {out_excel.name}")
print("==========================================================================================")
