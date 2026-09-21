import pandas as pd
import pathlib
import sys
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

sys.stdout.reconfigure(encoding='utf-8')

# ------------------- Configuration -------------------
BASE_DIR = pathlib.Path('D:/behaviour analysis/OI_DATA')   # OI/Futures data root
WORKSPACE_ROOT = pathlib.Path('d:/behaviour analysis')    # where quarter Excel files live
OUTPUT_PATH = pathlib.Path('D:/behaviour analysis/Future_Synthetic_Master.xlsx')

# ------------------- Helper functions -------------------

def latest_parquet(folder: pathlib.Path):
    """Return the newest .parquet file in a folder, or None."""
    files = sorted(folder.glob('*.parquet'))
    return files[-1] if files else None

def get_close_price(df: pd.DataFrame) -> float:
    """Return the closing price at 15:29:59 if present, else the last row."""
    row = df[df['Time'] == '15:29:59']
    if row.empty:
        row = df.sort_values('Time').tail(1)
    return float(row.iloc[0]['Close'])

def load_price(stock: str, date_str: str):
    """Load futures price for a stock on a given date.
    Prefers futures_parquet; falls back to spot_parquet.
    Returns (price, expiry_date) – expiry_date may be empty string.
    """
    stock_dir = BASE_DIR / stock
    # Try futures first
    fut_dir = stock_dir / 'futures_parquet'
    price = None
    expiry = ''
    if fut_dir.exists():
        fpath = fut_dir / f"{date_str}.parquet"
        if fpath.exists():
            df = pd.read_parquet(fpath)
            price = get_close_price(df)
            if 'ExpiryDate' in df.columns:
                expiry = sorted(df['ExpiryDate'].unique())[0] if df['ExpiryDate'].nunique() else ''
    # Fall back to spot
    if price is None:
        spot_dir = stock_dir / 'spot_parquet'
        fpath = spot_dir / f"{date_str}.parquet"
        if fpath.exists():
            df = pd.read_parquet(fpath)
            price = get_close_price(df)
    return price, expiry

# ------------------- Load trade data -------------------
quarter_files = sorted(WORKSPACE_ROOT.glob('*_Combined_Best_Capital_Utilisation.xlsx'))
if not quarter_files:
    print('⚠️ No quarter Excel files found.')
    sys.exit(1)

trade_frames = []
for qf in quarter_files:
    try:
        df = pd.read_excel(qf, sheet_name='Detailed_Trades')
    except Exception as e:
        print(f'⚠️ Could not read Detailed_Trades from {qf.name}: {e}')
        continue
    required = ['Symbol','Strategy','Entry Date','Exit Date','Expected Return (%)','Return We Get (%)','Quarter']
    missing = [c for c in required if c not in df.columns]
    if missing:
        print(f'⚠️ Missing columns in {qf.name}: {missing}')
        continue
    trade_frames.append(df[required])

if not trade_frames:
    print('❌ No valid trade data loaded – aborting.')
    sys.exit(1)

trades_df = pd.concat(trade_frames, ignore_index=True)

# ------------------- Build master rows -------------------
rows = []
for _, row in trades_df.iterrows():
    stock = str(row['Symbol']).strip()
    strategy = str(row['Strategy']).strip().upper()   # LONG / SHORT
    entry_date = pd.to_datetime(row['Entry Date']).strftime('%Y-%m-%d')
    exit_date = pd.to_datetime(row['Exit Date']).strftime('%Y-%m-%d')
    quarter = str(row['Quarter'])
    exp_ret = row.get('Expected Return (%)')
    ret_we_get = row.get('Return We Get (%)')

    entry_price, entry_expiry = load_price(stock, entry_date)
    exit_price, exit_expiry = load_price(stock, exit_date)
    if entry_price is None or exit_price is None:
        # skip if price data unavailable
        continue
    direction = 1 if 'LONG' in strategy else -1
    booked_pnl = direction * (exit_price - entry_price)
    # Use the expiry from the entry side if available, otherwise exit side
    expiry_date = entry_expiry or exit_expiry
    # Quarterly result date – we use the exit date as a placeholder
    quarterly_result_date = exit_date

    rows.append([
        quarter,
        stock,
        strategy,
        entry_date,
        exit_date,
        entry_price,
        exit_price,
        expiry_date,
        exp_ret,
        ret_we_get,
        booked_pnl,
        quarterly_result_date
    ])

# ------------------- Write to Excel -------------------
wb = Workbook()
ws = wb.active
ws.title = 'Futures_Synthetic_Master'

headers = [
    'Quarter',
    'Symbol',
    'Position',            # LONG / SHORT
    'Entry Date',
    'Exit Date',
    'Entry Futures Price (₹)',
    'Exit Futures Price (₹)',
    'Futures Expiry Date',
    'Expected Return (%)',
    'Return We Get (%)',
    'Booked P&L (₹)',
    'Quarterly Result Date'
]
ws.append(headers)
# Header styling (dark teal background, white bold text)
header_font = Font(bold=True, color='FFFFFFFF')
header_fill = PatternFill(start_color='FF006688', end_color='FF006688', fill_type='solid')
for col in range(1, len(headers)+1):
    cell = ws.cell(row=1, column=col)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center')

for r in rows:
    ws.append(r)

# Auto‑size columns for readability
for col in ws.columns:
    max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
    ws.column_dimensions[col[0].column_letter].width = max_len + 3

# Save with fallback versions similar to earlier scripts
def safe_save(wb_obj, path):
    try:
        wb_obj.save(path)
        return True
    except PermissionError:
        for i in range(2, 6):
            versioned = path.with_name(f"{path.stem}_v{i}{path.suffix}")
            try:
                wb_obj.save(versioned)
                print(f'⚠️ Permission error on {path.name}; saved as {versioned.name}')
                return True
            except PermissionError:
                continue
        print(f'❌ Could not save {path.name} after multiple attempts.')
        return False

saved = safe_save(wb, OUTPUT_PATH)
if saved:
    print(f'✅ Futures‑synthetic master Excel created at: {OUTPUT_PATH}')
    print(f'Rows written (valid trades): {len(rows)}')
else:
    print('❌ Failed to write the master workbook.')
