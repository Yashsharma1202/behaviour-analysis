import openpyxl, pandas as pd, sys

sys.stdout.reconfigure(errors='replace')

summary_path = r'D:/behaviour analysis/Nifty50_12_Quarters_Futures_Combined_Turnover_0.5PCT_Cost_Summary.xlsx'
wb = openpyxl.load_workbook(summary_path, data_only=True)
ws = wb['Futures_Comb_Turnover_0.5_Sum']

print("==========================================================================================")
print("VERIFYING COMBINED BUY+SELL TURNOVER (0.5% COST) SUMMARY WORKBOOK")
print("==========================================================================================")
print("Sheet Names:", wb.sheetnames)

print("\nSummary Table Rows 7-21:")
for r in range(7, 22):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 13)]
    print(f"  Row {r:2d}: {vals}")
print("==========================================================================================")
