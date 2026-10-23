"""
Top 5 / Top 10 / Top 15 stocks per calendar date (22-Oct..31-Oct) using the project 8x8 grid logic, plus a walk-forward
backtest for each portfolio size (stocks picked from PRIOR years only, traded in the test year, net of COST per trade).
Ranking = project rule: best 8x8 rule per stock (win rate, then avg return), stock needs >= TOP_MIN_EVENTS events and avg>0,
stocks ranked by (win %, avg %). Top 5 is nested inside Top 10 inside Top 15.
Output: Calendar_Top5_10_15_22Oct_31Oct.xlsx (new file).
"""
import sys, glob, os, pathlib
import numpy as np, pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
sys.path.insert(0, 'D:/behaviour analysis'); sys.stdout.reconfigure(encoding='utf-8')
import build_calendar_backtest_8x8 as b

OUT = b.ROOT / 'Calendar_Top5_10_15_22Oct_31Oct.xlsx'
SIZES = [5, 10, 15]
COST, TRAIN_MIN, FIRST, TOP_MIN_EVENTS = 0.15, 10, 2008, 15

px = {}
for f in glob.glob(b.DATA.as_posix() + '/*.csv'):
    s = os.path.basename(f)[:-4]
    if s in ('FINNIFTY', 'MIDCPNIFTY', 'SENSEX', 'TATAMOTORS') or s in b.INDICES: continue
    try: p = b.load(s)
    except Exception: continue
    if len(p) >= 250: px[s] = p

def side_txt(sd): return 'LONG' if sd == 1 else 'SHORT'

def analyse(m, d, label):
    per = {}
    for s, p in px.items():
        sm = b.samples(p, m, d, s, label)
        if sm: per[s] = sm
    # ---- current ranking (all history) ----
    cur = []
    for s, (ys, R, cals, _) in per.items():
        if len(ys) < TOP_MIN_EVENTS: continue
        k, win, avg, sd = b.best(R)
        if avg <= 0: continue
        r = R[:, k] * sd
        ins = ys <= b.WF_SPLIT
        wfw = wfa = np.nan
        if ins.sum() >= b.MIN_EVENTS and (~ins).sum():
            bw = b.best(R[ins])
            if bw:
                ro = R[~ins, bw[0]] * bw[3]; wfw, wfa = (ro > 0).mean() * 100, ro.mean()
        cur.append(dict(Symbol=s, Side=side_txt(sd), N=int(b.PN[k]), M=int(b.PM[k]), Window=f'T-{b.PN[k]} to T+{b.PM[k]}',
                        Events=len(ys), Win=win, Avg=avg, Worst=r.min(), Best=r.max(), WFwin=wfw, WFavg=wfa))
    cur.sort(key=lambda x: (-x['Win'], -x['Avg']))
    # ---- walk-forward ----
    wf = []
    for Y in range(FIRST, 2026):
        cand = []
        for s, (ys, R, cals, _) in per.items():
            idx = np.where(ys == Y)[0]
            if len(idx) == 0: continue
            tr = ys < Y
            if tr.sum() < TRAIN_MIN: continue
            k, win, avg, sd = b.best(R[tr])
            if avg <= 0: continue
            cand.append((win, avg, s, R[idx[0], k] * sd))
        if len(cand) < max(SIZES): continue
        cand.sort(key=lambda x: (-x[0], -x[1]))
        row = {'Year': Y}
        for n in SIZES:
            rr = [c[3] for c in cand[:n]]
            row[f'Top{n}_net'] = np.mean(rr) - COST
            row[f'Top{n}_hit'] = np.mean([x > 0 for x in rr]) * 100
        row['Top10_picks'] = ', '.join(f"{c[2]}({c[3]:+.1f})" for c in cand[:10])
        wf.append(row)
    return cur, pd.DataFrame(wf)

def main():
    res = {}
    for d in range(22, 32):
        label = f'{d:02d}-Oct'
        res[label] = analyse(10, d, label)
        print(label, 'ranked stocks:', len(res[label][0]), '| WF years:', len(res[label][1]), flush=True)
    hdr = PatternFill('solid', fgColor='1F3864'); hf = Font(bold=True, color='FFFFFF')
    GREEN = {5: 'A9D08E', 10: 'C6EFCE', 15: 'E2F0D9'}
    sec = PatternFill('solid', fgColor='D9E1F2')
    def style_header(ws, row, n):
        for c in range(1, n + 1):
            x = ws.cell(row, c); x.fill = hdr; x.font = hf; x.alignment = Alignment(wrap_text=True, vertical='center')
    # summary of walk-forward by size
    S = []
    for label, (cur, wf) in res.items():
        for n in SIZES:
            v = wf[f'Top{n}_net']
            S.append(dict(Date=label, Portfolio=f'Top {n}', TestYears=len(v), WinYearsPct=round((v > 0).mean() * 100, 1), AvgNetPct=round(v.mean(), 2),
                          MedianNetPct=round(v.median(), 2), WorstYearPct=round(v.min(), 2), BestYearPct=round(v.max(), 2),
                          StockHitPct=round(wf[f'Top{n}_hit'].mean(), 1), Last5yrAvgPct=round(v.tail(5).mean(), 2)))
    S = pd.DataFrame(S)
    with pd.ExcelWriter(OUT, engine='openpyxl') as xw:
        pd.DataFrame({'Top 5 / Top 10 / Top 15 stocks per date, 22-Oct..31-Oct': [
            'Per stock: best of the project 8x8 rules (buy T-N, sell T+M, N,M=1..8; LONG or SHORT), T = first trading day on/after the calendar date.',
            f'Ranking = win % then avg %, stock needs >= {TOP_MIN_EVENTS} events and positive avg. Top 5 is the first 5 of the Top 10 / Top 15 (nested).',
            'The date-sheet lists use ALL history, so their Win % / Avg % are in-sample (flattering).',
            f'Walk-forward backtest (the honest part): for each year {FIRST}-2025 stocks and rules are chosen from PRIOR years only (>= {TRAIN_MIN} prior events),',
            f'the top 5/10/15 are traded equal-weight in that year, net of {COST}% per position. Compare the sizes in Summary_WalkForward.',
            'Price moves only (no slippage/tax); dates overlap in time so the 10 dates are not independent evidence. Early Yahoo glitches (>60% moves) are excluded.']}).to_excel(xw, sheet_name='ReadMe', index=False)
        S.to_excel(xw, sheet_name='Summary_WalkForward', index=False)
        ws = xw.sheets['Summary_WalkForward']; style_header(ws, 1, S.shape[1]); ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
        for c in range(1, S.shape[1] + 1): ws.column_dimensions[get_column_letter(c)].width = 16
        # all lists in one sheet
        L = []
        for label, (cur, wf) in res.items():
            for i, t in enumerate(cur[:15], 1):
                L.append(dict(Date=label, Rank=i, In_Top5='Y' if i <= 5 else '', In_Top10='Y' if i <= 10 else '', In_Top15='Y', Symbol=t['Symbol'], Side=t['Side'],
                              Buy_N_days_before=t['N'], Sell_M_days_after=t['M'], Window=t['Window'], Events=t['Events'], Win_pct=round(t['Win'], 1),
                              Avg_pct=round(t['Avg'], 2), Worst_pct=round(t['Worst'], 2),
                              WF_win_2019plus=None if t['WFwin'] != t['WFwin'] else round(t['WFwin'], 1),
                              WF_avg_2019plus=None if t['WFavg'] != t['WFavg'] else round(t['WFavg'], 2)))
        L = pd.DataFrame(L); L.to_excel(xw, sheet_name='All_Lists', index=False)
        ws = xw.sheets['All_Lists']; style_header(ws, 1, L.shape[1]); ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for c in range(1, L.shape[1] + 1): ws.column_dimensions[get_column_letter(c)].width = 15
        for label, (cur, wf) in res.items():
            ws = xw.book.create_sheet(label); xw.sheets[label] = ws; r = 1
            ws.cell(r, 1, f'{label}  -  Top 5 / Top 10 / Top 15 stocks  (T = first trading day on/after {label}; buy T-N, sell T+M)').font = Font(bold=True, size=14); r += 2
            ws.cell(r, 1, 'WALK-FORWARD BACKTEST (stocks chosen from prior years only, net of cost)').font = Font(bold=True)
            for c in range(1, 12): ws.cell(r, c).fill = sec
            r += 1
            for i, h in enumerate(['Portfolio', 'Test years', 'Winning years %', 'Avg net % / yr', 'Median %', 'Worst year %', 'Best year %', 'Single-stock hit %', 'Last-5-yr avg %'], 1): ws.cell(r, i, h)
            style_header(ws, r, 9); r += 1
            for _, x in S[S.Date == label].iterrows():
                for i, v in enumerate([x.Portfolio, x.TestYears, x.WinYearsPct, x.AvgNetPct, x.MedianNetPct, x.WorstYearPct, x.BestYearPct, x.StockHitPct, x.Last5yrAvgPct], 1): ws.cell(r, i, v)
                r += 1
            r += 1
            for n in SIZES:
                ws.cell(r, 1, f'TOP {n} STOCKS for {label}  (rank by Win % then Avg %; in-sample on all history)').font = Font(bold=True, size=12)
                for c in range(1, 12): ws.cell(r, c).fill = PatternFill('solid', fgColor=GREEN[n])
                r += 1
                th = ['#', 'Symbol', 'Side', 'Buy N days BEFORE', 'Sell M days AFTER', 'Window', 'Events', 'Win %', 'Avg %', 'Worst %', 'WF win % 2019+', 'WF avg % 2019+']
                for i, h in enumerate(th, 1): ws.cell(r, i, h)
                style_header(ws, r, len(th)); r += 1
                for i, t in enumerate(cur[:n], 1):
                    vals = [i, t['Symbol'], t['Side'], t['N'], t['M'], t['Window'], t['Events'], round(t['Win'], 1), round(t['Avg'], 2), round(t['Worst'], 2),
                            None if t['WFwin'] != t['WFwin'] else round(t['WFwin'], 1), None if t['WFavg'] != t['WFavg'] else round(t['WFavg'], 2)]
                    for j, v in enumerate(vals, 1):
                        x = ws.cell(r, j, v); x.fill = PatternFill('solid', fgColor=GREEN[n])
                        if i <= 5: x.font = Font(bold=True)
                    r += 1
                r += 1
            ws.cell(r, 1, 'YEAR-BY-YEAR WALK-FORWARD RETURNS (net %, equal-weight)').font = Font(bold=True)
            for c in range(1, 12): ws.cell(r, c).fill = sec
            r += 1
            for i, h in enumerate(['Year', 'Top 5 %', 'Top 10 %', 'Top 15 %', 'Top 10 stock hit %', 'Top 10 picks (return % that year)'], 1): ws.cell(r, i, h)
            style_header(ws, r, 6); r += 1
            for _, x in wf.iterrows():
                for i, v in enumerate([int(x.Year), round(x.Top5_net, 2), round(x.Top10_net, 2), round(x.Top15_net, 2), round(x.Top10_hit, 0), x.Top10_picks], 1): ws.cell(r, i, v)
                r += 1
            ws.column_dimensions['A'].width = 16
            for c in range(2, 13): ws.column_dimensions[get_column_letter(c)].width = 15
    print('saved', OUT)
    pd.set_option('display.width', 250)
    print(S.pivot(index='Date', columns='Portfolio', values='AvgNetPct').to_string())
    print(S.pivot(index='Date', columns='Portfolio', values='WinYearsPct').to_string())
    print(S.pivot(index='Date', columns='Portfolio', values='WorstYearPct').to_string())
    print(S.groupby('Portfolio')[['AvgNetPct', 'WinYearsPct', 'WorstYearPct', 'StockHitPct']].mean().round(2))
    t = pd.DataFrame(res['30-Oct'][0][:15])[['Symbol', 'Side', 'Window', 'Events', 'Win', 'Avg']].round(1)
    print(t.to_string(index=False))

if __name__ == '__main__':
    main()
