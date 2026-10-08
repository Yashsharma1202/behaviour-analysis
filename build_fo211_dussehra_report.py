"""
ALL 211 F&O STOCKS — DUSSEHRA, YEAR-WISE, RANK-WISE REPORT (Excel + PDF).
Builds the final deliverable from the validated intermediate results computed
by build_fo211_dussehra_yearwise.py (run that first).
ADDITIVE: reads the intermediate pickle; writes Excel + HTML/PDF report files.
"""
import json, pickle, os, subprocess
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

import os as _os
BASE = r'D:\behaviour analysis'
OUT_XLSX = _os.environ.get('FO211_OUT_XLSX', BASE + r'\FO211_Dussehra_YearWise_Ranked_Report.xlsx')
OUT_HTML = BASE + r'\FO211_Dussehra_YearWise_Ranked_Report.html'
OUT_PDF = BASE + r'\FO211_Dussehra_YearWise_Ranked_Report.pdf'

data = pickle.load(open(BASE + r'\scraped_parquet\_tmp_fo211_dussehra.pkl', 'rb'))
results, meta = data['results'], data['meta']
qualified = {s: r for s, r in results.items() if r['pick']}
insufficient = {s: r for s, r in results.items() if not r['pick']}  # n=0, truly no data
ranked = sorted(qualified.items(), key=lambda kv: -(kv[1]['pick_avg'] if kv[1]['pick_avg'] is not None else -999))
ALL_YEARS = sorted(set(y for r in qualified.values() for y in r['year_rets']))
n_low_conf = sum(1 for _, r in ranked if r.get('low_confidence'))

# ---- upcoming Dussehra 2026 position-taking dates (same calendar/shift method
# the live dashboard was built with) \u2014 for the 50 Nifty-50 overlap, reuse the
# live dashboard's own stored dates directly (zero risk of divergence); for the
# other 161, compute via the identical trading-day-shift algorithm. ----
NSE_2026_HOLIDAYS = {
    # '2026-03-04' removed (not a real NSE holiday -- leftover from Holi's date
    # once being wrongly typed as 04-Mar; real Holi holiday is 03-Mar, kept
    # below). '2026-11-08'/'2026-11-10' (Diwali) added -- were missing entirely.
    '2026-01-26', '2026-03-03', '2026-03-26', '2026-03-31', '2026-04-03',
    '2026-04-14', '2026-05-01', '2026-05-28', '2026-06-26', '2026-09-14', '2026-10-02',
    '2026-10-20', '2026-11-08', '2026-11-10', '2026-11-24', '2026-12-25',
}
def is_trading_day(d):
    return d.weekday() < 5 and d.isoformat() not in NSE_2026_HOLIDAYS
def shift_trading_days(anchor, n):
    d = anchor
    step = 1 if n >= 0 else -1
    cnt = abs(n)
    while cnt > 0:
        d += timedelta(days=step)
        if is_trading_day(d): cnt -= 1
    return d
DUSSEHRA_ANCHOR_2026 = date(2026, 10, 20)

_live = json.load(open(BASE + r'\dashboard_data\event_dashboard_data.json', encoding='utf-8'))
_live_dussehra = [h for h in _live['holidays'] if h['id'] == 'dussehra_2026'][0]
LIVE_POSITIONS = {s['symbol']: s for s in _live_dussehra['stocks']}

def position_for(sym, r):
    if sym in LIVE_POSITIONS:
        s = LIVE_POSITIONS[sym]
        return dict(direction=s['direction'], entry=s['entry_date'], exit=s['exit_date'], lot=s.get('lot'), src='live-dashboard')
    en = shift_trading_days(DUSSEHRA_ANCHOR_2026, -r['lead'])
    ex = shift_trading_days(DUSSEHRA_ANCHOR_2026, r['hold'])
    return dict(direction=r['pick'], entry=en.strftime('%d-%b-%Y'), exit=ex.strftime('%d-%b-%Y'), lot=r.get('lot'), src='computed')

NAVY = PatternFill('solid', fgColor='1A202C'); WHITE = Font(bold=True, color='FFFFFF', size=9)
TITLE = Font(bold=True, color='1A202C', size=13); SUB = Font(color='4A5568', size=10, italic=True)
GF = PatternFill('solid', fgColor='E6F4EA'); RF = PatternFill('solid', fgColor='FCE8E6')
AM = PatternFill('solid', fgColor='FEF3C7'); LBL = PatternFill('solid', fgColor='F1F5F9')
GOOD = Font(color='137333', bold=True); BAD = Font(color='C5221F', bold=True); BOLD = Font(bold=True)
THIN = Border(*[Side(style='thin', color='D9E2EC')] * 4)
CEN = Alignment(horizontal='center'); RIGHT = Alignment(horizontal='right'); WRAP = Alignment(wrap_text=True, vertical='top')

wb = Workbook()

# ============ Sheet 1: Ranked Summary ============
rk = wb.active; rk.title = 'Ranked_Summary'
rk.merge_cells('A1:K1')
rk.cell(1, 1, 'ALL F&O STOCKS \u2014 DUSSEHRA, RANKED BY AVG RETURN (best direction per stock)').font = TITLE
rk.merge_cells('A2:L2')
rk.cell(2, 1, f'{len(qualified)} of {len(results)} F&O stocks have real data ({n_low_conf} low-confidence, n<3 yrs \u2014 highlighted amber). {len(insufficient)} have zero data (see Insufficient_Data sheet).').font = SUB
hdr = ['Rank', 'Symbol', 'Company', 'Sector', 'Direction', 'Win Rate %', 'Avg Return %', 'Years', 'Confidence', 'Data Source', 'Window', 'Lot Size']
for j, h in enumerate(hdr, 1):
    c = rk.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
row = 5
for i, (sym, r) in enumerate(ranked, 1):
    m = meta.get(sym, {})
    avg = r['pick_avg']
    low = r.get('low_confidence')
    vals = [i, sym, m.get('name', sym), m.get('sector', ''), r['pick'], r['pick_wr'],
            (round(avg, 2) if avg is not None else '\u2014'), r['n'],
            ('LOW (n<3)' if low else 'OK'), r['source'],
            f"T-{r['lead']} to T+{r['hold']}", r.get('lot', '')]
    fill = AM if low else (GF if (avg is None or avg >= 0) else RF)
    for j, v in enumerate(vals, 1):
        c = rk.cell(row, j, v); c.border = THIN; c.fill = fill
        if j in (1, 5, 8, 9): c.alignment = CEN
        if j in (6, 7, 12): c.alignment = RIGHT
        if j == 7 and isinstance(v, (int, float)): c.font = GOOD if v >= 0 else BAD
        if j == 9 and low: c.font = Font(color='92400E', bold=True, size=9)
    row += 1
for j, w in enumerate([6, 12, 26, 22, 10, 11, 13, 7, 12, 30, 13, 9], 1):
    rk.column_dimensions[get_column_letter(j)].width = w
rk.freeze_panes = 'A5'

# ============ Sheet 2: Year-by-Year detail (Stock x Year) ============
yw = wb.create_sheet('Year_by_Year_Detail')
yw.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4 + len(ALL_YEARS))
yw.cell(1, 1, 'STOCK x YEAR \u2014 Dussehra return %% using each stock\u2019s chosen direction').font = TITLE
yw.merge_cells(start_row=2, start_column=1, end_row=2, end_column=5 + len(ALL_YEARS))
yw.cell(2, 1, 'Blank cell = no data that year for that stock (varies by listing date / data source). Amber row = low-confidence (n<3 years). Ranked by avg return.').font = SUB
hdr = ['Rank', 'Symbol', 'Direction', 'Avg %', 'Yrs'] + [str(y) for y in ALL_YEARS]
for j, h in enumerate(hdr, 1):
    c = yw.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
row = 5
for i, (sym, r) in enumerate(ranked, 1):
    sign = 1 if r['pick'] == 'LONG' else -1
    avg = r['pick_avg']; low = r.get('low_confidence')
    vals = [i, sym, r['pick'], (round(avg, 2) if avg is not None else '\u2014'), r['n']]
    for y in ALL_YEARS:
        raw = r['year_rets'].get(y)
        vals.append(round(raw * sign, 2) if raw is not None else '')
    rowfill = AM if low else None
    for j, v in enumerate(vals, 1):
        c = yw.cell(row, j, v); c.border = THIN
        if rowfill: c.fill = rowfill
        if j in (1, 3, 5): c.alignment = CEN
        elif j >= 4:
            c.alignment = RIGHT
            if isinstance(v, (int, float)) and j != 5:
                c.font = GOOD if v >= 0 else BAD
    row += 1
for j, w in enumerate([6, 12, 9, 8, 6] + [8] * len(ALL_YEARS), 1):
    yw.column_dimensions[get_column_letter(j)].width = w
yw.freeze_panes = 'E5'

# ============ Sheet 3: Insufficient-data stocks (transparency) ============
ins = wb.create_sheet('Insufficient_Data_Stocks')
ins.merge_cells('A1:D1')
ins.cell(1, 1, f'{len(insufficient)} F&O stocks with ZERO Dussehra-window price data \u2014 genuinely nothing to compute (not ranked, not fabricated)').font = TITLE
hdr = ['Symbol', 'Company', 'Years Available', 'Reason']
for j, h in enumerate(hdr, 1):
    c = ins.cell(3, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
row = 4
for sym, r in sorted(insufficient.items()):
    m = meta.get(sym, {})
    vals = [sym, m.get('name', sym), r['n'], 'Recently listed F&O contract \u2014 no trading days overlap any historical Dussehra window yet']
    for j, v in enumerate(vals, 1):
        c = ins.cell(row, j, v); c.border = THIN
        if j == 3: c.alignment = CEN
    row += 1
for j, w in enumerate([13, 28, 15, 55], 1):
    ins.column_dimensions[get_column_letter(j)].width = w

# ============ Sheet 4: Upcoming Dussehra 2026 Position Schedule (all 211) ============
psch = wb.create_sheet('Position_Schedule_2026')
psch.merge_cells('A1:J1')
psch.cell(1, 1, 'DUSSEHRA 20-Oct-2026 — Upcoming Position-Taking Schedule, ALL 211 F&O Stocks').font = TITLE
psch.merge_cells('A2:J2')
psch.cell(2, 1, 'Entry/exit dates on the NSE-2026 trading calendar. 50 Nifty-50 stocks use the exact live-dashboard dates (0 mismatch); the other 161 are computed with the identical trading-day-shift method. Sorted by entry date.').font = SUB
hdr = ['Symbol', 'Company', 'Direction', 'Window', 'Entry Date', 'Exit Date', 'Lot Size', 'Win Rate %', 'Confidence', 'Status']
for j, hh in enumerate(hdr, 1):
    c = psch.cell(4, j, hh); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN

sched_rows = []
for sym, r in qualified.items():
    m = meta.get(sym, {})
    pos = position_for(sym, r)
    sort_key = __import__('datetime').datetime.strptime(pos['entry'].split(' (')[0].strip(), '%d-%b-%Y')
    sched_rows.append((sort_key, sym, m.get('name', sym), pos['direction'], f"T-{r['lead']} to T+{r['hold']}",
                        pos['entry'], pos['exit'], pos['lot'], r['pick_wr'], r.get('low_confidence'), 'ACTIVE'))
import datetime as _dt
for sym, r in insufficient.items():
    m = meta.get(sym, {})
    sched_rows.append((_dt.datetime.max, sym, m.get('name', sym), '—', '—', '—', '—', '—', '—', False, 'NO POSITION — insufficient data'))
sched_rows.sort(key=lambda t: (t[0], t[1]))

row = 5
for _, sym, name, direc, window, entry, exit_, lot, wr, low, status in sched_rows:
    vals = [sym, name, direc, window, entry, exit_, lot, wr, ('LOW (n<3)' if low else ('OK' if status == 'ACTIVE' else '—')), status]
    fill = AM if low else (RF if status != 'ACTIVE' else (GF if direc == 'LONG' else (RF if direc == 'SHORT' else None)))
    for j, v in enumerate(vals, 1):
        c = psch.cell(row, j, v); c.border = THIN
        if fill: c.fill = fill
        if j in (1, 3, 4, 5, 6, 9): c.alignment = CEN
        if j in (7, 8): c.alignment = RIGHT
    row += 1
for j, w in enumerate([13, 28, 10, 15, 13, 13, 9, 11, 12, 32], 1):
    psch.column_dimensions[get_column_letter(j)].width = w
psch.freeze_panes = 'A5'

# ============ Sheet 5: Overview (front) ============
ov = wb.create_sheet('Overview', 0)
ov.merge_cells('A1:B1')
ov.cell(1, 1, 'ALL F&O STOCKS \u2014 DUSSEHRA YEAR-WISE RANKED REPORT').font = TITLE
notes = [
    ('Universe', f'ALL {len(results)} F&O stocks (fo_stocks_211.json) are included. {len(qualified)} have real, usable Dussehra-window price data and are ranked; only {len(insufficient)} have genuinely ZERO data (too recently listed \u2014 listed on their own Insufficient_Data sheet, not fabricated).'),
    ('Confidence flagging', f'Of the {len(qualified)} ranked stocks, {n_low_conf} have fewer than 3 years of data (some just 1 year) \u2014 these are highlighted AMBER throughout and marked "LOW (n<3)" in the Confidence column. A "100% win rate" on n=1 is one data point, not a robust edge \u2014 read the Years column before trusting any number.'),
    ('Method', 'For each stock, its own T-n/T+m window is used. Both LONG and SHORT are backtested split-robust (corporate-action days >20%% excluded from compounding); the direction with the higher win rate is chosen.'),
    ('Data source per stock', '55 stocks (mostly Nifty-50) use 26-year real cash-equity history (2000-2025); 140 use near-month F&O futures history (2019-2025, up to 7 years, some as little as 1 year). Data source is shown per stock \u2014 do not compare a 26yr figure to a 1yr figure as if equally reliable.'),
    ('Consistency with the live dashboard', 'The 50 Nifty-50 stocks here use the SAME window and direction already verified and live on the dashboard for Dussehra (0 mismatches confirmed) \u2014 this report is a superset, not a conflicting alternate version.'),
    ('Upcoming position schedule', 'See the Position_Schedule_2026 sheet for the full entry/exit dates for the 20-Oct-2026 Dussehra event, all 211 stocks, sorted by entry date \u2014 includes the 16 zero-data stocks marked "NO POSITION".'),
    ('CAVEAT', 'Picking the single best direction per stock fits its own history closely but is prone to overfitting \u2014 in this project\u2019s own out-of-sample testing, a single uniform index-backed direction has repeatedly outperformed this per-stock method going forward. Treat rankings here as historical fit, not a guaranteed edge \u2014 and treat low-confidence (amber) rows with even more caution.'),
]
r0 = 3
for lbl, txt in notes:
    a = ov.cell(r0, 1, lbl); a.font = BOLD; a.fill = AM; a.border = THIN; a.alignment = WRAP
    b = ov.cell(r0, 2, txt); b.alignment = WRAP; b.border = THIN
    ov.row_dimensions[r0].height = 75; r0 += 1
ov.column_dimensions['A'].width = 24; ov.column_dimensions['B'].width = 105

wb.save(OUT_XLSX)
print('Saved Excel:', OUT_XLSX)

# ============ PDF (ranked summary, top/bottom + methodology) ============
def esc(s): return str(s).replace('&', '&amp;').replace('<', '&lt;')
top20 = ranked[:20]
bot10 = ranked[-10:]
def pdf_row(rank, s, r):
    low = r.get('low_confidence')
    conf = '<span style="color:#b45309;font-weight:700;">LOW</span>' if low else 'OK'
    rowstyle = ' style="background:#fffbeb;"' if low else ''
    return (f'<tr{rowstyle}><td>{rank}</td><td><b>{esc(s)}</b></td><td>{esc(meta.get(s,{}).get("name",s))[:24]}</td>'
            f'<td class="c">{r["pick"]}</td><td class="r">{r["pick_wr"]:.1f}%</td>'
            f'<td class="r {"pos" if (r["pick_avg"] or 0)>=0 else "neg"}">{(r["pick_avg"] or 0):+.2f}%</td>'
            f'<td class="c">{r["n"]}</td><td class="c">{conf}</td></tr>')
rows_top = ''.join(pdf_row(i + 1, s, r) for i, (s, r) in enumerate(top20))
rows_bot = ''.join(pdf_row(len(ranked) - 9 + i, s, r) for i, (s, r) in enumerate(bot10))

html = f"""<!doctype html><html><head><meta charset="utf-8">
<title>F&amp;O-211 Dussehra Ranked Report</title>
<style>
 *{{box-sizing:border-box}} body{{font-family:'Segoe UI',Arial,sans-serif;margin:0;color:#0f172a;background:#fff}}
 .wrap{{max-width:1000px;margin:0 auto;padding:26px 34px}}
 h1{{font-size:21px;margin:0 0 2px}} .sub{{color:#64748b;font-size:12px;margin-bottom:14px}}
 h2{{font-size:15px;margin:20px 0 8px;border-left:4px solid #2563eb;padding-left:9px}}
 table{{width:100%;border-collapse:collapse;font-size:11.5px}}
 th,td{{padding:5px 7px;text-align:right;border-bottom:1px solid #eef2f7}} th{{background:#f1f5f9;font-size:9.5px;text-transform:uppercase;color:#475569}}
 td.c,th.c{{text-align:center}} .pos{{color:#059669;font-weight:700}} .neg{{color:#dc2626;font-weight:700}}
 .note{{background:#fff7ed;border:1px solid #fed7aa;border-radius:8px;padding:11px 14px;font-size:11px;color:#7c2d12;margin-top:8px;line-height:1.5}}
 .cards{{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;margin:10px 0}}
 .mc{{background:#fff;border:1px solid #e2e8f0;border-radius:8px;padding:10px 12px}}
 .mcl{{font-size:9.5px;text-transform:uppercase;color:#64748b;font-weight:700}} .mcv{{font-size:19px;font-weight:800}}
 .foot{{color:#94a3b8;font-size:9.5px;margin-top:18px}} @page{{size:A4;margin:12mm}}
</style></head><body><div class="wrap">
 <h1>All F&amp;O Stocks &mdash; Dussehra Year-Wise, Ranked</h1>
 <div class="sub">ALL {len(results)} F&amp;O stocks covered &bull; {len(qualified)} ranked (best direction per stock, split-robust backtest) &bull; full year-by-year detail for every stock in the companion Excel.</div>
 <div class="cards">
  <div class="mc"><div class="mcl">Ranked stocks</div><div class="mcv">{len(qualified)}</div></div>
  <div class="mc"><div class="mcl">Low-confidence (n&lt;3yr)</div><div class="mcv" style="color:#b45309">{n_low_conf}</div></div>
  <div class="mc"><div class="mcl">Zero data</div><div class="mcv">{len(insufficient)}</div></div>
  <div class="mc"><div class="mcl">Best avg return</div><div class="mcv pos">{ranked[0][1]['pick_avg']:+.1f}%</div></div>
 </div>
 <h2>Top 20 by Average Return</h2>
 <table><thead><tr><th>Rank</th><th class="c">Symbol</th><th class="c">Company</th><th class="c">Dir</th><th>Win%</th><th>Avg Ret</th><th class="c">Yrs</th><th class="c">Conf.</th></tr></thead>
 <tbody>{rows_top}</tbody></table>
 <h2>Bottom 10 by Average Return</h2>
 <table><thead><tr><th>Rank</th><th class="c">Symbol</th><th class="c">Company</th><th class="c">Dir</th><th>Win%</th><th>Avg Ret</th><th class="c">Yrs</th><th class="c">Conf.</th></tr></thead>
 <tbody>{rows_bot}</tbody></table>
 <div class="note"><b>Caveat:</b> picking the best direction per stock fits its own history closely but has repeatedly underperformed a single uniform index-backed direction out-of-sample in this project's prior testing &mdash; treat as historical fit, not a guaranteed forward edge. Rows marked <b style="color:#b45309;">LOW</b> confidence have under 3 years of data (some just 1) &mdash; a "100% win rate" there is 1-2 data points, not a robust statistic. Full 211-stock year-by-year detail (every stock, every available year) is in the companion Excel (Year_by_Year_Detail sheet); the {len(insufficient)} stocks with zero data are listed on its Insufficient_Data sheet.</div>
 <div class="foot">Split-robust backtest (corporate-action days excluded) &bull; 55 stocks on 26yr cash history, 140 on up to 7yr F&amp;O futures history &bull; consistent with the live dashboard's Nifty-50 Dussehra positions (0 mismatches verified).</div>
</div></body></html>"""
open(OUT_HTML, 'w', encoding='utf-8').write(html)
print('Saved HTML:', OUT_HTML)

for exe in [r'C:\Program Files\Google\Chrome\Application\chrome.exe', r'C:\Program Files (x86)\Google\Chrome\Application\chrome.exe']:
    if os.path.exists(exe):
        try:
            subprocess.run([exe, '--headless=new', '--disable-gpu', '--no-pdf-header-footer',
                            '--print-to-pdf=' + OUT_PDF, 'file:///' + OUT_HTML.replace('\\', '/')], timeout=60, capture_output=True)
            print('Saved PDF:', OUT_PDF, 'exists:', os.path.exists(OUT_PDF))
        except Exception as e:
            print('PDF render failed:', e)
        break
