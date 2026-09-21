import openpyxl
import pathlib
import sys

sys.stdout.reconfigure(errors='replace')

ROOT = pathlib.Path('D:/behaviour analysis')
fpath = ROOT / 'Nifty211_All_Stocks_15_Quarters_WinRate_Ranking_Master_v2.xlsx'

wb = openpyxl.load_workbook(fpath, data_only=True)
ws1 = wb['All 211 Stocks 15Q Ranking']
ws2 = wb['Qualified Stocks (69)']

print("==========================================================================================")
print("TOP 15 STOCKS IN 15-QUARTER RANKING WITH EXPLICIT N/A HANDLING")
print("==========================================================================================")

for r in range(17, 32):
    row_vals = [ws1.cell(row=r, column=c).value for c in range(1, 14)]
    rank, sym, strat, status, n_q, avail_stat, win_str, wr, ratio_str, l_c, avg_r, margin, pnl = row_vals
    print(f"Rank #{rank:<2} | {sym:<12} | {strat:<13} | {status:<20} | {avail_stat:<35} | 15Q Win Rate: {wr:<7} | Win Ratio: {ratio_str}")

print("\nTOP 15 QUALIFIED STOCKS (>=12 QTRS) IN 15-QUARTER RANKING WITH EXPLICIT N/A HANDLING")
print("=" * 110)
for r in range(17, 32):
    row_vals = [ws2.cell(row=r, column=c).value for c in range(1, 14)]
    rank, sym, strat, status, n_q, avail_stat, win_str, wr, ratio_str, l_c, avg_r, margin, pnl = row_vals
    print(f"Rank #{rank:<2} | {sym:<12} | {strat:<13} | {n_q:<12} | {avail_stat:<35} | 15Q Win Rate: {wr:<7} | Win Ratio: {ratio_str}")

print("==========================================================================================")
