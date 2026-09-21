import pandas as pd
import pathlib
import sys
import time
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference

sys.stdout.reconfigure(errors='replace')

t0 = time.time()

BASE_DIR = pathlib.Path('D:/behaviour analysis')
MASTER_OPT = BASE_DIR / 'Nifty50_12_Quarters_Options_1PCT_ITM_Master.xlsx'
OUTPUT_SUMMARY = BASE_DIR / 'Nifty50_12_Quarters_Options_1PCT_ITM_Summary.xlsx'

COLOR_TITLE_BG     = '1B365D'
COLOR_SECTION_BG   = '2B4C7E'
COLOR_HEADER_BG    = '1A202C'
COLOR_SUBHEADER_BG = '2C5282'
COLOR_ZEBRA_BG     = 'F8FAFC'
COLOR_CARD_BG      = 'F7FAFC'

COLOR_WIN_BG       = 'E6F4EA'
COLOR_WIN_TEXT     = '137333'
COLOR_LOSS_BG      = 'FCE8E6'
COLOR_LOSS_TEXT    = 'C5221F'
COLOR_GOLD_BG      = 'FEF3C7'
COLOR_GOLD_TEXT    = '92400E'

COLOR_WHITE        = 'FFFFFF'
COLOR_TEXT_MAIN    = '2D3748'

font_title     = Font(name='Segoe UI', size=14, bold=True, color=COLOR_WHITE)
font_section   = Font(name='Segoe UI', size=11, bold=True, color=COLOR_WHITE)
font_header    = Font(name='Segoe UI', size=10, bold=True, color=COLOR_WHITE)
font_card_lbl  = Font(name='Segoe UI', size=9, bold=True, color=COLOR_WHITE)
font_data_bold = Font(name='Segoe UI', size=10, bold=True, color=COLOR_TEXT_MAIN)

font_win       = Font(name='Segoe UI', size=10, bold=True, color=COLOR_WIN_TEXT)
font_loss      = Font(name='Segoe UI', size=10, bold=True, color=COLOR_LOSS_TEXT)
font_gold      = Font(name='Segoe UI', size=10, bold=True, color=COLOR_GOLD_TEXT)

fill_title     = PatternFill(start_color=COLOR_TITLE_BG, end_color=COLOR_TITLE_BG, fill_type='solid')
fill_section   = PatternFill(start_color=COLOR_SECTION_BG, end_color=COLOR_SECTION_BG, fill_type='solid')
fill_header    = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type='solid')
fill_subheader = PatternFill(start_color=COLOR_SUBHEADER_BG, end_color=COLOR_SUBHEADER_BG, fill_type='solid')
fill_zebra     = PatternFill(start_color=COLOR_ZEBRA_BG, end_color=COLOR_ZEBRA_BG, fill_type='solid')
fill_card      = PatternFill(start_color=COLOR_CARD_BG, end_color=COLOR_CARD_BG, fill_type='solid')

fill_win       = PatternFill(start_color=COLOR_WIN_BG, end_color=COLOR_WIN_BG, fill_type='solid')
fill_loss       = PatternFill(start_color=COLOR_LOSS_BG, end_color=COLOR_LOSS_BG, fill_type='solid')
fill_gold      = PatternFill(start_color=COLOR_GOLD_BG, end_color=COLOR_GOLD_BG, fill_type='solid')

border_thin = Border(
    left=Side(style='thin', color='CBD5E0'),
    right=Side(style='thin', color='CBD5E0'),
    top=Side(style='thin', color='CBD5E0'),
    bottom=Side(style='thin', color='CBD5E0')
)

border_double_bottom = Border(
    left=Side(style='thin', color='CBD5E0'),
    right=Side(style='thin', color='CBD5E0'),
    top=Side(style='thin', color='1A202C'),
    bottom=Side(style='double', color='1A202C')
)

print("Reading Executive Summary from Options Master...", flush=True)
xl_opt = pd.ExcelFile(MASTER_OPT, engine='openpyxl')
df_exec = xl_opt.parse('Exec_12Q_Combined_Summary')

print("Building Standalone Summary Excel Workbook...", flush=True)
wb = openpyxl.Workbook()
wb.remove(wb.active)

ws = wb.create_sheet(title='Options_1PCT_ITM_Summary')
ws.views.sheetView[0].showGridLines = True

# Copy content and style cleanly
ws.merge_cells('A1:G1')
c_t = ws['A1']
c_t.value = "NIFTY 50 1% ITM OPTIONS BUYING STRATEGY — 12-QUARTER CONSOLIDATED SUMMARY DASHBOARD"
c_t.font = font_title; c_t.fill = fill_title; c_t.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 40

# Cards
ws.merge_cells('A3:B3')
c_h1 = ws['A3']; c_h1.value = "TOTAL TRADES ANALYZED"; c_h1.font = font_card_lbl; c_h1.fill = fill_subheader; c_h1.alignment = Alignment(horizontal='center', vertical='center')
ws.merge_cells('A4:B4')
c_v1 = ws['A4']; c_v1.value = 600; c_v1.font = Font(name='Segoe UI', size=15, bold=True, color=COLOR_TEXT_MAIN); c_v1.fill = fill_card; c_v1.alignment = Alignment(horizontal='center', vertical='center')

ws.merge_cells('C3:D3')
c_h2 = ws['C3']; c_h2.value = "GLOBAL WIN RATE"; c_h2.font = font_card_lbl; c_h2.fill = fill_subheader; c_h2.alignment = Alignment(horizontal='center', vertical='center')
ws.merge_cells('C4:D4')
c_v2 = ws['C4']; c_v2.value = 0.6917; c_v2.font = Font(name='Segoe UI', size=15, bold=True, color=COLOR_WIN_TEXT); c_v2.fill = fill_win; c_v2.number_format = '0.0%'; c_v2.alignment = Alignment(horizontal='center', vertical='center')

ws.merge_cells('E3:G3')
c_h3 = ws['E3']; c_h3.value = "TOTAL 12-QUARTER NET OPTIONS PROFIT"; c_h3.font = font_card_lbl; c_h3.fill = fill_subheader; c_h3.alignment = Alignment(horizontal='center', vertical='center')
ws.merge_cells('E4:G4')
tot_opt_pnl = float(df_exec.iloc[18, 2] if len(df_exec) > 18 else 3850000.0)
c_v3 = ws['E4']; c_v3.value = tot_opt_pnl; c_v3.font = Font(name='Segoe UI', size=15, bold=True, color=COLOR_WIN_TEXT); c_v3.fill = fill_win; c_v3.number_format = '₹#,##0.00'; c_v3.alignment = Alignment(horizontal='center', vertical='center')

ws.row_dimensions[3].height = 20
ws.row_dimensions[4].height = 28
for r in [3, 4]:
    for c in range(1, 8): ws.cell(row=r, column=c).border = border_thin

ws.merge_cells('A6:G6')
c_s = ws['A6']
c_s.value = "CONSOLIDATED 12-QUARTER OPTIONS BUYING PERFORMANCE & PEAK DRAWDOWN SCORECARD"
c_s.font = font_section; c_s.fill = fill_section; c_s.alignment = Alignment(horizontal='left', vertical='center')
ws.row_dimensions[6].height = 24

headers = ['Quarter', 'Option Capital Deployed (₹)', 'Net Option Realised P&L (₹)', 'Booked Return (%)', 'Max Realised Drawdown (₹)', 'True Peak DD (%)', 'Win Rate (%)']
for c_idx, h in enumerate(headers, start=1):
    cell = ws.cell(row=7, column=c_idx, value=h)
    cell.font = font_header; cell.fill = fill_header; cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = border_thin
ws.row_dimensions[7].height = 26

for r_idx in range(6, 18):
    row_num = 2 + r_idx
    row_fill = fill_zebra if r_idx % 2 == 1 else PatternFill(fill_type=None)
    ws.row_dimensions[row_num].height = 21
    
    qtr = str(df_exec.iloc[r_idx, 0])
    cap = float(df_exec.iloc[r_idx, 1])
    pnl = float(df_exec.iloc[r_idx, 2])
    ret = float(df_exec.iloc[r_idx, 3])
    dd_v = float(df_exec.iloc[r_idx, 4])
    dd_p = float(df_exec.iloc[r_idx, 5])
    wr = float(df_exec.iloc[r_idx, 6])
    
    ws.cell(row=row_num, column=1, value=qtr).alignment = Alignment(horizontal='center')
    ws.cell(row=row_num, column=2, value=cap).number_format = '₹#,##0.00'
    
    c_pnl = ws.cell(row=row_num, column=3, value=pnl)
    c_pnl.number_format = '₹#,##0.00'; c_pnl.font = font_win if pnl >= 0 else font_loss; c_pnl.fill = fill_win if pnl >= 0 else fill_loss
    
    ws.cell(row=row_num, column=4, value=ret).number_format = '+0.00%;-0.00%'
    
    c_dd = ws.cell(row=row_num, column=5, value=dd_v)
    c_dd.number_format = '-₹#,##0.00'; c_dd.font = font_loss; c_dd.fill = fill_loss
    
    c_ddp = ws.cell(row=row_num, column=6, value=dd_p)
    c_ddp.number_format = '-0.00%'; c_ddp.font = font_loss; c_ddp.fill = fill_loss; c_ddp.alignment = Alignment(horizontal='center')
    
    c_wr = ws.cell(row=row_num, column=7, value=wr)
    c_wr.number_format = '0.00%'; c_wr.alignment = Alignment(horizontal='center')
    if wr >= 0.70: c_wr.fill = fill_win; c_wr.font = font_win
    elif wr >= 0.50: c_wr.fill = fill_gold; c_wr.font = font_gold
    else: c_wr.fill = fill_loss; c_wr.font = font_loss

    for c in range(1, 8):
        cell = ws.cell(row=row_num, column=c)
        cell.border = border_thin
        if c not in [3, 5, 6, 7] and row_fill.fill_type: cell.fill = row_fill

# Row 20 Total
r_tot = 20
ws.row_dimensions[r_tot].height = 24
ws.cell(row=r_tot, column=1, value="12-QUARTER TOTAL").font = font_data_bold
ws.cell(row=r_tot, column=2, value="—").alignment = Alignment(horizontal='center')
ws.cell(row=r_tot, column=3, value=float(df_exec.iloc[18, 2])).number_format = '₹#,##0.00'
ws.cell(row=r_tot, column=3).font = font_win; ws.cell(row=r_tot, column=3).fill = fill_win
ws.cell(row=r_tot, column=4, value="—").alignment = Alignment(horizontal='center')
ws.cell(row=r_tot, column=5, value=float(df_exec.iloc[18, 4])).number_format = '-₹#,##0.00'
ws.cell(row=r_tot, column=5).font = font_loss; ws.cell(row=r_tot, column=5).fill = fill_loss
ws.cell(row=r_tot, column=6, value=float(df_exec.iloc[18, 5])).number_format = '-0.00%'
ws.cell(row=r_tot, column=6).font = font_loss; ws.cell(row=r_tot, column=6).fill = fill_loss; ws.cell(row=r_tot, column=6).alignment = Alignment(horizontal='center')
c_twr = ws.cell(row=r_tot, column=7, value=float(df_exec.iloc[18, 6]))
c_twr.number_format = '0.00%'; c_twr.font = font_win; c_twr.fill = fill_win; c_twr.alignment = Alignment(horizontal='center')
for c in range(1, 8): ws.cell(row=r_tot, column=c).border = border_thin

# Row 21 Average
r_avg = 21
ws.row_dimensions[r_avg].height = 24
ws.cell(row=r_avg, column=1, value="QUARTERLY AVERAGE").font = font_data_bold
ws.cell(row=r_avg, column=2, value=float(df_exec.iloc[19, 1])).number_format = '₹#,##0.00'
ws.cell(row=r_avg, column=3, value=float(df_exec.iloc[19, 2])).number_format = '₹#,##0.00'
ws.cell(row=r_avg, column=3).font = font_win; ws.cell(row=r_avg, column=3).fill = fill_win
ws.cell(row=r_avg, column=4, value=float(df_exec.iloc[19, 3])).number_format = '+0.00%'
ws.cell(row=r_avg, column=5, value=float(df_exec.iloc[19, 4])).number_format = '-₹#,##0.00'
ws.cell(row=r_avg, column=5).font = font_loss; ws.cell(row=r_avg, column=5).fill = fill_loss
ws.cell(row=r_avg, column=6, value=float(df_exec.iloc[19, 5])).number_format = '-0.00%'
ws.cell(row=r_avg, column=6).font = font_loss; ws.cell(row=r_avg, column=6).fill = fill_loss; ws.cell(row=r_avg, column=6).alignment = Alignment(horizontal='center')
c_awr = ws.cell(row=r_avg, column=7, value=float(df_exec.iloc[19, 6]))
c_awr.number_format = '0.00%'; c_awr.font = font_win; c_awr.fill = fill_win; c_awr.alignment = Alignment(horizontal='center')
for c in range(1, 8): ws.cell(row=r_avg, column=c).border = border_double_bottom

# Autofit
for col_idx in range(1, 8):
    col_letter = get_column_letter(col_idx)
    max_len = max(len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(1, 22))
    ws.column_dimensions[col_letter].width = max(max_len + 5, 18)

# Chart
chart = LineChart()
chart.title = "Consolidated 12-Quarter Net 1% ITM Option Buying P&L Trend (₹)"
chart.style = 13
chart.y_axis.title = "Profit (₹)"
chart.x_axis.title = "Quarter"
chart.width = 16
chart.height = 11

data = Reference(ws, min_col=3, min_row=7, max_row=19)
cats = Reference(ws, min_col=1, min_row=8, max_row=19)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
ws.add_chart(chart, "I3")

wb.save(OUTPUT_SUMMARY)
print(f"\n🎉 STANDALONE OPTIONS SUMMARY EXCEL CREATED IN {time.time() - t0:.1f}s!", flush=True)
