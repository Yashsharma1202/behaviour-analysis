import openpyxl
import pathlib
import sys

sys.stdout.reconfigure(errors='replace')

ROOT = pathlib.Path('D:/behaviour analysis')
fpath = ROOT / 'Nifty211_Quarterly_Performance_Master_v2.xlsx'

wb = openpyxl.load_workbook(fpath, data_only=True)
print(f"Sheets in workbook: {wb.sheetnames}")

for sname in wb.sheetnames:
    ws = wb[sname]
    headers = [ws.cell(row=7, column=c).value for c in range(1, 17)]
    print(f"\nSheet: {sname}")
    print(f"Headers (16 cols): {headers}")
    
    # Print first 5 rows sample data
    for r in range(8, 13):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 17)]
        print(f"Row {r-7}: Sym={row_vals[1]}, Sector={row_vals[2]}, StartQ={row_vals[3]}, FirstDate={row_vals[4]}, Strat={row_vals[5]}, TotalQ={row_vals[7]}")
