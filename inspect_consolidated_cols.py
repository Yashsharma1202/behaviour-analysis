import openpyxl, sys

sys.stdout.reconfigure(errors='replace')

path = r'D:/behaviour analysis/Nifty50_12_Quarters_Consolidated.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb['Q3 2023-24']

print("Row 8 Headers:")
for c in range(1, ws.max_column + 1):
    val = ws.cell(row=8, column=c).value
    print(f"  Col {c}: {val}")

print("\nRow 9 Values:")
for c in range(1, ws.max_column + 1):
    val = ws.cell(row=9, column=c).value
    print(f"  Col {c}: {val}")
