import openpyxl
import pandas as pd
import numpy as np
import pathlib
import sys
import time
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference

sys.stdout.reconfigure(errors='replace')

t0 = time.time()

# ------------------- Configuration & Paths -------------------
BASE_DIR = pathlib.Path('D:/behaviour analysis')
TARGET_FILE = BASE_DIR / 'Nifty50_12_Quarters_Consolidated.xlsx'
OUTPUT_FILE = BASE_DIR / 'Nifty50_12_Quarters_Consolidated_v2.xlsx'

print(f"Loading {TARGET_FILE.name} for Drawdown (DD) Enhancement...", flush=True)
wb = openpyxl.load_workbook(TARGET_FILE)

# Colors & Fonts
COLOR_TITLE_BG     = '1B365D'
COLOR_SECTION_BG   = '2B4C7E'
COLOR_HEADER_BG    = '1A202C'
COLOR_SUBHEADER_BG = '2C5282'

COLOR_WIN_BG       = 'E6F4EA'
COLOR_WIN_TEXT     = '137333'
COLOR_LOSS_BG      = 'FCE8E6'
COLOR_LOSS_TEXT    = 'C5221F'

COLOR_WHITE        = 'FFFFFF'
COLOR_TEXT_MAIN    = '2D3748'

font_header    = Font(name='Segoe UI', size=10, bold=True, color=COLOR_WHITE)
font_subheader = Font(name='Segoe UI', size=9, bold=True, color=COLOR_WHITE)
font_data_bold = Font(name='Segoe UI', size=10, bold=True, color=COLOR_TEXT_MAIN)

font_loss      = Font(name='Segoe UI', size=10, bold=True, color=COLOR_LOSS_TEXT)

fill_header    = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type='solid')
fill_subheader = PatternFill(start_color=COLOR_SUBHEADER_BG, end_color=COLOR_SUBHEADER_BG, fill_type='solid')
fill_loss       = PatternFill(start_color=COLOR_LOSS_BG, end_color=COLOR_LOSS_BG, fill_type='solid')

border_thin = Border(
    left=Side(style='thin', color='CBD5E0'),
    right=Side(style='thin', color='CBD5E0'),
    top=Side(style='thin', color='CBD5E0'),
    bottom=Side(style='thin', color='CBD5E0')
)

quarters_order = [
    'Q3 2023-24', 'Q4 2023-24', 'Q1 2024-25', 'Q2 2024-25',
    'Q3 2024-25', 'Q4 2024-25', 'Q1 2025-26', 'Q2 2025-26',
    'Q3 2025-26', 'Q4 2025-26', 'Q1 2026-27', 'Q2 2026-27'
]

quarter_dd_summary = {}

# Process each Quarter Sheet
for qtr in quarters_order:
    if qtr not in wb.sheetnames:
        continue
        
    ws = wb[qtr]
    
    # Read trades from row 9 onwards
    trades = []
    r_idx = 9
    while True:
        trade_no = ws.cell(row=r_idx, column=1).value
        sym = ws.cell(row=r_idx, column=2).value
        if trade_no is None or str(trade_no).strip() == '' or sym is None or str(sym).strip() == '':
            break
            
        entry_dt = str(ws.cell(row=r_idx, column=7).value or '')
        exit_dt  = str(ws.cell(row=r_idx, column=8).value or '')
        pnl_val  = float(ws.cell(row=r_idx, column=13).value or 0.0)  # Col 13 is Realised P&L
        
        trades.append({
            'row_idx': r_idx,
            'trade_no': int(trade_no) if str(trade_no).isdigit() else r_idx-8,
            'symbol': str(sym).strip(),
            'entry_dt': entry_dt,
            'pnl': pnl_val
        })
        r_idx += 1
        
    df_q = pd.DataFrame(trades)
    if df_q.empty:
        continue
        
    # Sort chronologically by Entry Date & Trade No
    df_q['entry_dt_parsed'] = pd.to_datetime(df_q['entry_dt'], errors='coerce')
    df_q = df_q.sort_values(by=['entry_dt_parsed', 'trade_no']).reset_index(drop=True)
    
    # Calculate Cumulative PnL & Drawdown
    df_q['cum_pnl'] = df_q['pnl'].cumsum()
    df_q['running_max'] = np.maximum.accumulate(df_q['cum_pnl'])
    df_q['drawdown'] = df_q['cum_pnl'] - df_q['running_max']
    
    max_dd_val = abs(df_q['drawdown'].min())
    
    # Get Fund Utilised from Cell A5
    fund_utilised = float(ws.cell(row=5, column=1).value or 1.0)
    max_dd_pct = -abs(max_dd_val / fund_utilised) if fund_utilised > 0 else 0.0
    
    quarter_dd_summary[qtr] = {
        'max_dd_val': max_dd_val,
        'max_dd_pct': max_dd_pct,
        'fund_utilised': fund_utilised,
        'count': len(df_q)
    }
    
    # Add Headers for Col 16 (Cumulative Profit) and Col 17 (Drawdown) at Row 8
    c16_hdr = ws.cell(row=8, column=16, value="Cumulative Profit (₹)")
    c16_hdr.font = font_header; c16_hdr.fill = fill_header; c16_hdr.alignment = Alignment(horizontal='center', vertical='center')
    c16_hdr.border = border_thin
    
    c17_hdr = ws.cell(row=8, column=17, value="Drawdown (₹)")
    c17_hdr.font = font_header; c17_hdr.fill = fill_header; c17_hdr.alignment = Alignment(horizontal='center', vertical='center')
    c17_hdr.border = border_thin
    
    # Write Cumulative P&L and Drawdown to Row 9+
    for idx, r in df_q.iterrows():
        orig_row = r['row_idx']
        
        cell_cpnl = ws.cell(row=orig_row, column=16, value=round(r['cum_pnl'], 2))
        cell_cpnl.number_format = '₹#,##0.00'
        cell_cpnl.font = font_data_bold
        cell_cpnl.border = border_thin
        
        cell_dd = ws.cell(row=orig_row, column=17, value=round(r['drawdown'], 2))
        cell_dd.number_format = '₹#,##0.00'
        if r['drawdown'] < 0:
            cell_dd.font = font_loss; cell_dd.fill = fill_loss
        cell_dd.border = border_thin
        
    # Update Summary Card Row 4 & 5 for Max Drawdown
    lbl_dd1 = ws.cell(row=4, column=10, value="Max Drawdown (₹)")
    lbl_dd1.font = font_subheader; lbl_dd1.fill = fill_subheader; lbl_dd1.alignment = Alignment(horizontal='center', vertical='center')
    lbl_dd1.border = border_thin
    
    val_dd1 = ws.cell(row=5, column=10, value=-abs(max_dd_val))
    val_dd1.number_format = '-₹#,##0.00'; val_dd1.font = font_loss; val_dd1.fill = fill_loss; val_dd1.alignment = Alignment(horizontal='center')
    val_dd1.border = border_thin
    
    lbl_dd2 = ws.cell(row=4, column=11, value="Max Drawdown (%)")
    lbl_dd2.font = font_subheader; lbl_dd2.fill = fill_subheader; lbl_dd2.alignment = Alignment(horizontal='center', vertical='center')
    lbl_dd2.border = border_thin
    
    val_dd2 = ws.cell(row=5, column=11, value=max_dd_pct)
    val_dd2.number_format = '-0.00%'; val_dd2.font = font_loss; val_dd2.fill = fill_loss; val_dd2.alignment = Alignment(horizontal='center')
    val_dd2.border = border_thin
    
    # Column dimensions
    ws.column_dimensions['P'].width = 22
    ws.column_dimensions['Q'].width = 18
    
    # Embed Native Excel Drawdown Chart at Column S6
    chart_q = LineChart()
    chart_q.title = f"{qtr} — Cumulative Realised Profit & Drawdown Curve (Max DD: -₹{max_dd_val:,.2f})"
    chart_q.style = 13
    chart_q.y_axis.title = "Rupees (₹)"
    chart_q.x_axis.title = "Trade #"
    chart_q.width = 16
    chart_q.height = 10

    max_r = 8 + len(df_q)
    data_q = Reference(ws, min_col=16, min_row=8, max_col=17, max_row=max_r)
    cats_q = Reference(ws, min_col=1, min_row=9, max_row=max_r)
    chart_q.add_data(data_q, titles_from_data=True)
    chart_q.set_categories(cats_q)
    ws.add_chart(chart_q, "S6")

# Save Workbook with fallback
def safe_save(wb_obj):
    try:
        print(f"Attempting save to {TARGET_FILE.name}...", flush=True)
        wb_obj.save(TARGET_FILE)
        print(f"✅ Successfully saved to {TARGET_FILE.name}", flush=True)
    except PermissionError:
        print(f"⚠️ {TARGET_FILE.name} is locked by Excel. Saving to {OUTPUT_FILE.name}...", flush=True)
        wb_obj.save(OUTPUT_FILE)
        print(f"✅ Successfully saved to {OUTPUT_FILE.name}", flush=True)

safe_save(wb)

print(f"\n🎉 DRAWDOWN ENHANCEMENT COMPLETED IN {time.time() - t0:.1f}s!", flush=True)
