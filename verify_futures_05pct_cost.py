import openpyxl, pandas as pd, sys

sys.stdout.reconfigure(errors='replace')

summary_path = r'D:/behaviour analysis/Nifty50_12_Quarters_Point_In_Time_Futures_0.5PCT_Cost_Summary.xlsx'
wb = openpyxl.load_workbook(summary_path, data_only=True)
ws = wb['Futures_0.5PCT_Cost_Summary']

print("==========================================================================================")
print("VERIFYING FUTURES MODEL (WITH 0.5% TRANSACTION COST) SUMMARY WORKBOOK")
print("==========================================================================================")
print("Sheet Names:", wb.sheetnames)

print("\nSummary Table Rows 7-21:")
for r in range(7, 22):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 12)]
    print(f"  Row {r:2d}: {vals}")
print("==========================================================================================")
