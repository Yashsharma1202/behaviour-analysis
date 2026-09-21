import openpyxl, pandas as pd, sys

sys.stdout.reconfigure(errors='replace')

path = r'D:/behaviour analysis/Nifty50_QuarterWise_And_YearWise_Best_Stock_Performance.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

print("==========================================================================================")
print("VERIFYING MASTER QUARTER-WISE & YEAR-WISE BEST STOCK WORKBOOK")
print("==========================================================================================")
print("Total Sheets:", len(wb.sheetnames))
print("Sheet Names :", wb.sheetnames)

for q_season in ['Q1', 'Q2', 'Q3', 'Q4']:
    sheet_name = f"{q_season}_Best_Stocks_Equity_&_Fut"
    df_s = pd.read_excel(path, sheet_name=sheet_name, engine='openpyxl')
    print(f"\n--- TOP 3 STOCKS IN {q_season} EARNINGS SEASON ---")
    print(df_s.iloc[2:5, [0, 1, 2, 5, 7, 8, 10]].to_string())

print("==========================================================================================")
