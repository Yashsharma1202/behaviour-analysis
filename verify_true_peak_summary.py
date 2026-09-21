import openpyxl, pandas as pd, sys

sys.stdout.reconfigure(errors='replace')

path = r'D:/behaviour analysis/Nifty50_12_Quarters_True_Peak_Drawdown_Summary.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb['True_Peak_Drawdown_Summary']

print("==========================================")
print("VERIFYING STANDALONE TRUE PEAK DRAWDOWN EXCEL WORKBOOK")
print("==========================================")
print("Sheet Names:", wb.sheetnames)
print(f"Total Charts: {len(ws._charts)}")

print("\nTable Rows 7-21:")
for r in range(7, 22):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 9)]
    print(f"  Row {r:2d}: {vals}")
print("==========================================")
