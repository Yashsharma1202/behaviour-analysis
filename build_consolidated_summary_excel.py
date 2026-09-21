import pandas as pd
import pathlib
import sys
import time
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference

sys.stdout.reconfigure(errors='replace')

t0 = time.time()

# ------------------- Configuration & Paths -------------------
BASE_DIR = pathlib.Path('D:/behaviour analysis')
OUTPUT_SUMMARY_EXCEL = BASE_DIR / 'Nifty50_Consolidated_12_Quarter_Drawdown_Summary.xlsx'

# Executive Color Tokens
COLOR_TITLE_BG     = '1B365D'  # Deep Imperial Navy
COLOR_SECTION_BG   = '2B4C7E'  # Royal Slate Accent
COLOR_HEADER_BG    = '1A202C'  # Rich Charcoal Navy
COLOR_SUBHEADER_BG = '2C5282'  # Deep Steel Teal
COLOR_ZEBRA_BG     = 'F8FAFC'  # Soft Slate Tint
COLOR_CARD_BG      = 'F7FAFC'  # Ice White

COLOR_WIN_BG       = 'E6F4EA'  # Mint Green Fill
COLOR_WIN_TEXT     = '137333'  # Dark Emerald Text

COLOR_LOSS_BG      = 'FCE8E6'  # Soft Peach Red Fill
COLOR_LOSS_TEXT    = 'C5221F'  # Dark Crimson Text

COLOR_GOLD_BG      = 'FEF3C7'  # Warm Gold Accent
COLOR_GOLD_TEXT    = '92400E'  # Dark Amber Text

COLOR_WHITE        = 'FFFFFF'
COLOR_TEXT_MAIN    = '2D3748'

font_title     = Font(name='Segoe UI', size=14, bold=True, color=COLOR_WHITE)
font_section   = Font(name='Segoe UI', size=11, bold=True, color=COLOR_WHITE)
font_header    = Font(name='Segoe UI', size=10, bold=True, color=COLOR_WHITE)
font_card_lbl  = Font(name='Segoe UI', size=9, bold=True, color=COLOR_WHITE)
font_data_bold = Font(name='Segoe UI', size=10, bold=True, color=COLOR_TEXT_MAIN)
font_data_reg  = Font(name='Segoe UI', size=10, color=COLOR_TEXT_MAIN)

font_win       = Font(name='Segoe UI', size=10, bold=True, color=COLOR_WIN_TEXT)
font_loss      = Font(name='Segoe UI', size=10, bold=True, color=COLOR_LOSS_TEXT)
font_gold      = Font(name='Segoe UI', size=10, bold=True, color=COLOR_GOLD_TEXT)

fill_title     = PatternFill(start_color=COLOR_TITLE_BG, end_color=COLOR_TITLE_BG, fill_type='solid')
fill_section   = PatternFill(start_color=COLOR_SECTION_BG, end_color=COLOR_SECTION_BG, fill_type='solid')
fill_header    = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type='solid')
fill_subheader = PatternFill(start_color=COLOR_SUBHEADER_BG, end_color=COLOR_SUBHEADER_BG, fill_type='solid')
fill_zebra     = PatternFill(start_color=COLOR_ZEBRA_BG, end_color=COLOR_ZEBRA_BG, fill_type='solid')
fill_card      = PatternFill(start_color=COLOR_CARD_BG, end_color=COLOR_CARD_BG, fill_type='solid')

fill_win       = PatternFill(start_color=COLOR_WIN_BG, end_color=COLOR_WIN_BG, fill_type='solid')
fill_loss       = PatternFill(start_color=COLOR_LOSS_BG, end_color=COLOR_LOSS_BG, fill_type='solid')
fill_gold      = PatternFill(start_color=COLOR_GOLD_BG, end_color=COLOR_GOLD_BG, fill_type='solid')

border_thin = Border(
    left=Side(style='thin', color='CBD5E0'),
    right=Side(style='thin', color='CBD5E0'),
    top=Side(style='thin', color='CBD5E0'),
    bottom=Side(style='thin', color='CBD5E0')
)

border_double_bottom = Border(
    left=Side(style='thin', color='CBD5E0'),
    right=Side(style='thin', color='CBD5E0'),
    top=Side(style='thin', color='1A202C'),
    bottom=Side(style='double', color='1A202C')
)

# ------------------- Table Data Exact Match -------------------
table_data = [
    {'Quarter': 'Q3 2023-24', 'Fund Utilised': 56125.63, 'Net Realised Profit': 2438.97, 'Booked Return': 0.0435, 'Max Drawdown ₹': -72.65, 'Max Drawdown %': -0.0013, 'Win Rate': 0.9400},
    {'Quarter': 'Q4 2023-24', 'Fund Utilised': 71393.52, 'Net Realised Profit': 5560.97, 'Booked Return': 0.0779, 'Max Drawdown ₹': -2.85, 'Max Drawdown %': -0.0000, 'Win Rate': 0.9608},
    {'Quarter': 'Q1 2024-25', 'Fund Utilised': 58151.51, 'Net Realised Profit': 1974.47, 'Booked Return': 0.0340, 'Max Drawdown ₹': -356.60, 'Max Drawdown %': -0.0061, 'Win Rate': 0.6667},
    {'Quarter': 'Q2 2024-25', 'Fund Utilised': 58676.80, 'Net Realised Profit': 2537.32, 'Booked Return': 0.0432, 'Max Drawdown ₹': -251.82, 'Max Drawdown %': -0.0043, 'Win Rate': 0.6800},
    {'Quarter': 'Q3 2024-25', 'Fund Utilised': 66124.85, 'Net Realised Profit': 1166.80, 'Booked Return': 0.0176, 'Max Drawdown ₹': -347.09, 'Max Drawdown %': -0.0052, 'Win Rate': 0.6600},
    {'Quarter': 'Q4 2024-25', 'Fund Utilised': 59487.75, 'Net Realised Profit': -812.25, 'Booked Return': -0.0137, 'Max Drawdown ₹': -1189.39, 'Max Drawdown %': -0.0200, 'Win Rate': 0.5200},
    {'Quarter': 'Q1 2025-26', 'Fund Utilised': 51045.10, 'Net Realised Profit': 1352.26, 'Booked Return': 0.0265, 'Max Drawdown ₹': -673.69, 'Max Drawdown %': -0.0132, 'Win Rate': 0.6400},
    {'Quarter': 'Q2 2025-26', 'Fund Utilised': 64804.02, 'Net Realised Profit': 300.92, 'Booked Return': 0.0046, 'Max Drawdown ₹': -417.42, 'Max Drawdown %': -0.0064, 'Win Rate': 0.5102},
    {'Quarter': 'Q3 2025-26', 'Fund Utilised': 52584.05, 'Net Realised Profit': 1836.82, 'Booked Return': 0.0349, 'Max Drawdown ₹': -436.32, 'Max Drawdown %': -0.0083, 'Win Rate': 0.6600},
    {'Quarter': 'Q4 2025-26', 'Fund Utilised': 69471.60, 'Net Realised Profit': 4407.57, 'Booked Return': 0.0634, 'Max Drawdown ₹': -179.42, 'Max Drawdown %': -0.0026, 'Win Rate': 0.7200},
    {'Quarter': 'Q1 2026-27', 'Fund Utilised': 62827.63, 'Net Realised Profit': 1990.95, 'Booked Return': 0.0317, 'Max Drawdown ₹': -207.40, 'Max Drawdown %': -0.0033, 'Win Rate': 0.6000},
    {'Quarter': 'Q2 2026-27', 'Fund Utilised': 94468.80, 'Net Realised Profit': 5162.26, 'Booked Return': 0.0546, 'Max Drawdown ₹': -96.00, 'Max Drawdown %': -0.0010, 'Win Rate': 0.8600}
]

df_table = pd.DataFrame(table_data)

print("Building Standalone Excel for Consolidated 12-Quarter Drawdown Summary Table...", flush=True)
wb = Workbook()
wb.remove(wb.active)

ws = wb.create_sheet(title='Drawdown_Summary')
ws.views.sheetView[0].showGridLines = True

# Row 1: Title Banner
ws.merge_cells('A1:G1')
c_title = ws['A1']
c_title.value = "NIFTY 50 STRATEGY — CONSOLIDATED 12-QUARTER DRAWDOWN SUMMARY TABLE"
c_title.font = font_title
c_title.fill = fill_title
c_title.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 40

# Row 3-4: Hero Summary Cards
ws.merge_cells('A3:B3')
c_h1 = ws['A3']; c_h1.value = "TOTAL TRADES ANALYZED"; c_h1.font = font_card_lbl; c_h1.fill = fill_subheader; c_h1.alignment = Alignment(horizontal='center', vertical='center')
ws.merge_cells('A4:B4')
c_v1 = ws['A4']; c_v1.value = 600; c_v1.font = Font(name='Segoe UI', size=15, bold=True, color=COLOR_TEXT_MAIN); c_v1.fill = fill_card; c_v1.alignment = Alignment(horizontal='center', vertical='center')

ws.merge_cells('C3:D3')
c_h2 = ws['C3']; c_h2.value = "GLOBAL WIN RATE"; c_h2.font = font_card_lbl; c_h2.fill = fill_subheader; c_h2.alignment = Alignment(horizontal='center', vertical='center')
ws.merge_cells('C4:D4')
c_v2 = ws['C4']; c_v2.value = 0.6967; c_v2.font = Font(name='Segoe UI', size=15, bold=True, color=COLOR_WIN_TEXT); c_v2.fill = fill_win; c_v2.number_format = '0.0%'; c_v2.alignment = Alignment(horizontal='center', vertical='center')

ws.merge_cells('E3:G3')
c_h3 = ws['E3']; c_h3.value = "TOTAL NET REALISED PROFIT"; c_h3.font = font_card_lbl; c_h3.fill = fill_subheader; c_h3.alignment = Alignment(horizontal='center', vertical='center')
ws.merge_cells('E4:G4')
c_v3 = ws['E4']; c_v3.value = df_table['Net Realised Profit'].sum(); c_v3.font = Font(name='Segoe UI', size=15, bold=True, color=COLOR_WIN_TEXT); c_v3.fill = fill_win; c_v3.number_format = '₹#,##0.00'; c_v3.alignment = Alignment(horizontal='center', vertical='center')

ws.row_dimensions[3].height = 20
ws.row_dimensions[4].height = 28

for r in [3, 4]:
    for c in range(1, 8):
        ws.cell(row=r, column=c).border = border_thin

# Row 6: Section Banner
ws.merge_cells('A6:G6')
c_sec = ws['A6']
c_sec.value = "CONSOLIDATED 12-QUARTER DRAWDOWN & CAPITAL PERFORMANCE TABLE"
c_sec.font = font_section
c_sec.fill = fill_section
c_sec.alignment = Alignment(horizontal='left', vertical='center')
ws.row_dimensions[6].height = 24

# Row 7: Table Headers
headers = ['Quarter', 'Fund Utilised (₹)', 'Net Realised Profit (₹)', 'Booked Return (%)', 'Max Drawdown (₹)', 'Max Drawdown (%)', 'Win Rate (%)']
for c_idx, h in enumerate(headers, start=1):
    cell = ws.cell(row=7, column=c_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = border_thin
ws.row_dimensions[7].height = 26

# Row 8-19: Data Rows
for r_idx, r_data in df_table.iterrows():
    row_num = 8 + r_idx
    row_fill = fill_zebra if r_idx % 2 == 1 else PatternFill(fill_type=None)
    ws.row_dimensions[row_num].height = 21
    
    ws.cell(row=row_num, column=1, value=r_data['Quarter']).alignment = Alignment(horizontal='center')
    
    c_fund = ws.cell(row=row_num, column=2, value=r_data['Fund Utilised'])
    c_fund.number_format = '₹#,##0.00'
    
    c_pnl = ws.cell(row=row_num, column=3, value=r_data['Net Realised Profit'])
    c_pnl.number_format = '₹#,##0.00'
    c_pnl.font = font_win if r_data['Net Realised Profit'] >= 0 else font_loss
    c_pnl.fill = fill_win if r_data['Net Realised Profit'] >= 0 else fill_loss
    
    c_ret = ws.cell(row=row_num, column=4, value=r_data['Booked Return'])
    c_ret.number_format = '+0.00%;-0.00%'
    
    c_dd = ws.cell(row=row_num, column=5, value=r_data['Max Drawdown ₹'])
    c_dd.number_format = '-₹#,##0.00'
    c_dd.font = font_loss; c_dd.fill = fill_loss
    
    c_ddp = ws.cell(row=row_num, column=6, value=r_data['Max Drawdown %'])
    c_ddp.number_format = '-0.00%'
    c_ddp.font = font_loss; c_ddp.fill = fill_loss; c_ddp.alignment = Alignment(horizontal='center')
    
    c_wr = ws.cell(row=row_num, column=7, value=r_data['Win Rate'])
    c_wr.number_format = '0.00%'
    c_wr.alignment = Alignment(horizontal='center')
    if r_data['Win Rate'] >= 0.70:
        c_wr.fill = fill_win; c_wr.font = font_win
    elif r_data['Win Rate'] >= 0.50:
        c_wr.fill = fill_gold; c_wr.font = font_gold
    else:
        c_wr.fill = fill_loss; c_wr.font = font_loss

    for c in range(1, 8):
        cell = ws.cell(row=row_num, column=c)
        cell.border = border_thin
        if c not in [3, 5, 6, 7] and row_fill.fill_type: cell.fill = row_fill

# Row 20: Summary Row
gt_row = 20
ws.row_dimensions[gt_row].height = 24
ws.cell(row=gt_row, column=1, value="12-Quarter Total / Summary").font = font_data_bold
ws.cell(row=gt_row, column=2, value=df_table['Fund Utilised'].mean()).number_format = '₹#,##0.00'
ws.cell(row=gt_row, column=3, value=df_table['Net Realised Profit'].sum()).number_format = '₹#,##0.00'
ws.cell(row=gt_row, column=3).font = font_win; ws.cell(row=gt_row, column=3).fill = fill_win

ws.cell(row=gt_row, column=4, value=df_table['Booked Return'].mean()).number_format = '+0.00%'
ws.cell(row=gt_row, column=5, value=df_table['Max Drawdown ₹'].min()).number_format = '-₹#,##0.00'
ws.cell(row=gt_row, column=5).font = font_loss; ws.cell(row=gt_row, column=5).fill = fill_loss

ws.cell(row=gt_row, column=6, value=df_table['Max Drawdown %'].min()).number_format = '-0.00%'
ws.cell(row=gt_row, column=6).font = font_loss; ws.cell(row=gt_row, column=6).fill = fill_loss; ws.cell(row=gt_row, column=6).alignment = Alignment(horizontal='center')

ws.cell(row=gt_row, column=7, value=0.6967).number_format = '0.00%'
ws.cell(row=gt_row, column=7).font = font_win; ws.cell(row=gt_row, column=7).fill = fill_win; ws.cell(row=gt_row, column=7).alignment = Alignment(horizontal='center')

for c in range(1, 8):
    cell = ws.cell(row=gt_row, column=c)
    cell.border = border_double_bottom
    if cell.font is None: cell.font = font_data_bold

# Column Autofit
for col_idx in range(1, 8):
    col_letter = get_column_letter(col_idx)
    max_len = max(len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(1, 21))
    ws.column_dimensions[col_letter].width = max(max_len + 5, 18)

# Add Embedded Line Chart at I3
chart = LineChart()
chart.title = "Consolidated 12-Quarter Net Realised Profit Trend (₹)"
chart.style = 13
chart.y_axis.title = "Profit (₹)"
chart.x_axis.title = "Quarter"
chart.width = 16
chart.height = 11

data = Reference(ws, min_col=3, min_row=7, max_row=19)
cats = Reference(ws, min_col=1, min_row=8, max_row=19)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
ws.add_chart(chart, "I3")

# Save Workbook
wb.save(OUTPUT_SUMMARY_EXCEL)
print(f"\n🎉 STANDALONE CONSOLIDATED SUMMARY EXCEL CREATED IN {time.time() - t0:.1f}s!", flush=True)
