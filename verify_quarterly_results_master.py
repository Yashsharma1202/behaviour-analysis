import openpyxl, sys

sys.stdout.reconfigure(errors='replace')

master_path = r'D:/behaviour analysis/Nifty50_12_Quarters_Quarterly_Results_Comparison_Master.xlsx'
wb = openpyxl.load_workbook(master_path, data_only=True)

print("==========================================================================================")
print("VERIFYING 14-SHEET QUARTERLY RESULTS COMPARISON MASTER WORKBOOK")
print("==========================================================================================")
print("Total Sheets Count:", len(wb.sheetnames))
print("Sheet Names       :", wb.sheetnames)

ws1 = wb['Quarterly_Results_Summary']
print(f"\nSheet 1 'Quarterly_Results_Summary' Rows 7-20:")
for r in range(7, 21):
    vals = [ws1.cell(row=r, column=c).value for c in range(1, 11)]
    print(f"  Row {r:2d}: {vals}")

ws_q3 = wb['Q3 2023-24']
print(f"\nSheet 3 'Q3 2023-24' Row 4 KPI Cards:")
cards = [ws_q3.cell(row=4, column=c).value for c in [1, 4, 7, 10, 13]]
print(f"  KPIs: {cards}")
print("First 3 Trade Rows in Q3 2023-24:")
for r in range(8, 11):
    vals = [ws_q3.cell(row=r, column=c).value for c in [1, 2, 5, 8, 9, 10, 12, 17, 19]]
    print(f"  Trade {r-7}: {vals}")

print("==========================================================================================")
