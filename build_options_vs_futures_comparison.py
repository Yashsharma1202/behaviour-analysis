"""
12-QUARTER COMPARISON: 1% ITM Nifty-50 Stock OPTIONS vs FUTURES.
Merges the two existing masters, per stock per quarter, into a side-by-side comparison
(exactly like the on-screen Q1 view but for all 12 quarters):
  Futures Realised P&L (1 lot)  vs  1% ITM Option P&L (1 lot)  ->  which vehicle won.

Sources (already built):
  GitHub_Dashboard_12_Quarters_Master.xlsx                 (futures Realised P&L)
  GitHub_Dashboard_12_Quarters_1PCT_ITM_Options_Master.xlsx (1% ITM Option P&L)
ADDITIVE: reads the two masters; writes a new comparison workbook.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

BASE = r'D:\behaviour analysis'
FUT = BASE + r'\GitHub_Dashboard_12_Quarters_Master.xlsx'
OPT = BASE + r'\GitHub_Dashboard_12_Quarters_1PCT_ITM_Options_Master.xlsx'
OUT = BASE + r'\GitHub_Dashboard_12_Quarters_Options_vs_Futures_Comparison.xlsx'


def read_master(path, pnl_label, ret_label):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for sn in wb.sheetnames:
        if sn == 'Dashboard_Summary':
            continue
        rows = list(wb[sn].iter_rows(values_only=True))
        hdr = rows[3]
        idx = {str(h).strip(): j for j, h in enumerate(hdr) if h is not None}
        def col(*names):
            for n in names:
                for k, j in idx.items():
                    if k.lower().startswith(n.lower()):
                        return j
            return None
        cS, cN, cB = col('Symbol'), col('Company'), col('Futures Bias', 'Option Bias', 'Bias')
        cP, cR = col(pnl_label), col(ret_label)
        d = {}
        for r in rows[4:]:
            if not r or r[cS] is None:
                continue
            sym = str(r[cS]).strip()
            d[sym] = dict(name=str(r[cN] or ''), dir=str(r[cB] or '').upper(),
                          pnl=r[cP], ret=r[cR])
        out[sn] = d
    return out


FUTD = read_master(FUT, 'Realised P&L', 'Return We Get')
OPTD = read_master(OPT, 'Option P&L', 'Option Return')
QUARTERS = [s for s in openpyxl.load_workbook(FUT, read_only=True).sheetnames if s != 'Dashboard_Summary']

# ---- styles ----
NAVY = PatternFill('solid', fgColor='1A202C'); WHITE = Font(bold=True, color='FFFFFF', size=10)
TITLE = Font(bold=True, color='1A202C', size=13); SUB = Font(color='4A5568', size=10, italic=True)
GREENF = PatternFill('solid', fgColor='E6F4EA'); REDF = PatternFill('solid', fgColor='FCE8E6')
BLUEF = PatternFill('solid', fgColor='E8F0FE')
GOOD = Font(color='137333', bold=True); BAD = Font(color='C5221F', bold=True)
THIN = Border(*[Side(style='thin', color='CBD5E0')] * 4)
CEN = Alignment(horizontal='center'); RIGHT = Alignment(horizontal='right')

HEADERS = ['#', 'Symbol', 'Company', 'Direction', 'Futures Realised P&L (1 Lot)',
           '1% ITM Option P&L (1 Lot)', 'Option \u2212 Futures', 'Better Vehicle',
           'Futures Ret %', 'Option Ret %']
WIDTHS = [5, 13, 28, 11, 26, 24, 18, 15, 13, 13]


def num(v):
    try:
        return float(v)
    except Exception:
        return None


out = openpyxl.Workbook()
summ = out.active; summ.title = 'Comparison_Summary'
qsum = []

for q in QUARTERS:
    fut = FUTD.get(q, {}); opt = OPTD.get(q, {})
    syms = list(fut.keys())  # keep futures order/universe (50 stocks)
    ws = out.create_sheet(q)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws.cell(1, 1, '%s  |  1%% ITM OPTIONS vs FUTURES  |  per-stock P&L comparison (1 lot)' % q).font = TITLE
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADERS))
    for j, h in enumerate(HEADERS, 1):
        c = ws.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN

    tot_fut = tot_opt = 0.0; opt_better = fut_better = 0; row = 5
    for i, sym in enumerate(syms, 1):
        f = fut[sym]; o = opt.get(sym, {})
        fp = num(f.get('pnl')); op = num(o.get('pnl'))
        fr = num(f.get('ret')); orr = num(o.get('ret'))
        diff = (op - fp) if (op is not None and fp is not None) else None
        if op is not None and fp is not None:
            better = 'OPTION' if op > fp else ('FUTURE' if fp > op else 'TIE')
            if better == 'OPTION': opt_better += 1
            elif better == 'FUTURE': fut_better += 1
        else:
            better = 'n/a'
        if fp is not None: tot_fut += fp
        if op is not None: tot_opt += op
        vals = [i, sym, f.get('name', ''), f.get('dir', ''),
                round(fp) if fp is not None else '\u2014',
                round(op) if op is not None else '\u2014',
                round(diff) if diff is not None else '\u2014', better,
                (round(fr * 100, 2) if fr is not None and abs(fr) < 5 else (round(fr, 2) if fr is not None else '\u2014')),
                (round(orr, 2) if orr is not None else '\u2014')]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row, j, v); c.border = THIN
            if j in (1, 4, 8): c.alignment = CEN
            if j in (5, 6, 7, 9, 10): c.alignment = RIGHT
            if j == 5 and fp is not None: c.font = GOOD if fp >= 0 else BAD
            if j == 6 and op is not None: c.font = GOOD if op >= 0 else BAD
            if j == 8: c.fill = GREENF if better == 'OPTION' else (BLUEF if better == 'FUTURE' else REDF)
        row += 1
    # totals row
    tr = ['', 'TOTAL', '%d stocks' % len(syms), '', round(tot_fut), round(tot_opt),
          round(tot_opt - tot_fut), 'OPT' if tot_opt > tot_fut else 'FUT', '', '']
    for j, v in enumerate(tr, 1):
        c = ws.cell(row, j, v); c.fill = NAVY; c.font = WHITE; c.border = THIN
        if j in (5, 6, 7): c.alignment = RIGHT
    ws.cell(2, 1, ('Futures net: Rs %s   |   1%% ITM Options net: Rs %s   |   Options won on %d/%d stocks, Futures on %d'
                   % (format(round(tot_fut), ','), format(round(tot_opt), ','), opt_better, len(syms), fut_better))).font = SUB
    for j, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.freeze_panes = 'A5'
    qsum.append((q, tot_fut, tot_opt, opt_better, fut_better, len(syms)))

# ---- summary sheet ----
summ.merge_cells('A1:H1'); summ.merge_cells('A2:H2')
summ.cell(1, 1, 'SMC GLOBAL \u2014 12 QUARTERS: 1% ITM OPTIONS vs FUTURES (per-stock P&L)').font = TITLE
summ.cell(2, 1, 'Same comparison as the live quarter view, across all 12 quarters. Net P&L, 1 lot per stock.').font = SUB
sh = ['Quarter', 'Futures Net P&L (Rs)', '1% ITM Options Net P&L (Rs)', 'Options \u2212 Futures',
      'Better Vehicle', 'Options Won (stocks)', 'Futures Won (stocks)', 'Total Stocks']
for j, h in enumerate(sh, 1):
    c = summ.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
tf = to = 0.0; tob = tfb = 0
for i, (q, fp, op, ob, fb, n) in enumerate(qsum, 5):
    tf += fp; to += op; tob += ob; tfb += fb
    vals = [q, round(fp), round(op), round(op - fp), 'OPTIONS' if op > fp else 'FUTURES', ob, fb, n]
    fill = GREENF if op > fp else BLUEF
    for j, v in enumerate(vals, 1):
        c = summ.cell(i, j, v); c.border = THIN; c.fill = fill
        if j in (2, 3, 4): c.alignment = RIGHT
        if j in (5, 6, 7, 8): c.alignment = CEN
r = 5 + len(qsum)
vals = ['CONSOLIDATED', round(tf), round(to), round(to - tf), 'OPTIONS' if to > tf else 'FUTURES', tob, tfb, sum(x[5] for x in qsum)]
for j, v in enumerate(vals, 1):
    c = summ.cell(r, j, v); c.fill = NAVY; c.font = WHITE; c.border = THIN
    if j in (2, 3, 4): c.alignment = RIGHT
    if j in (5, 6, 7, 8): c.alignment = CEN
for j, w in enumerate([16, 22, 26, 18, 15, 18, 18, 12], 1):
    summ.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
summ.freeze_panes = 'A5'

out.save(OUT)
print('Saved:', OUT)
print('12-quarter totals: Futures Rs %s  vs  1%% ITM Options Rs %s  (diff %s)'
      % (format(round(tf), ','), format(round(to), ','), format(round(to - tf), ',')))
print('Options beat futures in %d of 12 quarters; stock-level: options won %d, futures won %d'
      % (sum(1 for x in qsum if x[2] > x[1]), tob, tfb))
