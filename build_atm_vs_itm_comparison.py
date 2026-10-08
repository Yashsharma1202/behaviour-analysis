"""
12-QUARTER STRIKE COMPARISON: FUTURES vs 1% ITM OPTIONS vs ATM OPTIONS.
Recomputes both option legs from the same entry/exit spot prices in the futures master
(synthetic model: 1% time-value entry / 0.5% exit, 0.15% cost). ATM strike = spot;
1% ITM strike = spot*0.99 (LONG) / spot*1.01 (SHORT). Per-stock per-quarter P&L (1 lot),
quarter totals, and a Rs 20L/quarter performance summary for each option strike.

CAVEAT (in the workbook): the model uses a FLAT time value for every strike. Real ATM
options carry MORE time value than ITM, so this understates ATM cost and flatters ATM.
ADDITIVE: reads the futures master; writes a new comparison workbook.
"""
import openpyxl, statistics as st, re
from datetime import date
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

BASE = r'D:\behaviour analysis'
OUT = BASE + r'\GitHub_Dashboard_Futures_vs_ITM_vs_ATM_Comparison.xlsx'
EXT, EXT_X, COST, CAP = 0.01, 0.005, 0.0015, 2_000_000
ORDER = ['FY24_Q3', 'FY24_Q4', 'FY25_Q1', 'FY25_Q2', 'FY25_Q3', 'FY25_Q4', 'FY26_Q1', 'FY26_Q2', 'FY26_Q3', 'FY26_Q4', 'FY27_Q1', 'FY27_Q2']
_MONS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']


def num(v):
    try: return float(v)
    except Exception: return None


def pdate(s):
    s = str(s)
    m = re.search(r'(\d{4})-(\d{2})-(\d{2})', s)
    if m: return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.search(r'(\d{1,2})-([A-Za-z]{3})-(\d{4})', s)
    if m: return date(int(m.group(3)), _MONS.index(m.group(2)) + 1, int(m.group(1)))
    return None


def dshort(d):
    return d.strftime('%d-%b-%Y') if d else '\u2014'


def opt_leg(side, en, ex, lot, mode):
    call = 'LONG' in side.upper()
    k = en if mode == 'ATM' else en * (0.99 if call else 1.01)
    ie = max(0, (en - k) if call else (k - en))
    ix = max(0, (ex - k) if call else (k - ex))
    pen = ie + en * EXT; pex = ix + ex * EXT_X
    if pen <= 0: return None, None, None
    ret = pex / pen - 1 - COST
    return ret * pen * lot, pen * lot, ret * 100   # pnl, premium, ret%


fw = openpyxl.load_workbook(BASE + r'\GitHub_Dashboard_12_Quarters_Master.xlsx', read_only=True, data_only=True)
DATA = {}
for q in ORDER:
    rows = list(fw[q].iter_rows(values_only=True)); hdr = rows[3]
    idx = {str(h).strip(): j for j, h in enumerate(hdr) if h}
    def C(n):
        for k, j in idx.items():
            if k.lower().startswith(n.lower()): return j
    cS, cN, cB, cE, cX, cL, cP = C('Symbol'), C('Company'), C('Futures Bias'), C('Entry Price'), C('Exit Price'), C('Lot Size'), C('Realised P&L')
    cED, cXD = C('Entry Date'), C('Exit Date')
    recs = []
    for r in rows[4:]:
        if not r or r[cS] is None: continue
        en, ex, lot, fpnl = num(r[cE]), num(r[cX]), num(r[cL]), num(r[cP])
        if None in (en, ex, lot): continue
        itm, itm_prem, _ = opt_leg(str(r[cB]), en, ex, lot, 'ITM')
        atm, atm_prem, _ = opt_leg(str(r[cB]), en, ex, lot, 'ATM')
        recs.append(dict(sym=str(r[cS]), name=str(r[cN] or ''), dir=str(r[cB] or ''),
                         fut=fpnl, itm=itm, atm=atm, itm_prem=itm_prem, atm_prem=atm_prem,
                         ed=str(r[cED] or ''), xd=str(r[cXD] or ''), edt=pdate(r[cED]), xdt=pdate(r[cXD])))
    DATA[q] = recs

# styles
NAVY = PatternFill('solid', fgColor='1A202C'); WHITE = Font(bold=True, color='FFFFFF', size=10)
TITLE = Font(bold=True, color='1A202C', size=13); SUB = Font(color='4A5568', size=10, italic=True)
GF = PatternFill('solid', fgColor='E6F4EA'); RF = PatternFill('solid', fgColor='FCE8E6')
BF = PatternFill('solid', fgColor='E8F0FE'); AM = PatternFill('solid', fgColor='FEF3C7'); LBL = PatternFill('solid', fgColor='F1F5F9')
GOOD = Font(color='137333', bold=True); BAD = Font(color='C5221F', bold=True); BOLD = Font(bold=True)
THIN = Border(*[Side(style='thin', color='CBD5E0')] * 4)
CEN = Alignment(horizontal='center'); RIGHT = Alignment(horizontal='right'); WRAP = Alignment(wrap_text=True, vertical='top')

wb = openpyxl.Workbook()
HEAD = ['#', 'Symbol', 'Company', 'Dir', 'Entry Date', 'Exit Date', 'Futures P&L', '1% ITM Option P&L', 'ATM Option P&L', 'Best Vehicle']
WID = [5, 13, 24, 8, 13, 13, 14, 17, 15, 13]
DAYIDX = {5, 6}          # date columns (centre)
PNLIDX = {7, 8, 9}       # P&L columns (right, coloured)
qtot = {}
for q in ORDER:
    recs = DATA[q]
    entries = [r['edt'] for r in recs if r['edt']]
    exits = [r['xdt'] for r in recs if r['xdt']]
    first_en = min(entries) if entries else None
    last_ex = max(exits) if exits else None
    dur = (last_ex - first_en).days if (first_en and last_ex) else None
    ws = wb.create_sheet(q)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEAD))
    ws.cell(1, 1, '%s  |  FUTURES vs 1%% ITM vs ATM OPTIONS  |  per-stock P&L (1 lot)' % q).font = TITLE
    # ROW 2 = quarter trading window (first entry -> last exit -> duration)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEAD))
    ws.cell(2, 1, 'QUARTER WINDOW:  First Entry %s   →   Last Exit (quarter close) %s   |   Duration: %s calendar days'
            % (dshort(first_en), dshort(last_ex), (dur if dur is not None else '—'))).font = Font(bold=True, color='2B6CB0', size=11)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(HEAD))
    for j, h in enumerate(HEAD, 1):
        c = ws.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
    tf = ti = ta = 0.0; row = 5
    for i, r in enumerate(recs, 1):
        f, it, at = r['fut'] or 0, r['itm'] or 0, r['atm'] or 0
        tf += f; ti += it; ta += at
        best = max([('Futures', f), ('1% ITM', it), ('ATM', at)], key=lambda x: x[1])[0]
        vals = [i, r['sym'], r['name'][:22], r['dir'].replace('BUY ', ''),
                dshort(r['edt']), dshort(r['xdt']), round(f), round(it), round(at), best]
        fill = GF if best != 'Futures' else BF
        for j, v in enumerate(vals, 1):
            c = ws.cell(row, j, v); c.border = THIN; c.fill = fill
            if j in (1, 4, 10) or j in DAYIDX: c.alignment = CEN
            if j in PNLIDX: c.alignment = RIGHT; c.font = GOOD if v >= 0 else BAD
        row += 1
    tr = ['', 'TOTAL', '', '', dshort(first_en), dshort(last_ex), round(tf), round(ti), round(ta),
          max([('FUT', tf), ('ITM', ti), ('ATM', ta)], key=lambda x: x[1])[0]]
    for j, v in enumerate(tr, 1):
        c = ws.cell(row, j, v); c.fill = NAVY; c.font = WHITE; c.border = THIN
        if j in PNLIDX or j in DAYIDX: c.alignment = RIGHT if j in PNLIDX else CEN
    ws.cell(3, 1, 'Net: Futures Rs %s  |  1%% ITM Rs %s  |  ATM Rs %s' % (format(round(tf), ','), format(round(ti), ','), format(round(ta), ','))).font = SUB
    for j, w in enumerate(WID, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.freeze_panes = 'A5'
    qtot[q] = dict(fut=tf, itm=ti, atm=ta,
                   itm_prem=sum(r['itm_prem'] or 0 for r in DATA[q]),
                   atm_prem=sum(r['atm_prem'] or 0 for r in DATA[q]))

# summary + 20L performance
cs = wb.active; cs.title = 'Summary_and_20L'
cs.merge_cells('A1:E1'); cs.cell(1, 1, 'FUTURES vs 1% ITM vs ATM OPTIONS \u2014 12 quarters + Rs 20L/quarter performance').font = TITLE
cs.merge_cells('A2:E2'); cs.cell(2, 1, 'Net P&L per quarter (1 lot each) and return on a Rs 20L/quarter options book. See Caveat sheet.').font = SUB
hdr = ['Quarter', 'Futures Net (Rs)', '1% ITM Net (Rs)', 'ATM Net (Rs)', 'Best']
for j, h in enumerate(hdr, 1):
    c = cs.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
tf = ti = ta = 0.0
for i, q in enumerate(ORDER, 5):
    d = qtot[q]; tf += d['fut']; ti += d['itm']; ta += d['atm']
    best = max([('FUT', d['fut']), ('ITM', d['itm']), ('ATM', d['atm'])], key=lambda x: x[1])[0]
    for j, v in enumerate([q, round(d['fut']), round(d['itm']), round(d['atm']), best], 1):
        c = cs.cell(i, j, v); c.border = THIN; c.fill = GF if best != 'FUT' else BF
        if j in (2, 3, 4): c.alignment = RIGHT
        if j == 5: c.alignment = CEN
r = 5 + 12
for j, v in enumerate(['CONSOLIDATED', round(tf), round(ti), round(ta), max([('FUT', tf), ('ITM', ti), ('ATM', ta)], key=lambda x: x[1])[0]], 1):
    c = cs.cell(r, j, v); c.fill = NAVY; c.font = WHITE; c.border = THIN
    if j in (2, 3, 4): c.alignment = RIGHT

# 20L performance for ITM and ATM (return on premium)
def perf(key, prem):
    rets = [qtot[q][key] / qtot[q][prem] * 100 if qtot[q][prem] else 0 for q in ORDER]
    return rets
itm_r = perf('itm', 'itm_prem'); atm_r = perf('atm', 'atm_prem')
r0 = r + 2
cs.cell(r0, 1, 'Rs 20 LAKH / QUARTER PERFORMANCE (return on premium deployed)').font = BOLD
cs.cell(r0 + 1, 1, 'Measure').font = BOLD; cs.cell(r0 + 1, 2, '1% ITM').font = BOLD; cs.cell(r0 + 1, 3, 'ATM').font = BOLD
for j in (1, 2, 3):
    cs.cell(r0 + 1, j).fill = NAVY; cs.cell(r0 + 1, j).font = WHITE
def m(rets):
    n = len(rets); std = st.pstdev(rets)
    return dict(total=sum(r / 100 * CAP for r in rets), avg=st.mean(rets),
                wr=100 * sum(1 for r in rets if r > 0) / n, best=max(rets), worst=min(rets),
                sharpe=st.mean(rets) / std * 2 if std else 0)
MI, MA = m(itm_r), m(atm_r)
prows = [('Total P&L (12 x Rs 20L)', 'Rs {:+,.0f}'.format(MI['total']), 'Rs {:+,.0f}'.format(MA['total'])),
         ('Avg return / quarter', '{:+.1f}%'.format(MI['avg']), '{:+.1f}%'.format(MA['avg'])),
         ('Quarter win rate', '{:.0f}%'.format(MI['wr']), '{:.0f}%'.format(MA['wr'])),
         ('Best / Worst quarter', '{:+.0f}% / {:+.0f}%'.format(MI['best'], MI['worst']), '{:+.0f}% / {:+.0f}%'.format(MA['best'], MA['worst'])),
         ('Sharpe (annualised)', '{:.2f}'.format(MI['sharpe']), '{:.2f}'.format(MA['sharpe']))]
for k, (lbl, vi, va) in enumerate(prows):
    rr = r0 + 2 + k
    cs.cell(rr, 1, lbl).fill = LBL; cs.cell(rr, 1).font = BOLD; cs.cell(rr, 1).border = THIN
    cs.cell(rr, 2, vi).border = THIN; cs.cell(rr, 3, va).border = THIN
    cs.cell(rr, 2).alignment = CEN; cs.cell(rr, 3).alignment = CEN
for j, w in enumerate([30, 18, 18, 16, 8], 1):
    cs.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w

# caveat
cv = wb.create_sheet('Caveat')
cv.merge_cells('A1:B1'); cv.cell(1, 1, 'CAVEAT \u2014 why ATM looks (too) good here').font = Font(bold=True, color='7C2D12', size=12)
notes = [('Flat time-value model', 'Both strikes use the SAME 1% entry / 0.5% exit time value. In reality ATM options carry the MOST time value (higher than ITM), so this UNDERSTATES ATM premium and FLATTERS ATM returns. Real ATM would be more expensive and win less than shown.'),
         ('Why ATM shows smaller losses', 'ATM premium here is ~1% of spot vs ~2% for 1% ITM, so less capital is at risk on losers - hence smaller losses. With realistic (higher) ATM time value that gap shrinks.'),
         ('Same optimism as before', '1% synthetic model + 2023-26 bull sample = optimistic ceiling. Treat absolute returns as indicative; use for RELATIVE strike comparison, not as tradeable P&L. Re-run at realistic per-strike IV for a true picture.')]
rr = 3
for lbl, txt in notes:
    a = cv.cell(rr, 1, lbl); a.font = BOLD; a.fill = AM; a.border = THIN; a.alignment = WRAP
    cv.cell(rr, 2, txt).alignment = WRAP; cv.cell(rr, 2).border = THIN
    cv.row_dimensions[rr].height = 66; rr += 1
cv.column_dimensions['A'].width = 24; cv.column_dimensions['B'].width = 105

wb.move_sheet('Summary_and_20L', -(len(wb.sheetnames) - 1))
wb.save(OUT)
print('Saved:', OUT)
print('12-qtr net: Futures Rs {:+,.0f} | 1% ITM Rs {:+,.0f} | ATM Rs {:+,.0f}'.format(tf, ti, ta))
print('Rs20L/qtr: ITM total Rs {:+,.0f} ({:+.0f}%/q) | ATM total Rs {:+,.0f} ({:+.0f}%/q)'.format(MI['total'], MI['avg'], MA['total'], MA['avg']))
