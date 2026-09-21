import openpyxl
import pandas as pd
import numpy as np
import pathlib
import sys
import time
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference

sys.stdout.reconfigure(errors='replace')

t0 = time.time()

# ------------------- Configuration & Paths -------------------
BASE_DIR = pathlib.Path('D:/behaviour analysis')
SOURCE_FUT = BASE_DIR / 'Nifty50_12_Quarters_Futures_OI_Master_v13.xlsx'
OUTPUT_MASTER = BASE_DIR / 'Nifty50_12_Quarters_Options_1PCT_ITM_Master.xlsx'
OUTPUT_SUMMARY = BASE_DIR / 'Nifty50_12_Quarters_Options_1PCT_ITM_Summary.xlsx'

# Executive Color Tokens
COLOR_TITLE_BG     = '1B365D'
COLOR_SECTION_BG   = '2B4C7E'
COLOR_HEADER_BG    = '1A202C'
COLOR_SUBHEADER_BG = '2C5282'
COLOR_ZEBRA_BG     = 'F8FAFC'
COLOR_CARD_BG      = 'F7FAFC'

COLOR_WIN_BG       = 'E6F4EA'
COLOR_WIN_TEXT     = '137333'
COLOR_LOSS_BG      = 'FCE8E6'
COLOR_LOSS_TEXT    = 'C5221F'
COLOR_GOLD_BG      = 'FEF3C7'
COLOR_GOLD_TEXT    = '92400E'

COLOR_WHITE        = 'FFFFFF'
COLOR_TEXT_MAIN    = '2D3748'

font_title     = Font(name='Segoe UI', size=14, bold=True, color=COLOR_WHITE)
font_section   = Font(name='Segoe UI', size=11, bold=True, color=COLOR_WHITE)
font_header    = Font(name='Segoe UI', size=10, bold=True, color=COLOR_WHITE)
font_card_lbl  = Font(name='Segoe UI', size=9, bold=True, color=COLOR_WHITE)
font_data_bold = Font(name='Segoe UI', size=10, bold=True, color=COLOR_TEXT_MAIN)
font_data_reg  = Font(name='Segoe UI', size=10, color=COLOR_TEXT_MAIN)

font_win       = Font(name='Segoe UI', size=10, bold=True, color=COLOR_WIN_TEXT)
font_loss      = Font(name='Segoe UI', size=10, bold=True, color=COLOR_LOSS_TEXT)
font_gold      = Font(name='Segoe UI', size=10, bold=True, color=COLOR_GOLD_TEXT)

fill_title     = PatternFill(start_color=COLOR_TITLE_BG, end_color=COLOR_TITLE_BG, fill_type='solid')
fill_section   = PatternFill(start_color=COLOR_SECTION_BG, end_color=COLOR_SECTION_BG, fill_type='solid')
fill_header    = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type='solid')
fill_subheader = PatternFill(start_color=COLOR_SUBHEADER_BG, end_color=COLOR_SUBHEADER_BG, fill_type='solid')
fill_zebra     = PatternFill(start_color=COLOR_ZEBRA_BG, end_color=COLOR_ZEBRA_BG, fill_type='solid')
fill_card      = PatternFill(start_color=COLOR_CARD_BG, end_color=COLOR_CARD_BG, fill_type='solid')

fill_win       = PatternFill(start_color=COLOR_WIN_BG, end_color=COLOR_WIN_BG, fill_type='solid')
fill_loss       = PatternFill(start_color=COLOR_LOSS_BG, end_color=COLOR_LOSS_BG, fill_type='solid')
fill_gold      = PatternFill(start_color=COLOR_GOLD_BG, end_color=COLOR_GOLD_BG, fill_type='solid')

border_thin = Border(
    left=Side(style='thin', color='CBD5E0'),
    right=Side(style='thin', color='CBD5E0'),
    top=Side(style='thin', color='CBD5E0'),
    bottom=Side(style='thin', color='CBD5E0')
)

border_double_bottom = Border(
    left=Side(style='thin', color='CBD5E0'),
    right=Side(style='thin', color='CBD5E0'),
    top=Side(style='thin', color='1A202C'),
    bottom=Side(style='double', color='1A202C')
)

# Function to round strike price to nearest standard NSE strike interval
def round_nse_strike(price):
    if price > 5000: interval = 50
    elif price > 1000: interval = 20
    elif price > 500: interval = 10
    else: interval = 5
    return round(price / interval) * interval

quarters_order = [
    'Q3 2023-24', 'Q4 2023-24', 'Q1 2024-25', 'Q2 2024-25',
    'Q3 2024-25', 'Q4 2024-25', 'Q1 2025-26', 'Q2 2025-26',
    'Q3 2025-26', 'Q4 2025-26', 'Q1 2026-27', 'Q2 2026-27'
]

print("Loading Master Futures Trade Log for Options Calculation...", flush=True)
xl_src = pd.ExcelFile(SOURCE_FUT, engine='openpyxl')
df_src_trades = xl_src.parse('All_12Q_Futures_Trades')

options_processed_trades = []

for qtr in quarters_order:
    q_tr = df_src_trades[df_src_trades['Quarter'] == qtr].copy()
    
    t_no_col = next(c for c in q_tr.columns if 'Trade' in c)
    q_tr['Entry Date DT'] = pd.to_datetime(q_tr['Entry Date'])
    q_tr = q_tr.sort_values(by=['Entry Date DT', t_no_col]).reset_index(drop=True)
    
    for idx, row in q_tr.iterrows():
        sym      = str(row['Symbol']).strip()
        lot_size = int(row['Lot Size (Qty)'])
        s_entry  = float(row['Entry Futures Price (₹)'])
        s_exit   = float(row['Exit Futures Price (₹)'])
        strat    = str(row['Strategy']).strip().upper()
        
        is_call  = 'LONG' in strat or 'BUY' in strat
        opt_type = '1% ITM CALL (CE)' if is_call else '1% ITM PUT (PE)'
        
        # Strike Selection: 1% ITM
        if is_call:
            raw_strike = s_entry * 0.99
            strike = round_nse_strike(raw_strike)
            intrinsic_entry = max(0.0, s_entry - strike)
            intrinsic_exit  = max(0.0, s_exit - strike)
        else:
            raw_strike = s_entry * 1.01
            strike = round_nse_strike(raw_strike)
            intrinsic_entry = max(0.0, strike - s_entry)
            intrinsic_exit  = max(0.0, strike - s_exit)
            
        extrinsic_entry = s_entry * 0.02
        p_entry = intrinsic_entry + extrinsic_entry
        
        extrinsic_exit = s_exit * 0.005
        p_exit = intrinsic_exit + extrinsic_exit
        
        opt_pnl = round((p_exit - p_entry) * lot_size, 2)
        capital_outlay = round(p_entry * lot_size, 2)
        opt_return = round((p_exit - p_entry) / p_entry, 4) if p_entry > 0 else 0.0
        
        trade_rec = {
            'Quarter': qtr,
            'FY': str(row.get('FY', '')),
            'Trade No': idx + 1,
            'Symbol': sym,
            'Company Name': str(row.get('Company Name', '')),
            'Sector': str(row.get('Sector', '')),
            'Option Strategy': opt_type,
            'Position Taking Window': str(row.get('Position Taking Window', '')),
            'Entry Date': str(row['Entry Date']),
            'Exit Date': str(row['Exit Date']),
            'Entry Spot Price (₹)': s_entry,
            'Exit Spot Price (₹)': s_exit,
            '1% ITM Strike Price (₹)': strike,
            'Option Entry Premium (₹)': round(p_entry, 2),
            'Option Exit Premium (₹)': round(p_exit, 2),
            'Lot Size (Qty)': lot_size,
            'Capital Outlay per Trade (₹)': capital_outlay,
            'Booked Option P&L (₹)': opt_pnl,
            'Option Return (%)': opt_return,
            'Assigned Slot': str(row.get('Assigned Slot', f"Slot {(idx%15)+1}")),
            'Re-entry Type': str(row.get('Re-entry Type', 'First Entry')),
            'Quarterly Result Date': str(row.get('Quarterly Result Date', ''))
        }
        options_processed_trades.append(trade_rec)

df_all_opt_trades = pd.DataFrame(options_processed_trades)

# ------------------- 2. Calculate Options Summaries & Drawdown -------------------
opt_summaries = []

for qtr in quarters_order:
    q_tr = df_all_opt_trades[df_all_opt_trades['Quarter'] == qtr].copy()
    q_tr['Entry Date DT'] = pd.to_datetime(q_tr['Entry Date'])
    q_tr = q_tr.sort_values(by=['Entry Date DT', 'Trade No']).reset_index(drop=True)
    
    t_count = len(q_tr)
    wins = len(q_tr[q_tr['Booked Option P&L (₹)'] > 0])
    losses = len(q_tr[q_tr['Booked Option P&L (₹)'] <= 0])
    win_rate = (wins / t_count) if t_count > 0 else 0.0
    
    tot_opt_pnl = round(q_tr['Booked Option P&L (₹)'].sum(), 2)
    avg_slot_outlay = sum(r['Capital Outlay per Trade (₹)'] for _, r in q_tr.iterrows()) / t_count if t_count > 0 else 22500.0
    
    # 15 Active Concurrent Slots Premium Capital Pool
    slots_deployed = 15
    tot_opt_capital = round(avg_slot_outlay * slots_deployed, 2)
    final_val_quarter = round(tot_opt_capital + tot_opt_pnl, 2)
    opt_booked_return = (tot_opt_pnl / tot_opt_capital) if tot_opt_capital > 0 else 0.0
    
    # Chronological Cumulative & Peak Drawdown
    q_tr['Cum PnL'] = q_tr['Booked Option P&L (₹)'].cumsum()
    q_tr['Running Max'] = np.maximum.accumulate(q_tr['Cum PnL'])
    q_tr['Drawdown'] = q_tr['Cum PnL'] - q_tr['Running Max']
    
    trough_idx = q_tr['Drawdown'].idxmin()
    max_dd_val = abs(q_tr.loc[trough_idx, 'Drawdown'])
    peak_cum_pnl = q_tr.loc[:trough_idx, 'Cum PnL'].max()
    peak_portfolio_val = tot_opt_capital + max(0.0, peak_cum_pnl)
    max_dd_pct_peak = -abs(max_dd_val / peak_portfolio_val) if peak_portfolio_val > 0 else 0.0
    
    opt_summaries.append({
        'Quarter': qtr,
        'FY': q_tr.iloc[0]['FY'] if 'FY' in q_tr.columns else '',
        'Total Trades': t_count,
        'Wins': wins,
        'Losses': losses,
        'Win Rate (%)': win_rate,
        'Option Capital Deployed (₹)': tot_opt_capital,
        'Net Option P&L (₹)': tot_opt_pnl,
        'Final Quarter Value (₹)': final_val_quarter,
        'Booked Return (%)': opt_booked_return,
        'Slots Deployed': slots_deployed,
        'Peak Cum P&L (₹)': round(peak_cum_pnl, 2),
        'Trough Cum P&L (₹)': round(q_tr.loc[trough_idx, 'Cum PnL'], 2),
        'Max Drawdown (₹)': round(max_dd_val, 2),
        'True Peak DD (%)': round(max_dd_pct_peak, 4),
        'Trough Stock': q_tr.loc[trough_idx, 'Symbol']
    })

df_opt_summary = pd.DataFrame(opt_summaries)

# ------------------- 3. Build Full Options Master Excel Workbook -------------------
print("Building Master Options 1% ITM Excel Workbook with Corrected 15-Slot Margin Pool...", flush=True)
wb = Workbook()
wb.remove(wb.active)

def style_table_headers(ws, start_row, headers, fill_style=fill_header):
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=c_idx, value=h)
        cell.font = font_header; cell.fill = fill_style
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border_thin

def autofit_columns(ws, max_cols=25):
    ws.views.sheetView[0].showGridLines = True
    for col_idx in range(1, max_cols + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in range(1, min(ws.max_row + 1, 100)):
            val = str(ws.cell(row=row, column=col_idx).value or '')
            if len(val) > max_len: max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

# --- SHEET 1: Exec_12Q_Combined_Summary ---
ws_exec = wb.create_sheet(title='Exec_12Q_Combined_Summary')

ws_exec.merge_cells('A1:G1')
c_t = ws_exec['A1']
c_t.value = "NIFTY 50 1% ITM OPTIONS BUYING STRATEGY — 12-QUARTER EXECUTIVE SCORECARD"
c_t.font = font_title; c_t.fill = fill_title; c_t.alignment = Alignment(horizontal='center', vertical='center')
ws_exec.row_dimensions[1].height = 38

global_tot_pnl = round(df_all_opt_trades['Booked Option P&L (₹)'].sum(), 2)
global_win_rate = len(df_all_opt_trades[df_all_opt_trades['Booked Option P&L (₹)'] > 0]) / len(df_all_opt_trades)

ws_exec.merge_cells('A3:B3')
c_h1 = ws_exec['A3']; c_h1.value = "TOTAL TRADES ANALYZED"; c_h1.font = font_card_lbl; c_h1.fill = fill_subheader; c_h1.alignment = Alignment(horizontal='center', vertical='center')
ws_exec.merge_cells('A4:B4')
c_v1 = ws_exec['A4']; c_v1.value = len(df_all_opt_trades); c_v1.font = Font(name='Segoe UI', size=15, bold=True, color=COLOR_TEXT_MAIN); c_v1.fill = fill_card; c_v1.alignment = Alignment(horizontal='center', vertical='center')

ws_exec.merge_cells('C3:D3')
c_h2 = ws_exec['C3']; c_h2.value = "GLOBAL WIN RATE"; c_h2.font = font_card_lbl; c_h2.fill = fill_subheader; c_h2.alignment = Alignment(horizontal='center', vertical='center')
ws_exec.merge_cells('C4:D4')
c_v2 = ws_exec['C4']; c_v2.value = global_win_rate; c_v2.font = Font(name='Segoe UI', size=15, bold=True, color=COLOR_WIN_TEXT); c_v2.fill = fill_win; c_v2.number_format = '0.0%'; c_v2.alignment = Alignment(horizontal='center', vertical='center')

ws_exec.merge_cells('E3:G3')
c_h3 = ws_exec['E3']; c_h3.value = "TOTAL 12-QUARTER NET OPTIONS PROFIT"; c_h3.font = font_card_lbl; c_h3.fill = fill_subheader; c_h3.alignment = Alignment(horizontal='center', vertical='center')
ws_exec.merge_cells('E4:G4')
c_v3 = ws_exec['E4']; c_v3.value = global_tot_pnl; c_v3.font = Font(name='Segoe UI', size=15, bold=True, color=COLOR_WIN_TEXT); c_v3.fill = fill_win; c_v3.number_format = '₹#,##0.00'; c_v3.alignment = Alignment(horizontal='center', vertical='center')

ws_exec.row_dimensions[3].height = 20
ws_exec.row_dimensions[4].height = 28

for r in [3, 4]:
    for c in range(1, 8): ws_exec.cell(row=r, column=c).border = border_thin

ws_exec.merge_cells('A6:G6')
c_sec = ws_exec['A6']
c_sec.value = "CONSOLIDATED 12-QUARTER OPTIONS PERFORMANCE & PEAK DRAWDOWN SCORECARD"
c_sec.font = font_section; c_sec.fill = fill_section; c_sec.alignment = Alignment(horizontal='left', vertical='center')
ws_exec.row_dimensions[6].height = 24

headers_summary = ['Quarter', '15-Slot Option Premium Pool (₹)', 'Net Option Realised P&L (₹)', 'Booked Return (%)', 'Max Realised Drawdown (₹)', 'True Peak DD (%)', 'Win Rate (%)']
style_table_headers(ws_exec, 7, headers_summary, fill_header)
ws_exec.row_dimensions[7].height = 26

for r_idx, row_data in df_opt_summary.iterrows():
    row_num = 8 + r_idx
    row_fill = fill_zebra if r_idx % 2 == 1 else PatternFill(fill_type=None)
    ws_exec.row_dimensions[row_num].height = 21
    
    ws_exec.cell(row=row_num, column=1, value=row_data['Quarter']).alignment = Alignment(horizontal='center')
    ws_exec.cell(row=row_num, column=2, value=row_data['Option Capital Deployed (₹)']).number_format = '₹#,##0.00'
    
    c_pnl = ws_exec.cell(row=row_num, column=3, value=row_data['Net Option P&L (₹)'])
    c_pnl.number_format = '₹#,##0.00'; c_pnl.fill = fill_win if row_data['Net Option P&L (₹)'] >= 0 else fill_loss
    c_pnl.font = font_win if row_data['Net Option P&L (₹)'] >= 0 else font_loss
    
    c_ret = ws_exec.cell(row=row_num, column=4, value=row_data['Booked Return (%)'])
    c_ret.number_format = '+0.00%;-0.00%'
    
    c_dd = ws_exec.cell(row=row_num, column=5, value=-abs(row_data['Max Drawdown (₹)']))
    c_dd.number_format = '-₹#,##0.00'; c_dd.font = font_loss; c_dd.fill = fill_loss
    
    c_ddp = ws_exec.cell(row=row_num, column=6, value=row_data['True Peak DD (%)'])
    c_ddp.number_format = '-0.00%'; c_ddp.font = font_loss; c_ddp.fill = fill_loss; c_ddp.alignment = Alignment(horizontal='center')
    
    c_wr = ws_exec.cell(row=row_num, column=7, value=row_data['Win Rate (%)'])
    c_wr.number_format = '0.00%'; c_wr.alignment = Alignment(horizontal='center')
    if row_data['Win Rate (%)'] >= 0.70: c_wr.fill = fill_win; c_wr.font = font_win
    elif row_data['Win Rate (%)'] >= 0.50: c_wr.fill = fill_gold; c_wr.font = font_gold
    else: c_wr.fill = fill_loss; c_wr.font = font_loss

    for c in range(1, 8):
        cell = ws_exec.cell(row=row_num, column=c)
        cell.border = border_thin
        if c not in [3, 5, 6, 7] and row_fill.fill_type: cell.fill = row_fill

# Row 20: 12-QUARTER TOTAL ROW
r_tot = 20
ws_exec.row_dimensions[r_tot].height = 24
ws_exec.cell(row=r_tot, column=1, value="12-QUARTER TOTAL").font = font_data_bold
ws_exec.cell(row=r_tot, column=2, value="—").alignment = Alignment(horizontal='center')
ws_exec.cell(row=r_tot, column=3, value=global_tot_pnl).number_format = '₹#,##0.00'
ws_exec.cell(row=r_tot, column=3).font = font_win; ws_exec.cell(row=r_tot, column=3).fill = fill_win
ws_exec.cell(row=r_tot, column=4, value="—").alignment = Alignment(horizontal='center')
ws_exec.cell(row=r_tot, column=5, value=-abs(df_opt_summary['Max Drawdown (₹)'].max())).number_format = '-₹#,##0.00'
ws_exec.cell(row=r_tot, column=5).font = font_loss; ws_exec.cell(row=r_tot, column=5).fill = fill_loss
ws_exec.cell(row=r_tot, column=6, value=df_opt_summary['True Peak DD (%)'].min()).number_format = '-0.00%'
ws_exec.cell(row=r_tot, column=6).font = font_loss; ws_exec.cell(row=r_tot, column=6).fill = fill_loss; ws_exec.cell(row=r_tot, column=6).alignment = Alignment(horizontal='center')
c_tot_wr = ws_exec.cell(row=r_tot, column=7, value=global_win_rate)
c_tot_wr.number_format = '0.00%'; c_tot_wr.font = font_win; c_tot_wr.fill = fill_win; c_tot_wr.alignment = Alignment(horizontal='center')

for c in range(1, 8): ws_exec.cell(row=r_tot, column=c).border = border_thin

# Row 21: QUARTERLY AVERAGE ROW
r_avg = 21
ws_exec.row_dimensions[r_avg].height = 24
ws_exec.cell(row=r_avg, column=1, value="QUARTERLY AVERAGE").font = font_data_bold
ws_exec.cell(row=r_avg, column=2, value=df_opt_summary['Option Capital Deployed (₹)'].mean()).number_format = '₹#,##0.00'
ws_exec.cell(row=r_avg, column=3, value=df_opt_summary['Net Option P&L (₹)'].mean()).number_format = '₹#,##0.00'
ws_exec.cell(row=r_avg, column=3).font = font_win; ws_exec.cell(row=r_avg, column=3).fill = fill_win
ws_exec.cell(row=r_avg, column=4, value=df_opt_summary['Booked Return (%)'].mean()).number_format = '+0.00%'
ws_exec.cell(row=r_avg, column=5, value=-abs(df_opt_summary['Max Drawdown (₹)'].mean())).number_format = '-₹#,##0.00'
ws_exec.cell(row=r_avg, column=5).font = font_loss; ws_exec.cell(row=r_avg, column=5).fill = fill_loss
ws_exec.cell(row=r_avg, column=6, value=df_opt_summary['True Peak DD (%)'].mean()).number_format = '-0.00%'
ws_exec.cell(row=r_avg, column=6).font = font_loss; ws_exec.cell(row=r_avg, column=6).fill = fill_loss; ws_exec.cell(row=r_avg, column=6).alignment = Alignment(horizontal='center')
c_avg_wr = ws_exec.cell(row=r_avg, column=7, value=global_win_rate)
c_avg_wr.number_format = '0.00%'; c_avg_wr.font = font_win; c_avg_wr.fill = fill_win; c_avg_wr.alignment = Alignment(horizontal='center')

for c in range(1, 8): ws_exec.cell(row=r_avg, column=c).border = border_double_bottom

autofit_columns(ws_exec, max_cols=7)

# Chart
chart_opt = LineChart()
chart_opt.title = "Master 12-Quarter Net Option Buying P&L Trend (₹)"
chart_opt.style = 13
chart_opt.y_axis.title = "Net Option P&L (₹)"
chart_opt.x_axis.title = "Quarter"
chart_opt.width = 16
chart_opt.height = 11

data_o = Reference(ws_exec, min_col=3, min_row=7, max_row=19)
cats_o = Reference(ws_exec, min_col=1, min_row=8, max_row=19)
chart_opt.add_data(data_o, titles_from_data=True)
chart_opt.set_categories(cats_o)
ws_exec.add_chart(chart_opt, "I3")

# --- SHEETS 2-13: INDIVIDUAL QUARTER SHEETS ---
trade_headers = [
    'Trade #', 'Symbol', 'Company Name', 'Sector', 'Option Strategy',
    'Position Taking Window', 'Entry Date', 'Exit Date',
    'Entry Spot Price (₹)', 'Exit Spot Price (₹)', '1% ITM Strike Price (₹)',
    'Option Entry Premium (₹)', 'Option Exit Premium (₹)', 'Lot Size (Qty)',
    'Capital Outlay per Trade (₹)', 'Booked Option P&L (₹)', 'Cumulative P&L (₹)', 'Drawdown (₹)',
    'Assigned Slot', 'Re-entry Type', 'Quarterly Result Date'
]

for qtr in quarters_order:
    ws_q = wb.create_sheet(title=qtr)
    
    q_tr = df_all_opt_trades[df_all_opt_trades['Quarter'] == qtr].copy()
    q_tr['Entry Date DT'] = pd.to_datetime(q_tr['Entry Date'])
    q_tr = q_tr.sort_values(by=['Entry Date DT', 'Trade No']).reset_index(drop=True)
    
    q_sum = df_opt_summary[df_opt_summary['Quarter'] == qtr].iloc[0]
    
    # Cumulative & Drawdown
    q_tr['Cumulative P&L (₹)'] = q_tr['Booked Option P&L (₹)'].cumsum()
    q_tr['Running Max'] = np.maximum.accumulate(q_tr['Cumulative P&L (₹)'])
    q_tr['Drawdown (₹)'] = q_tr['Cumulative P&L (₹)'] - q_tr['Running Max']
    
    # Title Banner
    ws_q.merge_cells('A1:U1')
    c_t = ws_q['A1']
    c_t.value = f"NIFTY 50 1% ITM OPTIONS BUYING MODEL — {qtr.upper()}"
    c_t.font = font_title; c_t.fill = fill_title; c_t.alignment = Alignment(horizontal='center', vertical='center')
    ws_q.row_dimensions[1].height = 36
    
    # Card Banners Row 2-4
    card_lbls = ['15-Slot Option Premium Pool (₹)', 'Final Value (₹)', 'Net Option P&L (₹)', 'Booked Return (%)', 'Win Rate (%)', 'Total Trades', 'Winning Trades', 'Losing Trades', 'Max Drawdown (₹)', 'True Peak DD (%)', 'Slots Deployed']
    ws_q.merge_cells('A2:K2')
    c_s = ws_q['A2']; c_s.value = "1% ITM OPTIONS PERFORMANCE & CAPITAL OVERVIEW"; c_s.font = font_section; c_s.fill = fill_section
    ws_q.row_dimensions[2].height = 24
    
    style_table_headers(ws_q, 3, card_lbls, fill_subheader)
    ws_q.row_dimensions[3].height = 22
    
    ws_q.cell(row=4, column=1, value=q_sum['Option Capital Deployed (₹)']).number_format = '₹#,##0.00'
    ws_q.cell(row=4, column=2, value=q_sum['Final Quarter Value (₹)']).number_format = '₹#,##0.00'
    c_pnl = ws_q.cell(row=4, column=3, value=q_sum['Net Option P&L (₹)'])
    c_pnl.number_format = '₹#,##0.00'; c_pnl.font = font_win if q_sum['Net Option P&L (₹)'] >= 0 else font_loss; c_pnl.fill = fill_win if q_sum['Net Option P&L (₹)'] >= 0 else fill_loss
    
    ws_q.cell(row=4, column=4, value=q_sum['Booked Return (%)']).number_format = '0.00%'
    c_wr = ws_q.cell(row=4, column=5, value=q_sum['Win Rate (%)'])
    c_wr.number_format = '0.00%'; c_wr.font = font_win; c_wr.fill = fill_win; c_wr.alignment = Alignment(horizontal='center')
    
    ws_q.cell(row=4, column=6, value=int(q_sum['Total Trades'])).alignment = Alignment(horizontal='center')
    ws_q.cell(row=4, column=7, value=int(q_sum['Wins'])).alignment = Alignment(horizontal='center')
    ws_q.cell(row=4, column=8, value=int(q_sum['Losses'])).alignment = Alignment(horizontal='center')
    
    c_dd1 = ws_q.cell(row=4, column=9, value=-abs(q_sum['Max Drawdown (₹)']))
    c_dd1.number_format = '-₹#,##0.00'; c_dd1.font = font_loss; c_dd1.fill = fill_loss
    
    c_dd2 = ws_q.cell(row=4, column=10, value=q_sum['True Peak DD (%)'])
    c_dd2.number_format = '-0.00%'; c_dd2.font = font_loss; c_dd2.fill = fill_loss; c_dd2.alignment = Alignment(horizontal='center')
    
    ws_q.cell(row=4, column=11, value=int(q_sum['Slots Deployed'])).alignment = Alignment(horizontal='center')
    ws_q.row_dimensions[4].height = 25
    
    for c in range(1, 12): ws_q.cell(row=4, column=c).border = border_thin
    
    # Trade Journal Row 6+
    ws_q.merge_cells('A6:U6')
    c_s2 = ws_q['A6']; c_s2.value = "1% ITM OPTIONS TRADE JOURNAL — ENTRY DATE CHRONOLOGICAL ORDER"; c_s2.font = font_section; c_s2.fill = fill_section
    ws_q.row_dimensions[6].height = 24
    
    style_table_headers(ws_q, 7, trade_headers, fill_header)
    ws_q.row_dimensions[7].height = 26
    
    for r_idx, t_row in q_tr.iterrows():
        row_num = 8 + r_idx
        row_fill = fill_zebra if r_idx % 2 == 1 else PatternFill(fill_type=None)
        ws_q.row_dimensions[row_num].height = 21
        
        ws_q.cell(row=row_num, column=1, value=r_idx + 1).alignment = Alignment(horizontal='center')
        ws_q.cell(row=row_num, column=2, value=t_row['Symbol']).alignment = Alignment(horizontal='center')
        ws_q.cell(row=row_num, column=3, value=t_row['Company Name'])
        ws_q.cell(row=row_num, column=4, value=t_row['Sector'])
        
        c_st = ws_q.cell(row=row_num, column=5, value=t_row['Option Strategy'])
        c_st.alignment = Alignment(horizontal='center')
        if 'CALL' in t_row['Option Strategy']: c_st.fill = fill_win; c_st.font = font_win
        else: c_st.fill = fill_loss; c_st.font = font_loss
        
        ws_q.cell(row=row_num, column=6, value=t_row['Position Taking Window']).alignment = Alignment(horizontal='center')
        ws_q.cell(row=row_num, column=7, value=t_row['Entry Date']).alignment = Alignment(horizontal='center')
        ws_q.cell(row=row_num, column=8, value=t_row['Exit Date']).alignment = Alignment(horizontal='center')
        
        ws_q.cell(row=row_num, column=9, value=t_row['Entry Spot Price (₹)']).number_format = '₹#,##0.00'
        ws_q.cell(row=row_num, column=10, value=t_row['Exit Spot Price (₹)']).number_format = '₹#,##0.00'
        ws_q.cell(row=row_num, column=11, value=t_row['1% ITM Strike Price (₹)']).number_format = '₹#,##0.00'
        
        ws_q.cell(row=row_num, column=12, value=t_row['Option Entry Premium (₹)']).number_format = '₹#,##0.00'
        ws_q.cell(row=row_num, column=13, value=t_row['Option Exit Premium (₹)']).number_format = '₹#,##0.00'
        ws_q.cell(row=row_num, column=14, value=int(t_row['Lot Size (Qty)'])).alignment = Alignment(horizontal='center')
        ws_q.cell(row=row_num, column=15, value=t_row['Capital Outlay per Trade (₹)']).number_format = '₹#,##0.00'
        
        c_op = ws_q.cell(row=row_num, column=16, value=t_row['Booked Option P&L (₹)'])
        c_op.number_format = '₹#,##0.00'
        c_op.font = font_win if t_row['Booked Option P&L (₹)'] >= 0 else font_loss
        c_op.fill = fill_win if t_row['Booked Option P&L (₹)'] >= 0 else fill_loss
        
        c_cp = ws_q.cell(row=row_num, column=17, value=t_row['Cumulative P&L (₹)'])
        c_cp.number_format = '₹#,##0.00'; c_cp.font = font_data_bold
        
        c_dd = ws_q.cell(row=row_num, column=18, value=t_row['Drawdown (₹)'])
        c_dd.number_format = '₹#,##0.00'
        if t_row['Drawdown (₹)'] < 0: c_dd.font = font_loss; c_dd.fill = fill_loss
        
        ws_q.cell(row=row_num, column=19, value=t_row['Assigned Slot']).alignment = Alignment(horizontal='center')
        ws_q.cell(row=row_num, column=20, value=t_row.get('Re-entry Type', 'First Entry')).alignment = Alignment(horizontal='center')
        ws_q.cell(row=row_num, column=21, value=t_row.get('Quarterly Result Date', '')).alignment = Alignment(horizontal='center')
        
        for c in range(1, 22):
            cell = ws_q.cell(row=row_num, column=c)
            cell.border = border_thin
            if c not in [5, 16, 18] and row_fill.fill_type: cell.fill = row_fill

    autofit_columns(ws_q, max_cols=21)
    ws_q.freeze_panes = 'A8'

    # Line Chart
    chart_q = LineChart()
    chart_q.title = f"{qtr} — 1% ITM Option Buying Cumulative P&L & Drawdown Curve (Max DD: -₹{q_sum['Max Drawdown (₹)']:,.2f})"
    chart_q.style = 13
    chart_q.y_axis.title = "Rupees (₹)"
    chart_q.x_axis.title = "Trade #"
    chart_q.width = 16
    chart_q.height = 10

    max_r = 7 + len(q_tr)
    data_q = Reference(ws_q, min_col=17, min_row=7, max_col=18, max_row=max_r)
    cats_q = Reference(ws_q, min_col=1, min_row=8, max_row=max_r)
    chart_q.add_data(data_q, titles_from_data=True)
    chart_q.set_categories(cats_q)
    ws_q.add_chart(chart_q, "W6")

# Save Master Workbook
print(f"Saving Options 1% ITM Master Workbook to {OUTPUT_MASTER.name}...", flush=True)
wb.save(OUTPUT_MASTER)

print(f"\n🎉 1% ITM OPTIONS MASTER WORKBOOK UPDATED IN {time.time() - t0:.1f}s!", flush=True)
