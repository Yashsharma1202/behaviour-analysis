"""
ALL F&O STOCKS — DUSSEHRA, RANKED BY WIN RATE % (percentage-wise report).
Companion to the avg-return-ranked report: same 211-stock universe, same
backtest, but primary sort is Win Rate % (ties broken by avg return) instead
of avg return \u2014 so you can see exactly who wins most often, not just who
returns most.

ADDITIVE: reads the existing intermediate pickle + live dashboard data; writes
new report files only.
"""
import json, pickle, os, subprocess, statistics as st
from datetime import date, timedelta, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

BASE = r'D:\behaviour analysis'
OUT_XLSX = os.environ.get('FO211_WR_OUT_XLSX', BASE + r'\FO211_Dussehra_WinRate_Ranked_Report.xlsx')
OUT_HTML = BASE + r'\FO211_Dussehra_WinRate_Ranked_Report.html'
OUT_PDF = BASE + r'\FO211_Dussehra_WinRate_Ranked_Report.pdf'

data = pickle.load(open(BASE + r'\scraped_parquet\_tmp_fo211_dussehra.pkl', 'rb'))
results, meta = data['results'], data['meta']
qualified = {s: r for s, r in results.items() if r['pick']}
insufficient = {s: r for s, r in results.items() if not r['pick']}
ranked = sorted(qualified.items(), key=lambda kv: (-kv[1]['pick_wr'], -(kv[1]['pick_avg'] if kv[1]['pick_avg'] is not None else -999)))
n_low_conf = sum(1 for _, r in ranked if r.get('low_confidence'))

# ---- upcoming Dussehra 2026 positions (same method as the other reports) ----
NSE_2026_HOLIDAYS = {
    # '2026-03-04' removed (not a real NSE holiday -- leftover from Holi's date
    # once being wrongly typed as 04-Mar; real Holi holiday is 03-Mar, kept
    # below). '2026-11-08'/'2026-11-10' (Diwali) added -- were missing entirely.
    '2026-01-26', '2026-03-03', '2026-03-26', '2026-03-31', '2026-04-03',
    '2026-04-14', '2026-05-01', '2026-05-28', '2026-06-26', '2026-09-14', '2026-10-02',
    '2026-10-20', '2026-11-08', '2026-11-10', '2026-11-24', '2026-12-25',
}
def is_trading_day(d):
    return d.weekday() < 5 and d.isoformat() not in NSE_2026_HOLIDAYS
def shift_trading_days(anchor, n):
    d = anchor
    step = 1 if n >= 0 else -1
    cnt = abs(n)
    while cnt > 0:
        d += timedelta(days=step)
        if is_trading_day(d): cnt -= 1
    return d
DUSSEHRA_ANCHOR_2026 = date(2026, 10, 20)
_live = json.load(open(BASE + r'\dashboard_data\event_dashboard_data.json', encoding='utf-8'))
_live_dussehra = [h for h in _live['holidays'] if h['id'] == 'dussehra_2026'][0]
LIVE_POSITIONS = {s['symbol']: s for s in _live_dussehra['stocks']}
def position_for(sym, r):
    if sym in LIVE_POSITIONS:
        s = LIVE_POSITIONS[sym]
        return dict(direction=s['direction'], entry=s['entry_date'], exit=s['exit_date'], lot=s.get('lot'))
    en = shift_trading_days(DUSSEHRA_ANCHOR_2026, -r['lead'])
    ex = shift_trading_days(DUSSEHRA_ANCHOR_2026, r['hold'])
    return dict(direction=r['pick'], entry=en.strftime('%d-%b-%Y'), exit=ex.strftime('%d-%b-%Y'), lot=r.get('lot'))

def perf_measures(r):
    sign = 1 if r['pick'] == 'LONG' else -1
    years_sorted = sorted(r['year_rets'].keys())
    signed = [r['year_rets'][y] * sign for y in years_sorted]
    n = len(signed)
    cum = 1.0
    for v in signed: cum *= (1 + v / 100)
    cagr = (cum ** (1 / n) - 1) * 100 if n > 0 and cum > 0 else None
    eq = [100.0]
    for v in signed: eq.append(eq[-1] * (1 + v / 100))
    peak = eq[0]; maxdd = 0.0
    for v in eq[1:]:
        peak = max(peak, v)
        dd = (v - peak) / peak * 100
        maxdd = min(maxdd, dd)
    sharpe = (st.mean(signed) / st.pstdev(signed)) if n > 1 and st.pstdev(signed) > 0 else None
    return dict(cagr=cagr, maxdd=maxdd, sharpe=sharpe, n=n)

NAVY = PatternFill('solid', fgColor='1A202C'); WHITE = Font(bold=True, color='FFFFFF', size=9)
TITLE = Font(bold=True, color='1A202C', size=13); SUB = Font(color='4A5568', size=10, italic=True)
GF = PatternFill('solid', fgColor='E6F4EA'); RF = PatternFill('solid', fgColor='FCE8E6')
AM = PatternFill('solid', fgColor='FEF3C7'); LBL = PatternFill('solid', fgColor='F1F5F9')
GOOD = Font(color='137333', bold=True); BAD = Font(color='C5221F', bold=True); BOLD = Font(bold=True)
THIN = Border(*[Side(style='thin', color='D9E2EC')] * 4)
CEN = Alignment(horizontal='center'); RIGHT = Alignment(horizontal='right'); WRAP = Alignment(wrap_text=True, vertical='top')

wb = Workbook()

# ============ Sheet 1: Ranked by Win Rate % ============
rk = wb.active; rk.title = 'Ranked_By_WinRate'
rk.merge_cells('A1:M1')
rk.cell(1, 1, 'ALL F&O STOCKS \u2014 DUSSEHRA, RANKED BY WIN RATE %% (percentage-wise)').font = TITLE
rk.merge_cells('A2:M2')
rk.cell(2, 1, f'{len(qualified)} of {len(results)} F&O stocks have real data ({n_low_conf} low-confidence, n<3 yrs \u2014 amber). Sorted by Win Rate % (ties broken by avg return). {len(insufficient)} have zero data.').font = SUB
hdr = ['Rank', 'Symbol', 'Company', 'Sector', 'Direction', 'Win Rate %', 'Avg Return %', 'CAGR %', 'Max DD %', 'Years', 'Confidence', 'Data Source', 'Window']
for j, hh in enumerate(hdr, 1):
    c = rk.cell(4, j, hh); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
row = 5
perf_cache = {}
for i, (sym, r) in enumerate(ranked, 1):
    m = meta.get(sym, {})
    p = perf_measures(r); perf_cache[sym] = p
    avg = r['pick_avg']; low = r.get('low_confidence')
    vals = [i, sym, m.get('name', sym), m.get('sector', ''), r['pick'], r['pick_wr'],
            (round(avg, 2) if avg is not None else '\u2014'),
            (round(p['cagr'], 2) if p['cagr'] is not None else '\u2014'), round(p['maxdd'], 2),
            r['n'], ('LOW (n<3)' if low else 'OK'), r['source'], f"T-{r['lead']} to T+{r['hold']}"]
    fill = AM if low else (GF if (avg is None or avg >= 0) else RF)
    for j, v in enumerate(vals, 1):
        c = rk.cell(row, j, v); c.border = THIN; c.fill = fill
        if j in (1, 5, 10, 11): c.alignment = CEN
        if j in (6, 7, 8, 9): c.alignment = RIGHT
        if j == 6: c.font = BOLD
        if j == 7 and isinstance(v, (int, float)): c.font = GOOD if v >= 0 else BAD
        if j == 9: c.font = BAD if p['maxdd'] < 0 else GOOD
        if j == 11 and low: c.font = Font(color='92400E', bold=True, size=9)
    row += 1
for j, w in enumerate([6, 12, 26, 20, 10, 11, 12, 9, 10, 7, 12, 30, 13], 1):
    rk.column_dimensions[get_column_letter(j)].width = w
rk.freeze_panes = 'A5'

# ============ Sheet 2: Win Rate Distribution (percentage bands) ============
db = wb.create_sheet('WinRate_Distribution')
db.merge_cells('A1:D1')
db.cell(1, 1, 'HOW MANY STOCKS FALL IN EACH WIN RATE BAND').font = TITLE
bands = [(100, 100, '100%'), (90, 99.9, '90-99%'), (80, 89.9, '80-89%'), (75, 79.9, '75-79%'),
         (60, 74.9, '60-74%'), (50, 59.9, '50-59%'), (0, 49.9, 'Below 50%')]
hdr = ['Win Rate Band', 'Stock Count', '% of Ranked Universe', 'OK Confidence / LOW Confidence']
for j, hh in enumerate(hdr, 1):
    c = db.cell(3, j, hh); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
row = 4
for lo, hi, label in bands:
    in_band = [r for _, r in ranked if lo <= r['pick_wr'] <= hi]
    ok_n = sum(1 for r in in_band if not r.get('low_confidence'))
    low_n = len(in_band) - ok_n
    pct = round(100 * len(in_band) / len(ranked), 1) if ranked else 0
    vals = [label, len(in_band), f'{pct}%', f'{ok_n} OK / {low_n} LOW']
    for j, v in enumerate(vals, 1):
        c = db.cell(row, j, v); c.border = THIN; c.alignment = CEN if j != 1 else Alignment(horizontal='left')
    row += 1
for j, w in enumerate([18, 14, 20, 28], 1):
    db.column_dimensions[get_column_letter(j)].width = w

# ============ Sheet 3: Insufficient-data stocks ============
ins = wb.create_sheet('Insufficient_Data_Stocks')
ins.merge_cells('A1:D1')
ins.cell(1, 1, f'{len(insufficient)} F&O stocks with ZERO Dussehra-window price data').font = TITLE
hdr = ['Symbol', 'Company', 'Years Available', 'Reason']
for j, hh in enumerate(hdr, 1):
    c = ins.cell(3, j, hh); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
row = 4
for sym, r in sorted(insufficient.items()):
    m = meta.get(sym, {})
    vals = [sym, m.get('name', sym), r['n'], 'Recently listed F&O contract \u2014 no trading days overlap any historical Dussehra window yet']
    for j, v in enumerate(vals, 1):
        c = ins.cell(row, j, v); c.border = THIN
        if j == 3: c.alignment = CEN
    row += 1
for j, w in enumerate([13, 28, 15, 55], 1):
    ins.column_dimensions[get_column_letter(j)].width = w

# ============ Sheet 4: Position Schedule 2026 ============
psch = wb.create_sheet('Position_Schedule_2026')
psch.merge_cells('A1:J1')
psch.cell(1, 1, 'DUSSEHRA 20-Oct-2026 \u2014 Position Schedule, ALL 211 F&O Stocks (Win-Rate-Ranked universe)').font = TITLE
psch.merge_cells('A2:J2')
psch.cell(2, 1, 'Entry/exit on the NSE-2026 trading calendar; Nifty-50 names use the exact live-dashboard dates. Sorted by entry date.').font = SUB
hdr = ['Symbol', 'Company', 'Direction', 'Window', 'Entry Date', 'Exit Date', 'Lot Size', 'Win Rate %', 'Years', 'Confidence']
for j, hh in enumerate(hdr, 1):
    c = psch.cell(4, j, hh); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
sched = []
for sym, r in qualified.items():
    pos = position_for(sym, r)
    m = meta.get(sym, {})
    sort_key = datetime.strptime(pos['entry'].split(' (')[0].strip(), '%d-%b-%Y')
    sched.append((sort_key, sym, m.get('name', sym), pos['direction'], f"T-{r['lead']} to T+{r['hold']}",
                  pos['entry'], pos['exit'], pos['lot'], r['pick_wr'], r['n'], r.get('low_confidence')))
for sym, r in insufficient.items():
    m = meta.get(sym, {})
    sched.append((datetime.max, sym, m.get('name', sym), '\u2014', '\u2014', '\u2014', '\u2014', '\u2014', '\u2014', '\u2014', False))
sched.sort(key=lambda t: (t[0], t[1]))
row = 5
for _, sym, name, direc, window, entry, exit_, lot, wr, n, low in sched:
    vals = [sym, name, direc, window, entry, exit_, lot, wr, n, ('LOW (n<3)' if low else ('OK' if direc != '\u2014' else '\u2014'))]
    fill = AM if low else (GF if direc == 'LONG' else (RF if direc == 'SHORT' else None))
    for j, v in enumerate(vals, 1):
        c = psch.cell(row, j, v); c.border = THIN
        if fill: c.fill = fill
        if j in (1, 3, 4, 5, 6, 9, 10): c.alignment = CEN
        if j in (7, 8): c.alignment = RIGHT
    row += 1
for j, w in enumerate([13, 28, 10, 15, 16, 16, 9, 11, 7, 12], 1):
    psch.column_dimensions[get_column_letter(j)].width = w
psch.freeze_panes = 'A5'

# ============ Sheet 5: Overview ============
ov = wb.create_sheet('Overview', 0)
ov.merge_cells('A1:B1')
ov.cell(1, 1, 'ALL F&O STOCKS \u2014 DUSSEHRA, WIN-RATE-RANKED REPORT').font = TITLE
notes = [
    ('Universe & ranking', f'ALL {len(results)} F&O stocks. {len(qualified)} have real, usable data and are ranked here by WIN RATE %% (not avg return \u2014 see the companion avg-return-ranked report for that view). Ties broken by avg return.'),
    ('Confidence \u2014 read first', f'{n_low_conf} of the {len(qualified)} ranked stocks are LOW confidence (n<3 years) \u2014 amber highlighted. A 100% win rate on 1 trade ranks #1 here on paper but is one data point, not a statistic. Check the Years/Confidence column before trusting any rank near the top.'),
    ('Win Rate Distribution', 'See the WinRate_Distribution sheet for a breakdown of how many stocks fall in each win-rate band, and how many of those are OK vs LOW confidence \u2014 useful for judging how real the overall edge is, not just the top few names.'),
    ('Method', 'For each stock, its own T-n/T+m window is used. Both LONG and SHORT are backtested split-robust (corporate-action days >20%% excluded from compounding); the direction with the higher win rate is chosen.'),
    ('Upcoming positions', 'See Position_Schedule_2026 for the 20-Oct-2026 entry/exit dates for all 211, sorted by entry date.'),
    ('CAVEAT', 'Best-direction-per-stock selection fits history closely but has repeatedly underperformed a single uniform index-backed direction out-of-sample in this project\u2019s own testing \u2014 treat ranks as historical fit, not guaranteed forward edge, especially for LOW-confidence (amber) rows.'),
]
r0 = 3
for lbl, txt in notes:
    a = ov.cell(r0, 1, lbl); a.font = BOLD; a.fill = AM; a.border = THIN; a.alignment = WRAP
    b = ov.cell(r0, 2, txt); b.alignment = WRAP; b.border = THIN
    ov.row_dimensions[r0].height = 68; r0 += 1
ov.column_dimensions['A'].width = 24; ov.column_dimensions['B'].width = 105

wb.save(OUT_XLSX)
print('Saved Excel:', OUT_XLSX)

# ============ PDF ============
def esc(s): return str(s).replace('&', '&amp;').replace('<', '&lt;')
top20 = ranked[:20]
def pdf_row(rank, sym, r, p):
    low = r.get('low_confidence')
    rowstyle = ' style="background:#fffbeb;"' if low else ''
    conf = '<span style="color:#b45309;font-weight:700;">LOW</span>' if low else 'OK'
    avg = r['pick_avg']; avg_str = f"{avg:+.2f}%" if avg is not None else '\u2014'
    return (f'<tr{rowstyle}><td>{rank}</td><td><b>{esc(sym)}</b></td><td>{esc(meta.get(sym,{}).get("name",sym))[:24]}</td>'
            f'<td class="c">{r["pick"]}</td><td class="r"><b>{r["pick_wr"]:.1f}%</b></td>'
            f'<td class="r {"pos" if (avg or 0)>=0 else "neg"}">{avg_str}</td>'
            f'<td class="c">{r["n"]}</td><td class="c">{conf}</td></tr>')
rows_top = ''.join(pdf_row(i + 1, s, r, perf_cache[s]) for i, (s, r) in enumerate(top20))
band_rows = ''
for lo, hi, label in bands:
    in_band = [r for _, r in ranked if lo <= r['pick_wr'] <= hi]
    ok_n = sum(1 for r in in_band if not r.get('low_confidence')); low_n = len(in_band) - ok_n
    band_rows += f'<tr><td>{label}</td><td class="c">{len(in_band)}</td><td class="c">{ok_n} OK / {low_n} LOW</td></tr>'

html = f"""<!doctype html><html><head><meta charset="utf-8">
<title>F&amp;O-211 Dussehra Win-Rate Ranked</title>
<style>
 *{{box-sizing:border-box}} body{{font-family:'Segoe UI',Arial,sans-serif;margin:0;color:#0f172a;background:#fff}}
 .wrap{{max-width:1000px;margin:0 auto;padding:26px 34px}}
 h1{{font-size:21px;margin:0 0 2px}} .sub{{color:#64748b;font-size:12px;margin-bottom:14px}}
 h2{{font-size:15px;margin:20px 0 8px;border-left:4px solid #7c3aed;padding-left:9px}}
 table{{width:100%;border-collapse:collapse;font-size:11.5px}}
 th,td{{padding:5px 7px;text-align:right;border-bottom:1px solid #eef2f7}} th{{background:#f1f5f9;font-size:9.5px;text-transform:uppercase;color:#475569}}
 td.c,th.c{{text-align:center}} .pos{{color:#059669;font-weight:700}} .neg{{color:#dc2626;font-weight:700}}
 .note{{background:#fff7ed;border:1px solid #fed7aa;border-radius:8px;padding:11px 14px;font-size:11px;color:#7c2d12;margin-top:8px;line-height:1.5}}
 .cards{{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;margin:10px 0}}
 .mc{{background:#fff;border:1px solid #e2e8f0;border-radius:8px;padding:10px 12px}}
 .mcl{{font-size:9.5px;text-transform:uppercase;color:#64748b;font-weight:700}} .mcv{{font-size:19px;font-weight:800}}
 .foot{{color:#94a3b8;font-size:9.5px;margin-top:18px}} @page{{size:A4;margin:12mm}}
</style></head><body><div class="wrap">
 <h1>All F&amp;O Stocks &mdash; Dussehra, Ranked by Win Rate %</h1>
 <div class="sub">ALL {len(results)} F&amp;O stocks covered &bull; {len(qualified)} ranked by Win Rate % (percentage-wise) &bull; companion to the avg-return-ranked report.</div>
 <div class="cards">
  <div class="mc"><div class="mcl">Ranked stocks</div><div class="mcv">{len(qualified)}</div></div>
  <div class="mc"><div class="mcl">Low-confidence (n&lt;3yr)</div><div class="mcv" style="color:#b45309">{n_low_conf}</div></div>
  <div class="mc"><div class="mcl">Zero data</div><div class="mcv">{len(insufficient)}</div></div>
  <div class="mc"><div class="mcl">Top win rate</div><div class="mcv pos">{ranked[0][1]['pick_wr']:.1f}%</div></div>
 </div>
 <h2>Top 20 by Win Rate</h2>
 <table><thead><tr><th>Rank</th><th class="c">Symbol</th><th class="c">Company</th><th class="c">Dir</th><th>Win%</th><th>Avg Ret</th><th class="c">Yrs</th><th class="c">Conf.</th></tr></thead>
 <tbody>{rows_top}</tbody></table>
 <h2>Win Rate Distribution</h2>
 <table><thead><tr><th>Band</th><th class="c">Count</th><th class="c">Confidence split</th></tr></thead><tbody>{band_rows}</tbody></table>
 <div class="note"><b>Caveat:</b> rows marked <b style="color:#b45309;">LOW</b> confidence have under 3 years of data (some just 1) &mdash; a "100% win rate" there is 1-2 data points, not a robust statistic; check the distribution table above for how much of the win-rate edge sits in thin samples. Best-direction-per-stock selection has repeatedly underperformed a uniform index-backed direction out-of-sample in this project's own testing.</div>
 <div class="foot">Full 211-stock detail in the companion Excel. Upcoming 20-Oct-2026 position schedule in Position_Schedule_2026.</div>
</div></body></html>"""
open(OUT_HTML, 'w', encoding='utf-8').write(html)
print('Saved HTML:', OUT_HTML)

for exe in [r'C:\Program Files\Google\Chrome\Application\chrome.exe', r'C:\Program Files (x86)\Google\Chrome\Application\chrome.exe']:
    if os.path.exists(exe):
        try:
            subprocess.run([exe, '--headless=new', '--disable-gpu', '--no-pdf-header-footer',
                            '--print-to-pdf=' + OUT_PDF, 'file:///' + OUT_HTML.replace('\\', '/')], timeout=60, capture_output=True)
            print('Saved PDF:', OUT_PDF, 'exists:', os.path.exists(OUT_PDF))
        except Exception as e:
            print('PDF render failed:', e)
        break
