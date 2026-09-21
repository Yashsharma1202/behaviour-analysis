import openpyxl, pandas as pd, sys

sys.stdout.reconfigure(errors='replace')

summary_path = r'D:/behaviour analysis/Nifty50_12_Quarters_Options_1PCT_ITM_40PCT_SL_Summary.xlsx'
wb = openpyxl.load_workbook(summary_path, data_only=True)
ws = wb['Options_40PCT_SL_Summary']

print("==========================================================================================")
print("VERIFYING STRICT -40% STOP LOSS OPTIONS SUMMARY WORKBOOK")
print("==========================================================================================")
print("Sheet Names:", wb.sheetnames)

print("\nSummary Table Rows 7-21:")
for r in range(7, 22):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 11)]
    print(f"  Row {r:2d}: {vals}")
print("==========================================================================================")
