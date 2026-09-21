"""
exp_bar_host.py  —  SEPARATE dashboard for the Expectation-Bar engine
===============================================================================
A standalone host (own port, default 8097) that renders the pre-event news
attribution for upcoming events. Kept OFF the main Nifty-50 dashboard until it
proves out — then we promote it.

    python exp_bar_host.py            # http://localhost:8097
    python exp_bar_host.py 9002       # custom port

Reads news.db (populated by news_ingest.py) via expectation_bar.score_batch.
It measures how HIGH the bar was set — never a prediction or trade view.
===============================================================================
"""
from __future__ import annotations

import datetime as dt
import json
import sys
import threading
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs

import expectation_bar as EB
import expectation_bar_backtest as BT
import news_effect_study as FX
import stock_server as S               # for the historical expected-move (behaviour engine)
import peer_backtest as PB             # peer read-across: predicted vs actual
import news_ingest                     # latest headlines for the detail view
import forward_track as FT             # snapshot predictions now, measure after

_LTP = None
_BEH: dict = {}


def _ltp():
    """symbol -> last traded price (from the Q1 sheet), for points conversion."""
    global _LTP
    if _LTP is None:
        _LTP = {}
        try:
            import openpyxl
            ws = openpyxl.load_workbook("Nifty50_Q1Results_BehaviourEngine.xlsx").active
            for r in range(3, ws.max_row + 1):
                sym = (ws.cell(r, 2).value or "").strip()
                v = ws.cell(r, 3).value
                if sym and isinstance(v, (int, float)):
                    _LTP[sym] = float(v)
        except Exception:                                  # noqa: BLE001
            pass
    return _LTP


_REAL = None


def _realised():
    """symbol -> realised Q1 move % (from the tracker), for peer read-across."""
    global _REAL
    if _REAL is None:
        _REAL = {}
        for fn in ("Nifty50_Q1Results_BehaviourEngine_UPDATED_06Aug_fixed.xlsx",
                   "UPDATED.xlsx"):
            try:
                import openpyxl
                ws = openpyxl.load_workbook(fn).active
                for r in range(3, ws.max_row + 1):
                    sym = (ws.cell(r, 2).value or "").strip()
                    o = ws.cell(r, 15).value
                    if sym and isinstance(o, (int, float)):
                        _REAL[sym] = float(o)
                if _REAL:
                    break
            except Exception:                              # noqa: BLE001
                continue
    return _REAL


def peer_expected_move(sym: str) -> dict | None:
    """Read-across prediction: average of sector peers' ACTUAL realised moves this
    quarter (data-grounded, not headline tone)."""
    real = _realised()
    sector = EB.SYM_SECTOR.get(sym)
    members = [(p, real[p]) for p in EB.SECTORS.get(sector, [])
               if p != sym and p in real]
    if not members:
        return None
    avg = sum(v for _, v in members) / len(members)
    ltp = _ltp().get(sym)
    return {"avg_pct": round(avg, 2), "n": len(members),
            "members": [{"peer": p, "ret": round(v, 2)} for p, v in members],
            "avg_points": round(avg / 100 * ltp, 1) if ltp else None}


def tracker_row(sym: str) -> dict | None:
    """The stock's full Performance-Review row from the tracker."""
    import openpyxl
    for fn in ("Nifty50_Q1Results_BehaviourEngine_UPDATED_06Aug_fixed.xlsx", "UPDATED.xlsx"):
        try:
            ws = openpyxl.load_workbook(fn).active
        except Exception:                                  # noqa: BLE001
            continue
        for r in range(3, ws.max_row + 1):
            if (ws.cell(r, 2).value or "").strip() == sym:
                g = lambda c: ws.cell(r, c).value
                dd = lambda v: v.date().isoformat() if isinstance(v, dt.datetime) else v
                o = g(15)
                status = ("realised" if isinstance(o, (int, float))
                          else "open" if isinstance(o, str) and "unrealised" in o
                          else "no-trade")
                return {"name": g(1), "ltp": g(3), "result": dd(g(4)),
                        "buy_before": g(5), "sell_after": g(6), "win": g(7),
                        "expected_return": g(9), "entry": dd(g(11)), "buy": g(12),
                        "exit": dd(g(13)), "sell": g(14), "realised": o, "status": status}
    return None


def stock_detail(sym: str) -> dict:
    tr = tracker_row(sym)
    det = {"symbol": sym, "name": (tr or {}).get("name") or sym,
           "sector": EB.SYM_SECTOR.get(sym), "tracker": tr,
           "expected_move": expected_move(sym),
           "peer_readacross": peer_expected_move(sym)}
    try:
        det["peer_bt"] = next((r for r in PB.compute()["rows"] if r["symbol"] == sym), None)
    except Exception:                                      # noqa: BLE001
        det["peer_bt"] = None
    try:
        rd = (tr or {}).get("result")
        when = dt.datetime.fromisoformat(rd) if rd else dt.datetime.now()
        sc = EB.score_batch(sym, when)
        det["news"] = {"bar": sc["expectation_bar_raw"], "no_coverage": sc["no_coverage"],
                       "counts": sc["counts"], "governance": sc["governance_flag"]}
    except Exception:                                      # noqa: BLE001
        det["news"] = None
    arts = news_ingest.articles_for(sym)[-6:][::-1]
    det["headlines"] = [{"title": a["title"][:120], "source": a["source"], "link": a["link"]}
                        for a in arts]
    return det


def expected_move(sym: str) -> dict | None:
    """Empirical move around this stock's events (NOT a forecast): avg/median/
    range/win-rate from the behaviour engine, plus a points estimate via LTP."""
    try:
        if sym not in _BEH:
            _BEH[sym] = S.build_behaviour(sym)
        b = _BEH[sym]
        types = b.get("types", {}) if b.get("available") else {}
        et = ("RESULTS" if "RESULTS" in types else
              "BOARD_MEETING" if "BOARD_MEETING" in types else
              (next(iter(types)) if types else None))
        if not et:
            return None
        best = types[et]["best"]
        ltp = _ltp().get(sym)
        avg = best.get("avg_return_pct")
        return {
            "event_type": et, "avg_pct": avg,
            "median_pct": best.get("median_pct"),
            "best_pct": best.get("best_case"), "worst_pct": best.get("worst_case"),
            "win_rate": best.get("win_rate_pct"), "n": best.get("n"),
            "buy_before": best.get("days_before"), "sell_after": best.get("days_after"),
            "ltp": ltp,
            "avg_points": round(avg / 100 * ltp, 1) if (ltp and avg is not None) else None,
            "worst_points": round(best["worst_case"] / 100 * ltp, 1) if (ltp and best.get("worst_case") is not None) else None,
            "best_points": round(best["best_case"] / 100 * ltp, 1) if (ltp and best.get("best_case") is not None) else None,
        }
    except Exception:                                      # noqa: BLE001
        return None

PAGE = r"""<!doctype html><html><head><meta charset="utf-8">
<title>Expectation Bar — Nifty 50</title><style>
:root{--bg:#0b0f17;--card:#111827;--line:#1f2937;--ink:#e6edf3;--mut:#8b98ad;--blue:#58a6ff;
 --pos:#3fb950;--neg:#f85149;--amber:#d29922;}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);
 font-family:-apple-system,"Segoe UI",system-ui,sans-serif}
.wrap{max-width:1100px;margin:0 auto;padding:26px}
h1{margin:0 0 2px;font-size:1.5rem}
.sub{color:var(--mut);font-size:.9rem;margin-bottom:6px}
.warn{background:#1a1330;border:1px solid #3b2d5e;border-radius:10px;
 padding:10px 14px;font-size:.82rem;margin:14px 0;color:#c7b9e6}
.warn b{color:#e6d9ff}
.row{display:flex;gap:10px;margin:14px 0;flex-wrap:wrap}
input,button{background:var(--card);border:1px solid var(--line);color:var(--ink);
 border-radius:8px;padding:9px 12px;font-size:.9rem}
button{background:var(--blue);color:#0b1220;font-weight:700;cursor:pointer;border:0}
.cards{display:grid;grid-template-columns:1fr 1fr;gap:14px}
@media(max-width:820px){.cards{grid-template-columns:1fr}}
.card{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:16px 18px}
.ctop{display:flex;justify-content:space-between;align-items:baseline}
.sym{font-weight:800;font-size:1.15rem}.date{color:var(--mut);font-size:.82rem}
.barwrap{margin:12px 0}
.barline{position:relative;height:12px;background:#0e1522;border:1px solid var(--line);border-radius:8px;overflow:hidden}
.zero{position:absolute;left:50%;top:0;bottom:0;width:1px;background:#3a465c}
.fill{position:absolute;top:0;bottom:0;border-radius:6px}
.barlab{display:flex;justify-content:space-between;font-size:.72rem;color:var(--mut);margin-top:3px}
.big{font-size:1.6rem;font-weight:800}
.emv{background:#0e1522;border:1px solid var(--line);border-left:3px solid var(--pos);border-radius:10px;padding:10px 13px;margin:6px 0 4px}
.tag{display:inline-block;font-size:.72rem;border:1px solid var(--line);border-radius:20px;
 padding:2px 10px;color:var(--mut);margin:2px 4px 2px 0}
.gov{background:#3a1d1d;border:1px solid #7a2b2b;color:#ffb4b4;border-radius:8px;padding:7px 11px;
 font-size:.8rem;margin:8px 0}
.null{color:var(--amber);font-weight:700}
.hl{font-size:.8rem;color:var(--mut);margin:3px 0}
.hl b{color:var(--ink);font-weight:600}
details summary{cursor:pointer;color:var(--blue);font-size:.82rem;margin-top:8px}
.mini{font-size:.76rem;color:var(--mut);margin:2px 0}
.pos{color:var(--pos)}.neg{color:var(--neg)}.mut{color:var(--mut)}
tr.clk{cursor:pointer}tr.clk:hover td{background:#182234}
.sec{font-weight:700;color:var(--blue);font-size:.9rem;margin:16px 0 6px;border-bottom:1px solid var(--line);padding-bottom:4px}
.grid2{display:grid;grid-template-columns:1fr 1fr;gap:6px 18px}
@media(max-width:640px){.grid2{grid-template-columns:1fr}}
.kv{display:flex;justify-content:space-between;gap:10px;border-bottom:1px dotted var(--line);padding:4px 0;font-size:.85rem}
.kk{color:var(--mut)}.kvv{font-weight:600;text-align:right}
.chartbox{background:#0e1522;border:1px solid var(--line);border-radius:10px;padding:10px 12px;margin:12px 0}
.chartttl{color:var(--ink);font-weight:700;font-size:.9rem;margin:2px 4px 6px}
.chartsub{color:var(--mut);font-size:.76rem;margin:0 4px 4px}
.foot{color:var(--mut);font-size:.75rem;margin-top:20px;line-height:1.5}
</style>
<script src="https://cdn.jsdelivr.net/npm/apexcharts"></script>
</head><body><div class="wrap">
<h1>📰 Expectation Bar</h1>
<div class="sub">Pre-event news attribution — how high was the bar set before each event.</div>
<div class="warn">⚠️ This measures the <b>bar</b>, not a prediction. A <b>high bar = good news already priced</b>,
 not "bullish". Experimental · headline-based (deterministic v1) · trustworthy only as daily ingestion accumulates.</div>
<div class="row">
  <input id="sym" placeholder="symbol e.g. APOLLOHOSP" style="flex:1;min-width:160px">
  <input id="ts" placeholder="decision date YYYY-MM-DD" style="width:190px">
  <button onclick="runOne()">Score event</button>
  <button onclick="loadUpcoming()" style="background:#1f2937;color:#e6edf3">↻ Upcoming</button>
  <button onclick="loadBacktest()" style="background:#1f2937;color:#e6edf3">📊 Backtest (bar vs move)</button>
  <button onclick="loadEffect()" style="background:#1f2937;color:#e6edf3">🔬 News features</button>
  <button onclick="loadPeers()" style="background:#1f2937;color:#e6edf3">🤝 Peer read-across</button>
  <button onclick="loadPeerBt()" style="background:#1f2937;color:#e6edf3">🎯 Peer: predicted vs actual</button>
  <button onclick="loadForward()" style="background:#1f2937;color:#e6edf3">🔮 Forward tracking</button>
</div>
<div id="out"></div>
<div class="foot">Bar range −1 … +1 · reactive articles excluded · coverage &lt;2 credible → <span class="null">null</span>
 (absence is a signal, not zero). Peer read-across shown where sector peers reported pre-event.
 Never a buy/sell view. Swap the deterministic engine for your LLM prompt (same schema) once an API key is wired.</div>
</div>
<script>
let lastView=null;                       // remembers the last list view for the ← back button
function backBtn(){return '<button onclick="(lastView||loadUpcoming)()" style="background:#1f2937;color:#e6edf3;margin-bottom:12px">← back</button>';}
function bar(v){
  if(v==null) return '<span class="null">null (no coverage)</span>';
  const pct=Math.abs(v)*50, col=v>=0?'var(--pos)':'var(--neg)';
  const left=v>=0?50:50-pct;
  return `<div class="big" style="color:${col}">${v>0?'+':''}${v.toFixed(2)}</div>
   <div class="barwrap"><div class="barline"><div class="zero"></div>
   <div class="fill" style="left:${left}%;width:${pct}%;background:${col}"></div></div>
   <div class="barlab"><span>−1 low bar</span><span>0</span><span>+1 high bar</span></div></div>`;
}
function card(r){
  const c=r.counts||{};
  const gov=r.governance_flag?`<div class="gov">⚠ governance: <b>${r.governance_flag.type}</b> — ${r.governance_flag.headline||''}</div>`:'';
  const peers=(r.peer_signals||[]).map(p=>`<span class="tag">${p.peer}: ${p.direction}${p.linked_to_subject?' ↔':''}</span>`).join('')||'<span class="mut" style="font-size:.78rem">none</span>';
  const hlink=a=>a.link?`<a href="${a.link}" target="_blank" rel="noopener" style="color:inherit;text-decoration:none;border-bottom:1px dotted var(--line)">${a.headline} ↗</a>`:a.headline;
  const scored=(r.articles||[]).filter(a=>a.bar_contrib).slice(0,6)
     .map(a=>`<div class="mini">${a.bar_contrib>0?'<span class=pos>+':'<span class=neg>'}${a.bar_contrib}</span> · ${a.classification} · ${hlink(a)}</div>`).join('');
  const recent=(r.articles||[]).slice(-6).reverse()
     .map(a=>`<div class="mini">${a.is_reactive?'<span class=mut>↩ reactive</span>':a.classification} · ${hlink(a)}</div>`).join('');
  const em=r.expected_move;
  const sgn=v=>v>0?'+':'';
  const emhtml = em ? `<div class="emv">
      <div class="hl"><b>📈 Historical move around ${em.event_type.toLowerCase().replace('_',' ')}</b> <span class="mut">(empirical — not a forecast)</span></div>
      <div class="big" style="color:${em.avg_pct>=0?'var(--pos)':'var(--neg)'}">${sgn(em.avg_pct)}${em.avg_pct}%${em.avg_points!=null?` <span style="font-size:1rem">≈ ${sgn(em.avg_points)}${em.avg_points} pts</span>`:''}</div>
      <div class="hl mut">median ${em.median_pct==null?'—':em.median_pct+'%'} · range ${em.worst_pct}%…${em.best_pct}%${em.worst_points!=null?` (${em.worst_points}…${sgn(em.best_points)}${em.best_points} pts)`:''} · win ${em.win_rate}% over ${em.n} past events</div>
      <div class="hl mut" style="font-size:.72rem">buy ${em.buy_before}c before → sell ${em.sell_after}c after · LTP ₹${em.ltp!=null?em.ltp:'—'}</div>
    </div>` : '<div class="hl mut">no behaviour history for expected move</div>';
  return `<div class="card">
    <div class="ctop"><span class="sym">${r.symbol}</span><span class="date">decision ${r.decision_ts.slice(0,10)}</span></div>
    ${emhtml}
    <div class="hl mut" style="margin:10px 0 2px">Expectation bar (news, context only — no proven edge):</div>
    ${bar(r.expectation_bar_raw)}
    ${gov}
    <div><span class="tag">pre ${c.pre_event||0}</span><span class="tag">reactive ${c.reactive||0}</span>
      <span class="tag">info ${c.informational||0}</span><span class="tag">anticip ${c.anticipatory||0}</span>
      <span class="tag">coverage: ${r.confidence?r.confidence.coverage:'?'}</span></div>
    <div class="hl mut" style="margin-top:8px;font-size:.76rem">Peer read-across → see the 🤝 tab</div>
    <button onclick="refresh('${r.symbol}','${r.decision_ts.slice(0,10)}')" style="margin-top:10px;background:#1f2937;color:#e6edf3;font-size:.78rem;padding:6px 10px">↻ refresh from Google News</button>
    <details><summary>▸ scored pre-event headlines (${(r.articles||[]).filter(a=>a.bar_contrib).length})</summary>${scored||'<div class=mini>—</div>'}</details>
    <details><summary>▸ latest Google News headlines</summary>${recent||'<div class=mini>—</div>'}</details>
  </div>`;
}
async function loadUpcoming(){
  lastView=loadUpcoming;
  const o=document.getElementById('out'); o.innerHTML='<div class="mut">scoring upcoming events…</div>';
  const d=await (await fetch('/api/upcoming')).json();
  if(!d.length){o.innerHTML='<div class="mut">No upcoming events in the feed window.</div>';return;}
  o.innerHTML='<div class="cards">'+d.map(card).join('')+'</div>';
}
async function refresh(sym,ts){
  const o=document.getElementById('out'); o.innerHTML='<div class="mut">pulling fresh Google News for '+sym+'…</div>';
  const d=await (await fetch('/api/ingest?sym='+encodeURIComponent(sym)+'&ts='+encodeURIComponent(ts))).json();
  o.innerHTML=backBtn()+'<div class="cards">'+card(d)+'</div>';
}
async function loadForward(){
  lastView=loadForward;
  const o=document.getElementById('out'); o.innerHTML='<div class="mut">capturing snapshots &amp; settling matured events…</div>';
  const d=await (await fetch('/api/forward')).json();
  const c=v=>v==null?'var(--mut)':(v>=0?'var(--pos)':'var(--neg)');
  const p=(v)=>v==null?'—':(v>0?'+':'')+v+'%';
  const stcol={realised:'var(--pos)',open:'var(--amber)',pending:'var(--mut)'};
  const rows=(d.rows||[]).map(r=>`<tr class="clk" onclick="openStock('${r.symbol}')" title="click for detail">
     <td style="text-align:left;font-weight:600;color:var(--blue)">${r.symbol} ›</td>
     <td class=mut>${r.decision_ts.slice(0,10)}</td>
     <td class=mut style="font-size:.74rem">${(r.snapshot||'').slice(0,10)}</td>
     <td style="text-align:right;color:${c(r.subj_pred)}">${p(r.subj_pred)}</td>
     <td style="text-align:right;color:${c(r.peer_pred)}">${p(r.peer_pred)}</td>
     <td style="text-align:right" class=mut>${r.bar==null?'—':r.bar}</td>
     <td style="text-align:right;color:${c(r.actual)};font-weight:700">${r.actual==null?'—':p(r.actual)}</td>
     <td style="text-align:left;color:${stcol[r.status]||'var(--mut)'};font-weight:600">${r.status}${r.status==='open'?' (running)':''}</td></tr>`).join('');
  o.innerHTML=`
    <div class="warn">🔮 <b>Forward tracking</b> — predictions are <b>snapshotted before</b> each event (locked with a timestamp), then the <b>actual</b> move is filled in after it reports. This is the clean, forward test that grows over time. Captured daily at 6 PM.</div>
    <div class="card">
      <div class="hl mut">${(d.rows||[]).length} tracked event(s) · realised are the finished ones (compare pred vs actual); open = entered &amp; running; pending = not entered yet.</div>
      <table style="width:100%;border-collapse:collapse;font-size:.82rem;margin-top:8px">
        <tr class=mut><th style="text-align:left">stock</th><th style="text-align:left">decision</th><th style="text-align:left">captured</th><th>subj pred</th><th>peer pred</th><th>bar</th><th>actual</th><th style="text-align:left">status</th></tr>
        ${rows||'<tr><td colspan=8 class=mini>no snapshots yet — captured automatically as events approach</td></tr>'}
      </table>
    </div>`;
}
function kv(k,v){return `<div class="kv"><span class="kk">${k}</span><span class="kvv">${v}</span></div>`;}
async function openStock(sym){
  const o=document.getElementById('out'); o.innerHTML='<div class="mut">loading '+sym+' detail…</div>';
  const d=await (await fetch('/api/stock_detail?sym='+encodeURIComponent(sym))).json();
  if(d.error){o.innerHTML='<div class="card">'+sym+' — '+d.error+'</div>';return;}
  const sg=v=>v>0?'+':''; const col=v=>v==null?'var(--mut)':(v>=0?'var(--pos)':'var(--neg)');
  const t=d.tracker||{}, em=d.expected_move, pr=d.peer_readacross, bt=d.peer_bt, nw=d.news;
  const realTxt = (t.realised==null)?'—':(typeof t.realised==='number'?`<span style="color:${col(t.realised)}">${sg(t.realised)}${t.realised}%</span>`:`<i>${t.realised}</i>`);
  // result / trade block
  const trade = t.result? `<div class="grid2">
     ${kv('Result date',t.result)} ${kv('Status',t.status)}
     ${kv('Rule',`buy ${t.buy_before}c before → sell ${t.sell_after}c after`)} ${kv('Win rate',t.win+'%')}
     ${kv('Entry',t.entry+'  @ ₹'+(t.buy??'—'))} ${kv('Exit',(t.exit??'—')+'  @ ₹'+(t.sell??'—'))}
     ${kv('Expected (history)',(t.expected_return!=null?sg(t.expected_return)+t.expected_return+'%':'—'))} ${kv('ACTUAL realised',realTxt)}
   </div>` : '<div class="hl mut">No Q1 result row for this stock.</div>';
  // predictions
  const emb = em?`${kv('Subject history',`${sg(em.avg_pct)}${em.avg_pct}% ≈ ${sg(em.avg_points)}${em.avg_points} pts · median ${em.median_pct}% · win ${em.win_rate}% (n=${em.n})`)}`:'';
  const prb = pr?`${kv('Peer read-across',`${sg(pr.avg_pct)}${pr.avg_pct}% ≈ ${sg(pr.avg_points)}${pr.avg_points} pts · ${pr.n} sector peers`)}`
    +`<div class="hl mut" style="font-size:.75rem;margin-top:2px">${pr.members.map(m=>m.peer+' '+sg(m.ret)+m.ret+'%').join(' · ')}</div>`:'';
  const btb = bt?`${kv('Peer PREDICTED vs ACTUAL',`pred ${bt.pred==null?'—':sg(bt.pred)+bt.pred+'%'} → actual ${sg(bt.actual)}${bt.actual}% ${bt.error==null?'':'(err '+sg(bt.error)+bt.error+')'} · basis: ${bt.basis||'—'}`)}`
    +(bt.members&&bt.members.length?`<div class="hl mut" style="font-size:.75rem">earlier sector reporters: ${bt.members.map(m=>m.peer+' '+sg(m.ret)+m.ret+'%').join(' · ')}</div>`
       :(bt.basis==='market'?`<div class="hl mut" style="font-size:.75rem">market fallback: avg of ${bt.n_peers} earlier reporters (no earlier sector peer)</div>`:''))
    +((bt.sector_peers&&bt.sector_peers.length)?`<div class="hl mut" style="font-size:.75rem">sector peer group: ${bt.sector_peers.map(p=>(bt.used_peers||[]).includes(p)?`<span style="color:var(--pos)">${p}</span>`:p).join(', ')}</div>`:'')  :'';
  // news
  const newsb = nw?`<div class="grid2">${kv('Expectation bar',nw.bar==null?'null (no pre-event coverage)':nw.bar)}
     ${kv('Coverage',nw.no_coverage?'none':'ok')}
     ${kv('Articles',`pre ${nw.counts.pre_event} · reactive ${nw.counts.reactive} · info ${nw.counts.informational}`)}
     ${kv('Governance',nw.governance?nw.governance.type:'—')}</div>
     <div class="hl mut" style="font-size:.72rem;margin-top:4px">Bar is usually null for past events (news ingested after the result) — that's expected.</div>`:'';
  const heads = (d.headlines||[]).map(h=>`<div class="mini">${h.link?`<a href="${h.link}" target="_blank" rel="noopener" style="color:inherit;border-bottom:1px dotted var(--line)">${h.title} ↗</a>`:h.title} <span class=mut>· ${h.source||''}</span></div>`).join('')||'<div class=mini>—</div>';
  o.innerHTML=`
    ${backBtn()}
    <div class="card">
      <div class="ctop"><span class="sym">${d.symbol}</span><span class="date">${d.name||''} · ${d.sector||'—'} · LTP ₹${t.ltp??'—'}</span></div>
      <div class="sec">📋 Result &amp; trade</div>${trade}
      <div class="sec">🎯 Predictions</div><div class="emv">${emb}${prb}${btb||'<div class=hl mut>no peer prediction (no earlier reporters)</div>'}</div>
      <div class="chartbox"><div class="chartttl">Predictions vs actual</div><div id="sdPred"></div></div>
      <div class="chartbox"><div class="chartttl">Sector peer moves</div><div id="sdPeers"></div></div>
      <div class="sec">📰 News / expectation bar</div>${newsb||'<div class=hl mut>no news scored</div>'}
      <div class="sec">🗞 Latest headlines</div>${heads}
    </div>`;
  renderStockCharts(d);
}
function renderStockCharts(d){
  if(typeof ApexCharts==='undefined')return;
  const em=d.expected_move, pr=d.peer_readacross, bt=d.peer_bt, t=d.tracker||{};
  const actual=(typeof t.realised==='number')?t.realised:null;
  const cats=[], vals=[];
  if(em){cats.push('Subject history');vals.push(em.avg_pct);}
  if(pr){cats.push('Peer read-across');vals.push(pr.avg_pct);}
  if(bt&&bt.pred!=null){cats.push('Peer (earlier)');vals.push(bt.pred);}
  if(actual!=null){cats.push('ACTUAL');vals.push(actual);}
  const el1=document.getElementById('sdPred');
  if(el1&&cats.length){ if(window._sd1)window._sd1.destroy();
    window._sd1=new ApexCharts(el1,{
      chart:{type:'bar',height:230,background:'transparent',toolbar:{show:false},fontFamily:'inherit'},
      theme:{mode:'dark'}, series:[{name:'move %',data:vals}],
      colors:cats.map((c,i)=>c==='ACTUAL'?'#58a6ff':(vals[i]>=0?'#3fb950':'#f85149')),
      plotOptions:{bar:{distributed:true,borderRadius:3,columnWidth:'55%'}},
      dataLabels:{enabled:true,formatter:v=>(v>0?'+':'')+v+'%',style:{colors:['#e6edf3']}},
      xaxis:{categories:cats,labels:{style:{colors:'#8b98ad'}}},
      yaxis:{labels:{style:{colors:'#8b98ad'}}},
      grid:{borderColor:'#1f2937'}, legend:{show:false},
      tooltip:{theme:'dark',y:{formatter:v=>(v>0?'+':'')+v+'%'}}
    }); window._sd1.render();
  }
  const mem=(pr&&pr.members)||[];
  const el2=document.getElementById('sdPeers');
  if(el2){ if(mem.length){ if(window._sd2)window._sd2.destroy();
      window._sd2=new ApexCharts(el2,{
        chart:{type:'bar',height:Math.max(150,mem.length*30+50),background:'transparent',toolbar:{show:false},fontFamily:'inherit'},
        theme:{mode:'dark'}, series:[{name:'peer move %',data:mem.map(m=>m.ret)}],
        colors:mem.map(m=>m.ret>=0?'#3fb950':'#f85149'),
        plotOptions:{bar:{distributed:true,horizontal:true,borderRadius:3,barHeight:'60%'}},
        dataLabels:{enabled:true,formatter:v=>(v>0?'+':'')+v+'%',style:{colors:['#e6edf3']}},
        xaxis:{categories:mem.map(m=>m.peer),labels:{style:{colors:'#8b98ad'}}},
        yaxis:{labels:{style:{colors:'#8b98ad'}}},
        grid:{borderColor:'#1f2937'}, legend:{show:false},
        tooltip:{theme:'dark',y:{formatter:v=>(v>0?'+':'')+v+'%'}}
      }); window._sd2.render();
    } else { el2.innerHTML='<div class="mini">no peer realised moves for this sector</div>'; }
  }
}
async function loadPeerBt(){
  lastView=loadPeerBt;
  const o=document.getElementById('out'); o.innerHTML='<div class="mut">back-testing peer read-across (predicted vs actual)…</div>';
  const d=await (await fetch('/api/peerbacktest')).json();
  const c=v=>v==null?'var(--mut)':(v>=0?'var(--pos)':'var(--neg)');
  const rows=d.rows.map(r=>`<tr class="clk" onclick="openStock('${r.symbol}')" title="click for full detail">
     <td style="text-align:left;font-weight:600;color:var(--blue)">${r.symbol} ›</td>
     <td style="text-align:left" class=mut>${r.sector||'—'}</td>
     <td class=mut>${r.result}</td>
     <td style="text-align:right;color:${c(r.pred)}">${r.pred==null?'—':(r.pred>0?'+':'')+r.pred+'%'}${r.basis==='market'?' <span class="mut" style="font-size:.68rem">mkt</span>':''}</td>
     <td style="text-align:right;color:${c(r.actual)};font-weight:700">${r.actual>0?'+':''}${r.actual}%</td>
     <td style="text-align:right" class=mut>${r.error==null?'—':(r.error>0?'+':'')+r.error}</td>
     <td style="text-align:left;white-space:normal;font-size:.74rem">${(r.sector_peers&&r.sector_peers.length)?r.sector_peers.map(p=>(r.used_peers||[]).includes(p)?`<span style="color:var(--pos);font-weight:600">${p}</span>`:`<span class=mut>${p}</span>`).join(', '):'<span class=mut>— (only one in sector)</span>'}</td></tr>`).join('');
  const cc=d.correlation==null?'var(--mut)':(Math.abs(d.correlation)<0.2?'var(--mut)':(d.correlation>0?'var(--pos)':'var(--neg)'));
  const strg=c=>c==null?'—':(Math.abs(c)<0.2?'no relationship':(Math.abs(c)<0.4?'weak':'moderate'));
  o.innerHTML=`
    <div class="warn">🎯 <b>Peer read-across — predicted vs actual</b> (stocks that already reported). PREDICTED = mean realised move of sector peers that reported <b>earlier</b>; if none, market fallback (tagged <span class=mut>mkt</span>). ${d.note}</div>
    <div class="card">
      <div class="big" style="color:${cc}">corr = ${d.correlation==null?'—':(d.correlation>0?'+':'')+d.correlation} <span style="font-size:.9rem;color:var(--mut)">(sector-only, n=${d.n_sector}) · ${strg(d.correlation)}</span></div>
      <div class="hl">incl. market fallback: corr <b>${d.correlation_all==null?'—':(d.correlation_all>0?'+':'')+d.correlation_all}</b> (n=${d.n}) · direction ${d.direction_hit}% (${d.direction_hit>=60?'some edge':'≈ coin flip'}) · avg error ${d.mae}%</div>
      <div class="chartbox"><div class="chartttl">Predicted vs Actual move</div><div class="chartsub">each dot = a stock · dashed line = perfect prediction · closer to the line = better · <b>click a dot for full detail</b></div><div id="pbScatter"></div></div>
      <div class="chartbox"><div class="chartttl">Average actual move by sector</div><div class="chartsub">how each sector moved around its Q1 results</div><div id="pbSector"></div></div>
      <details open><summary>▸ per-stock: predicted vs actual</summary>
        <table style="width:100%;border-collapse:collapse;font-size:.8rem;margin-top:6px">
        <tr class=mut><th style="text-align:left">stock</th><th style="text-align:left">sector</th><th style="text-align:left">result</th><th>pred</th><th>actual</th><th>err</th><th style="text-align:left">peer stocks (green = reported earlier)</th></tr>
        ${rows}</table></details>
    </div>`;
  renderPBCharts(d);
}
function renderPBCharts(d){
  if(typeof ApexCharts==='undefined'){
    document.getElementById('pbScatter').innerHTML='<div class="mini">charts need internet (ApexCharts CDN) — table still works</div>';return;}
  const pts=d.rows.filter(r=>r.pred!=null).map(r=>({x:r.pred,y:r.actual,sym:r.symbol,basis:r.basis}));
  const xs=pts.map(p=>p.x).concat(pts.map(p=>p.y)); const lo=Math.floor(Math.min(...xs)), hi=Math.ceil(Math.max(...xs));
  if(window._sc)window._sc.destroy();
  window._sc=new ApexCharts(document.getElementById('pbScatter'),{
    chart:{height:360,background:'transparent',toolbar:{show:false},fontFamily:'inherit',
      events:{dataPointSelection:function(e,ctx,cfg){if(cfg.seriesIndex===0){openStock(pts[cfg.dataPointIndex].sym);}},
              markerClick:function(e,ctx,cfg){if(cfg.seriesIndex===0){openStock(pts[cfg.dataPointIndex].sym);}}}},
    theme:{mode:'dark'},
    series:[{name:'stocks',type:'scatter',data:pts.map(p=>({x:p.x,y:p.y}))},
            {name:'perfect (pred = actual)',type:'line',data:[{x:lo,y:lo},{x:hi,y:hi}]}],
    colors:['#58a6ff','#8b98ad'], stroke:{width:[0,2],dashArray:[0,6]},
    markers:{size:[6,0],strokeWidth:0},
    xaxis:{title:{text:'Predicted move %  (peer read-across)',style:{color:'#8b98ad'}},min:lo,max:hi,tickAmount:8,decimalsInFloat:0,labels:{style:{colors:'#8b98ad'}}},
    yaxis:{title:{text:'Actual move %',style:{color:'#8b98ad'}},decimalsInFloat:0,labels:{style:{colors:'#8b98ad'}}},
    grid:{borderColor:'#1f2937'},
    legend:{show:true,labels:{colors:'#8b98ad'}},
    tooltip:{theme:'dark',custom:function({seriesIndex,dataPointIndex}){
      if(seriesIndex!==0)return''; const p=pts[dataPointIndex];
      return '<div style="padding:6px 10px"><b>'+p.sym+'</b> ('+(p.basis||'—')+')<br>pred '+(p.x>0?'+':'')+p.x+'% · actual '+(p.y>0?'+':'')+p.y+'%</div>';}}
  }); window._sc.render();
  // sector averages
  const bs={}; d.rows.forEach(r=>{if(r.sector){(bs[r.sector]=bs[r.sector]||[]).push(r.actual);}});
  const secs=Object.keys(bs).sort((a,b)=>(bs[b].reduce((x,y)=>x+y,0)/bs[b].length)-(bs[a].reduce((x,y)=>x+y,0)/bs[a].length));
  const vals=secs.map(s=>+(bs[s].reduce((x,y)=>x+y,0)/bs[s].length).toFixed(2));
  if(window._sb)window._sb.destroy();
  window._sb=new ApexCharts(document.getElementById('pbSector'),{
    chart:{type:'bar',height:Math.max(220,secs.length*28+60),background:'transparent',toolbar:{show:false},fontFamily:'inherit'},
    theme:{mode:'dark'},
    series:[{name:'avg move %',data:vals}],
    colors:vals.map(v=>v>=0?'#3fb950':'#f85149'),
    plotOptions:{bar:{distributed:true,horizontal:true,borderRadius:3,barHeight:'62%'}},
    dataLabels:{enabled:true,formatter:v=>(v>0?'+':'')+v+'%',style:{colors:['#e6edf3'],fontSize:'11px'}},
    xaxis:{categories:secs,labels:{style:{colors:'#8b98ad'}}},
    yaxis:{labels:{style:{colors:'#8b98ad'}}},
    grid:{borderColor:'#1f2937'}, legend:{show:false},
    tooltip:{theme:'dark',y:{formatter:v=>(v>0?'+':'')+v+'%'}}
  }); window._sb.render();
}
async function loadPeers(){
  lastView=loadPeers;
  const o=document.getElementById('out'); o.innerHTML='<div class="mut">building peer read-across…</div>';
  const d=await (await fetch('/api/peers')).json();
  const dc={up:'var(--pos)',down:'var(--neg)',neutral:'var(--mut)'};
  const sgn=v=>v>0?'+':'';
  const pct=(v,pts)=>`<span style="color:${v>=0?'var(--pos)':'var(--neg)'};font-weight:700">${sgn(v)}${v}%</span>${pts!=null?` <span class=mut>≈ ${sgn(pts)}${pts} pts</span>`:''}`;
  const cards=d.map(e=>{
    if(e.error) return `<div class="card"><b>${e.symbol}</b> — ${e.error}</div>`;
    const em=e.expected_move, pem=e.peer_expected_move;
    const emhtml = em?`<div class="emv"><div class="hl"><b>📈 Prediction · subject history:</b> ${pct(em.avg_pct,em.avg_points)} <span class=mut>· median ${em.median_pct}% · win ${em.win_rate}% (n=${em.n})</span></div></div>`:'';
    const pemhtml = pem?`<div class="emv" style="border-left-color:var(--blue)"><div class="hl"><b>🤝 Prediction · peer read-across:</b> ${pct(pem.avg_pct,pem.avg_points)} <span class=mut>· avg of ${pem.n} sector peers reported this quarter</span></div><div class="hl mut" style="font-size:.75rem">${pem.members.map(m=>m.peer+' '+(m.ret>0?'+':'')+m.ret+'%').join(' · ')}</div></div>`:'<div class="hl mut">no peer realised moves yet for read-across</div>';
    const rows=(e.peers||[]).map(p=>{
      const s=p.signal;
      if(!s) return `<tr><td style="text-align:left">${p.peer}</td><td class=mut>—</td><td></td><td class=mut style="font-size:.75rem">no pre-event results signal</td></tr>`;
      const hl=s.link?`<a href="${s.link}" target="_blank" rel="noopener" style="color:inherit;border-bottom:1px dotted var(--line)">${s.headline} ↗</a>`:s.headline;
      return `<tr><td style="text-align:left;font-weight:600">${p.peer}</td>
        <td style="color:${dc[s.direction]};font-weight:700">${s.direction}</td>
        <td class=mut>${s.linked?'↔ linked':''}</td>
        <td style="text-align:left" class=mut style="font-size:.78rem">${hl}</td></tr>`;
    }).join('');
    const netcol=e.net==='up'?'var(--pos)':(e.net==='down'?'var(--neg)':'var(--mut)');
    return `<div class="card">
      <div class="ctop"><span class="sym">${e.symbol}</span><span class="date">${e.sector||'—'} · decision ${e.decision_ts.slice(0,10)}</span></div>
      ${emhtml}
      ${pemhtml}
      <div class="hl" style="margin:8px 0 2px"><b>${e.reported}</b> of ${e.peers.length} peers reported pre-event · <span style="color:var(--pos)">${e.up} up</span> / <span style="color:var(--neg)">${e.down} down</span> · net <b style="color:${netcol}">${e.net}</b></div>
      <table style="width:100%;border-collapse:collapse;font-size:.82rem;margin-top:4px">
        <tr class=mut><th style="text-align:left">peer</th><th style="text-align:left">direction</th><th style="text-align:left">link</th><th style="text-align:left">headline</th></tr>
        ${rows}</table></div>`;
  }).join('');
  o.innerHTML='<div class="warn">🤝 <b>Peer read-across</b> — sector peers that reported <b>before</b> each event. A peer beat/miss resets the bar even when the subject has no coverage of its own. Direction from headline tone (proxy).</div>'
    +'<div class="cards">'+cards+'</div>';
}
async function loadEffect(){
  lastView=loadEffect;
  const o=document.getElementById('out'); o.innerHTML='<div class="mut">testing past-news features vs the actual move…</div>';
  const d=await (await fetch('/api/effect')).json();
  const col=c=>c==null?'var(--mut)':(Math.abs(c)<0.2?'var(--mut)':(c>0?'var(--pos)':'var(--neg)'));
  const rows=d.features.map(f=>`<tr>
     <td style="text-align:left">${f.feature}</td>
     <td style="text-align:right;color:${col(f.corr)};font-weight:700">${f.corr==null?'—':(f.corr>0?'+':'')+f.corr.toFixed(3)}</td>
     <td style="text-align:left" class=mut>${f.strength}</td>
     <td style="text-align:left" class=mut style="font-size:.74rem">${f.desc}</td></tr>`).join('');
  const g=d.governance, b=d.buzz_split, mv=v=>v==null?'—':(v>0?'+':'')+v+'%';
  o.innerHTML=`
    <div class="warn">🔬 <b>Past-news features vs the move</b> — do any news signals relate to the realised move? ${d.note}</div>
    <div class="card">
      <div class="hl">${d.n} realised events · avg move ${mv(d.avg_move)}</div>
      <table style="width:100%;border-collapse:collapse;font-size:.82rem;margin-top:8px">
        <tr class=mut><th style="text-align:left">feature</th><th style="text-align:right">corr</th><th style="text-align:left">&nbsp;strength</th><th style="text-align:left">what it is</th></tr>
        ${rows}
      </table>
      <div class="hl" style="margin-top:12px"><b>Governance flag:</b> present n=${g.present.n} avg ${mv(g.present.avg)} · absent n=${g.absent.n} avg ${mv(g.absent.avg)} <span class=mut>(no real difference)</span></div>
      <div class="hl"><b>Attention:</b> |move| when coverage LOW (n=${b.low.n}) <b>${b.low.abs_move}%</b> vs HIGH (n=${b.high.n}) <b>${b.high.abs_move}%</b> <span class=mut>(thinly-covered names swung more — weak, one quarter)</span></div>
      <div class="hl mut" style="margin-top:10px">Verdict: no reliable effect. The one flicker (pos_frac) is weak and likely momentum/anticipation, not edge.</div>
    </div>`;
}
async function loadBacktest(){
  lastView=loadBacktest;
  const o=document.getElementById('out'); o.innerHTML='<div class="mut">running proxy backtest (bar vs realised move)…</div>';
  const d=await (await fetch('/api/backtest')).json();
  const b=d.buckets, col=v=>v==null?'var(--mut)':(v>=0?'var(--pos)':'var(--neg)');
  const bk=(k,lab)=>`<div class="tag">${lab}: n=${b[k].n} · avg <span style="color:${col(b[k].avg)}">${b[k].avg==null?'—':(b[k].avg>0?'+':'')+b[k].avg+'%'}</span></div>`;
  const rows=d.rows.map(r=>`<tr>
     <td style="text-align:left">${r.symbol}</td><td class=mut>${r.result}</td>
     <td style="text-align:right;color:${r.bar==null?'var(--amber)':'var(--ink)'}">${r.bar==null?'null':(r.bar>0?'+':'')+r.bar.toFixed(2)}</td>
     <td style="text-align:right" class=mut>${r.arts}</td>
     <td style="text-align:right;color:${col(r.realised)}">${r.realised>0?'+':''}${r.realised.toFixed(2)}%</td></tr>`).join('');
  const verdict = (d.correlation==null)?'—':
     (Math.abs(d.correlation)<0.2?`<b class=mut>no relationship</b> (|corr| &lt; 0.2)`:
      (d.correlation<0?`<b>negative</b> — weakly supports "sell the news"`:`<b>positive</b>`));
  o.innerHTML=`
    <div class="warn">📊 <b>Proxy backtest</b> — does the bar relate to the actual move? ${d.note}</div>
    <div class="card">
      <div class="big" style="color:${col(d.correlation)}">corr = ${d.correlation==null?'—':(d.correlation>0?'+':'')+d.correlation}</div>
      <div class="hl">over n=${d.n} realised events · ${verdict}</div>
      <div style="margin:10px 0">${bk('high','HIGH ≥+0.34')}${bk('neutral','NEUTRAL')}${bk('low','LOW ≤−0.34')}${bk('null','NO COVERAGE')}</div>
      <details open><summary>▸ per-event: bar vs realised move</summary>
        <table style="width:100%;border-collapse:collapse;font-size:.8rem;margin-top:6px">
        <tr class=mut style="text-align:right"><th style="text-align:left">stock</th><th style="text-align:left">result</th><th>bar</th><th>arts</th><th>move</th></tr>
        ${rows}</table></details>
    </div>`;
}
async function runOne(){
  const s=document.getElementById('sym').value.trim().toUpperCase();
  const t=document.getElementById('ts').value.trim();
  if(!s){return;}
  const o=document.getElementById('out'); o.innerHTML='<div class="mut">scoring '+s+'…</div>';
  const d=await (await fetch('/api/bar?sym='+encodeURIComponent(s)+(t?'&ts='+encodeURIComponent(t):''))).json();
  o.innerHTML=backBtn()+'<div class="cards">'+card(d)+'</div>';
}
function openTab(){const h=location.hash.slice(1);
  if(h.startsWith('stock='))openStock(decodeURIComponent(h.slice(6)).toUpperCase());
  else if(h==='peerbt')loadPeerBt(); else if(h==='peers')loadPeers();
  else if(h==='backtest')loadBacktest(); else if(h==='features')loadEffect();
  else if(h==='forward')loadForward();
  else loadUpcoming();}
window.addEventListener('hashchange',openTab);
openTab();
</script></body></html>"""


def api(name: str, q: dict):
    """Route an api name (path after /api/) to its JSON-able result. Reusable so
    the main dashboard can mount these under a prefix. Returns None if unknown."""
    if name == "upcoming":
        out = []
        for sym, ts in EB.upcoming_events():
            try:
                r = EB.score_batch(sym, ts); r["expected_move"] = expected_move(sym); out.append(r)
            except Exception as e:                          # noqa: BLE001
                out.append({"symbol": sym, "decision_ts": ts.isoformat(), "error": str(e)})
        return out
    if name == "forward":
        try:
            FT.capture(); FT.settle(); return {"rows": FT.rows()}
        except Exception as e:                              # noqa: BLE001
            return {"error": str(e), "rows": FT.rows()}
    if name == "stock_detail":
        sym = (q.get("sym", [""])[0]).upper()
        try:
            return stock_detail(sym)
        except Exception as e:                              # noqa: BLE001
            return {"symbol": sym, "error": str(e)}
    if name == "peerbacktest":
        try:
            return PB.compute()
        except Exception as e:                              # noqa: BLE001
            return {"error": str(e)}
    if name == "peers":
        out = []
        for sym, ts in EB.upcoming_events():
            try:
                r = EB.peer_readacross(sym, ts)
                r["expected_move"] = expected_move(sym)
                r["peer_expected_move"] = peer_expected_move(sym)
                out.append(r)
            except Exception as e:                          # noqa: BLE001
                out.append({"symbol": sym, "error": str(e)})
        return out
    if name == "effect":
        try:
            return FX.compute(fetch_missing=False)
        except Exception as e:                              # noqa: BLE001
            return {"error": str(e)}
    if name == "backtest":
        try:
            return BT.compute(fetch_missing=False)
        except Exception as e:                              # noqa: BLE001
            return {"error": str(e)}
    if name in ("ingest", "bar"):
        sym = (q.get("sym", [""])[0]).upper()
        ts = q.get("ts", [""])[0]
        when = dt.datetime.fromisoformat(ts) if ts else dt.datetime.now() + dt.timedelta(days=1)
        try:
            if name == "ingest":
                import news_ingest
                news_ingest.ingest([sym]); _BEH.pop(sym, None)
            r = EB.score_batch(sym, when); r["expected_move"] = expected_move(sym)
            return r
        except Exception as e:                              # noqa: BLE001
            return {"symbol": sym, "error": str(e)}
    return None


class H(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def _send(self, code, body, ctype="application/json"):
        b = body.encode("utf-8") if isinstance(body, str) else body
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(b)))
        self.end_headers()
        self.wfile.write(b)

    def do_GET(self):
        u = urlparse(self.path)
        q = parse_qs(u.query)
        if u.path in ("/", "/index.html"):
            return self._send(200, PAGE, "text/html; charset=utf-8")
        if u.path.startswith("/api/"):
            obj = api(u.path[len("/api/"):], q)
            if obj is not None:
                return self._send(200, json.dumps(obj, ensure_ascii=False))
        self._send(404, json.dumps({"error": "not found"}))


def main():
    port = 8097
    for a in sys.argv[1:]:
        if a.isdigit():
            port = int(a)
    srv = ThreadingHTTPServer(("127.0.0.1", port), H)
    print("=" * 60)
    print(f"  EXPECTATION-BAR host  →  http://localhost:{port}")
    print("  (separate from the main Nifty-50 dashboard)")
    print("  Ctrl+C to stop.")
    print("=" * 60)
    if "--no-open" not in sys.argv:
        threading.Timer(0.6, lambda: webbrowser.open(f"http://localhost:{port}")).start()
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        srv.shutdown()


if __name__ == "__main__":
    main()
