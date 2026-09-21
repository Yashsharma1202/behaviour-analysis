import openpyxl, pandas as pd, sys

sys.stdout.reconfigure(errors='replace')

path = r'D:/behaviour analysis/Nifty50_12_Quarters_Consolidated_v2.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

print("==========================================")
print("VERIFYING DRAWDOWN IN CONSOLIDATED WORKBOOK")
print("==========================================")
print("Sheet Names:", wb.sheetnames)

charts_count = 0
for sheet in wb.sheetnames:
    ws = wb[sheet]
    num_charts = len(ws._charts)
    if num_charts > 0:
        charts_count += num_charts
        print(f"  Sheet '{sheet}': {num_charts} embedded chart(s)")

print(f"\nTotal Embedded Charts: {charts_count}")

ws_q3 = wb['Q3 2023-24']
print("\nQ3 2023-24 Summary Cards Row 4-5:")
print(f"  Col 10 (Max DD ₹): {ws_q3.cell(row=4, column=10).value} -> {ws_q3.cell(row=5, column=10).value}")
print(f"  Col 11 (Max DD %): {ws_q3.cell(row=4, column=11).value} -> {ws_q3.cell(row=5, column=11).value}")

print("\nQ3 2023-24 Trade Row 9 (Head):")
print(f"  Col 16 (Cum Profit): {ws_q3.cell(row=8, column=16).value} -> {ws_q3.cell(row=9, column=16).value}")
print(f"  Col 17 (Drawdown)  : {ws_q3.cell(row=8, column=17).value} -> {ws_q3.cell(row=9, column=17).value}")
print("==========================================")
