"""
DUSSEHRA 2025 - TOP 50 STOCKS FROM THE POINT-IN-TIME DETAILED REPORT.
Reads Dussehra_2025_PIT_Detailed_Report.xlsx (built by build_dussehra_2025_pit_detailed_report.py).
Each stock's 2025 window+direction is point-in-time (prior years only). Top 50 is ranked by
2025 realised P&L, the same ranking as the earlier 2025 top-50 report (hindsight selection).
ADDITIVE: writes a new workbook only.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

BASE = r'D:\behaviour analysis'
SRC = BASE + r'\Dussehra_2025_PIT_Detailed_Report.xlsx'
OUT = BASE + r'\Dussehra_2025_PIT_Top50_LotFixed.xlsx'

trades = pd.read_excel(SRC, sheet_name='PIT_2025_Trades', header=2)
trades = trades[pd.to_numeric(trades['PnL Rank'], errors='coerce').notna()].copy()
top = trades.sort_values('P&L Rs', ascending=False).head(50).reset_index(drop=True)
top['Top50 Rank'] = range(1, len(top) + 1)
top_syms = list(top['Symbol'])

grid = pd.read_excel(SRC, sheet_name='Daily_Cum_PnL', header=2)
grid['Date'] = pd.to_datetime(grid['Date']).dt.date
keep = ['Date'] + [s for s in top_syms if s in grid.columns]
g50 = grid[keep].copy()
g50['TOP50 TOTAL'] = g50[[c for c in keep if c != 'Date']].sum(axis=1).round(0)
port = g50['TOP50 TOTAL']
max_dd = round(min(0, (port - port.cummax()).min()), 0)

cols = ['Top50 Rank', 'Symbol', 'Name', 'Sector', 'Direction', 'Window', 'Entry Date', 'Entry Px (fut close)',
        'Exit Date', 'Exit Px (fut close)', 'Direction-Adj Return %', 'Pickle OOS Ret % (cross-check)',
        'Lot Size', 'Notional per Lot', 'Indicative Margin (20%)', 'P&L Rs', 'Max DD Rs', 'Daily Sharpe',
        'Prior-yr OOS n', 'Low Confidence']
top = top[cols]

summary = [
    ('Stocks in top 50', len(top)),
    ('Ranking basis', '2025 realised P&L, 1 lot each (hindsight, same as earlier 2025 top-50 report)'),
    ('Long / Short split', f"{(top.Direction == 'LONG').sum()} / {(top.Direction == 'SHORT').sum()}"),
    ('Winning trades (P&L > 0)', f"{(top['P&L Rs'] > 0).sum()} of {len(top)}"),
    ('Total P&L, 1 lot each (Rs)', round(top['P&L Rs'].sum(), 0)),
    ('Total indicative margin, 1 lot each (Rs)', round(top['Indicative Margin (20%)'].sum(), 0)),
    ('Combined max drawdown, top 50 (Rs)', max_dd),
    ('Low-confidence stocks in top 50', int((top['Low Confidence'] == 'YES').sum())),
]

wb = Workbook()
ws = wb.active; ws.title = 'Overview'
ws.append(['DUSSEHRA 2025 - TOP 50 STOCKS (POINT-IN-TIME WINDOWS, RANKED BY 2025 P&L)'])
ws['A1'].font = Font(bold=True, size=13)
ws.append([])
for k, v in summary:
    ws.append([k, v])
ws.append([])
ws.append(['Windows and directions are point-in-time (chosen from prior years only). Ranking uses realised 2025 P&L, so the top-50 total is optimistic. Lot sizes and margin are the same caveats as the detailed report.'])
ws.column_dimensions['A'].width = 48; ws.column_dimensions['B'].width = 70


def write_table(wsx, frame, title):
    wsx.append([title]); wsx['A1'].font = Font(bold=True, size=12); wsx.append([])
    wsx.append(list(frame.columns))
    hr = wsx.max_row
    for c in wsx[hr]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78')
        c.alignment = Alignment(horizontal='center', wrap_text=True)
    for r in frame.itertuples(index=False):
        wsx.append([None if pd.isna(v) else v for v in r])
    wsx.freeze_panes = wsx.cell(row=hr + 1, column=3)
    for i, col in enumerate(frame.columns, 1):
        wsx.column_dimensions[wsx.cell(row=hr, column=i).column_letter].width = max(12, min(28, len(str(col)) + 2))


write_table(wb.create_sheet('Top50_Trades'), top, 'TOP 50 - 2025 POINT-IN-TIME TRADES')
write_table(wb.create_sheet('Top50_Daily_Cum_PnL'), g50, 'TOP 50 - DAILY CUMULATIVE P&L (Rs, 1 lot each)')
wb.save(OUT)

print(top[['Top50 Rank', 'Symbol', 'Direction', 'Window', 'P&L Rs']].to_string(index=False))
for k, v in summary:
    print(f'{k}: {v}')
