import openpyxl, sys

sys.stdout.reconfigure(errors='replace')

path = r'D:/behaviour analysis/Nifty50_12_Quarters_Futures_OI_Master_v11.xlsx'

wb = openpyxl.load_workbook(path, data_only=True)
print("==========================================")
print(f"VERIFYING V11 MASTER WORKBOOK WITH 12 CHARTS")
print("==========================================")
print("Total Sheets:", len(wb.sheetnames))
print("Sheet Names :", wb.sheetnames)

charts_count = 0
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    num_charts = len(ws._charts)
    if num_charts > 0:
        charts_count += num_charts
        print(f"  Sheet '{sheet_name}': {num_charts} embedded chart(s) -> Title: '{ws._charts[0].title}'")

print(f"\nTotal Embedded Charts across Workbook: {charts_count}")
