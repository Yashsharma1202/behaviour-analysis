import pandas as pd
import pathlib
import sys
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

sys.stdout.reconfigure(encoding='utf-8')

# ---------- Configuration ----------
BASE_DIR = pathlib.Path('D:/behaviour analysis/OI_DATA')  # OI data root
WORKSPACE_ROOT = pathlib.Path('d:/behaviour analysis')   # where quarter Excel files live
OUTPUT_PATH = pathlib.Path('D:/behaviour analysis/Future_Master_With_Trades.xlsx')

# ---------- Helpers ----------
def latest_parquet(folder: pathlib.Path):
    files = sorted(folder.glob('*.parquet'))
    return files[-1] if files else None

def get_close_price(df: pd.DataFrame):
    """Return the close price at market close (15:29:59) or the last row if missing."""
    row = df[df['Time'] == '15:29:59']
    if row.empty:
        row = df.sort_values('Time').tail(1)
    return float(row.iloc[0]['Close'])

def load_futures_price(stock: str, date_str: str):
    """Return futures close price for a given stock on a given date.
    Uses futures_parquet if present, otherwise falls back to spot parquet.
    Returns (price, expiry_date) where expiry_date may be empty string.
    """
    stock_dir = BASE_DIR / stock
    # First try futures_parquet
    fut_dir = stock_dir / 'futures_parquet'
    date_file = None
    if fut_dir.exists():
        candidate = fut_dir / f"{date_str}.parquet"
        if candidate.exists():
            date_file = candidate
    # If not found, fall back to spot_parquet (proxy for futures price)
    if not date_file:
        spot_dir = stock_dir / 'spot_parquet'
        candidate = spot_dir / f"{date_str}.parquet"
        if candidate.exists():
            date_file = candidate
    if not date_file:
        return None, ''
    df = pd.read_parquet(date_file)
    price = get_close_price(df)
    # Try to pull an expiry date from futures data if available
    expiry = ''
    if fut_dir.exists() and date_file.parent.name == 'futures_parquet':
        if 'ExpiryDate' in df.columns:
            expiry = sorted(df['ExpiryDate'].unique())[0] if len(df['ExpiryDate'].unique()) else ''
    return price, expiry

# ---------- Load quarterly trade data ----------
trade_frames = []
quarter_files = sorted(WORKSPACE_ROOT.glob('*_Combined_Best_Capital_Utilisation.xlsx'))
if not quarter_files:
    print('⚠️ No quarter files found in workspace.')
else:
    for qf in quarter_files:
        try:
            df_q = pd.read_excel(qf, sheet_name='Detailed_Trades')
        except Exception as e:
            print(f'⚠️ Could not read {qf.name}: {e}')
            continue
        # Ensure required columns exist; rename if needed
        required = ['Symbol', 'Strategy', 'Entry Date', 'Exit Date',
                    'Expected Return (%)', 'Return We Get (%)', 'Quarter']
        missing = [c for c in required if c not in df_q.columns]
        if missing:
            print(f'⚠️ Missing columns in {qf.name}: {missing}')
            continue
        # Add a column with the source file name (optional)
        df_q['SourceFile'] = qf.name
        trade_frames.append(df_q)

if not trade_frames:
    print('❌ No trade data loaded – aborting.')
    sys.exit(1)

trades_df = pd.concat(trade_frames, ignore_index=True)

# ---------- Build master rows ----------
rows = []
for idx, row in trades_df.iterrows():
    stock = str(row['Symbol']).strip()
    strategy = str(row['Strategy']).strip()  # LONG or SHORT
    entry_dt = pd.to_datetime(row['Entry Date']).strftime('%Y-%m-%d')
    exit_dt = pd.to_datetime(row['Exit Date']).strftime('%Y-%m-%d')
    quarter = str(row['Quarter'])
    exp_ret = row.get('Expected Return (%)', None)
    ret_we_get = row.get('Return We Get (%)', None)
    # Load entry & exit futures/spot prices
    entry_price, entry_expiry = load_futures_price(stock, entry_dt)
    exit_price, exit_expiry = load_futures_price(stock, exit_dt)
    if entry_price is None or exit_price is None:
        # Skip if we cannot obtain price data
        continue
    # Compute booked P&L (simple price diff * direction)
    direction = 1 if 'LONG' in strategy.upper() else -1
    booked_pnl = direction * (exit_price - entry_price)
    # Quarterly result date – placeholder: use exit date (can be adjusted later)
    quarterly_result_date = exit_dt
    # Determine which expiry to display – use entry expiry if available else exit expiry
    expiry_date = entry_expiry or exit_expiry

    rows.append([
        stock,
        strategy,
        entry_dt,
        exit_dt,
        entry_price,
        exit_price,
        expiry_date,
        exp_ret,
        ret_we_get,
        booked_pnl,
        quarterly_result_date,
        quarter
    ])

# ---------- Write to Excel ----------
wb = Workbook()
ws = wb.active
ws.title = 'Futures_Trades_Master'

headers = [
    'Stock',
    'Position',            # LONG / SHORT
    'Entry Date',
    'Exit Date',
    'Entry Futures Price (₹)',
    'Exit Futures Price (₹)',
    'Futures Expiry Date',
    'Expected Return (%)',
    'Return We Get (%)',
    'Booked P&L (₹)',
    'Quarterly Result Date',
    'Quarter'
]
ws.append(headers)
# Header styling
header_font = Font(bold=True, color='FFFFFFFF')
header_fill = PatternFill(start_color='FF006699', end_color='FF006699', fill_type='solid')
for col in range(1, len(headers)+1):
    cell = ws.cell(row=1, column=col)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center')

for r in rows:
    ws.append(r)

# Auto‑size columns
for col in ws.columns:
    max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
    ws.column_dimensions[col[0].column_letter].width = max_len + 3

wb.save(OUTPUT_PATH)
print(f'✅ Futures‑with‑trades master Excel created at: {OUTPUT_PATH}')
print(f'Rows written (valid trades): {len(rows)}')
