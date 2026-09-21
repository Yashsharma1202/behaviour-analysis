import openpyxl, pathlib, sys

sys.stdout.reconfigure(errors='replace')

b_dir = pathlib.Path('D:/behaviour analysis')

master_path = b_dir / 'Nifty50_12_Quarters_Point_In_Time_Futures_0.5PCT_Cost_Master.xlsx'
summary_path = b_dir / 'Nifty50_12_Quarters_Point_In_Time_Futures_0.5PCT_Cost_Summary.xlsx'

print("==========================================================================================")
print("FINAL WORKBOOK VERIFICATION")
print("==========================================================================================")

if master_path.exists():
    wb_m = openpyxl.load_workbook(master_path, read_only=True)
    print(f"✅ Master Workbook Exists: {master_path.name}")
    print(f"   Size: {master_path.stat().st_size / 1024:.1f} KB")
    print(f"   Sheets Count: {len(wb_m.sheetnames)}")
    print(f"   Sheets: {wb_m.sheetnames}")

if summary_path.exists():
    wb_s = openpyxl.load_workbook(summary_path, read_only=True)
    print(f"\n✅ Summary Workbook Exists: {summary_path.name}")
    print(f"   Size: {summary_path.stat().st_size / 1024:.1f} KB")
    print(f"   Sheets Count: {len(wb_s.sheetnames)}")

print("==========================================================================================")
