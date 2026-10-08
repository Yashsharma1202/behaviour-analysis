"""
ALL 211 F&O STOCKS \u2014 DUSSEHRA, WALK-FORWARD RANKED (supersedes the earlier
in-sample FO211 Dussehra reports built today). Ranks stocks by genuine
out-of-sample (OOS) average return from the walk-forward 8x8 grid search
(fix_dussehra_position_window_walkforward.py) \u2014 point-in-time validated,
no lookahead, same method as the Nifty index options model.

ADDITIVE: reads the walk-forward pickle; writes new report files.
"""
import json, pickle, os, subprocess
from datetime import date, timedelta, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

BASE = r'D:\behaviour analysis'
OUT_XLSX = BASE + r'\FO211_Dussehra_WalkForward_Ranked_Report.xlsx'
OUT_HTML = BASE + r'\FO211_Dussehra_WalkForward_Ranked_Report.html'
OUT_PDF = BASE + r'\FO211_Dussehra_WalkForward_Ranked_Report.pdf'

results = pickle.load(open(BASE + r'\scraped_parquet\_tmp_dussehra_walkforward.pkl', 'rb'))
fo = json.load(open(BASE + r'\dashboard_data\fo_stocks_211.json', encoding='utf-8'))
meta = {f['symbol']: f for f in fo}

qualified = {s: r for s, r in results.items() if r['oos_n'] >= 1}
insufficient = {s: r for s, r in results.items() if r['oos_n'] == 0}
ranked = sorted(qualified.items(), key=lambda kv: -kv[1]['oos_avg'])
n_low_conf = sum(1 for _, r in ranked if r['low_confidence'])

NSE_2026_HOLIDAYS = {
    # '2026-03-04' removed (not a real NSE holiday -- leftover from Holi's date
    # once being wrongly typed as 04-Mar; real Holi holiday is 03-Mar, kept
    # below). '2026-11-08'/'2026-11-10' (Diwali) added -- were missing entirely.
    '2026-01-26', '2026-03-03', '2026-03-26', '2026-03-31', '2026-04-03',
    '2026-04-14', '2026-05-01', '2026-05-28', '2026-06-26', '2026-09-14', '2026-10-02',
    '2026-10-20', '2026-11-08', '2026-11-10', '2026-11-24', '2026-12-25',
}
def is_trading_day(d):
    return d.weekday() < 5 and d.isoformat() not in NSE_2026_HOLIDAYS
def shift_trading_days(anchor, n):
    d = anchor
    step = 1 if n >= 0 else -1
    cnt = abs(n)
    while cnt > 0:
        d += timedelta(days=step)
        if is_trading_day(d): cnt -= 1
    return d
ANCHOR = date(2026, 10, 20)

NAVY = PatternFill('solid', fgColor='1A202C'); WHITE = Font(bold=True, color='FFFFFF', size=9)
TITLE = Font(bold=True, color='1A202C', size=13); SUB = Font(color='4A5568', size=10, italic=True)
GF = PatternFill('solid', fgColor='E6F4EA'); RF = PatternFill('solid', fgColor='FCE8E6')
AM = PatternFill('solid', fgColor='FEF3C7'); LBL = PatternFill('solid', fgColor='F1F5F9')
GOOD = Font(color='137333', bold=True); BAD = Font(color='C5221F', bold=True); BOLD = Font(bold=True)
THIN = Border(*[Side(style='thin', color='D9E2EC')] * 4)
CEN = Alignment(horizontal='center'); RIGHT = Alignment(horizontal='right'); WRAP = Alignment(wrap_text=True, vertical='top')

wb = Workbook()

# ============ Sheet 1: Ranked Summary (walk-forward) ============
rk = wb.active; rk.title = 'WalkForward_Ranked'
rk.merge_cells('A1:L1')
rk.cell(1, 1, 'ALL F&O STOCKS \u2014 DUSSEHRA, WALK-FORWARD RANKED (point-in-time, no lookahead)').font = TITLE
rk.merge_cells('A2:L2')
rk.cell(2, 1, f'{len(qualified)} of {len(results)} stocks have >=1 genuine out-of-sample trade ({n_low_conf} low-confidence, n<3 OOS trades). Ranked by OOS avg return. This SUPERSEDES the earlier in-sample FO211 Dussehra reports.').font = SUB
hdr = ['Rank', 'Symbol', 'Company', 'Sector', 'Upcoming Direction', 'OOS Win Rate %', 'OOS Avg Return %', 'OOS Trades (n)', 'Confidence', 'Window', 'Data Source']
for j, h in enumerate(hdr, 1):
    c = rk.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
row = 5
for i, (sym, r) in enumerate(ranked, 1):
    m = meta.get(sym, {})
    low = r['low_confidence']
    vals = [i, sym, m.get('name', sym), m.get('sector', ''), r['upcoming_side'], r['oos_wr'], r['oos_avg'],
            r['oos_n'], ('LOW (n<3)' if low else 'OK'), f"T-{r['upcoming_bo']} to T+{r['upcoming_so']}", r['source']]
    fill = AM if low else (GF if r['oos_avg'] >= 0 else RF)
    for j, v in enumerate(vals, 1):
        c = rk.cell(row, j, v); c.border = THIN; c.fill = fill
        if j in (1, 5, 8, 9): c.alignment = CEN
        if j in (6, 7): c.alignment = RIGHT
        if j == 7: c.font = GOOD if r['oos_avg'] >= 0 else BAD
        if j == 9 and low: c.font = Font(color='92400E', bold=True, size=9)
    row += 1
for j, w in enumerate([6, 12, 26, 20, 12, 13, 13, 10, 12, 15, 13], 1):
    rk.column_dimensions[get_column_letter(j)].width = w
rk.freeze_panes = 'A5'

# ============ Sheet 2: Insufficient history ============
ins = wb.create_sheet('Insufficient_History')
ins.merge_cells('A1:D1')
ins.cell(1, 1, f'{len(insufficient)} F&O stocks with ZERO walk-forward trades possible (need >=4yr history to clear the 3-prior-year minimum)').font = TITLE
hdr = ['Symbol', 'Company', 'Years of Underlying Data', 'Reason']
for j, h in enumerate(hdr, 1):
    c = ins.cell(3, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
row = 4
for sym, r in sorted(insufficient.items()):
    m = meta.get(sym, {})
    reason = 'No price history for any Dussehra date' if r.get('source') is None else 'Has some history, but fewer than 3 PRIOR years exist before its first possible trade \u2014 walk-forward needs a 3-year warm-up, in-sample does not (this is why it appeared ranked in the earlier report but not here)'
    vals = [sym, m.get('name', sym), r.get('source') or '\u2014', reason]
    for j, v in enumerate(vals, 1):
        c = ins.cell(row, j, v); c.border = THIN; c.alignment = WRAP if j == 4 else None
    row += 1
for j, w in enumerate([13, 28, 20, 70], 1):
    ins.column_dimensions[get_column_letter(j)].width = w

# ============ Sheet 3: Position Schedule 2026 ============
ps = wb.create_sheet('Position_Schedule_2026')
ps.merge_cells('A1:I1')
ps.cell(1, 1, 'DUSSEHRA 20-Oct-2026 \u2014 Walk-Forward Position Schedule (qualified stocks only)').font = TITLE
ps.merge_cells('A2:I2')
ps.cell(2, 1, 'Entry/exit on the NSE-2026 trading calendar. Sorted by entry date. Window/direction chosen via walk-forward 8x8 grid search using ALL history available up to today.').font = SUB
hdr = ['Symbol', 'Company', 'Direction', 'Window', 'Entry Date', 'Exit Date', 'OOS Win Rate %', 'OOS Trades', 'Confidence']
for j, h in enumerate(hdr, 1):
    c = ps.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
sched = []
for sym, r in qualified.items():
    m = meta.get(sym, {})
    en = shift_trading_days(ANCHOR, -r['upcoming_bo'])
    ex = shift_trading_days(ANCHOR, r['upcoming_so'])
    sched.append((en, sym, m.get('name', sym), r['upcoming_side'], f"T-{r['upcoming_bo']} to T+{r['upcoming_so']}",
                  en.strftime('%d-%b-%Y (%a)'), ex.strftime('%d-%b-%Y (%a)'), r['oos_wr'], r['oos_n'], r['low_confidence']))
sched.sort(key=lambda t: (t[0], t[1]))
row = 5
for _, sym, name, direc, window, entry, exit_, wr, n, low in sched:
    vals = [sym, name, direc, window, entry, exit_, wr, n, ('LOW (n<3)' if low else 'OK')]
    fill = AM if low else (GF if direc == 'LONG' else RF)
    for j, v in enumerate(vals, 1):
        c = ps.cell(row, j, v); c.border = THIN; c.fill = fill
        if j in (1, 3, 4, 5, 6, 9): c.alignment = CEN
        if j in (7, 8): c.alignment = RIGHT
    row += 1
for j, w in enumerate([13, 28, 10, 15, 18, 18, 13, 10, 12], 1):
    ps.column_dimensions[get_column_letter(j)].width = w
ps.freeze_panes = 'A5'

# ============ Sheet 4: Overview ============
ov = wb.create_sheet('Overview', 0)
ov.merge_cells('A1:B1')
ov.cell(1, 1, 'DUSSEHRA WALK-FORWARD RANKED REPORT \u2014 OVERVIEW').font = TITLE
n_long = sum(1 for _, r in ranked if r['upcoming_side'] == 'LONG')
n_short = len(ranked) - n_long
notes = [
    ('What changed', 'The earlier FO211 Dussehra reports (Ranked_Summary, Top50, High-WinRate, WinRate-Ranked) picked each stock\u2019s best direction using ALL available years at once \u2014 in-sample, no held-out data. This report uses the SAME walk-forward 8x8 grid search already used for the Nifty index options model: for every year, only years strictly BEFORE it are used to choose that year\u2019s window+direction, which is then applied forward and the real result recorded. This is genuine point-in-time validation.'),
    ('Result', f'{len(qualified)} of {len(results)} F&O stocks have at least 1 genuine out-of-sample trade. Direction split: {n_long} LONG / {n_short} SHORT. Most Nifty-50 names show OOS win rates much closer to 50% than the in-sample reports showed \u2014 this is the honest number.'),
    ('Why some stocks dropped out', f'{len(insufficient)} stocks that appeared ranked in the earlier in-sample report are NOT here \u2014 walk-forward needs 3 PRIOR years before the first trade can even be taken (a warm-up period), so stocks with only 1-3 years of total history can\u2019t produce a single OOS trade, even though they could produce an in-sample "result" using all 1-3 years at once. See Insufficient_History.'),
    ('Confidence', f'{n_low_conf} of the {len(qualified)} qualified stocks have fewer than 3 OOS trades and are flagged LOW confidence \u2014 amber. Everything else has a real walk-forward track record (many Nifty-50 names show n=23, i.e. tested across 23 separate Dussehra years).'),
    ('Standout', 'HINDUNILVR is the strongest genuinely-validated result found: 82.6% OOS win rate across 23 real walk-forward trades, SHORT direction.'),
    ('Live dashboard', 'This report\u2019s numbers now match what is live on the dashboard\u2019s Dussehra 2026 card for the 50 Nifty-50 names (BSE and JIOFIN excepted \u2014 insufficient history, kept on their prior fallback values).'),
]
r0 = 3
for lbl, txt in notes:
    a = ov.cell(r0, 1, lbl); a.font = BOLD; a.fill = AM; a.border = THIN; a.alignment = WRAP
    b = ov.cell(r0, 2, txt); b.alignment = WRAP; b.border = THIN
    ov.row_dimensions[r0].height = 85; r0 += 1
ov.column_dimensions['A'].width = 22; ov.column_dimensions['B'].width = 105

wb.save(OUT_XLSX)
print('Saved Excel:', OUT_XLSX)
print(f'{len(qualified)} qualified / {len(insufficient)} insufficient. Direction: {n_long} LONG / {n_short} SHORT')

# ============ PDF ============
def esc(s): return str(s).replace('&', '&amp;').replace('<', '&lt;')
top20 = ranked[:20]
bot10 = ranked[-10:]
def pdf_row(rank, sym, r):
    low = r['low_confidence']
    rowstyle = ' style="background:#fffbeb;"' if low else ''
    conf = '<span style="color:#b45309;font-weight:700;">LOW</span>' if low else 'OK'
    return (f'<tr{rowstyle}><td>{rank}</td><td><b>{esc(sym)}</b></td><td>{esc(meta.get(sym,{}).get("name",sym))[:24]}</td>'
            f'<td class="c">{r["upcoming_side"]}</td><td class="r">{r["oos_wr"]:.1f}%</td>'
            f'<td class="r {"pos" if r["oos_avg"]>=0 else "neg"}">{r["oos_avg"]:+.2f}%</td>'
            f'<td class="c">{r["oos_n"]}</td><td class="c">{conf}</td></tr>')
rows_top = ''.join(pdf_row(i + 1, s, r) for i, (s, r) in enumerate(top20))
rows_bot = ''.join(pdf_row(len(ranked) - 9 + i, s, r) for i, (s, r) in enumerate(bot10))

html = f"""<!doctype html><html><head><meta charset="utf-8">
<title>F&amp;O-211 Dussehra Walk-Forward Ranked</title>
<style>
 *{{box-sizing:border-box}} body{{font-family:'Segoe UI',Arial,sans-serif;margin:0;color:#0f172a;background:#fff}}
 .wrap{{max-width:1000px;margin:0 auto;padding:26px 34px}}
 h1{{font-size:21px;margin:0 0 2px}} .sub{{color:#64748b;font-size:12px;margin-bottom:14px}}
 h2{{font-size:15px;margin:20px 0 8px;border-left:4px solid #0f766e;padding-left:9px}}
 table{{width:100%;border-collapse:collapse;font-size:11.5px}}
 th,td{{padding:5px 7px;text-align:right;border-bottom:1px solid #eef2f7}} th{{background:#f1f5f9;font-size:9.5px;text-transform:uppercase;color:#475569}}
 td.c,th.c{{text-align:center}} .pos{{color:#059669;font-weight:700}} .neg{{color:#dc2626;font-weight:700}}
 .note{{background:#f0fdfa;border:1px solid #99f6e4;border-radius:8px;padding:11px 14px;font-size:11px;color:#134e4a;margin-top:8px;line-height:1.5}}
 .cards{{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;margin:10px 0}}
 .mc{{background:#fff;border:1px solid #e2e8f0;border-radius:8px;padding:10px 12px}}
 .mcl{{font-size:9.5px;text-transform:uppercase;color:#64748b;font-weight:700}} .mcv{{font-size:19px;font-weight:800}}
 .foot{{color:#94a3b8;font-size:9.5px;margin-top:18px}} @page{{size:A4;margin:12mm}}
</style></head><body><div class="wrap">
 <h1>All F&amp;O Stocks &mdash; Dussehra, Walk-Forward Ranked</h1>
 <div class="sub">Genuine out-of-sample (OOS) results &bull; point-in-time 8x8 grid search, same method as the Nifty index options model &bull; supersedes the earlier in-sample FO211 Dussehra reports.</div>
 <div class="cards">
  <div class="mc"><div class="mcl">OOS-qualified</div><div class="mcv">{len(qualified)}</div></div>
  <div class="mc"><div class="mcl">Low-confidence (n&lt;3)</div><div class="mcv" style="color:#b45309">{n_low_conf}</div></div>
  <div class="mc"><div class="mcl">Insufficient history</div><div class="mcv">{len(insufficient)}</div></div>
  <div class="mc"><div class="mcl">LONG / SHORT</div><div class="mcv">{n_long} / {n_short}</div></div>
 </div>
 <h2>Top 20 by OOS Avg Return</h2>
 <table><thead><tr><th>Rank</th><th class="c">Symbol</th><th class="c">Company</th><th class="c">Dir</th><th>OOS Win%</th><th>OOS Avg</th><th class="c">n</th><th class="c">Conf.</th></tr></thead>
 <tbody>{rows_top}</tbody></table>
 <h2>Bottom 10 by OOS Avg Return</h2>
 <table><thead><tr><th>Rank</th><th class="c">Symbol</th><th class="c">Company</th><th class="c">Dir</th><th>OOS Win%</th><th>OOS Avg</th><th class="c">n</th><th class="c">Conf.</th></tr></thead>
 <tbody>{rows_bot}</tbody></table>
 <div class="note"><b>Why this is different from earlier reports today:</b> every number here is a genuine out-of-sample result \u2014 the window and direction for each historical trade were chosen using ONLY years strictly before it, never future data. Most Nifty-50 win rates here are much closer to 50% than the in-sample reports showed; that gap IS the overfitting this project has repeatedly flagged. HINDUNILVR (82.6% WR, n=23) is the strongest genuinely-validated result found.</div>
 <div class="foot">Full 211-stock detail, insufficient-history list, and the 2026 position schedule are in the companion Excel.</div>
</div></body></html>"""
open(OUT_HTML, 'w', encoding='utf-8').write(html)
print('Saved HTML:', OUT_HTML)

for exe in [r'C:\Program Files\Google\Chrome\Application\chrome.exe', r'C:\Program Files (x86)\Google\Chrome\Application\chrome.exe']:
    if os.path.exists(exe):
        try:
            subprocess.run([exe, '--headless=new', '--disable-gpu', '--no-pdf-header-footer',
                            '--print-to-pdf=' + OUT_PDF, 'file:///' + OUT_HTML.replace('\\', '/')], timeout=60, capture_output=True)
            print('Saved PDF:', OUT_PDF, 'exists:', os.path.exists(OUT_PDF))
        except Exception as e:
            print('PDF render failed:', e)
        break
