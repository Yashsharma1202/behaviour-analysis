import openpyxl, pandas as pd, sys

sys.stdout.reconfigure(errors='replace')

path = r'D:/behaviour analysis/Nifty50_12_Quarters_Equity_Consolidated_Summary.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

print("==========================================")
print("VERIFYING FULL 12-SHEET CONSOLIDATED WORKBOOK")
print("==========================================")
print("Total Sheets:", len(wb.sheetnames))
print("Sheet Names :", wb.sheetnames)

charts_count = 0
for sheet in wb.sheetnames:
    ws = wb[sheet]
    num_charts = len(ws._charts)
    if num_charts > 0:
        charts_count += num_charts
        print(f"  Sheet '{sheet}': {num_charts} embedded chart(s)")

print(f"\nTotal Embedded Charts: {charts_count}")

df_exec = pd.read_excel(path, sheet_name='Executive_Drawdown_Summary', engine='openpyxl')
print("\nExecutive Summary Table:")
print(df_exec.iloc[5:18, :7].to_string())
print("==========================================")
