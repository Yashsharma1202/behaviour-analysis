"""
Calendar-date backtest, 22-Oct .. 31-Oct, using the project's standard 8x8 grid logic
(Buy_Before_Sell_After_Logic.pdf / build_walkforward_windows_real_dates.py):
  T      = first trading day on/after the calendar date (Sat/Sun dates included; positions only on trading days)
  entry  = close[T - N], N = 1..8 trading days before
  exit   = close[T + M], M = 1..8 trading days after
  64 rules x {LONG, SHORT}; need >= MIN_EVENTS events; prefer rules with avg > 0; rank by (win rate, avg return).
Extra honesty columns: walk-forward plan (chosen on years <= WF_SPLIT only) and its result on later years.
Output: Calendar_Backtest_22Oct_31Oct_8x8.xlsx (new file, additive).
"""
import sys, glob, os, pathlib, datetime
import numpy as np, pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter
sys.stdout.reconfigure(encoding='utf-8')

ROOT = pathlib.Path('D:/behaviour analysis')
DATA = ROOT / 'Stock_Data_213' / 'daily_ohlc'
OUT = ROOT / 'Calendar_Backtest_22Oct_31Oct_8x8.xlsx'
TARGETS = [(10, d) for d in range(22, 32)]
NS = list(range(1, 9)); MS = list(range(1, 9))
PN = np.array([n for n in NS for m in MS]); PM = np.array([m for n in NS for m in MS])
MIN_EVENTS, WF_SPLIT, TOP_MIN_YEARS = 5, 2018, 15
OUTLIER = 60.0   # drop stock-years with any |return| > 60% inside the 16-day window (Yahoo early-data glitches)
EXCLUDED = []
INDICES = {'NIFTY', 'BANKNIFTY'}
NIFTY50 = ['RELIANCE','HDFCBANK','ICICIBANK','INFY','TCS','AXISBANK','SBIN','BHARTIARTL','BAJFINANCE','KOTAKBANK',
 'LT','M&M','MARUTI','TATACONSUM','SUNPHARMA','TATASTEEL','HINDUNILVR','NTPC','POWERGRID','ITC','COALINDIA',
 'JSWSTEEL','HINDALCO','GRASIM','EICHERMOT','BAJAJ-AUTO','CIPLA','DRREDDY','APOLLOHOSP','ASIANPAINT','TITAN',
 'DIVISLAB','NESTLEIND','BRITANNIA','ADANIENT','ADANIPORTS','ONGC','BPCL','BEL','ULTRACEMCO','TRENT','HCLTECH',
 'TECHM','WIPRO','SHRIRAMFIN','BAJAJFINSV','INDIGO','JIOFIN','HDFCLIFE','ETERNAL','TMPV','SBILIFE']

def load(sym):
    d = pd.read_csv(DATA / f'{sym}.csv', usecols=['Date', 'Adj Close']).dropna()
    d['Date'] = pd.to_datetime(d['Date'])
    return d.sort_values('Date').drop_duplicates('Date').set_index('Date')['Adj Close']

def samples(px, m, d, sym='', label=''):
    dates = px.index.values; c = px.values; recs = []
    for y in range(px.index[0].year, px.index[-1].year + 1):
        try: cal = datetime.date(y, m, d)
        except ValueError: continue
        t = np.datetime64(cal); a = int(np.searchsorted(dates, t))
        if a >= len(dates) or a - 8 < 0 or a + 8 >= len(dates): continue
        if dates[a] - t > np.timedelta64(6, 'D'): continue
        rr = (c[a + PM] / c[a - PN] - 1) * 100
        if np.abs(rr).max() > OUTLIER:
            EXCLUDED.append(dict(Date=label, Symbol=sym, Year=y, MaxAbsMovePct=round(float(np.abs(rr).max()), 1))); continue
        recs.append((y, cal, a, rr))
    if not recs: return None
    return np.array([r[0] for r in recs]), np.array([r[3] for r in recs]), [r[1] for r in recs], [r[2] for r in recs]

def pick(R, side):
    """project selection: need >= MIN_EVENTS, prefer avg>0, rank by (win, avg). returns (k, win, avg) or None"""
    n = len(R)
    if n < MIN_EVENTS: return None
    r = R * side
    win = (r > 0).mean(0) * 100; avg = r.mean(0)
    cand = np.where(avg > 0)[0]
    if len(cand) == 0: cand = np.arange(len(PN))
    k = max(cand, key=lambda i: (round(win[i], 6), avg[i]))
    return int(k), win[k], avg[k]

def best(R):
    L, S = pick(R, 1), pick(R, -1)
    if L is None: return None
    if S is None or (L[1], L[2]) >= (S[1], S[2]): return (*L, 1)
    return (*S, -1)

def plan(ys, R, wkend):
    b = best(R)
    if not b: return None
    k, win, avg, sd = b; r = R[:, k] * sd
    wf = None
    ins = ys <= WF_SPLIT
    if ins.sum() >= MIN_EVENTS and (~ins).sum() > 0:
        bw = best(R[ins])
        if bw:
            kw, _, _, sw = bw; ro = R[~ins, kw] * sw
            wf = dict(WF_N=int(PN[kw]), WF_M=int(PM[kw]), WF_Side='LONG' if sw == 1 else 'SHORT', WF_N_oos=int(len(ro)),
                      WF_Avg=ro.mean(), WF_Win=(ro > 0).mean() * 100)
    if wf is None: wf = dict(WF_N=np.nan, WF_M=np.nan, WF_Side='', WF_N_oos=0, WF_Avg=np.nan, WF_Win=np.nan)
    def st(mask):
        n = int(mask.sum()); return (n, r[mask].mean() if n else np.nan, (r[mask] > 0).mean() * 100 if n else np.nan)
    wdn, wda, wdw = st(~wkend); wen, wea, wew = st(wkend)
    return dict(k=k, Side='LONG' if sd == 1 else 'SHORT', N=int(PN[k]), M=int(PM[k]), Window=f'T-{PN[k]} to T+{PM[k]}',
                Events=len(ys), Win=win, Avg=avg, Worst=r.min(), Best=r.max(),
                WeekdayN=wdn, WeekdayAvg=wda, WeekdayWin=wdw, WeekendN=wen, WeekendAvg=wea, WeekendWin=wew, **wf)

def top10(rows):
    c = [r for r in rows if r['Type'] != 'Index' and r['Events'] >= TOP_MIN_YEARS and r['Avg'] > 0]
    c.sort(key=lambda r: (-r['Win'], -r['Avg']))
    return c[:10]

def main():
    px = {}
    for f in sorted(glob.glob(str(DATA / '*.csv'))):
        s = os.path.basename(f)[:-4]
        if s in ('FINNIFTY', 'MIDCPNIFTY', 'SENSEX'): continue
        try: p = load(s)
        except Exception: continue
        if len(p) >= 250: px[s] = p
    if 'TATAMOTORS' in px and 'TMPV' in px: del px['TATAMOTORS']
    n50 = [s for s in NIFTY50 if s in px]
    print('universe', len(px), 'nifty50 members with data', len(n50), flush=True)
    sheets, summary = {}, []
    for (m, d) in TARGETS:
        label = f'{d:02d}-Oct'; per, rows = {}, []
        for s, p in px.items():
            sm = samples(p, m, d, s, label)
            if not sm: continue
            per[s] = sm; ys, R, cals, _ = sm
            pl = plan(ys, R, np.array([c.weekday() >= 5 for c in cals]))
            if pl: rows.append(dict(Symbol=s, Type='Index' if s in INDICES else ('Nifty50' if s in n50 else 'Stock'), **pl))
        def basket(members):
            yr = {}
            for s in members:
                if s in per:
                    ys, R, cals, _ = per[s]
                    for i, y in enumerate(ys): yr.setdefault(y, []).append((R[i], cals[i]))
            yy = np.array(sorted(y for y, v in yr.items() if len(v) >= 3))
            if len(yy) < MIN_EVENTS: return None
            R = np.array([np.mean([x[0] for x in yr[y]], axis=0) for y in yy]); cals = [yr[y][0][1] for y in yy]
            wk = np.array([c.weekday() >= 5 for c in cals]); pl = plan(yy, R, wk)
            return (pl, yy, R, cals) if pl else None
        B50 = basket(n50); BALL = basket([s for s in px if s not in INDICES])
        sheets[label] = dict(B50=B50, BALL=BALL, rows=rows, top=top10(rows), m=m, d=d)
        for nm, B in (('Nifty 50 basket', B50), ('All-stock basket', BALL)):
            if B: summary.append(dict(Date=label, Universe=nm, **{k: v for k, v in B[0].items() if k != 'k'}))
        for r in rows:
            if r['Symbol'] in INDICES: summary.append(dict(Date=label, Universe=r['Symbol'], **{k: v for k, v in r.items() if k not in ('k', 'Symbol', 'Type')}))
        print(label, 'symbols with plan:', len(rows), flush=True)

    hdr = PatternFill('solid', fgColor='1F3864'); hf = Font(bold=True, color='FFFFFF')
    sec = PatternFill('solid', fgColor='D9E1F2'); GREEN = PatternFill('solid', fgColor='C6EFCE')
    def style_header(ws, row, ncol):
        for c in range(1, ncol + 1):
            x = ws.cell(row, c); x.fill = hdr; x.font = hf; x.alignment = Alignment(wrap_text=True, vertical='center')
    def band(ws, r, text, fill=sec, color='000000', size=11):
        ws.cell(r, 1, text).font = Font(bold=True, color=color, size=size)
        for c in range(1, 14): ws.cell(r, c).fill = fill
    scols = ['Date','Universe','Side','N','M','Window','Events','Win','Avg','Worst','Best','WeekdayN','WeekdayAvg','WeekdayWin',
             'WeekendN','WeekendAvg','WeekendWin','WF_Side','WF_N','WF_M','WF_N_oos','WF_Win','WF_Avg']
    snames = ['Date','Universe','Side','Buy N days BEFORE (T-N)','Sell M days AFTER (T+M)','Window','Events (years)','Win %','Avg return %','Worst %','Best %',
              'Weekday-date yrs','Weekday avg %','Weekday win %','Weekend-date yrs','Weekend avg %','Weekend win %',
              'WF plan side (chosen <=2018)','WF N','WF M','WF yrs 2019+','WF win % 2019+','WF avg % 2019+']
    S = pd.DataFrame(summary)[scols]; S.columns = snames
    with pd.ExcelWriter(OUT, engine='openpyxl') as xw:
        pd.DataFrame({'Calendar-date backtest 22-Oct..31-Oct - project 8x8 grid logic': [
            'Same logic as Buy_Before_Sell_After_Logic.pdf: T = first trading day on/after the calendar date (Sat/Sun dates included, positions only on trading days).',
            'Buy at close of T-N, sell at close of T+M, N=1..8 trading days before, M=1..8 after: 64 rules, each as LONG and SHORT.',
            f'Selection: need >= {MIN_EVENTS} events, prefer rules with positive average, rank by win rate then average return (best LONG vs best SHORT -> higher wins).',
            'Data: Yahoo Finance via yfinance, Adj Close, full history per symbol, up to 15-Oct-2026. Price moves only (no costs). "Nifty 50 basket" = the project nifty50_list (52 names, same list as the holiday builds).',
            'IN-SAMPLE: the winning rule is chosen from 64x2 tries on the same history it is scored on, so Win % / Avg % flatter. Read with the Events count.',
            f'WF columns = walk-forward check: the rule re-chosen using years <= {WF_SPLIT} only, then scored on later years (honest out-of-sample).',
            f'Data guard: any stock-year with a move >{OUTLIER:.0f}% inside the T-8..T+8 window is dropped as a probable Yahoo data glitch (listed in Excluded_Outliers).',
            f'Top 10 = stocks (no indices) with >= {TOP_MIN_YEARS} events and Avg>0, ranked by Win % then Avg %.']}).to_excel(xw, sheet_name='ReadMe', index=False)
        S.round(2).to_excel(xw, sheet_name='Summary', index=False)
        T = pd.DataFrame([dict(Date=lb, Rank=i, Symbol=t['Symbol'], Side=t['Side'], Buy_N_days_before=t['N'], Sell_M_days_after=t['M'], Window=t['Window'],
                               Events=t['Events'], Win_pct=round(t['Win'], 1), Avg_pct=round(t['Avg'], 2), Worst_pct=round(t['Worst'], 2),
                               WF_win_pct_2019plus=round(t['WF_Win'], 1) if t['WF_Win'] == t['WF_Win'] else None,
                               WF_avg_pct_2019plus=round(t['WF_Avg'], 2) if t['WF_Avg'] == t['WF_Avg'] else None)
                          for lb, sh in sheets.items() for i, t in enumerate(sh['top'], 1)])
        T.to_excel(xw, sheet_name='Top10_All_Dates', index=False)
        wt = xw.sheets['Top10_All_Dates']; style_header(wt, 1, T.shape[1]); wt.freeze_panes = 'A2'; wt.auto_filter.ref = wt.dimensions
        for c in range(1, T.shape[1] + 1): wt.column_dimensions[get_column_letter(c)].width = 17
        for rr in range(2, wt.max_row + 1):
            for c in range(1, T.shape[1] + 1): wt.cell(rr, c).fill = GREEN
        X = pd.DataFrame(EXCLUDED); X.to_excel(xw, sheet_name='Excluded_Outliers', index=False)
        style_header(xw.sheets['Excluded_Outliers'], 1, X.shape[1])
        for c in range(1, X.shape[1] + 1): xw.sheets['Excluded_Outliers'].column_dimensions[get_column_letter(c)].width = 18
        print('excluded stock-years:', len(X), '| unique symbol-years:', len(X.drop_duplicates(['Symbol','Year'])))
        print(X.drop_duplicates(['Symbol','Year']).sort_values('MaxAbsMovePct', ascending=False).head(12).to_string(index=False))
        for label, sh in sheets.items():
            ws = xw.book.create_sheet(label); xw.sheets[label] = ws; r = 1
            ws.cell(r, 1, f'{label}  -  T = first trading day on/after {label};  buy T-N, sell T+M  (N,M = 1..8)').font = Font(bold=True, size=14); r += 2
            band(ws, r, f'TOP 10 STOCKS for {label}  (>= {TOP_MIN_YEARS} events, Avg>0, ranked by Win % then Avg %)', GREEN, '006100', 12); r += 1
            th = ['#','Symbol','Side','Buy N days BEFORE','Sell M days AFTER','Window','Events','Win %','Avg %','Worst %','WF win % 2019+','WF avg % 2019+']
            for i, h in enumerate(th, 1): ws.cell(r, i, h)
            style_header(ws, r, len(th)); r += 1
            for n, t in enumerate(sh['top'], 1):
                vals = [n, t['Symbol'], t['Side'], t['N'], t['M'], t['Window'], t['Events'], round(t['Win'], 1), round(t['Avg'], 2), round(t['Worst'], 2),
                        None if t['WF_Win'] != t['WF_Win'] else round(t['WF_Win'], 1), None if t['WF_Avg'] != t['WF_Avg'] else round(t['WF_Avg'], 2)]
                for i, v in enumerate(vals, 1):
                    x = ws.cell(r, i, v); x.fill = GREEN; x.font = Font(bold=True)
                r += 1
            r += 1
            for title, B in (('NIFTY 50 BASKET (equal-weight)', sh['B50']), ('ALL-STOCK BASKET (equal-weight, ex-index)', sh['BALL'])):
                band(ws, r, title); r += 1
                if not B: ws.cell(r, 1, 'insufficient data'); r += 2; continue
                p = B[0]
                wd = f"{p['WeekdayN']} yrs avg {p['WeekdayAvg']:.2f}% win {p['WeekdayWin']:.0f}%"
                we = f"{p['WeekendN']} yrs avg {p['WeekendAvg']:.2f}% win {p['WeekendWin']:.0f}%" if p['WeekendN'] else 'no weekend-date years'
                wf = (f"{p['WF_Side']} T-{int(p['WF_N'])} to T+{int(p['WF_M'])} -> {p['WF_N_oos']} yrs 2019+: win {p['WF_Win']:.0f}%, avg {p['WF_Avg']:.2f}%"
                      if p['WF_N_oos'] else 'n/a')
                for a, b in [('Recommended plan', f"{p['Side']}: buy {p['N']} trading days BEFORE T, sell {p['M']} days AFTER T  ({p['Window']})"),
                             ('Events / win % / avg %', f"{p['Events']} yrs / {p['Win']:.1f}% / {p['Avg']:.2f}%"),
                             ('Worst / best year', f"{p['Worst']:.2f}% / {p['Best']:.2f}%"),
                             ('Weekday-date vs weekend-date yrs', f'{wd}  |  {we}'),
                             ('Walk-forward check (honest OOS)', wf)]:
                    ws.cell(r, 1, a).font = Font(bold=True); ws.cell(r, 2, b); r += 1
                r += 1
            # 8x8 grids for the Nifty 50 basket (LONG)
            B = sh['B50']
            if B:
                p, yy, R, cals = B
                for gname, mat in (('LONG win % grid - Nifty 50 basket  (rows = buy N days before, columns = sell M days after)', (R > 0).mean(0) * 100),
                                   ('LONG avg return % grid - Nifty 50 basket', R.mean(0))):
                    band(ws, r, gname); r += 1
                    ws.cell(r, 1, 'N \\ M')
                    for j, mm in enumerate(MS, 2): ws.cell(r, j, f'T+{mm}')
                    style_header(ws, r, 9); r += 1; g0 = r
                    for ni, nn in enumerate(NS):
                        ws.cell(r, 1, f'T-{nn}').font = Font(bold=True)
                        for mi in range(8): ws.cell(r, mi + 2, round(float(mat[ni * 8 + mi]), 2))
                        r += 1
                    ws.conditional_formatting.add(f'B{g0}:I{r-1}', ColorScaleRule(start_type='min', start_color='F8696B', mid_type='percentile', mid_value=50,
                                                                              mid_color='FFFFFF', end_type='max', end_color='63BE7B'))
                    r += 1
                band(ws, r, 'YEAR-BY-YEAR TRADES - Nifty 50 basket plan (side-adjusted returns)'); r += 1
                heads = ['Year','Calendar date','Weekday','T (event day)','Entry date','Exit date','Basket return %','Result']
                for i, h in enumerate(heads, 1): ws.cell(r, i, h)
                style_header(ws, r, len(heads)); r += 1
                cpx = px['RELIANCE']
                dts = cpx.index.values; sg = 1 if p['Side'] == 'LONG' else -1
                for i, y in enumerate(yy):
                    cd = cals[i]; a = int(np.searchsorted(dts, np.datetime64(cd)))
                    if a - 8 < 0 or a + 8 >= len(dts): continue
                    ret = R[i, p['k']] * sg
                    ws.append([int(y), cd.strftime('%d-%b-%Y'), cd.strftime('%a'), str(dts[a])[:10], str(dts[a - p['N']])[:10],
                               str(dts[a + p['M']])[:10], round(float(ret), 2), 'WIN' if ret > 0 else 'LOSS']); r += 1
                r += 1
            band(ws, r, 'EVERY SYMBOL - best 8x8 rule per symbol (green = Top 10)'); r += 1
            df = pd.DataFrame(sh['rows'])
            if len(df):
                df = df.sort_values(['Win', 'Avg'], ascending=False)
                df = df[['Symbol','Type','Side','N','M','Window','Events','Win','Avg','Worst','Best','WeekdayN','WeekdayAvg','WeekendN','WeekendAvg','WF_Side','WF_N','WF_M','WF_N_oos','WF_Win','WF_Avg']].round(2)
                df.columns = ['Symbol','Type','Side','Buy N days BEFORE','Sell M days AFTER','Window','Events','Win %','Avg %','Worst %','Best %',
                              'Weekday yrs','Weekday avg %','Weekend yrs','Weekend avg %','WF side','WF N','WF M','WF yrs 2019+','WF win % 2019+','WF avg % 2019+']
                for i, h in enumerate(df.columns, 1): ws.cell(r, i, h)
                style_header(ws, r, len(df.columns)); hr = r; r += 1
                topset = {t['Symbol'] for t in sh['top']}
                for row in df.itertuples(index=False):
                    ws.append([None if (isinstance(v, float) and v != v) else v for v in row])
                    if row[0] in topset:
                        for c in range(1, len(df.columns) + 1): ws.cell(r, c).fill = GREEN
                    r += 1
                ws.auto_filter.ref = f'A{hr}:{get_column_letter(len(df.columns))}{r-1}'
            ws.column_dimensions['A'].width = 36; ws.column_dimensions['B'].width = 40
            for c in range(3, 22): ws.column_dimensions[get_column_letter(c)].width = 14
        ws = xw.sheets['Summary']; style_header(ws, 1, S.shape[1]); ws.freeze_panes = 'C2'; ws.row_dimensions[1].height = 48
        for c in range(1, S.shape[1] + 1): ws.column_dimensions[get_column_letter(c)].width = 14
        ws.auto_filter.ref = ws.dimensions
        xw.sheets['ReadMe'].column_dimensions['A'].width = 170
    print('saved', OUT)
    pd.set_option('display.width', 250)
    q = S[S.Universe.isin(['Nifty 50 basket', 'All-stock basket'])]
    print(q[['Date','Universe','Side','Window','Events (years)','Win %','Avg return %','WF win % 2019+','WF avg % 2019+']].round(2).to_string(index=False))
    print(T[T.Date.isin(['22-Oct', '27-Oct'])].to_string(index=False))
    print(T.Symbol.value_counts().head(8))

if __name__ == '__main__':
    main()
