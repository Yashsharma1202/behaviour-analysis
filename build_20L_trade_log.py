"""
DETAILED TRADE LOG (12 quarters) — 1% ITM OPTIONS on Rs 20 Lakh / quarter.
Every individual option trade (50 stocks x 12 quarters = 600), ordered by exit date,
with each trade's P&L scaled to a Rs 20L/quarter book, plus running equity and drawdown,
and a full performance-measures sheet. Same detailed-log style as the earlier masters.

Source: GitHub_Dashboard_12_Quarters_1PCT_ITM_Options_Master.xlsx.
ADDITIVE: writes a new workbook.
"""
import openpyxl, statistics as st, math, re
from datetime import date
from collections import defaultdict
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

BASE = r'D:\behaviour analysis'
OUT = BASE + r'\GitHub_Dashboard_20L_Options_Trade_Log.xlsx'
CAP = 2_000_000
MONS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']


def pdate(s):
    s = str(s)
    m = re.search(r'(\d{4})-(\d{2})-(\d{2})', s)
    if m: return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.search(r'(\d{1,2})-([A-Za-z]{3})-(\d{4})', s)
    if m: return date(int(m.group(3)), MONS.index(m.group(2)) + 1, int(m.group(1)))
    return None


wb0 = openpyxl.load_workbook(BASE + r'\GitHub_Dashboard_12_Quarters_1PCT_ITM_Options_Master.xlsx', read_only=True, data_only=True)
T = []; premq = defaultdict(float)
for sn in wb0.sheetnames:
    if sn == 'Dashboard_Summary': continue
    rows = list(wb0[sn].iter_rows(values_only=True)); hdr = rows[3]
    idx = {str(h).strip(): j for j, h in enumerate(hdr) if h}
    def C(n):
        for k, j in idx.items():
            if k.lower().startswith(n.lower()): return j
    cS, cN, cB, cW = C('Symbol'), C('Company'), C('Option Bias'), C('Position Window')
    cED, cXD, cK, cEP, cXP = C('Entry Date'), C('Exit Date'), C('Strike'), C('Entry Premium'), C('Exit Premium')
    cR, cP, cPr, cO = C('Option Return'), C('Option P&L'), C('Premium Paid'), C('Outcome')
    for r in rows[4:]:
        if not r or r[cS] is None: continue
        try: pnl = float(r[cP]); prem = float(r[cPr])
        except Exception: continue
        T.append(dict(q=sn, sym=str(r[cS]), name=str(r[cN] or ''), dir=str(r[cB] or ''), win=str(r[cW] or ''),
                      ed=str(r[cED] or ''), xd=str(r[cXD] or ''), xdt=pdate(r[cXD]),
                      strike=r[cK], ep=r[cEP], xp=r[cXP], ret=r[cR], pnl=pnl, prem=prem, out=str(r[cO] or '')))
        premq[sn] += prem
for t in T:
    t['scale'] = CAP / premq[t['q']] if premq[t['q']] else 0
    t['spnl'] = t['pnl'] * t['scale']
T.sort(key=lambda t: (t['xdt'] or date(1900, 1, 1), t['q']))

# running equity + drawdown (continuous over 600 trades)
cum = peak = 0.0
for t in T:
    cum += t['spnl']; peak = max(peak, cum)
    t['cum'] = cum; t['dd'] = cum - peak

# performance measures (per-quarter returns)
byq = defaultdict(float)
for t in T: byq[t['q']] += t['spnl']
ORDER = ['FY24_Q3', 'FY24_Q4', 'FY25_Q1', 'FY25_Q2', 'FY25_Q3', 'FY25_Q4', 'FY26_Q1', 'FY26_Q2', 'FY26_Q3', 'FY26_Q4', 'FY27_Q1', 'FY27_Q2']
rets = [byq[q] / CAP * 100 for q in ORDER]
mdd = min(t['dd'] for t in T); mdd_pct = mdd / CAP * 100
trade_win = 100 * sum(1 for t in T if t['spnl'] > 0) / len(T)
std = st.pstdev(rets)
pos = [r for r in rets if r > 0]; neg = [r for r in rets if r < 0]

# ---- styles ----
NAVY = PatternFill('solid', fgColor='1A202C'); WHITE = Font(bold=True, color='FFFFFF', size=9)
TITLE = Font(bold=True, color='1A202C', size=13); SUB = Font(color='4A5568', size=10, italic=True)
GREENF = PatternFill('solid', fgColor='E6F4EA'); REDF = PatternFill('solid', fgColor='FCE8E6')
AMBERF = PatternFill('solid', fgColor='FEF3C7'); LBL = PatternFill('solid', fgColor='F1F5F9')
GOOD = Font(color='137333', bold=True); BAD = Font(color='C5221F', bold=True); BOLD = Font(bold=True)
THIN = Border(*[Side(style='thin', color='D9E2EC')] * 4)
CEN = Alignment(horizontal='center'); RIGHT = Alignment(horizontal='right'); WRAP = Alignment(wrap_text=True, vertical='top')

wb = openpyxl.Workbook()
tl = wb.active; tl.title = 'Trade_Log'
tl.merge_cells('A1:R1'); tl.cell(1, 1, 'DETAILED OPTION TRADE LOG (12 quarters) — 1% ITM, Rs 20 Lakh / quarter book').font = TITLE
tl.merge_cells('A2:R2'); tl.cell(2, 1, 'All 600 option trades, ordered by exit date. P&L scaled to a fresh Rs 20L each quarter; running equity & drawdown shown.').font = SUB
HEAD = ['#', 'Quarter', 'Symbol', 'Company', 'Dir', 'Window', 'Entry', 'Exit', 'Strike',
        'Entry Prem', 'Exit Prem', 'Option Ret %', 'P&L 1 Lot (Rs)', 'Scaled P&L 20L (Rs)',
        'Outcome', 'Equity (Rs)', 'Drawdown (Rs)', 'DD %']
WID = [5, 9, 12, 24, 9, 12, 13, 13, 10, 11, 11, 12, 15, 17, 9, 15, 15, 8]
for j, h in enumerate(HEAD, 1):
    c = tl.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
row = 5
for i, t in enumerate(T, 1):
    dcol = '137333' if 'LONG' in t['dir'].upper() or 'CALL' in t['dir'].upper() else 'C5221F'
    win = t['spnl'] > 0
    vals = [i, t['q'], t['sym'], t['name'][:22], t['dir'].replace('BUY ', ''), t['win'],
            t['ed'][:11], t['xd'][:11], t['strike'],
            round(t['ep'], 2) if isinstance(t['ep'], (int, float)) else t['ep'],
            round(t['xp'], 2) if isinstance(t['xp'], (int, float)) else t['xp'],
            round(t['ret'], 1) if isinstance(t['ret'], (int, float)) else t['ret'],
            round(t['pnl']), round(t['spnl']), 'WIN' if win else 'LOSS',
            round(t['cum']), round(t['dd']), round(t['dd'] / CAP * 100, 1)]
    fill = GREENF if win else REDF
    for j, v in enumerate(vals, 1):
        c = tl.cell(row, j, v); c.border = THIN; c.fill = fill
        if j in (1, 2, 5, 6, 7, 8, 15): c.alignment = CEN
        if j in (9, 10, 11, 12, 13, 14, 16, 17, 18): c.alignment = RIGHT
        if j == 14: c.font = GOOD if win else BAD
        if j in (17, 18) and t['dd'] < 0: c.font = BAD
    row += 1
for j, w in enumerate(WID, 1):
    tl.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
tl.freeze_panes = 'A5'

# performance sheet
ps = wb.create_sheet('Performance', 0)
ps.merge_cells('A1:C1'); ps.cell(1, 1, 'PERFORMANCE MEASURES — 1% ITM Options, Rs 20 Lakh / Quarter (12 quarters)').font = TITLE
ps.merge_cells('A2:C2'); ps.cell(2, 1, 'From the detailed trade log (600 trades). Drawdown is trade-level (real).').font = SUB
hdr = ['Quarter', 'Return %', 'Rs P&L on 20L']
for j, h in enumerate(hdr, 1):
    c = ps.cell(4, j, h); c.fill = NAVY; c.font = WHITE; c.border = THIN; c.alignment = CEN
for i, q in enumerate(ORDER, 5):
    r = byq[q] / CAP * 100
    for j, v in enumerate([q, round(r, 1), round(byq[q])], 1):
        c = ps.cell(i, j, v); c.border = THIN; c.fill = GREENF if r >= 0 else REDF
        if j >= 2: c.alignment = RIGHT
        if j == 2: c.font = GOOD if r >= 0 else BAD
r0 = 5 + 12 + 1
ps.cell(r0, 1, 'SUMMARY').font = BOLD
meas = [
    ('Total P&L (12 x Rs 20L)', 'Rs {:+,.0f}'.format(sum(byq.values()))),
    ('Overall return (Rs 2.4 Cr deployed)', '{:+.1f}%'.format(sum(byq.values()) / (CAP * 12) * 100)),
    ('Quarter win rate', '{:.0f}%  ({}/12 quarters)'.format(100 * sum(1 for r in rets if r > 0) / 12, sum(1 for r in rets if r > 0))),
    ('Trade-level win rate', '{:.1f}%  ({} of {} trades)'.format(trade_win, sum(1 for t in T if t['spnl'] > 0), len(T))),
    ('Avg / Median return per quarter', '{:+.1f}% / {:+.1f}%'.format(st.mean(rets), st.median(rets))),
    ('Best / Worst quarter', '{:+.1f}% / {:+.1f}%'.format(max(rets), min(rets))),
    ('Sharpe (annualised)', '{:.2f}'.format(st.mean(rets) / std * math.sqrt(4) if std else 0)),
    ('Profit factor (quarterly)', ('{:.2f}'.format(sum(pos) / abs(sum(neg))) if neg else 'n/a (no losing quarter)')),
    ('>> MAX DRAWDOWN (real, trade-level)', 'Rs {:,.0f}  ({:.1f}% of a 20L book)'.format(mdd, mdd_pct)),
    ('Biggest single winner', 'Rs {:+,.0f}  ({})'.format(max(t['spnl'] for t in T), max(T, key=lambda t: t['spnl'])['sym'])),
    ('Biggest single loser', 'Rs {:+,.0f}  ({})'.format(min(t['spnl'] for t in T), min(T, key=lambda t: t['spnl'])['sym'])),
]
for k, (lbl, val) in enumerate(meas):
    rr = r0 + 1 + k
    a = ps.cell(rr, 1, lbl); a.fill = (AMBERF if lbl.startswith('>>') else LBL); a.font = BOLD; a.border = THIN
    b = ps.cell(rr, 2, val); b.border = THIN; ps.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=3)
note = ps.cell(r0 + 1 + len(meas) + 1, 1, 'NOTE: 1% synthetic option model (not real fills) + a 12-quarter 2023-26 bull sample where every quarter won. Returns are an optimistic ceiling; the -{:.0f}% trade-level drawdown and {:.0f}% trade win rate are the honest risk. Re-run at 3% time value for a conservative view.'.format(abs(mdd_pct), trade_win))
note.font = Font(color='7C2D12', size=9, italic=True); note.alignment = WRAP
ps.merge_cells(start_row=r0 + 1 + len(meas) + 1, start_column=1, end_row=r0 + 3 + len(meas), end_column=3)
for j, w in enumerate([36, 22, 16], 1):
    ps.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w

wb.save(OUT)
print('Saved:', OUT)
print('%d trades logged | total Rs %+,.0f on 20L/qtr | trade-win %.1f%% | maxDD %.1f%% (Rs %,.0f)'
      % (len(T), sum(byq.values()), trade_win, mdd_pct, mdd))
