import openpyxl, pandas as pd, sys

sys.stdout.reconfigure(errors='replace')

master_path = r'D:/behaviour analysis/Nifty50_12_Quarters_Options_1PCT_ITM_Master.xlsx'
summary_path = r'D:/behaviour analysis/Nifty50_12_Quarters_Options_1PCT_ITM_Summary.xlsx'

wb_m = openpyxl.load_workbook(master_path, data_only=True)
ws_m = wb_m['Exec_12Q_Combined_Summary']

wb_s = openpyxl.load_workbook(summary_path, data_only=True)
ws_s = wb_s['Options_1PCT_ITM_Summary']

print("==========================================================================================")
print("VERIFYING 1% ITM OPTIONS STRATEGY MASTER & SUMMARY WORKBOOKS")
print("==========================================================================================")
print("Master Sheet Names:", wb_m.sheetnames)
print(f"Master Sheet Count: {len(wb_m.sheetnames)}")

print("\n--- MASTER EXECUTIVE SCORECARD ROWS 7-21 ---")
for r in range(7, 22):
    vals = [ws_m.cell(row=r, column=c).value for c in range(1, 8)]
    print(f"  Row {r:2d}: {vals}")

print("\n--- STANDALONE SUMMARY ROWS 7-21 ---")
for r in range(7, 22):
    vals = [ws_s.cell(row=r, column=c).value for c in range(1, 8)]
    print(f"  Row {r:2d}: {vals}")
print("==========================================================================================")
