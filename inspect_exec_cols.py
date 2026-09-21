import openpyxl, sys

sys.stdout.reconfigure(errors='replace')

path = r'D:/behaviour analysis/Nifty50_12_Quarters_Equity_Master_v2.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb['Exec_12Q_Combined_Summary']

print("Row 5 Headers:")
for c in range(1, ws.max_column + 1):
    val = ws.cell(row=5, column=c).value
    print(f"  Col {c}: {val}")

print("\nRow 6 Data (Q3 2023-24):")
for c in range(1, ws.max_column + 1):
    val = ws.cell(row=6, column=c).value
    print(f"  Col {c}: {val}")
