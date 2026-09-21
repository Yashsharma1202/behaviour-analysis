import openpyxl, pathlib, sys

sys.stdout.reconfigure(errors='replace')

b_dir = pathlib.Path('D:/behaviour analysis')

master_path  = b_dir / 'Nifty50_12_Quarters_Futures_0.5PCT_Total_Turnover_Master.xlsx'
summary_path = b_dir / 'Nifty50_12_Quarters_Futures_0.5PCT_Total_Turnover_Summary.xlsx'

print("==========================================================================================")
print("VERIFYING GENERATED EXCEL WORKBOOKS FOR DELIVERABLE")
print("==========================================================================================")

if master_path.exists():
    wb_m = openpyxl.load_workbook(master_path, read_only=True)
    print(f"✅ Master 13-Sheet Workbook: {master_path.name}")
    print(f"   Full Absolute Path: {master_path.resolve()}")
    print(f"   Size: {master_path.stat().st_size / 1024:.1f} KB")
    print(f"   Sheets: {len(wb_m.sheetnames)} -> {wb_m.sheetnames[:4]} ...")

if summary_path.exists():
    wb_s = openpyxl.load_workbook(summary_path, read_only=True)
    print(f"\n✅ Summary Workbook: {summary_path.name}")
    print(f"   Full Absolute Path: {summary_path.resolve()}")
    print(f"   Size: {summary_path.stat().st_size / 1024:.1f} KB")

print("==========================================================================================")
