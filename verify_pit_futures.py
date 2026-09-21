import openpyxl, pandas as pd, sys

sys.stdout.reconfigure(errors='replace')

path = r'D:/behaviour analysis/Nifty50_12_Quarters_Point_In_Time_Futures_Master.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

print("==========================================================================================")
print("VERIFYING POINT-IN-TIME FUTURES MASTER WORKBOOK")
print("==========================================================================================")
print("Total Sheets:", len(wb.sheetnames))
print("Sheet Names :", wb.sheetnames)

df_exec = pd.read_excel(path, sheet_name='Exec_12Q_Combined_Summary', engine='openpyxl')
print("\nExecutive Summary Scorecard Table:")
print(df_exec.iloc[5:20, :7].to_string())
print("==========================================================================================")
