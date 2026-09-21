"""
news_effect_study.py  —  do ANY past-news features relate to the event move?
===============================================================================
The single 'bar' saturated, so this widens the test: for each realised Q1 event
it derives several PRE-EVENT (published_at < decision_ts) news features from the
articles we already hold, and correlates each with the realised move from the
tracker. Purely diagnostic — still a PROXY (published_at), one quarter.

    python news_effect_study.py
===============================================================================
"""
from __future__ import annotations
import datetime as dt
import statistics as st
import sys
from email.utils import parsedate_to_datetime
from pathlib import Path

import openpyxl
sys.stdout.reconfigure(encoding="utf-8")
import expectation_bar as EB
import news_ingest

ROOT = Path(__file__).resolve().parent
XLSX = ROOT / "Nifty50_Q1Results_BehaviourEngine_UPDATED_06Aug_fixed.xlsx"


def realised_events():
    ws = openpyxl.load_workbook(XLSX).active
    out = []
    for r in range(3, ws.max_row + 1):
        sym = (ws.cell(r, 2).value or "").strip()
        D = ws.cell(r, 4).value
        o = ws.cell(r, 15).value
        if isinstance(D, dt.datetime) and isinstance(o, (int, float)):
            out.append((sym, D, float(o)))
    return out


def features(sym, D):
    nm = EB.names(); name = nm.get(sym, sym)
    n_pre = n_react = pos = neg = gov = 0
    tot_words = 0
    for a in news_ingest.articles_for(sym):
        try:
            pub = parsedate_to_datetime(a["published_at"]).replace(tzinfo=None)
        except Exception:                                  # noqa: BLE001
            continue
        if pub >= D:
            continue
        n_pre += 1
        t = a["title"]
        if EB.REACTIVE.search(t):
            n_react += 1
            continue
        if EB._attribution(sym, name, t) != "high":
            continue
        tot_words += 1
        if EB.POS.search(t):
            pos += 1
        if EB.NEG.search(t):
            neg += 1
        for rx in EB.GOVERNANCE.values():
            if rx.search(t):
                gov += 1
    net_tilt = (pos - neg) / max(1, pos + neg)             # de-saturated: only opinionated
    pos_frac = pos / tot_words if tot_words else 0.0
    react_frac = n_react / n_pre if n_pre else 0.0
    return {"buzz": n_pre, "react_frac": round(react_frac, 3),
            "net_tilt": round(net_tilt, 3), "pos_frac": round(pos_frac, 3),
            "gov": 1 if gov else 0}


def corr(xs, ys):
    if len(xs) < 3:
        return None
    mx, my = st.mean(xs), st.mean(ys)
    cov = sum((x - mx) * (y - my) for x, y in zip(xs, ys))
    dx = sum((x - mx) ** 2 for x in xs) ** .5
    dy = sum((y - my) ** 2 for y in ys) ** .5
    return round(cov / (dx * dy), 3) if dx and dy else None


FEATURE_DESC = {
    "pos_frac": "positive-word fraction of company-specific headlines",
    "react_frac": "share of coverage that is reactive (price-move) news",
    "net_tilt": "de-saturated sentiment tilt (pos−neg among opinionated)",
    "buzz": "pre-event news volume (attention)",
    "gov": "governance flag present (0/1)",
}


def compute(fetch_missing: bool = True) -> dict:
    events = realised_events()
    if fetch_missing:
        missing = sorted({s for s, _, _ in events if not news_ingest.articles_for(s)})
        if missing:
            news_ingest.ingest(missing)
    data = [(sym, features(sym, D), ret) for sym, D, ret in events]
    moves = [ret for *_, ret in data]

    ranked = []
    for f in ["buzz", "react_frac", "net_tilt", "pos_frac", "gov"]:
        c = corr([d[1][f] for d in data], moves)
        strength = ("—" if c is None else "none" if abs(c) < 0.2
                    else "weak" if abs(c) < 0.4 else "moderate")
        ranked.append({"feature": f, "corr": c, "strength": strength,
                       "desc": FEATURE_DESC[f]})
    ranked.sort(key=lambda x: -(abs(x["corr"]) if x["corr"] is not None else 0))

    gpos = [ret for _, fe, ret in data if fe["gov"]]
    gneg = [ret for _, fe, ret in data if not fe["gov"]]
    ds = sorted(data, key=lambda d: d[1]["buzz"])
    k = len(ds) // 3 or 1
    lowb = [abs(r) for *_, r in ds[:k]]
    hib = [abs(r) for *_, r in ds[-k:]]
    return {
        "n": len(data), "avg_move": round(st.mean(moves), 2),
        "features": ranked,
        "governance": {"present": {"n": len(gpos), "avg": round(st.mean(gpos), 2) if gpos else None},
                       "absent": {"n": len(gneg), "avg": round(st.mean(gneg), 2) if gneg else None}},
        "buzz_split": {"low": {"n": len(lowb), "abs_move": round(st.mean(lowb), 2)},
                       "high": {"n": len(hib), "abs_move": round(st.mean(hib), 2)}},
        "note": ("PROXY (published_at) · one quarter · headline-based. Best flicker (pos_frac) "
                 "is weak and likely momentum/anticipation, not tradeable edge."),
    }


def main():
    d = compute()
    print(f"analysed {d['n']} realised events · avg move {d['avg_move']:+.2f}%\n")
    print(f"{'feature':12}{'corr w/ move':>14}   interpretation")
    print("-" * 60)
    for f in d["features"]:
        cs = "—" if f["corr"] is None else f"{f['corr']:+.3f}"
        print(f"{f['feature']:12}{cs:>14}   {f['strength']}")
    g = d["governance"]
    print(f"\ngovernance present: n={g['present']['n']} avg {g['present']['avg']}%  |  "
          f"absent: n={g['absent']['n']} avg {g['absent']['avg']}%")
    b = d["buzz_split"]
    print(f"|move| buzz LOW (n={b['low']['n']}): {b['low']['abs_move']}%  |  "
          f"buzz HIGH (n={b['high']['n']}): {b['high']['abs_move']}%")


if __name__ == "__main__":
    main()
